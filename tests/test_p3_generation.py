from __future__ import annotations

import hashlib
import json
import os
import sys
import subprocess
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "tests"))

from au_hikari_conditions.p2 import (
    calculate_workbook,
    compute_record_hash,
    write_report,
)
from au_hikari_conditions.p3 import (
    P3GenerationError,
    generate_distribution_candidates,
    generate_agency_documents,
)
from test_p2_integration import AGENCY_ID, _build_fixture


def _fixture_with_result(
    directory: Path,
    *,
    unknown_tax: bool = False,
    withhold: bool = False,
    missing_isp: bool = False,
) -> tuple[Path, Path]:
    input_path = directory / "approved-input.xlsx"
    result_path = directory / "p2-result.xlsx"
    _build_fixture(input_path)
    if unknown_tax or withhold or missing_isp:
        workbook = load_workbook(input_path)
        if unknown_tax:
            worksheet = workbook["粗利原価"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(3, headers.index("direct_cost_tax_basis") + 1, "UNKNOWN")
        if withhold:
            worksheet = workbook["粗利原価"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(3, headers.index("display_policy") + 1, "WITHHOLD")
        if missing_isp:
            worksheet = workbook["条件対応"]
            headers = [cell.value for cell in worksheet[1]]
            mapping_headers = [cell.value for cell in worksheet[1]]
            mapping_row = next(
                row_number
                for row_number in range(2, worksheet.max_row + 1)
                if worksheet.cell(
                    row_number,
                    mapping_headers.index("mapping_id") + 1,
                ).value
                == "mapping-2026-11"
            )
            worksheet.cell(mapping_row, headers.index("isp") + 1).value = None
            mapping_headers = [cell.value for cell in worksheet[1]]
            mapping_record = {
                header: worksheet.cell(mapping_row, index + 1).value
                for index, header in enumerate(mapping_headers)
                if header
            }
            approvals = workbook["承認管理"]
            approval_headers = [cell.value for cell in approvals[1]]
            for row_number in range(2, approvals.max_row + 1):
                if approvals.cell(row_number, approval_headers.index("record_id") + 1).value == "mapping-2026-11":
                    approvals.cell(
                        row_number,
                        approval_headers.index("record_hash") + 1,
                        compute_record_hash("条件対応", mapping_record),
                    )
        if unknown_tax or withhold:
            worksheet = workbook["粗利原価"]
            gross_headers = [cell.value for cell in worksheet[1]]
            gross_record = {
                header: worksheet.cell(3, index + 1).value
                for index, header in enumerate(gross_headers)
                if header
            }
            approvals = workbook["承認管理"]
            approval_headers = [cell.value for cell in approvals[1]]
            for row_number in range(2, approvals.max_row + 1):
                if approvals.cell(row_number, approval_headers.index("record_id") + 1).value == "gross-2026-11":
                    approvals.cell(
                        row_number,
                        approval_headers.index("record_hash") + 1,
                        compute_record_hash("粗利原価", gross_record),
                    )
        workbook.save(input_path)
        workbook.close()
    report = calculate_workbook(input_path, "2026-11")
    write_report(report, result_path)
    return input_path, result_path


class P3GenerationTest(unittest.TestCase):
    def test_creates_clean_review_candidate_from_structured_p2_data(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory)
            candidate_dir = directory / "approval-candidates"

            report = generate_distribution_candidates(
                approved_input_path=input_path,
                p2_result_path=p2_path,
                target_period="2026-11",
                output_directory=candidate_dir,
            )

            self.assertEqual(
                report.status,
                "APPROVAL_CANDIDATES_NOT_APPROVED_FOR_DELIVERY",
            )
            artifact = report.artifacts[0]
            workbook = load_workbook(
                artifact.spreadsheet_path,
                read_only=True,
                data_only=True,
            )
            try:
                sheet = workbook["条件表"]
                self.assertIsNone(sheet["A1"].value)
                self.assertNotIn("INTERNAL_REVIEW_ONLY", str(sheet["A4"].value))
                self.assertEqual(sheet["A3"].value, "2026-11 条件表")
            finally:
                workbook.close()
            word_text = "\n".join(paragraph.text for paragraph in Document(artifact.word_path).paragraphs)
            self.assertNotIn("内部確認用", word_text)
            self.assertNotIn("未承認", word_text)

    def test_generates_and_reads_back_anonymous_excel_and_word_forms(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory)
            output = directory / "deliverables"

            report = generate_agency_documents(
                approved_input_path=input_path,
                p2_result_path=p2_path,
                target_period="2026-11",
                output_directory=output,
            )

            self.assertEqual(report.status, "INTERNAL_REVIEW_NOT_APPROVED_FOR_DELIVERY")
            self.assertEqual(len(report.artifacts), 1)
            artifact = report.artifacts[0]
            self.assertEqual(artifact.agency_id, AGENCY_ID)
            self.assertEqual(artifact.condition_count, 1)
            self.assertTrue(artifact.spreadsheet_path.exists())
            self.assertTrue(artifact.word_path.exists())

            workbook = load_workbook(
                artifact.spreadsheet_path, read_only=False, data_only=True
            )
            try:
                sheet = workbook["条件表"]
                self.assertIn("内部確認用", sheet["A1"].value)
                self.assertIn("Anonymous Agency", sheet["A2"].value)
                self.assertEqual(sheet["A3"].value, "2026-11 条件表")
                self.assertEqual(sheet["A7"].value, "anon-item-01")
                self.assertEqual(sheet["D7"].value, 70)
                self.assertEqual(sheet["E7"].value, "掲載")
                self.assertIn("ISP: synthetic-isp", sheet["C7"].value)
                self.assertEqual(sheet["D7"].number_format, "#,##0.##;[Red]-#,##0.##")
                self.assertEqual(sheet.page_setup.fitToWidth, 1)
            finally:
                workbook.close()

            word = Document(artifact.word_path)
            word_text = "\n".join(paragraph.text for paragraph in word.paragraphs)
            self.assertIn("内部確認用", word_text)
            self.assertIn("Anonymous Agency", word_text)
            self.assertEqual(len(word.tables), 1)
            rows = {
                row.cells[0].text: row.cells[1].text
                for row in word.tables[0].rows
            }
            self.assertEqual(rows["ISP"], "synthetic-isp")
            self.assertEqual(rows["代理店報酬（税抜）"], "70 円")
            self.assertEqual(rows["代理店向け掲載可否"], "掲載")

    def test_held_p2_results_block_all_artifact_generation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory, unknown_tax=True)
            output = directory / "deliverables"

            with self.assertRaises(P3GenerationError):
                generate_agency_documents(
                    approved_input_path=input_path,
                    p2_result_path=p2_path,
                    target_period="2026-11",
                    output_directory=output,
                )
            self.assertFalse(output.exists())

    def test_modified_p2_report_is_rejected_before_output(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory)
            workbook = load_workbook(p2_path)
            worksheet = workbook["代理店別計算"]
            worksheet.cell(2, 9, 999)
            workbook.save(p2_path)
            workbook.close()
            output = directory / "deliverables"

            with self.assertRaises(P3GenerationError):
                generate_agency_documents(
                    approved_input_path=input_path,
                    p2_result_path=p2_path,
                    target_period="2026-11",
                    output_directory=output,
                )
            self.assertFalse(output.exists())

    def test_withheld_compensation_is_not_rendered_as_an_amount(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory, withhold=True)
            output = directory / "deliverables"
            report = generate_agency_documents(
                approved_input_path=input_path,
                p2_result_path=p2_path,
                target_period="2026-11",
                output_directory=output,
            )
            artifact = report.artifacts[0]

            workbook = load_workbook(
                artifact.spreadsheet_path, read_only=True, data_only=True
            )
            try:
                self.assertEqual(
                    workbook["条件表"]["D7"].value, "非提示（社内条件）"
                )
                self.assertEqual(workbook["条件表"]["E7"].value, "非提示")
            finally:
                workbook.close()
            word = Document(artifact.word_path)
            reward_row = next(
                row
                for row in word.tables[0].rows
                if row.cells[0].text == "代理店報酬（税抜）"
            )
            self.assertEqual(reward_row.cells[1].text, "非提示（社内条件）")

    def test_missing_approved_condition_display_field_blocks_generation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory, missing_isp=True)

            with self.assertRaises(P3GenerationError):
                generate_agency_documents(
                    approved_input_path=input_path,
                    p2_result_path=p2_path,
                    target_period="2026-11",
                    output_directory=directory / "deliverables",
                )

    def test_existing_output_files_are_never_overwritten(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory)
            output = directory / "deliverables"
            output.mkdir()
            filename_key = AGENCY_ID + "_" + hashlib.sha256(
                AGENCY_ID.encode("utf-8")
            ).hexdigest()[:8]
            existing = output / f"2026-11_{filename_key}_conditions.xlsx"
            existing.write_bytes(b"keep")

            with self.assertRaises(FileExistsError):
                generate_agency_documents(
                    approved_input_path=input_path,
                    p2_result_path=p2_path,
                    target_period="2026-11",
                    output_directory=output,
                )
            self.assertEqual(existing.read_bytes(), b"keep")

    def test_command_line_generates_forms_outside_the_repository(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            directory = Path(temporary_directory)
            input_path, p2_path = _fixture_with_result(directory)
            output = directory / "deliverables"
            completed = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "scripts" / "run_p3_generation.py"),
                    "--input",
                    str(input_path),
                    "--p2-result",
                    str(p2_path),
                    "--target-period",
                    "2026-11",
                    "--output-directory",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
                encoding="utf-8",
                env={**os.environ, "PYTHONIOENCODING": "utf-8"},
            )
            self.assertEqual(completed.returncode, 0, completed.stderr)
            summary = json.loads(completed.stdout)
            self.assertEqual(summary["status"], "INTERNAL_REVIEW_NOT_APPROVED_FOR_DELIVERY")
            self.assertEqual(summary["agency_count"], 1)
            self.assertEqual(len(list(output.glob("*.xlsx"))), 1)
            self.assertEqual(len(list(output.glob("*.docx"))), 1)


if __name__ == "__main__":
    unittest.main()
