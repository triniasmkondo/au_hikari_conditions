from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

import prepare_input_workbooks as workbooks
from au_hikari_conditions.document_reader import EvidenceBlock
from au_hikari_conditions.normalization import ConditionDimension, NormalizedField


class InputWorkbookTest(unittest.TestCase):
    def test_approved_input_template_has_required_sheets_and_no_default_amounts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "input.xlsx"
            workbooks.create_approved_input_template(path)
            workbook = load_workbook(path, data_only=False)
            try:
                self.assertEqual(
                    set(workbook.sheetnames),
                    {
                        "使い方",
                        "粗利原価",
                        "条件対応",
                        "代理店別例外",
                        "前回条件",
                        "承認管理",
                    },
                )
                margin = workbook["粗利原価"]
                headers = [cell.value for cell in margin[1]]
                self.assertIn("approved_target_gross_margin", headers)
                self.assertIn("attributable_direct_cost", headers)
                self.assertIn("display_policy", headers)
                self.assertIsNone(
                    margin.cell(2, headers.index("upstream_amount") + 1).value
                )
                self.assertIsNone(
                    margin.cell(2, headers.index("approved_target_gross_margin") + 1).value
                )
                self.assertIsNone(
                    margin.cell(2, headers.index("attributable_direct_cost") + 1).value
                )
                mapping_headers = [
                    cell.value for cell in workbook["条件対応"][1]
                ]
                self.assertIn("period_rule", mapping_headers)
                previous_headers = [
                    cell.value for cell in workbook["前回条件"][1]
                ]
                self.assertIn("previous_agency_reward", previous_headers)
                self.assertGreaterEqual(len(margin.data_validations.dataValidation), 5)
                self.assertTrue(workbook.calculation.fullCalcOnLoad)
            finally:
                workbook.close()

    def test_delivery_template_contains_only_nonroutable_fictional_examples(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "delivery.xlsx"
            workbooks.create_delivery_template(path)
            workbook = load_workbook(path, data_only=False)
            try:
                self.assertEqual(
                    set(workbook.sheetnames),
                    {
                        "使い方",
                        "送付先マスター",
                        "本文テンプレート",
                        "月次承認",
                        "送信履歴",
                        "承認失効ルール",
                    },
                )
                recipient = workbook["送付先マスター"]
                self.assertEqual(recipient["C2"].value, "SAMPLE_ONLY_NOT_APPROVED")
                self.assertTrue(recipient["D2"].value.endswith("@example.invalid"))
                self.assertTrue(recipient["G2"].value.endswith("@example.invalid"))
                self.assertEqual(workbook["月次承認"]["D2"].value, "SAMPLE_ONLY_NOT_APPROVED")
                self.assertEqual(workbook["送信履歴"]["F2"].value, "NOT_SENT_SAMPLE_ONLY")
                self.assertEqual(
                    workbook["承認失効ルール"].max_row,
                    len(workbooks.DELIVERY_RULES) + 2,
                )
            finally:
                workbook.close()

    def test_extraction_candidates_keep_evidence_and_unapproved_status(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            agency_root = root / "03_条件表" / "2026-10"
            upstream_root = root / "01_上位原本" / "2026-11"
            agency_root.mkdir(parents=True)
            upstream_root.mkdir(parents=True)
            agency_file = agency_root / "【202610】【Agency A 御中】条件.xlsx"
            upstream_file = upstream_root / "November notice.pdf"
            agency_file.write_bytes(b"synthetic agency")
            upstream_file.write_bytes(b"synthetic upstream")

            def read_mock(path: Path, *, project_root: Path) -> list[EvidenceBlock]:
                return [
                    EvidenceBlock(
                        source_path=path.relative_to(project_root).as_posix(),
                        location="page:1" if path.suffix == ".pdf" else "sheet:Terms;cells:A1",
                        text="報奨金額 18,000円/件",
                    )
                ]

            def normalize_mock(blocks: list[EvidenceBlock]) -> list[NormalizedField]:
                block = blocks[0]
                return [
                    NormalizedField(
                        item_id="p1.commission_amount",
                        raw_value=block.text,
                        normalized_values=("18000/件",),
                        source_path=block.source_path,
                        location=block.location,
                        status="hold",
                        hold_reasons=("未承認候補です。",),
                        condition_dimensions=(
                            ConditionDimension("isp", ("BIGLOBE",)),
                        ),
                    )
                ]

            with (
                patch.object(workbooks, "read_document", side_effect=read_mock),
                patch.object(workbooks, "normalize_blocks", side_effect=normalize_mock),
            ):
                output = root / "extraction.xlsx"
                counts = workbooks.create_extraction_workbook(root, output)

            workbook = load_workbook(output, read_only=True, data_only=True)
            try:
                candidates = workbook["抽出候補"]
                candidate_rows = list(
                    candidates.iter_rows(min_row=2, values_only=True)
                )
                agency_candidate = next(
                    row for row in candidate_rows if row[1] == "Agency A"
                )
                self.assertEqual(counts["source_document_count"], 2)
                self.assertEqual(counts["agency_document_count"], 1)
                self.assertEqual(counts["agency_count"], 1)
                self.assertEqual(counts["candidate_count"], 2)
                self.assertEqual(counts["read_error_count"], 0)
                self.assertEqual(agency_candidate[1], "Agency A")
                self.assertEqual(agency_candidate[7], '["18000/件"]')
                self.assertIn("sheet:Terms;cells:A1", agency_candidate[9])
                self.assertIn("18,000円", agency_candidate[10])
                self.assertEqual(agency_candidate[11], "NOT_REVIEWED")
                self.assertEqual(agency_candidate[12], "UNAPPROVED")
                self.assertIsNone(agency_candidate[13])
                self.assertIsNone(agency_candidate[14])
                self.assertEqual(counts["unknown_item_count"], 8)
                unknown_rows = list(
                    workbook["判断待ち一覧"].iter_rows(min_row=2, values_only=True)
                )
                self.assertTrue(all(row[2] == "NEEDS_REVIEW" for row in unknown_rows))
                self.assertTrue(all(row[5] == "UNAPPROVED" for row in unknown_rows))
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
