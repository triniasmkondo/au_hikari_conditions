from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

import prepare_input_workbooks as templates
from au_hikari_conditions.p2 import (
    P2InputError,
    calculate_workbook,
    compute_record_hash,
    write_report,
)


APPROVED_AT = "2026-10-29T13:00:00+09:00"
AGENCY_ID = "anon-agency-01"
ITEM_ID = "anon-item-01"


def _mapping(mapping_id: str, month: str, trigger: str = "activation") -> dict[str, object]:
    year, month_number = (int(part) for part in month.split("-"))
    from calendar import monthrange

    last_day = monthrange(year, month_number)[1]
    return {
        "mapping_id": mapping_id,
        "source_document": "synthetic-conditions.pdf",
        "source_location": "page:1",
        "source_candidate_id": f"candidate-{mapping_id}",
        "upstream_item_id": "anon-upstream-item",
        "agency_id": AGENCY_ID,
        "agency_item_id": ITEM_ID,
        "product_family": "internet",
        "plus_type": "plus1",
        "isp": "synthetic-isp",
        "dwelling_type": "apartment",
        "speed": "1g",
        "access_type": "new",
        "trigger": trigger,
        "option_code": "NOT_APPLICABLE",
        "campaign_code": "NOT_APPLICABLE",
        "effective_from": f"{year:04d}-{month_number:02d}-01",
        "effective_until": f"{year:04d}-{month_number:02d}-{last_day:02d}",
        "period_rule": "monthly",
        "payment_timing": "following_month",
        "exclusions": "duplicate_application",
        "reversal_terms": "early_cancellation",
        "equivalence_review_status": "REVIEWED_EQUIVALENT",
        "approval_status": "APPROVED",
        "approved_by": "synthetic-approver",
        "approved_at": APPROVED_AT,
        "review_reason": "synthetic test approval",
    }


def _gross_row(record_id: str, month: str, upstream: str) -> dict[str, object]:
    year, month_number = (int(part) for part in month.split("-"))
    from calendar import monthrange

    last_day = monthrange(year, month_number)[1]
    return {
        "record_id": record_id,
        "agency_id": AGENCY_ID,
        "agency_name": "Anonymous Agency",
        "condition_item_id": ITEM_ID,
        "condition_version": month,
        "mapping_id": f"mapping-{month}",
        "upstream_amount": upstream,
        "upstream_amount_state": "VALUE",
        "upstream_tax_basis": "EXCLUSIVE",
        "approved_target_gross_margin": "30",
        "gross_margin_state": "VALUE",
        "gross_margin_tax_basis": "EXCLUSIVE",
        "display_policy": "PRESENT",
        "direct_cost_component": "synthetic-direct-cost",
        "direct_cost_description": "anonymous regression fixture",
        "cost_allocation_basis": "per_application",
        "attributable_direct_cost": "10",
        "direct_cost_state": "VALUE",
        "direct_cost_tax_basis": "EXCLUSIVE",
        "minimum_gross_margin": "20",
        "valid_from": f"{year:04d}-{month_number:02d}-01",
        "valid_until": f"{year:04d}-{month_number:02d}-{last_day:02d}",
        "source_document": f"synthetic-{month}-gross.pdf",
        "evidence_location": "page:1",
        "approval_status": "APPROVED",
        "approved_by": "synthetic-approver",
        "approved_at": APPROVED_AT,
        "approval_reason": "synthetic test approval",
    }


def _previous_row(record_id: str, target_period: str, previous_reward: str) -> dict[str, object]:
    year, month = (int(part) for part in target_period.split("-"))
    previous_year, previous_month = (year - 1, 12) if month == 1 else (year, month - 1)
    return {
        "record_id": record_id,
        "target_period": target_period,
        "agency_id": AGENCY_ID,
        "condition_item_id": ITEM_ID,
        "mapping_id": f"mapping-{previous_year:04d}-{previous_month:02d}",
        "condition_version": "synthetic-previous",
        "previous_agency_reward": previous_reward,
        "amount_state": "VALUE",
        "tax_basis": "EXCLUSIVE",
        "source_document": f"synthetic-previous-{target_period}.pdf",
        "evidence_location": "page:1",
        "approval_status": "APPROVED",
        "approved_by": "synthetic-approver",
        "approved_at": APPROVED_AT,
        "approval_reason": "synthetic test approval",
    }


def _write_record(worksheet, record: dict[str, object]) -> int:
    headers = [cell.value for cell in worksheet[1]]
    row_number = worksheet.max_row + 1
    for column, header in enumerate(headers, 1):
        if header in record:
            worksheet.cell(row_number, column, record[header])
    return row_number


def _add_approval(worksheet, record_type: str, record: dict[str, object]) -> None:
    record_id = (
        record.get("record_id")
        or record.get("mapping_id")
        or record.get("exception_id")
    )
    approval_record = {
        "approval_id": f"approval-{record_id}",
        "batch_id": "synthetic-batch",
        "agency_id": AGENCY_ID,
        "record_type": record_type,
        "record_id": record_id,
        "record_hash": compute_record_hash(record_type, record),
        "approval_status": "APPROVED",
        "requested_by": "synthetic-requester",
        "requested_at": APPROVED_AT,
        "approved_by": record["approved_by"],
        "approved_at": record["approved_at"],
        "decision_reason": "synthetic test approval",
        "supersedes_approval_id": None,
    }
    _write_record(worksheet, approval_record)


def _build_fixture(path: Path) -> None:
    templates.create_approved_input_template(path)
    workbook = load_workbook(path)
    approvals = workbook["承認管理"]
    records: list[tuple[str, dict[str, object]]] = []

    for month, upstream in (("2026-10", "100"), ("2026-11", "110")):
        record = _gross_row(f"gross-{month}", month, upstream)
        _write_record(workbook["粗利原価"], record)
        records.append(("粗利原価", record))
    for month in ("2026-09", "2026-10", "2026-11"):
        record = _mapping(f"mapping-{month}", month)
        _write_record(workbook["条件対応"], record)
        records.append(("条件対応", record))
    for target_period, previous_amount in (("2026-10", "40"), ("2026-11", "60")):
        record = _previous_row(f"previous-{target_period}", target_period, previous_amount)
        _write_record(workbook["前回条件"], record)
        records.append(("前回条件", record))
    for record_type, record in records:
        _add_approval(approvals, record_type, record)
    workbook.save(path)
    workbook.close()


class P2WorkbookIntegrationTest(unittest.TestCase):
    def test_anonymous_october_reproduction_and_november_update(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "approved-input.xlsx"
            _build_fixture(input_path)

            october = calculate_workbook(input_path, "2026-10")
            november = calculate_workbook(input_path, "2026-11")

            self.assertEqual(october.proposed_count, 1)
            self.assertEqual(october.hold_count, 0)
            self.assertEqual(october.results[0].proposed_agency_reward, 60)
            self.assertEqual(october.results[0].proposed_difference, 20)
            self.assertEqual(november.proposed_count, 1)
            self.assertEqual(november.hold_count, 0)
            self.assertEqual(november.results[0].proposed_agency_reward, 70)
            self.assertEqual(november.results[0].previous_agency_reward, 60)
            self.assertEqual(november.results[0].proposed_difference, 10)

            result_path = Path(temporary_directory) / "p2-result.xlsx"
            write_report(november, result_path)
            result_book = load_workbook(result_path, read_only=True, data_only=True)
            try:
                row = next(result_book["代理店別計算"].iter_rows(min_row=2, values_only=True))
                self.assertEqual(row[3], "proposed")
                self.assertEqual(row[8], 70)
                self.assertEqual(row[9], 10)
                self.assertEqual(result_book["保留理由"].max_row, 1)
            finally:
                result_book.close()

    def test_unapproved_record_never_exposes_a_proposed_amount(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "unapproved.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            worksheet = workbook["粗利原価"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(3, headers.index("approval_status") + 1, "PENDING")
            workbook.save(input_path)
            workbook.close()

            report = calculate_workbook(input_path, "2026-11")
            result = report.results[0]
            self.assertEqual(result.status, "hold")
            self.assertIsNone(result.proposed_agency_reward)
            self.assertIsNone(result.proposed_difference)
            self.assertIsNone(result.upstream_reward)
            self.assertTrue(any("未承認" in reason for reason in result.hold_reasons))

    def test_unknown_tax_basis_holds_without_emitting_money(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "unknown-tax.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            worksheet = workbook["粗利原価"]
            headers = [cell.value for cell in worksheet[1]]
            worksheet.cell(3, headers.index("direct_cost_tax_basis") + 1, "UNKNOWN")
            workbook.save(input_path)
            workbook.close()

            result = calculate_workbook(input_path, "2026-11").results[0]
            self.assertEqual(result.status, "hold")
            self.assertIsNone(result.direct_cost)
            self.assertIsNone(result.proposed_agency_reward)
            self.assertTrue(any("税区分" in reason for reason in result.hold_reasons))

    def test_approved_payment_prohibition_exception_holds_calculation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "exception.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            exception = {
                "exception_id": "exception-payment-prohibited",
                "agency_id": AGENCY_ID,
                "condition_item_id_or_rule_group": ITEM_ID,
                "priority": 1,
                "rule_type": "PAYMENT_PROHIBITED",
                "rule_value": "true",
                "display_policy": "",
                "reason": "synthetic contract exception",
                "source_document": "synthetic-exception.pdf",
                "evidence_location": "page:1",
                "approval_status": "APPROVED",
                "approved_by": "synthetic-approver",
                "approved_at": APPROVED_AT,
                "valid_from": "2026-11-01",
                "valid_until": "2026-11-30",
            }
            _write_record(workbook["代理店別例外"], exception)
            _add_approval(workbook["承認管理"], "代理店別例外", exception)
            workbook.save(input_path)
            workbook.close()

            result = calculate_workbook(input_path, "2026-11").results[0]
            self.assertEqual(result.status, "hold")
            self.assertIsNone(result.proposed_agency_reward)
            self.assertTrue(any("支払禁止" in reason for reason in result.hold_reasons))

    def test_direct_cost_components_are_summed_after_individual_approval(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "cost-components.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            additional = _gross_row("gross-2026-11-cost-2", "2026-11", "110")
            additional["direct_cost_component"] = "synthetic-second-cost"
            additional["attributable_direct_cost"] = "5"
            _write_record(workbook["粗利原価"], additional)
            _add_approval(workbook["承認管理"], "粗利原価", additional)
            workbook.save(input_path)
            workbook.close()

            result = calculate_workbook(input_path, "2026-11").results[0]
            self.assertEqual(result.status, "proposed")
            self.assertEqual(result.direct_cost, 15)
            self.assertEqual(result.proposed_agency_reward, 65)
            self.assertEqual(result.proposed_difference, 5)

    def test_condition_mismatch_and_changed_approval_hash_hold(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "mismatch.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            mapping_sheet = workbook["条件対応"]
            mapping_headers = [cell.value for cell in mapping_sheet[1]]
            mapping_row = next(
                row
                for row in range(2, mapping_sheet.max_row + 1)
                if mapping_sheet.cell(row, mapping_headers.index("mapping_id") + 1).value
                == "mapping-2026-11"
            )
            mapping_sheet.cell(
                mapping_row,
                mapping_headers.index("trigger") + 1,
                "application",
            )
            workbook.save(input_path)
            workbook.close()

            report = calculate_workbook(input_path, "2026-11")
            result = report.results[0]
            self.assertEqual(result.status, "hold")
            self.assertIsNone(result.proposed_agency_reward)
            self.assertTrue(
                any("ハッシュ" in reason for reason in result.hold_reasons)
            )
            self.assertTrue(
                any("一致しません" in reason for reason in result.hold_reasons)
            )

    def test_unconfirmed_condition_value_holds_even_with_matching_approval_hash(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "unconfirmed-condition.xlsx"
            _build_fixture(input_path)
            workbook = load_workbook(input_path)
            mapping_sheet = workbook["条件対応"]
            mapping_headers = [cell.value for cell in mapping_sheet[1]]
            isp_column = mapping_headers.index("isp") + 1
            mapping_row = next(
                row
                for row in range(2, mapping_sheet.max_row + 1)
                if mapping_sheet.cell(row, mapping_headers.index("mapping_id") + 1).value
                == "mapping-2026-11"
            )
            mapping_sheet.cell(mapping_row, isp_column, "UNKNOWN")
            mapping_record = {
                header: mapping_sheet.cell(mapping_row, column).value
                for column, header in enumerate(mapping_headers, 1)
                if header
            }

            approval_sheet = workbook["承認管理"]
            approval_headers = [cell.value for cell in approval_sheet[1]]
            approval_row = next(
                row
                for row in range(2, approval_sheet.max_row + 1)
                if approval_sheet.cell(
                    row, approval_headers.index("record_id") + 1
                ).value
                == "mapping-2026-11"
            )
            approval_sheet.cell(
                approval_row,
                approval_headers.index("record_hash") + 1,
                compute_record_hash("条件対応", mapping_record),
            )
            workbook.save(input_path)
            workbook.close()

            result = calculate_workbook(input_path, "2026-11").results[0]
            self.assertEqual(result.status, "hold")
            self.assertIsNone(result.proposed_agency_reward)
            self.assertTrue(any("未確定値" in reason for reason in result.hold_reasons))

    def test_missing_p2_schema_is_an_explicit_input_error(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "old-template.xlsx"
            templates.create_approved_input_template(input_path)
            workbook = load_workbook(input_path)
            workbook.remove(workbook["前回条件"])
            workbook.save(input_path)
            workbook.close()

            with self.assertRaises(P2InputError):
                calculate_workbook(input_path, "2026-11")

    def test_report_writer_refuses_to_overwrite(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "result.xlsx"
            path.write_bytes(b"existing")
            with self.assertRaises(FileExistsError):
                write_report(
                    calculate_workbook(
                        self._fixture_input(Path(temporary_directory)),
                        "2026-11",
                    ),
                    path,
                )

    def test_command_line_runs_from_approved_input_to_result_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            input_path = root / "approved-input.xlsx"
            output_path = root / "p2-result.xlsx"
            _build_fixture(input_path)

            completed = subprocess.run(
                [
                    sys.executable,
                    str(PROJECT_ROOT / "scripts" / "run_p2_calculation.py"),
                    "--input",
                    str(input_path),
                    "--target-period",
                    "2026-11",
                    "--output",
                    str(output_path),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(output_path.exists())
            summary = json.loads(completed.stdout)
            self.assertEqual(summary["proposed_count"], 1)
            self.assertEqual(summary["hold_count"], 0)

    @staticmethod
    def _fixture_input(directory: Path) -> Path:
        path = directory / "input.xlsx"
        _build_fixture(path)
        return path


if __name__ == "__main__":
    unittest.main()
