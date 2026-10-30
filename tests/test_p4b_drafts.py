from __future__ import annotations

import sqlite3
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from au_hikari_conditions.p3 import GeneratedAgencyArtifacts
from au_hikari_conditions.p4b import (
    AgencyMailSettings,
    DraftHistoryStore,
    DraftPreparationError,
    DraftPreview,
    DuplicateDraftError,
    _file_hash,
    _message_hash,
    _recipient_snapshot_hash,
    prepare_draft_preview,
    save_draft,
)
from au_hikari_conditions.workflow import (
    AgencyWorkflow,
    BatchStatus,
    DeliverySnapshot,
)


AGENCY_ID = "anonymous-agency-01"
AGENCY_NAME = "匿名代理店"
PERIOD = "2026-11"
FROM = "sender@example.invalid"
TO = ("recipient@example.invalid",)
SUBJECT_TEMPLATE = "{target_period} 条件表のご案内"
BODY_TEMPLATE = "{agency_name} 御中\n{target_period}分の条件表をご確認ください。"


class FakeMailClient:
    def __init__(
        self,
        *,
        save_error: Exception | None = None,
        upload_error: Exception | None = None,
        after_upload=None,
        attachment_name_override: str | None = None,
    ) -> None:
        self.save_error = save_error
        self.upload_error = upload_error
        self.after_upload = after_upload
        self.attachment_name_override = attachment_name_override
        self.uploads: list[tuple[str, bytes, str]] = []
        self.saved_payloads: list[dict[str, object]] = []

    def upload_draft_attachment(
        self, *, file_name: str, content: bytes, content_type: str
    ) -> dict[str, str]:
        if self.upload_error is not None:
            raise self.upload_error
        self.uploads.append((file_name, content, content_type))
        if self.after_upload is not None:
            callback, self.after_upload = self.after_upload, None
            callback()
        return {
            "storeName": f"store-{len(self.uploads)}",
            "attachmentName": self.attachment_name_override or file_name,
            "attachmentPath": f"/attachments/{len(self.uploads)}",
        }

    def save_draft(self, payload: dict[str, object]) -> dict[str, object]:
        self.saved_payloads.append(payload)
        if self.save_error is not None:
            raise self.save_error
        return {"status": {"code": 200}}


class P4BDraftTest(unittest.TestCase):
    def setUp(self) -> None:
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.artifacts = self._create_artifacts()
        self.settings = AgencyMailSettings(
            agency_id=AGENCY_ID,
            from_address=FROM,
            to_addresses=TO,
            cc_addresses=("cc@example.invalid",),
            subject_template=SUBJECT_TEMPLATE,
            body_template=BODY_TEMPLATE,
        )
        self.snapshot = self._create_snapshot()
        self.workflow = self._approved_workflow()
        self.history = DraftHistoryStore(self.root / "history.sqlite3")

    def tearDown(self) -> None:
        self.directory.cleanup()

    def _create_artifacts(self) -> GeneratedAgencyArtifacts:
        spreadsheet_path = self.root / "approved-conditions.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "条件表"
        worksheet["A1"] = f"{AGENCY_NAME} 御中"
        worksheet["A2"] = f"{PERIOD} 条件表"
        worksheet["A3"] = f"代理店ID: {AGENCY_ID}"
        worksheet["A5"] = "匿名条件行"
        workbook.save(spreadsheet_path)
        workbook.close()

        word_path = self.root / "approved-conditions.docx"
        document = Document()
        document.add_paragraph(f"{PERIOD} 条件表")
        document.add_paragraph(f"{AGENCY_NAME} 御中")
        document.add_paragraph(f"代理店ID: {AGENCY_ID}")
        document.add_paragraph("匿名条件行")
        document.save(word_path)
        return GeneratedAgencyArtifacts(
            agency_id=AGENCY_ID,
            agency_name=AGENCY_NAME,
            spreadsheet_path=spreadsheet_path,
            word_path=word_path,
            condition_count=1,
        )

    def _create_snapshot(self) -> DeliverySnapshot:
        subject = SUBJECT_TEMPLATE.format(target_period=PERIOD)
        body = BODY_TEMPLATE.format(agency_name=AGENCY_NAME, target_period=PERIOD)
        return DeliverySnapshot(
            batch_id="anonymous-batch-2026-11",
            agency_id=AGENCY_ID,
            target_period=PERIOD,
            source_master_hash="anonymous-master-hash",
            template_hash="anonymous-template-hash",
            artifact_hashes=(
                _file_hash(self.artifacts.spreadsheet_path),
                _file_hash(self.artifacts.word_path),
            ),
            recipient_snapshot_hash=_recipient_snapshot_hash(self.settings),
            message_hash=_message_hash(FROM, subject, body),
        )

    def _approved_workflow(self) -> AgencyWorkflow:
        workflow = AgencyWorkflow(self.snapshot.batch_id, AGENCY_ID)
        for status in (
            BatchStatus.PARSED,
            BatchStatus.VALIDATED,
            BatchStatus.PROPOSED,
            BatchStatus.GENERATED,
            BatchStatus.QA_PASSED,
            BatchStatus.PENDING_APPROVAL,
            BatchStatus.APPROVED,
        ):
            workflow.transition(
                status,
                actor_id="anonymous-approver",
                snapshot=self.snapshot if status == BatchStatus.APPROVED else None,
            )
        return workflow

    def _preview(self) -> DraftPreview:
        return prepare_draft_preview(
            artifacts=self.artifacts,
            target_period=PERIOD,
            settings=self.settings,
            workflow=self.workflow,
            snapshot=self.snapshot,
        )

    def test_preview_and_mock_draft_save_with_two_approved_attachments(self) -> None:
        preview = self._preview()
        self.assertIn(PERIOD, preview.subject)
        self.assertEqual(preview.target_period, PERIOD)
        self.assertEqual(len(preview.attachments), 2)
        self.assertNotIn("recipient@example.invalid", repr(preview))
        self.assertNotIn(AGENCY_ID, repr(self.settings))

        client = FakeMailClient()
        result = save_draft(
            preview=preview,
            mail_client=client,
            history=self.history,
        )

        self.assertEqual(result.status, "CREATED")
        self.assertEqual(len(client.uploads), 2)
        self.assertEqual(len(client.saved_payloads), 1)
        payload = client.saved_payloads[0]
        self.assertEqual(payload["mode"], "draft")
        self.assertEqual(payload["toAddress"], TO[0])
        self.assertEqual(len(payload["attachments"]), 2)
        self.assertEqual(self.history.get_status(AGENCY_ID, PERIOD), "CREATED")

        connection = sqlite3.connect(self.history.path)
        try:
            (stored_scope,) = connection.execute(
                "SELECT scope_hash FROM draft_attempts"
            ).fetchone()
        finally:
            connection.close()
        self.assertNotIn(AGENCY_ID, stored_scope)
        with self.assertRaises(DuplicateDraftError):
            save_draft(
                preview=preview,
                mail_client=client,
                history=self.history,
            )
        self.assertEqual(len(client.saved_payloads), 1)

    def test_p3_internal_or_unapproved_artifacts_are_rejected(self) -> None:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "条件表"
        worksheet["A1"] = "内部確認用・配信用成果物ではありません（未承認）"
        worksheet["A2"] = AGENCY_NAME
        worksheet["A3"] = AGENCY_ID
        worksheet["A4"] = PERIOD
        workbook.save(self.artifacts.spreadsheet_path)
        workbook.close()

        with self.assertRaisesRegex(DraftPreparationError, "internal or unapproved"):
            self._preview()

    def test_unapproved_marker_in_word_footer_is_rejected(self) -> None:
        document = Document(self.artifacts.word_path)
        document.sections[0].footer.paragraphs[0].text = "未承認"
        document.save(self.artifacts.word_path)

        with self.assertRaisesRegex(DraftPreparationError, "internal or unapproved"):
            self._preview()

    def test_workflow_approval_is_mandatory(self) -> None:
        workflow = AgencyWorkflow(self.snapshot.batch_id, AGENCY_ID)
        with self.assertRaisesRegex(DraftPreparationError, "current approval"):
            prepare_draft_preview(
                artifacts=self.artifacts,
                target_period=PERIOD,
                settings=self.settings,
                workflow=workflow,
                snapshot=self.snapshot,
            )

    def test_changed_recipient_invalidates_approved_snapshot(self) -> None:
        changed_settings = AgencyMailSettings(
            agency_id=AGENCY_ID,
            from_address=FROM,
            to_addresses=("changed@example.invalid",),
            subject_template=SUBJECT_TEMPLATE,
            body_template=BODY_TEMPLATE,
        )
        with self.assertRaisesRegex(DraftPreparationError, "recipients"):
            prepare_draft_preview(
                artifacts=self.artifacts,
                target_period=PERIOD,
                settings=changed_settings,
                workflow=self.workflow,
                snapshot=self.snapshot,
            )

    def test_missing_to_recipient_is_rejected(self) -> None:
        missing_to = AgencyMailSettings(
            agency_id=AGENCY_ID,
            from_address=FROM,
            to_addresses=(),
            subject_template=SUBJECT_TEMPLATE,
            body_template=BODY_TEMPLATE,
        )
        with self.assertRaisesRegex(DraftPreparationError, "To recipient"):
            prepare_draft_preview(
                artifacts=self.artifacts,
                target_period=PERIOD,
                settings=missing_to,
                workflow=self.workflow,
                snapshot=self.snapshot,
            )

    def test_changed_attachment_is_rejected_before_upload(self) -> None:
        preview = self._preview()
        preview.attachments[0].path.write_bytes(b"tampered attachment")
        client = FakeMailClient()
        with self.assertRaisesRegex(DraftPreparationError, "changed after preview"):
            save_draft(
                preview=preview,
                mail_client=client,
                history=self.history,
            )
        self.assertEqual(client.uploads, [])
        self.assertIsNone(self.history.get_status(AGENCY_ID, PERIOD))

    def test_non_draft_mode_stops_before_history_or_network(self) -> None:
        client = FakeMailClient()
        with self.assertRaisesRegex(DraftPreparationError, 'mode="draft"'):
            save_draft(
                preview=self._preview(),
                mail_client=client,
                history=self.history,
                mode="template",
            )
        self.assertEqual(client.uploads, [])
        self.assertEqual(client.saved_payloads, [])
        self.assertIsNone(self.history.get_status(AGENCY_ID, PERIOD))

    def test_upload_error_is_recorded_and_does_not_save_a_draft(self) -> None:
        client = FakeMailClient(upload_error=RuntimeError("sensitive detail"))
        with self.assertRaisesRegex(DraftPreparationError, "attachment upload failed") as raised:
            save_draft(
                preview=self._preview(),
                mail_client=client,
                history=self.history,
            )
        self.assertNotIn("sensitive detail", str(raised.exception))
        self.assertEqual(self.history.get_status(AGENCY_ID, PERIOD), "FAILED")
        self.assertEqual(client.saved_payloads, [])

    def test_save_error_is_outcome_unknown_and_retry_is_blocked(self) -> None:
        client = FakeMailClient(save_error=RuntimeError("sensitive detail"))
        with self.assertRaisesRegex(DraftPreparationError, "result is unknown"):
            save_draft(
                preview=self._preview(),
                mail_client=client,
                history=self.history,
            )
        self.assertEqual(
            self.history.get_status(AGENCY_ID, PERIOD),
            "OUTCOME_UNKNOWN",
        )
        with self.assertRaises(DuplicateDraftError):
            save_draft(
                preview=self._preview(),
                mail_client=FakeMailClient(),
                history=self.history,
            )

    def test_approval_expiring_after_upload_stops_before_draft_post(self) -> None:
        def expire_approval() -> None:
            self.workflow.transition(BatchStatus.HOLD, actor_id="anonymous-reviewer")

        client = FakeMailClient(after_upload=expire_approval)
        with self.assertRaisesRegex(DraftPreparationError, "Approval changed or expired"):
            save_draft(
                preview=self._preview(),
                mail_client=client,
                history=self.history,
            )
        self.assertEqual(self.history.get_status(AGENCY_ID, PERIOD), "FAILED")
        self.assertEqual(client.saved_payloads, [])

    def test_mismatched_uploaded_filename_fails_closed(self) -> None:
        client = FakeMailClient(attachment_name_override="wrong-file.docx")
        with self.assertRaisesRegex(DraftPreparationError, "did not match"):
            save_draft(
                preview=self._preview(),
                mail_client=client,
                history=self.history,
            )
        self.assertEqual(self.history.get_status(AGENCY_ID, PERIOD), "FAILED")
        self.assertEqual(client.saved_payloads, [])

    def test_malformed_template_placeholder_is_rejected(self) -> None:
        invalid_settings = AgencyMailSettings(
            agency_id=AGENCY_ID,
            from_address=FROM,
            to_addresses=TO,
            subject_template="{agency_name.upper}",
            body_template=BODY_TEMPLATE,
        )
        with self.assertRaisesRegex(DraftPreparationError, "unsupported"):
            prepare_draft_preview(
                artifacts=self.artifacts,
                target_period=PERIOD,
                settings=invalid_settings,
                workflow=self.workflow,
                snapshot=self.snapshot,
            )


class CompanySubjectRuleTest(unittest.TestCase):
    def test_company_subject_exactly_matches_confirmed_rule(self):
        from au_hikari_conditions.locked_delivery import render_company_subject
        self.assertEqual(
            render_company_subject("2026-11", "keioh株式会社"),
            "【202611】【keioh株式会社様】auひかり(AL商流)/auひかりプラス_条件表通知表",
        )

    def test_individual_and_existing_honorifics_never_duplicate(self):
        from au_hikari_conditions.locked_delivery import render_company_subject
        for name in ("高橋大地", "高橋大地様", "高橋大地 様", "高橋大地様様"):
            with self.subTest(name=name):
                self.assertEqual(
                    render_company_subject("2026-11", name),
                    "【202611】【高橋大地様】auひかり(AL商流)/auひかりプラス_条件表通知表",
                )
        self.assertEqual(
            render_company_subject("2026-11", "株式会社J&C 御中"),
            "【202611】【株式会社J&C様】auひかり(AL商流)/auひかりプラス_条件表通知表",
        )

if __name__ == "__main__":
    unittest.main()
