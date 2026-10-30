import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook,load_workbook
from test_phase2_triage import fixture
from au_hikari_conditions.phase2_triage import build_triage
from au_hikari_conditions.phase2_november import write_november
from au_hikari_conditions.phase2_review import _write_sheet,input_hash

class NovemberTest(unittest.TestCase):
    def run_policy(self,result,current,path=None):
        groups,members,_=build_triage([result],'2026-10','2026-11')
        book=Workbook()
        stats=write_november(book,[result],current,groups,members,'snapshot',path,_write_sheet)
        return book,stats

    def test_same_new_end_and_unknown_finance(self):
        p,_,r=fixture()
        b,s=self.run_policy(r,[p.upper])
        self.assertEqual(s['categories'],{'同等項目・既存粗利/個別条件引継ぎ':1})
        self.assertEqual(s['price_tasks'],1)
        self.assertIsNone(r['proposed'])
        b,s=self.run_policy(r,[])
        self.assertEqual(s['categories'],{'10月のみ・終了':1})
        p.old.terms='個別約定'
        b,s=self.run_policy(r,[])
        self.assertEqual(s['categories'],{'10月のみ・終了':1})
        self.assertEqual(s['remaining_individual'],1)

    def test_common_approval_applies_and_tampering_is_rejected(self):
        p,_,r=fixture();p.old.tax='不明'
        b,s=self.run_policy(r,[p.upper])
        sheet=b['共通一括判断入力'];headers=[c.value for c in sheet[1]]
        record=dict(zip(headers,[c.value for c in sheet[2]]))
        record.update(decision='CONFIRMED',resolution='当該対象は税抜',source_document='根拠',
                      evidence_location='1頁',approval_reason='確認済',approval_id='A1',
                      approval_status='APPROVED',approved_by='担当者',approved_at='2026-10-30T12:00:00+09:00')
        record['approved_hash']=input_hash('共通一括判断入力',record)
        for i,h in enumerate(headers,1):sheet.cell(2,i,record.get(h))
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'input.xlsx';b.save(path)
            _,s=self.run_policy(r,[p.upper],path)
            self.assertEqual(s['common_resolved'],1)
            sheet.cell(2,headers.index('resolution')+1,'改変');b.save(path)
            _,s=self.run_policy(r,[p.upper],path)
            self.assertEqual(s['common_resolved'],0)
