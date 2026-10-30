from __future__ import annotations

from dataclasses import replace
from datetime import datetime, timezone
from decimal import Decimal
import hashlib
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "scripts"))
import prepare_input_workbooks
import prepare_campaign_conditions
from au_hikari_conditions.calculation import ApprovalEvidence, ApprovalStatus
from au_hikari_conditions.campaign_import import (
    build_campaign_candidates, extract_campaign_document, write_campaign_candidate_workbook,
)
from au_hikari_conditions.campaign_workbook import (
    CAMPAIGN_HEADERS, CAMPAIGN_SHEET, evaluate_campaign_workbook, load_campaign_schedules,
)
from au_hikari_conditions.p2 import P2InputError, calculate_workbook, compute_record_hash
from au_hikari_conditions.reward_design import AgencyRewardProfile, VolumeScenario

# Explicit synthetic fixtures, matching the documented November tier shapes.
DATA = (
    ("ボリュームインセンティブ（auひかり既存・auひかりプラスタイプ①）", (500,750,1000,1500,2000,2500), (6500,9000,11500,14000,16500,17500), (5500,8000,10500,13000,15500,16500)),
    ("販促支援金（auひかり既存・auひかりプラスタイプ①）", (1800,2000,2300), (3000,5000,6500), (1000,3000,4500)),
    ("ボリュームインセンティブ（auひかりプラスタイプ②・③）", (220,440), (3000,5000), (1000,3000)),
    ("特別調整金（auひかりプラスタイプ③）", (1,), (10000,), (8000,)),
    ("ボリュームインセンティブ（auでんき）", (30,50,80), (3000,4000,5000), (3000,4000,5000)),
)


def _document(path, upper=False, previous=False):
    doc = Document()
    period = "2026年10月1日～2026年10月31日" if previous else "2026年11月1日～2026年11月30日"
    doc.add_paragraph(period)
    doc.add_paragraph("設定期間合計数が受付目標に達したとき、対応する支払いレートを適用します。")
    doc.add_paragraph("登録月内に取消・解約された案件は除きます。")
    for name, thresholds, upstream, agency in DATA:
        doc.add_paragraph(name)
        doc.add_paragraph("キャリアエントリ及び登録での商流全体ボリューム達成時にて、貴社開通案件へ増額")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "サービス名"
        table.rows[0].cells[1].text = "金額(税別)"
        for threshold, amount in zip(thresholds, upstream if upper else agency):
            if previous and name.startswith("販促"):
                amount = {1800:1000, 2000:2000, 2300:3000}[threshold]
            row = table.add_row()
            row.cells[0].text = f"{threshold:,}件～"
            row.cells[1].text = f"{amount:,}円／件"
    doc.add_paragraph("後確認手数料")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 1).text = "金額(税別)"
    table.cell(1, 0).text = "１件～"
    table.cell(1, 1).text = "－２，０００円／件"
    doc.save(path)


class CampaignWorkbookTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.upper, self.agency, self.previous = (self.root / name for name in ("upper.docx", "agency.docx", "previous.docx"))
        _document(self.upper, upper=True)
        _document(self.agency)
        _document(self.previous, previous=True)
        self.input = self.root / "input.xlsx"
        prepare_input_workbooks.create_approved_input_template(self.input)
        self.candidates = build_campaign_candidates(self.upper, self.agency, agency_id="agency-test", target_period="2026-11", previous_agency_path=self.previous)
        self.output = self.root / "candidate.xlsx"
        write_campaign_candidate_workbook(self.input, self.output, self.candidates)
        self.rows = [dict(row) for row in self.candidates.rows]
        for row in self.rows:
            group = next(i for i, (name, *_rest) in enumerate(DATA) if name.replace("（", "(").replace("）", ")").replace("①", "1").replace("②", "2").replace("③", "3") == row["campaign_name"])
            row.update(campaign_code=f"fixture-campaign-{group}", condition_item_id=f"fixture-item-{group}-{row['threshold_value']}",
                eligible_opening_classes="5G;FTTH", payment_timing="承認済み契約の回収後", exclusions="登録月取消・解約除外",
                attributable_direct_cost=0, direct_cost_state="ZERO", direct_cost_tax_basis="EXCLUSIVE",
                cost_source_document="synthetic-cost.xlsx", cost_evidence_location="sheet:cost;row:2")
        self.profile = AgencyRewardProfile("agency-test", "2026-11", Decimal(0), Decimal(0), True, True,
            ApprovalEvidence(ApprovalStatus.APPROVED, "fixture-collection", "fixture-contract", "clause:1", "fixture-reviewer", datetime(2026,10,29,tzinfo=timezone.utc)))

    def _save(self, approve=True):
        wb = load_workbook(self.output)
        ws = wb[CAMPAIGN_SHEET]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        approvals = wb["承認管理"]
        if approvals.max_row > 1:
            approvals.delete_rows(2, approvals.max_row - 1)
        headers = [c.value for c in approvals[1]]
        for row in self.rows:
            if approve:
                row.update(approval_status="APPROVED", approved_by="fixture-reviewer", approved_at="2026-10-29T12:00:00+09:00")
            ws.append([row.get(key) for key in CAMPAIGN_HEADERS])
            if approve:
                approval = dict(approval_id="approve-"+row["record_id"], agency_id="agency-test", record_type=CAMPAIGN_SHEET,
                    record_id=row["record_id"], record_hash=compute_record_hash(CAMPAIGN_SHEET,row), approval_status="APPROVED",
                    approved_by=row["approved_by"], approved_at=row["approved_at"])
                approvals.append([approval.get(key) for key in headers])
        wb.save(self.output)
        wb.close()

    def test_extracts_five_schedules_fifteen_tiers_and_separate_deductions(self):
        self.assertEqual(len(self.candidates.rows), 15)
        self.assertEqual(len({r["campaign_name"] for r in self.candidates.rows}), 5)
        self.assertEqual(self.candidates.deductions[0]["agency_unit_amount"], -2000)
        self.assertTrue(all(r["approval_status"] == "UNREVIEWED" for r in self.candidates.rows))
        self.assertTrue(all(not r["condition_item_id"] and not r["campaign_code"] for r in self.candidates.rows))
        self.assertTrue(all(r["evidence_location"] and r["agency_evidence_location"] for r in self.candidates.rows))
        self.assertTrue(all("取消" in r["eligibility_terms"] for r in self.candidates.rows))

    def test_unreviewed_candidates_cannot_enter_engine_or_normal_p2(self):
        result = evaluate_campaign_workbook(self.output, self.profile)
        self.assertEqual(result.status, "hold")
        self.assertFalse(result.campaign_proposals)
        report = calculate_workbook(self.output, "2026-11")
        self.assertEqual(report.hold_count, 15)
        self.assertEqual(report.proposed_count, 0)

    def test_confirmed_engine_validation_preserves_rates_and_margins(self):
        self._save()
        result = evaluate_campaign_workbook(self.output, self.profile)
        self.assertEqual(result.status, "candidate", result.hold_reasons)
        by_code = {}
        for row in result.campaign_proposals:
            by_code.setdefault(row.campaign_code, set()).add(row.retained_gross_margin_per_unit)
        self.assertEqual(by_code, {"fixture-campaign-0":{Decimal(1000)}, "fixture-campaign-1":{Decimal(2000)},
            "fixture-campaign-2":{Decimal(2000)}, "fixture-campaign-3":{Decimal(2000)}, "fixture-campaign-4":{Decimal(0)}})
        # Even approved tiers cannot be flattened into ordinary P2/P3 payments.
        self.assertEqual(calculate_workbook(self.output,"2026-11").hold_count, 15)

    def test_commerce_count_and_agency_count_remain_separate(self):
        self._save()
        scenario = VolumeScenario("fixture-volume", "fixture-campaign-0", 750, 4, "5G", True, True)
        result = evaluate_campaign_workbook(self.output, self.profile, volume_scenarios=(scenario,))
        outcome = next(s for s in result.scenarios if s.scenario_id == "fixture-volume")
        self.assertEqual(outcome.agency_payment, Decimal(32000))
        self.assertEqual(outcome.retained_gross_margin, Decimal(4000))
        self.assertEqual(outcome.commerce_count, 750)
        self.assertEqual(outcome.agency_eligible_count, 4)

    def test_confirmed_prices_are_not_reduced_to_hide_margin_shortfall(self):
        self._save()
        result = evaluate_campaign_workbook(self.output, replace(self.profile,target_gross_margin=Decimal(1000), minimum_gross_margin=Decimal(1000)))
        self.assertEqual(result.status, "hold")
        self.assertTrue(any("単価変更不可" in reason for reason in result.hold_reasons))

    def test_design_uses_previous_month_rather_than_confirmed_current_rates(self):
        self._save()
        result = evaluate_campaign_workbook(self.output, self.profile, mode="design")
        self.assertEqual(result.status, "candidate", result.hold_reasons)
        tier = next(p for p in result.campaign_proposals if p.campaign_code == "fixture-campaign-1" and p.threshold == 2000)
        self.assertEqual(tier.proposed_agency_unit_reward, Decimal(2000))

    def test_hash_binds_scopes_prices_terms_and_costs(self):
        self._save()
        wb = load_workbook(self.output)
        sheet = wb[CAMPAIGN_SHEET]
        sheet.cell(2, CAMPAIGN_HEADERS.index("counting_scope")+1, "altered-commerce-scope")
        wb.save(self.output)
        wb.close()
        loaded = load_campaign_schedules(self.output,"agency-test","2026-11")
        self.assertEqual(loaded.status, "hold")
        self.assertFalse(loaded.campaigns)
        self.assertTrue(any("ハッシュ" in r for r in loaded.hold_reasons))

    def test_changed_source_file_is_held_even_with_unchanged_approval_hash(self):
        self._save()
        self.upper.write_bytes(self.upper.read_bytes()+b"changed")
        loaded=load_campaign_schedules(self.output,"agency-test","2026-11")
        self.assertTrue(any("SHA-256" in r for r in loaded.hold_reasons))
        self.assertFalse(loaded.campaigns)

    def test_missing_or_duplicate_tier_does_not_silently_select_lower_rate(self):
        self.rows.pop(1)
        self._save()
        loaded=load_campaign_schedules(self.output,"agency-test","2026-11")
        self.assertTrue(any("閾値一覧" in r for r in loaded.hold_reasons))
        self.assertFalse(loaded.campaigns)

    def test_contradictory_scopes_are_held_even_if_individually_approved(self):
        self.rows[0]["payout_scope"] = "different-scope"
        self._save()
        loaded=load_campaign_schedules(self.output,"agency-test","2026-11")
        self.assertTrue(any("矛盾" in r for r in loaded.hold_reasons))
        self.assertFalse(loaded.campaigns)

    def test_missing_cost_unknown_mode_tax_and_fractional_threshold_are_held(self):
        for key,value in (("attributable_direct_cost",None),("threshold_application","unknown"),
                          ("upstream_tax_basis","INCLUSIVE"),("threshold_value",500.5),
                          ("valid_from","2026-11-02")):
            with self.subTest(key=key):
                old=self.rows[0][key]
                self.rows[0][key]=value
                self._save()
                self.assertEqual(load_campaign_schedules(self.output,"agency-test","2026-11").status,"hold")
                self.rows[0][key]=old

    def test_missing_previous_data_blocks_design_but_not_current_verification(self):
        self.rows[0]["previous_agency_unit_amount"]=None
        self._save()
        self.assertEqual(evaluate_campaign_workbook(self.output,self.profile).status,"candidate")
        self.assertEqual(evaluate_campaign_workbook(self.output,self.profile,mode="design").status,"hold")

    def test_candidate_copy_preserves_original_and_refuses_overwrite(self):
        before=hashlib.sha256(self.input.read_bytes()).hexdigest()
        with self.assertRaises(FileExistsError):
            write_campaign_candidate_workbook(self.input,self.output,self.candidates)
        self.assertEqual(hashlib.sha256(self.input.read_bytes()).hexdigest(),before)
        with self.assertRaises(P2InputError):
            write_campaign_candidate_workbook(self.input,ROOT/"never-write-campaign.xlsx",self.candidates)

    def test_optional_template_and_cli_keep_old_schema_compatible(self):
        extended=self.root/"extended.xlsx"
        prepare_input_workbooks.create_approved_input_template(extended,include_campaigns=True)
        wb=load_workbook(extended)
        self.assertEqual(tuple(c.value for c in wb[CAMPAIGN_SHEET][1]),CAMPAIGN_HEADERS)
        wb.close()
        cli_output=self.root/"cli.xlsx"
        with patch.object(sys,"argv",["prepare_campaign_conditions", "--input",str(self.input),
            "--upstream",str(self.upper),"--agency",str(self.agency),"--agency-id","agency-test",
            "--target-period","2026-11","--output",str(cli_output)]):
            self.assertEqual(prepare_campaign_conditions.main(),2)
        self.assertTrue(cli_output.exists())


if __name__ == "__main__":
    unittest.main()
