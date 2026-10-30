from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook
from test_phase2_triage import fixture
from au_hikari_conditions.phase2_review import _write_sheet,input_hash
from au_hikari_conditions.phase2_pricing import write_pricing,identity_errors


def approve(sheet,row=2,**values):
    headers=[c.value for c in sheet[1]]
    r=dict(zip(headers,[c.value for c in sheet[row]]))
    r.update(values,source_document='根拠資料',evidence_location='1頁',approval_reason='同等性・対象確認',
             approval_id='ID1',approval_status='APPROVED',approved_by='担当',approved_at='2026-10-30T12:00:00+09:00')
    r['approved_hash']=input_hash(sheet.title,r)
    for i,h in enumerate(headers,1):sheet.cell(row,i).value=r.get(h)


class PricingTest(unittest.TestCase):
    def run_pricing(self,r,u,path=None):
        b=Workbook();s,a=write_pricing(b,[r],[u],'snapshot',path,_write_sheet,'2026-10','2026-11')
        return b,s,a

    def test_a_b_exception_and_approval_tamper(self):
        p,_,r=fixture()
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'review.xlsx'
            b,s,a=self.run_pricing(r,p.upper)
            self.assertEqual(s['possible'],0);self.assertEqual(s['C'],1)
            approve(b['単価同等性入力'],equivalence_status='CONFIRMED',resolution='支払・除外・戻入・個別約定が同一')
            approve(b['単価基本方針入力'],method='A');b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(a[0]['amount'],'22000');self.assertEqual(s['A'],1)
            approve(b['方式別単価承認入力']);b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(s['approved'],1)
            approve(b['成果項目例外入力'],method='B');b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(a[0]['amount'],'23000');self.assertEqual(s['B'],1)
            self.assertEqual(a[0]['status'],'APPROVAL_INVALID')
            self.assertIn('粗利保証なし',b['方式別11月単価案'].cell(2,15).value)
            self.assertIsNone(r['proposed'])

    def test_invalid_axes_unknown_negative_and_payment_prohibition(self):
        p,_,r=fixture()
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'review.xlsx'
            b,_,_=self.run_pricing(r,p.upper)
            approve(b['単価同等性入力'],equivalence_status='CONFIRMED',resolution='同等')
            approve(b['単価基本方針入力'],method='B');b.save(path)
            p.upper.amount=1000
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(s['possible'],0)
            p.upper.amount=25000;p.old.tax='不明'
            self.assertTrue(identity_errors(p,p.upper,'2026-10','2026-11'))
            with self.assertRaises(ValueError):self.run_pricing(r,p.upper,path)
            p.old.tax='税別';p.upper.access='事業者変更'
            self.assertTrue(identity_errors(p,p.upper,'2026-10','2026-11'))
            with self.assertRaises(ValueError):self.run_pricing(r,p.upper,path)

    def test_c_manual_and_policy_edit_invalidates(self):
        p,_,r=fixture()
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'review.xlsx'
            b,_,_=self.run_pricing(r,p.upper)
            approve(b['単価基本方針入力'],method='C')
            approve(b['単価同等性入力'],equivalence_status='CONFIRMED',resolution='適用条件を個別設定')
            b['方式別単価承認入力'].cell(2,3,21000);b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(a[0]['amount'],'21000');self.assertEqual(s['C'],1)
            approve(b['方式別単価承認入力']);b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(s['approved'],1)
            b['単価基本方針入力'].cell(2,2,'A');b.save(path)
            _,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(a[0]['status'],'APPROVAL_INVALID')

    def test_b_decrease_and_source_scope_checks(self):
        p,_,r=fixture();p.upper.amount=23000
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'review.xlsx'
            b,_,_=self.run_pricing(r,p.upper)
            approve(b['単価基本方針入力'],method='B')
            approve(b['単価同等性入力'],equivalence_status='CONFIRMED',resolution='同等')
            b.save(path)
            b,s,a=self.run_pricing(r,p.upper,path)
            self.assertEqual(a[0]['amount'],'21000')
            self.assertEqual(s['B'],1)
            b['単価基本方針入力'].cell(2,3,'別資料');b.save(path)
            with self.assertRaises(ValueError):self.run_pricing(r,p.upper,path)

    def test_payment_prohibition_and_ambiguous_old_rows(self):
        p,_,r=fixture('keioh株式会社');p.upper.item='ftth_open'
        self.assertIn('代理店支払禁止',identity_errors(p,p.upper,'2026-10','2026-11'))
        p,_,r=fixture();p.initial_identity_collision=True
        self.assertTrue(identity_errors(p,p.upper,'2026-10','2026-11'))
        p.initial_identity_collision=False;p.upper.state='DASH'
        self.assertTrue(identity_errors(p,p.upper,'2026-10','2026-11'))

    def test_common_confirmation_applies_only_to_bound_agencies(self):
        p,_,r=fixture('代理店一');_,_,second=fixture('代理店二')
        def run(path=None,rs=None):
            b=Workbook();s,a=write_pricing(b,rs or [r,second],[p.upper],'snapshot',path,_write_sheet,'2026-10','2026-11')
            return b,s,a
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'review.xlsx';b,_,_=run()
            approve(b['単価共通同等性入力'],equivalence_status='CONFIRMED',resolution='記載された両代理店の契約・適用条件同等')
            b.save(path)
            b,s,a=run(path)
            self.assertEqual(s['A'],2);self.assertEqual(s['possible'],2)
            self.assertTrue(all(x['status']=='未承認案' for x in a))
            with self.assertRaises(ValueError):run(path,[r])
