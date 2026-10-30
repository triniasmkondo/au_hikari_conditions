from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from au_hikari_conditions.condition_ledger import Condition
from au_hikari_conditions.upper_conditions import UpperCondition, extract_upper_pdf
from au_hikari_conditions.phase2_review import (
    Proposal, INPUT_SCHEMAS, input_hash, empty_input, evaluate_proposal, make_proposals,
    old_axes, classification, build_phase2, read_inputs, read_phase1)

ROOT=Path(__file__).resolve().parents[1]
LEDGER=ROOT/'04_成果物/Phase1/2026-10_11社共通条件台帳_監査修正版.xlsx'


def approved(name,record):
    record.update(source_document='approved-internal-register',evidence_location='record:1',
        approval_reason='担当者が根拠・条件・対象月を照合',approval_id='approval-'+name,
        approval_status='APPROVED',approved_by='reviewer',approved_at='2026-10-30T12:00:00+09:00')
    record['approved_hash']=input_hash(name,record)
    return record


class ReviewCalculationTest(unittest.TestCase):
    def fixture(self,hold='',tax='税別'):
        old=Condition('架空社','通常報酬','auひかり','回線','auひかりホーム',amount=22000,
            unit='円/件',tax=tax,period='2026-10-01/2026-10-31',hold=hold)
        upper=UpperCondition('auひかり','base','net_open',dwelling='ホーム',amount=25000,
            period='2026-11-01/2027-03-31',source='new.pdf',location='page:5',digest='source-hash')
        prior=replace(upper,amount=24000,period='2026-10-01/2026-10-31')
        p=make_proposals([old],[prior],[upper],'2026-11')[0]
        inputs={name:empty_input(headers,p,'snapshot','2026-11') for name,headers in INPUT_SCHEMAS.items()}
        inputs['条件対応入力'].update(decision='CONTINUE',equivalence_status='CONFIRMED',scope_review='CONFIRMED',
            period_rule='月次登録・6ヶ月以内開通',payment_timing='条件充足の翌月末',payout_trigger='開通',
            exclusions='同月解約除外',reversal_terms='6ヶ月内解約・原本根拠に従う')
        inputs['粗利承認入力'].update(approved_gross_margin=2000,gross_margin_tax_basis='exclusive',
            direct_cost=500,direct_cost_tax_basis='exclusive',display_policy='present')
        approved('条件対応入力',inputs['条件対応入力']);approved('粗利承認入力',inputs['粗利承認入力'])
        return p,inputs

    def evaluate(self,p,inputs):return evaluate_proposal(p,inputs,'snapshot','2026-11')

    def test_margin_preserved_and_unapproved_never_final(self):
        p,inputs=self.fixture();r=self.evaluate(p,inputs)
        self.assertEqual(r['proposed'],Decimal(22500))
        self.assertEqual(r['retained'],Decimal(2000))
        self.assertEqual(r['status'],'PENDING')
        self.assertIsNone(r['approved_amount'])
        decision=inputs['最終承認入力'];decision['candidate_hash']=r['candidate_hash']
        approved('最終承認入力',decision)
        r=self.evaluate(p,inputs)
        self.assertEqual(r['status'],'APPROVED')
        self.assertEqual(r['approved_amount'],Decimal(22500))

    def test_unknown_cost_margin_or_tax_and_phase1_hold_are_individual_blocks(self):
        for field in ('approved_gross_margin','direct_cost'):
            p,inputs=self.fixture();inputs['粗利承認入力'][field]=None
            approved('粗利承認入力',inputs['粗利承認入力'])
            self.assertIsNone(self.evaluate(p,inputs)['proposed'])
        p,inputs=self.fixture('原本特典コード空欄','原本に税区分明示なし')
        r=self.evaluate(p,inputs);self.assertIsNone(r['proposed']);self.assertEqual(r['issues'],['PHASE1_HOLD','TAX_UNKNOWN'])
        inputs['保留解消入力'].update(resolved_issues='PHASE1_HOLD;TAX_UNKNOWN',tax_basis='exclusive',resolution='原本発行元へ照会し税抜と適用対象を確認')
        approved('保留解消入力',inputs['保留解消入力'])
        self.assertEqual(self.evaluate(p,inputs)['proposed'],Decimal(22500))
        inputs['保留解消入力']['resolution']='変更'
        self.assertIsNone(self.evaluate(p,inputs)['proposed'])

    def test_final_amount_cannot_reduce_margin_or_be_invalid(self):
        for amount in (23000,'bad',True,-1):
            p,inputs=self.fixture();r=self.evaluate(p,inputs)
            d=inputs['最終承認入力'];d.update(candidate_hash=r['candidate_hash'],approved_amount=amount)
            approved('最終承認入力',d)
            self.assertEqual(self.evaluate(p,inputs)['status'],'APPROVAL_INVALID',str(amount))
        p,inputs=self.fixture();r=self.evaluate(p,inputs);d=inputs['最終承認入力']
        d.update(candidate_hash=r['candidate_hash'],approved_amount=22000);approved('最終承認入力',d)
        self.assertEqual(self.evaluate(p,inputs)['approved_retained'],Decimal(2500))

    def test_tamper_mapping_cost_source_period_and_flags_invalidates(self):
        for mutate in (
            lambda p,i:i['粗利承認入力'].update(direct_cost=100),
            lambda p,i:i['条件対応入力'].update(exclusions='取消除外なし'),
            lambda p,i:setattr(p.candidates[0],'digest','changed-source'),
            lambda p,i:setattr(p.candidates[0],'period','2026-12-01/2027-03-31'),
        ):
            p,inputs=self.fixture();r=self.evaluate(p,inputs)
            inputs['最終承認入力']['candidate_hash']=r['candidate_hash'];approved('最終承認入力',inputs['最終承認入力'])
            mutate(p,inputs)
            self.assertNotEqual(self.evaluate(p,inputs)['status'],'APPROVED')

    def test_classification_and_dash_is_not_zero_or_end(self):
        p,inputs=self.fixture()
        self.assertEqual(classification(p.old,p.previous,p.upper)[0],'変更')
        self.assertEqual(classification(p.old,p.upper,p.upper)[0],'継承')
        self.assertEqual(classification(None,None,p.upper)[0],'新規')
        self.assertEqual(classification(p.old,p.previous,None)[0],'要確認')
        self.assertEqual(classification(p.old,p.previous,replace(p.upper,state='ENDED',amount=None))[0],'終了')
        p.candidates[0].state='DASH';p.candidates[0].amount=None
        self.assertIsNone(self.evaluate(p,inputs)['proposed'])
        self.assertEqual(classification(p.old,p.previous,p.upper)[0],'要確認')

    def test_changed_type_code_or_scope_does_not_fuzzy_match(self):
        p,inputs=self.fixture();a=p.axes.copy();a.update(code='KOF',type='1')
        upper=replace(p.upper,code='KCM',type='3')
        from au_hikari_conditions.phase2_review import related_axes
        self.assertFalse(related_axes(a,upper.axes))

    def test_type_is_not_inferred_from_price_digits_and_duplicate_claims_block(self):
        from au_hikari_conditions.upper_conditions import types
        from au_hikari_conditions.phase2_review import find_conflicts
        self.assertEqual(types('auひかりプラス タイプ① - 5,000 2,000 25,000'),'1')
        self.assertEqual(types('auひかりプラス タイプ③ - 5,000 2,000 20,000'),'3')
        p,inputs=self.fixture()
        second=replace(p,id='second-proposal')
        mapping1=inputs['条件対応入力'];mapping2=deepcopy(mapping1)
        mapping2['proposal_id']=second.id;approved('条件対応入力',mapping2)
        conflicts=find_conflicts([p,second],{'条件対応入力':{p.id:mapping1,second.id:mapping2}},'snapshot')
        self.assertEqual(set(conflicts),{p.id,second.id})
        self.assertIsNone(evaluate_proposal(p,inputs,'snapshot','2026-11',conflicts[p.id])['proposed'])

    def test_keioh_ftth_receipt_cannot_be_adopted_as_agency_payment(self):
        p,inputs=self.fixture()
        p.agency='keioh株式会社';p.candidates[0].item='ftth_open';p.candidates[0].amount=2000
        inputs['条件対応入力']['selected_upper_id']=p.candidates[0].id
        approved('条件対応入力',inputs['条件対応入力'])
        r=self.evaluate(p,inputs)
        self.assertIsNone(r['proposed'])
        self.assertTrue(any('代理店支払禁止' in b for b in r['blocks']))


class KeiohSourceReviewTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.old,_=read_phase1(ROOT,LEDGER)
        cls.uppers={}
        for month in ('2026-10','2026-11'):
            data=[]
            for p in (ROOT/'01_上位原本'/month).glob('*.pdf'):
                records,_=extract_upper_pdf(p,ROOT);data.extend(records)
            cls.uppers[month]=data

    def test_keioh_campaign_and_changed_scope(self):
        old=[c for c in self.old if c.agency=='keioh株式会社']
        proposals=make_proposals(old,self.uppers['2026-10'],self.uppers['2026-11'],'2026-11')
        self.assertEqual(sum(p.old is not None for p in proposals),90)
        self.assertEqual(len({p.id for p in proposals}),len(proposals))
        campaign=[p for p in proposals if p.old and p.old.kind=='キャンペーン']
        self.assertEqual(len(campaign),15)
        main=[p for p in campaign if p.old.item=='ボリュームインセンティブ' and p.old.type=='au ひかり既存・au ひかりプラス タイプ1']
        # Assert by source threshold and target axes instead of accepting a total candidate count.
        main=[p for p in campaign if old_axes(p.old)['type']=='1' and p.old.item=='ボリュームインセンティブ']
        self.assertEqual([p.upper.amount for p in main],[6500,9000,11500,14000,16500,17500])
        self.assertTrue(all(p.category=='継承' for p in main))
        support=[p for p in campaign if p.old.item=='販促支援金']
        self.assertEqual([p.upper.amount for p in support],[3000,5000,6500])
        self.assertEqual([p.category for p in support],['継承','変更','変更'])
        adjustment=next(p for p in campaign if p.old.item=='特別調整金')
        self.assertFalse(adjustment.exact);self.assertEqual(adjustment.upper.type,'3')
        self.assertEqual(adjustment.upper.amount,10000);self.assertEqual(adjustment.category,'変更')
        self.assertEqual(old_axes(adjustment.old)['type'],'2,3')
        native=[p for p in proposals if p.old and p.old.item in ('5G開通手数料','ネット開通手数料')]
        self.assertEqual(len(native),6)
        for p in native:
            self.assertIsNotNone(p.upper)
            self.assertEqual(p.upper.type,old_axes(p.old)['type'])
        call=next(p for p in proposals if p.old and p.old.item=='電話開通手数料' and p.old.type=='タイプ1')
        self.assertEqual(call.upper.state,'DASH');self.assertIsNone(call.upper.amount)
        wifi=next(p for p in proposals if p.old and p.old.item=='おうちどこでもWi-Fi')
        self.assertEqual((wifi.previous.amount,wifi.upper.amount,wifi.category),(2500,4500,'変更'))
        incentive=next(p for p in proposals if p.old and p.old.item.startswith('10Gインセンティブ') and p.old.product=='auひかり')
        self.assertEqual((incentive.previous.amount,incentive.upper.amount),(5000,20000))
        for p in [p for p in proposals if p.old and p.old.item=='回線サービス' and p.old.code=='KWE' and p.old.type=='タイプ3']:
            self.assertGreaterEqual(len(p.candidates),2)
            self.assertEqual(p.category,'要確認')

    def test_all_phase1_flags_and_no_guessed_upstream_prices(self):
        self.assertEqual(sum(bool(c.hold) for c in self.old),120)
        self.assertEqual(sum(c.tax!='税別' for c in self.old),222)
        for u in self.uppers['2026-11']:
            if u.family=='nifty_unresolved':
                self.assertIsNone(u.amount)
                self.assertEqual(u.state,'UNRESOLVED')
        new=[u for u in self.uppers['2026-11'] if u.product=='auひかりプラス' and u.family=='isp_line']
        self.assertEqual(len(new),14)
        code=[u for u in new if u.type=='3' and u.code=='KWE' and u.speed=='1G']
        self.assertEqual({u.access:u.amount for u in code},{'回線新規':13000,'事業者変更':0})
        self.assertTrue(all(u.period.startswith('2026-11') for u in self.uppers['2026-11']))

    def test_real_workbook_roundtrip_edit_display_does_not_approve(self):
        with tempfile.TemporaryDirectory() as folder:
            first=Path(folder)/'first.xlsx';edited=Path(folder)/'edited.xlsx';second=Path(folder)/'second.xlsx'
            stats=build_phase2(ROOT,LEDGER,first,agency_filter='keioh株式会社')
            self.assertEqual(stats['inherited_phase1_holds'],6)
            self.assertEqual(stats['inherited_unknown_tax'],21)
            self.assertEqual(stats['calculation_proposals'],0)
            w=load_workbook(first);w['11月条件案']['AI2']='APPROVED';w['11月条件案']['S2']=999999
            w.save(edited);w.close()
            result=build_phase2(ROOT,LEDGER,second,agency_filter='keioh株式会社',review_input=edited)
            self.assertEqual(result['approval_states'],{'HELD':result['proposal_count']})
            inputs,manifest=read_inputs(second)
            # A copied approval state with no provenance or matching approved_hash is rejected per condition.
            self.assertTrue(all(r['approval_status']=='PENDING' for r in inputs['最終承認入力'].values()))
            with self.assertRaises(FileExistsError):build_phase2(ROOT,LEDGER,second,agency_filter='keioh株式会社')

class ApprovalWorkbookIntegrationTest(unittest.TestCase):
    def test_edit_recompute_approve_partial_and_reject_tamper(self):
        import hashlib
        from unittest.mock import patch
        from openpyxl import Workbook
        from au_hikari_conditions.condition_ledger import HEADERS,row_values
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source=root/'03_条件表/2026-10/source.docx'
            source.parent.mkdir(parents=True);source.write_bytes(b'synthetic-agency-source')
            digest=hashlib.sha256(source.read_bytes()).hexdigest()
            old=Condition('架空社','通常報酬','auひかり','回線','auひかりホーム',amount=22000,
                unit='円/件',tax='税別',period='2026-10-01/2026-10-31',source=source.relative_to(root).as_posix(),
                locations=['table:1;row:2;cell:2'],digest=digest)
            unknown=replace(old,item='auひかりマンション',hold='元条件未確定',tax='原本に税区分明示なし')
            book=Workbook();s=book.active;s.title='共通条件台帳';s.append(HEADERS);s.append(row_values(old));s.append(row_values(unknown))
            s=book.create_sheet('資料一覧');s.append(['代理店','根拠資料','SHA256']);s.append([old.agency,old.source,digest])
            ledger=root/'ledger.xlsx';book.save(ledger)
            for month in ('2026-10','2026-11'):
                path=root/'01_上位原本'/month/'source.pdf';path.parent.mkdir(parents=True);path.write_bytes(month.encode())
            def upper_reader(path,project):
                month=path.parent.name
                return [UpperCondition('auひかり','base','net_open',dwelling=dw,amount=24000 if month=='2026-10' else 25000,
                    period=month+'-01/'+('2026-10-31' if month=='2026-10' else '2027-03-31'),
                    source=path.relative_to(project).as_posix(),location='page:1;row:2',digest=hashlib.sha256(path.read_bytes()).hexdigest())
                    for dw in ('ホーム','マンション')],[[path.relative_to(project).as_posix(),1,'架空上位条件']]
            with patch('au_hikari_conditions.phase2_review.extract_upper_pdf',side_effect=upper_reader):
                first=root/'first.xlsx';build_phase2(root,ledger,first)
                records,_=read_inputs(first)
                pid=next(pid for pid,r in records['条件対応入力'].items() if r['condition_id']==old.id)
                mapping=records['条件対応入力'][pid]
                mapping.update(decision='CONTINUE',equivalence_status='CONFIRMED',scope_review='CONFIRMED',
                    period_rule='月次登録',payment_timing='翌月末',payout_trigger='開通',exclusions='同月解約除外',reversal_terms='原本の6ヶ月戻入')
                margin=records['粗利承認入力'][pid]
                margin.update(approved_gross_margin=2000,gross_margin_tax_basis='exclusive',direct_cost=500,
                              direct_cost_tax_basis='exclusive',display_policy='present')
                for name in ('条件対応入力','粗利承認入力'):approved(name,records[name][pid])
                def write_inputs(original,target,inputs):
                    w=load_workbook(original)
                    for name,headers in INPUT_SCHEMAS.items():
                        s=w[name]
                        for ri,row in enumerate(s.iter_rows(min_row=2),2):
                            data=inputs[name][row[0].value]
                            for ci,key in enumerate(headers,1):s.cell(ri,ci).value=data.get(key)
                    w.save(target);w.close()
                edited=root/'edited.xlsx';write_inputs(first,edited,records)
                second=root/'second.xlsx';stats=build_phase2(root,ledger,second,review_input=edited)
                self.assertEqual(stats['calculation_proposals'],1)
                records,_=read_inputs(second)
                d=records['最終承認入力'][pid];approved('最終承認入力',d)
                final_in=root/'final-in.xlsx';write_inputs(second,final_in,records)
                final=root/'final.xlsx';stats=build_phase2(root,ledger,final,review_input=final_in)
                self.assertEqual(stats['approval_states'],{'APPROVED':1,'HELD':1})
                import json
                journal=json.loads(final.with_suffix('.approvals.json').read_text())
                approved_result=next(r for r in journal['decisions'] if r['status']=='APPROVED')
                self.assertEqual(approved_result['previous_condition_id'],old.id)
                self.assertEqual(approved_result['approved_amount'],'22500')
                # Editing a reviewed input without matching approval fingerprint revokes final status.
                records,_=read_inputs(final);records['粗利承認入力'][pid]['direct_cost']=1000
                tamper=root/'tamper.xlsx';write_inputs(final,tamper,records)
                changed=build_phase2(root,ledger,root/'changed.xlsx',review_input=tamper)
                self.assertNotIn('APPROVED',changed['approval_states'])
                # Source replacement is a batch-level stop, not a silently reused approval.
                source.write_bytes(b'changed')
                with self.assertRaisesRegex(ValueError,'ハッシュ'):
                    build_phase2(root,ledger,root/'must-not-write.xlsx',review_input=final)
