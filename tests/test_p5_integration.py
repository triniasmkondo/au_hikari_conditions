from __future__ import annotations

import contextlib
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "tests"))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from au_hikari_conditions.p4b import DraftHistoryStore, DuplicateDraftError
from au_hikari_conditions.p5 import (
    P5Error,
    compute_mail_template_hash,
    create_approved_zoho_drafts,
    finalize_approved_monthly_batch,
    load_approved_mail_settings,
    prepare_monthly_batch,
)
from au_hikari_conditions.monthly_bridge import build_monthly_comparison_workbook
import run_p5
import prepare_monthly_comparison
from prepare_input_workbooks import create_delivery_template
from test_p2_integration import AGENCY_ID, _build_fixture


PERIOD = "2026-11"
BATCH_ID = "anonymous-batch-2026-11"


class MockDraftClient:
    def __init__(self) -> None:
        self.uploaded: list[str] = []
        self.payloads: list[dict[str, object]] = []

    def upload_draft_attachment(
        self, *, file_name: str, content: bytes, content_type: str
    ) -> dict[str, str]:
        self.uploaded.append(file_name)
        return {
            "storeName": f"store-{len(self.uploaded)}",
            "attachmentName": file_name,
            "attachmentPath": f"/attachment/{len(self.uploaded)}",
        }

    def save_draft(self, payload: dict[str, object]) -> dict[str, object]:
        self.payloads.append(payload)
        return {"status": {"code": 200}}


class P5IntegrationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name)
        self.input_path = self.root / "approved-input.xlsx"
        self.p2_path = self.root / "p2-result.xlsx"
        self.master_path = self.root / "delivery-master.xlsx"
        self.approval_path = self.root / "monthly-approval.xlsx"
        self.candidate_directory = self.root / "approval-candidates"
        self.final_directory = self.root / "approved-distribution"
        self.comparison_path = self.root / "monthly-comparison.xlsx"
        _build_fixture(self.input_path)
        self.project_root = self._create_monthly_comparison()
        self._create_delivery_master()

    def _create_monthly_comparison(self) -> Path:
        source_root = self.root / "comparison-sources"
        source_folders = (
            source_root / "01_上位原本" / "2026-10",
            source_root / "03_条件表" / "2026-10",
            source_root / "01_上位原本" / "2026-11",
        )
        contents = (
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 120円/件",
        )
        for index, (folder, content) in enumerate(zip(source_folders, contents)):
            folder.mkdir(parents=True)
            document = Document()
            document.add_paragraph(content)
            document.save(folder / f"fictional-source-{index}.docx")
        build_monthly_comparison_workbook(
            project_root=source_root,
            previous_period="2026-10",
            target_period=PERIOD,
            output_path=self.comparison_path,
        )
        return source_root

    def test_fictional_template_master_is_not_treated_as_approved(self) -> None:
        create_delivery_template(self.master_path)
        with self.assertRaises(P5Error):
            load_approved_mail_settings(
                self.master_path,
                agency_id=AGENCY_ID,
                target_period=PERIOD,
            )

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _create_delivery_master(self) -> None:
        create_delivery_template(self.master_path)
        workbook = load_workbook(self.master_path)
        recipients = workbook["送付先マスター"]
        recipients.delete_rows(2, 1)
        recipient_headers = [cell.value for cell in recipients[1]]
        recipient = {
            "agency_id": AGENCY_ID,
            "agency_name": "Anonymous Agency",
            "recipient_status": "APPROVED",
            "to_addresses": "recipient@example.invalid",
            "cc_addresses": "copy@example.invalid",
            "bcc_addresses": "",
            "sender_address": "sender@example.invalid",
            "salutation": "Anonymous Agency ご担当者様",
            "recipient_source": "synthetic-approved-master",
            "approved_by": "anonymous-approver",
            "approved_at": "2026-10-28T10:00:00+09:00",
            "valid_from": "2026-11-01",
            "valid_until": "2026-11-30",
            "last_reviewed_at": "2026-10-28T10:00:00+09:00",
        }
        recipients.append([recipient.get(header) for header in recipient_headers])

        templates = workbook["本文テンプレート"]
        templates.delete_rows(2, 1)
        template_headers = [cell.value for cell in templates[1]]
        template = {
            "template_id": "anonymous-conditions-v1",
            "template_version": "1",
            "doc_type": "conditions",
            "subject_template": "{target_period} conditions",
            "body_template": "{salutation}\n{target_period} conditions attached.",
            "allowed_placeholders": "salutation,target_period",
            "approval_status": "APPROVED",
            "approved_by": "anonymous-approver",
            "approved_at": "2026-10-28T10:00:00+09:00",
            "valid_from": "2026-11-01",
            "valid_until": "2026-11-30",
        }
        template["template_hash"] = compute_mail_template_hash(template)
        templates.append([template.get(header) for header in template_headers])
        workbook.save(self.master_path)
        workbook.close()

    def _prepare(self):
        with patch(
            "au_hikari_conditions.p5.build_p1_report",
            return_value={"p1_gate": "met"},
        ):
            return prepare_monthly_batch(
                project_root=self.project_root,
                approved_input_path=self.input_path,
                p2_result_path=self.p2_path,
                monthly_comparison_path=self.comparison_path,
                target_period=PERIOD,
                batch_id=BATCH_ID,
                delivery_workbook_path=self.master_path,
                approval_workbook_path=self.approval_path,
                candidate_directory=self.candidate_directory,
                requested_by="anonymous-operator",
            )

    def _approve(self) -> None:
        workbook = load_workbook(self.approval_path)
        worksheet = workbook["月次承認"]
        headers = [cell.value for cell in worksheet[1]]
        for row_number in range(2, worksheet.max_row + 1):
            values = {
                header: worksheet.cell(row_number, column + 1).value
                for column, header in enumerate(headers)
            }
            if (
                values.get("batch_id") == BATCH_ID
                and values.get("agency_id") == AGENCY_ID
            ):
                for field, value in {
                    "approval_status": "APPROVED",
                    "qa_status": "PASSED",
                    "approved_by": "anonymous-approver",
                    "approved_at": datetime.now(timezone.utc).isoformat(),
                    "approval_reason": "Anonymous QA approval",
                }.items():
                    worksheet.cell(
                        row_number,
                        headers.index(field) + 1,
                        value,
                    )
                break
        else:
            self.fail("P5 monthly approval request was not written.")
        workbook.save(self.approval_path)
        workbook.close()

    def _finalize(self):
        with patch(
            "au_hikari_conditions.p5.build_p1_report",
            return_value={"p1_gate": "met"},
        ):
            return finalize_approved_monthly_batch(
                project_root=self.project_root,
                approved_input_path=self.input_path,
                p2_result_path=self.p2_path,
                monthly_comparison_path=self.comparison_path,
                target_period=PERIOD,
                batch_id=BATCH_ID,
                delivery_workbook_path=self.master_path,
                approval_workbook_path=self.approval_path,
                candidate_directory=self.candidate_directory,
                final_directory=self.final_directory,
            )

    def test_p2_to_p5_candidate_approval_finalize_and_mock_draft(self) -> None:
        prepared = self._prepare()
        self.assertEqual(prepared.p1_status, "met")
        self.assertEqual(prepared.agency_count, 1)
        self.assertEqual(len(prepared.requests[0].snapshot.artifact_hashes), 2)
        workbook = load_workbook(self.approval_path, read_only=False)
        approval_table = next(iter(workbook["月次承認"].tables.values()))
        self.assertIn(str(workbook["月次承認"].max_row), approval_table.ref)
        workbook.close()

        with self.assertRaisesRegex(P5Error, "completed QA and approval"):
            self._finalize()

        self._approve()
        packages = self._finalize()
        self.assertEqual(len(packages), 1)
        package = packages[0]
        self.assertEqual(package.workflow.approval_fingerprint, package.snapshot.fingerprint())
        self.assertTrue(package.artifacts.spreadsheet_path.exists())
        self.assertTrue(package.artifacts.word_path.exists())
        self.assertEqual(package.preview.target_period, PERIOD)

        client = MockDraftClient()
        history = DraftHistoryStore(self.root / "draft-history.sqlite3")
        with self.assertRaisesRegex(P5Error, "Company authorization"):
            create_approved_zoho_drafts(
                packages=packages,
                mail_client=client,
                history=history,
                company_authorized=False,
            )
        self.assertEqual(client.payloads, [])

        result = create_approved_zoho_drafts(
            packages=packages,
            mail_client=client,
            history=history,
            company_authorized=True,
        )
        self.assertEqual(result, ("CREATED",))
        self.assertEqual(len(client.uploaded), 2)
        self.assertEqual(client.payloads[0]["mode"], "draft")
        self.assertEqual(history.get_status(AGENCY_ID, PERIOD), "CREATED")
        with self.assertRaises(DuplicateDraftError):
            create_approved_zoho_drafts(
                packages=packages,
                mail_client=client,
                history=history,
                company_authorized=True,
            )
        self.assertEqual(len(client.payloads), 1)

    def test_cli_prepare_human_approval_finalize_with_mock_draft(self) -> None:
        synthetic_project = self.root / "synthetic-project"
        previous_upper = synthetic_project / "01_上位原本" / "2026-10"
        previous_agency = synthetic_project / "03_条件表" / "2026-10"
        current_upper = synthetic_project / "01_上位原本" / "2026-11"
        previous_upper.mkdir(parents=True)
        previous_agency.mkdir(parents=True)
        current_upper.mkdir(parents=True)
        (synthetic_project / "01_上位原本" / "2026-11" / "fictional-source.txt").write_text(
            "Synthetic source placeholder; not parsed by this test.",
            encoding="utf-8",
        )
        for folder, filename, content in (
            (
                previous_upper,
                "fictional-previous-upper.docx",
                "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
            ),
            (
                previous_agency,
                "fictional-agency-conditions.docx",
                "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
            ),
            (
                current_upper,
                "fictional-current-upper.docx",
                "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 120円/件",
            ),
        ):
            document = Document()
            document.add_paragraph(content)
            document.save(folder / filename)

        comparison_path = self.root / "monthly-comparison-cli.xlsx"
        comparison_stdout = io.StringIO()
        with (
            patch.object(prepare_monthly_comparison, "PROJECT_ROOT", synthetic_project),
            contextlib.redirect_stdout(comparison_stdout),
        ):
            comparison_exit = prepare_monthly_comparison.main(
                [
                    "--previous-period",
                    "2026-10",
                    "--target-period",
                    PERIOD,
                    "--output",
                    str(comparison_path),
                ]
            )
        self.assertEqual(comparison_exit, 2)
        comparison = json.loads(comparison_stdout.getvalue())
        self.assertEqual(comparison["status"], "REVIEW_REQUIRED")
        self.assertTrue(comparison_path.is_file())

        prepare_stdout = io.StringIO()
        with (
            patch.object(run_p5, "PROJECT_ROOT", synthetic_project),
            patch(
                "au_hikari_conditions.p5.build_p1_report",
                return_value={"p1_gate": "met"},
            ),
            contextlib.redirect_stdout(prepare_stdout),
        ):
            self.assertEqual(
                run_p5.main(
                    [
                        "prepare",
                        "--target-period",
                        PERIOD,
                        "--batch-id",
                        BATCH_ID,
                        "--approved-input",
                        str(self.input_path),
                        "--p2-result",
                        str(self.p2_path),
                        "--monthly-comparison",
                        str(comparison_path),
                        "--delivery-workbook",
                        str(self.master_path),
                        "--approval-workbook",
                        str(self.approval_path),
                        "--candidate-directory",
                        str(self.candidate_directory),
                        "--requested-by",
                        "anonymous-operator",
                    ]
                ),
                0,
            )
        self.assertEqual(
            json.loads(prepare_stdout.getvalue())["phase"],
            "AWAITING_HUMAN_QA_AND_APPROVAL",
        )
        self.assertTrue(self.p2_path.is_file())
        self.assertTrue(self.approval_path.is_file())
        self.assertEqual(len(list(self.candidate_directory.glob("*"))), 2)

        self._approve()
        draft_client = MockDraftClient()
        finalize_stdout = io.StringIO()
        with (
            patch.object(run_p5, "PROJECT_ROOT", synthetic_project),
            patch(
                "au_hikari_conditions.p5.build_p1_report",
                return_value={"p1_gate": "met"},
            ),
            patch.object(
                run_p5.ZohoMailConfig,
                "from_secret_provider",
                return_value=object(),
            ),
            patch.object(run_p5, "ZohoMailClient", return_value=draft_client),
            patch.dict(
                os.environ,
                {
                    "P5_ZOHO_DRAFTS_APPROVED": "YES",
                    "P5_DRAFT_HISTORY_DB": str(self.root / "draft-history.sqlite3"),
                },
            ),
            contextlib.redirect_stdout(finalize_stdout),
        ):
            self.assertEqual(
                run_p5.main(
                    [
                        "finalize",
                        "--target-period",
                        PERIOD,
                        "--batch-id",
                        BATCH_ID,
                        "--approved-input",
                        str(self.input_path),
                        "--p2-result",
                        str(self.p2_path),
                        "--monthly-comparison",
                        str(comparison_path),
                        "--delivery-workbook",
                        str(self.master_path),
                        "--approval-workbook",
                        str(self.approval_path),
                        "--candidate-directory",
                        str(self.candidate_directory),
                        "--final-directory",
                        str(self.final_directory),
                        "--save-drafts",
                    ]
                ),
                0,
            )
        final_result = json.loads(finalize_stdout.getvalue())
        self.assertEqual(final_result["phase"], "DRAFTS_CREATED")
        self.assertEqual(final_result["agency_count"], 1)
        self.assertEqual(final_result["draft_count"], 1)
        self.assertEqual(len(list(self.final_directory.glob("*"))), 2)
        self.assertEqual(len(draft_client.uploaded), 2)
        self.assertEqual(len(draft_client.payloads), 1)
        self.assertEqual(draft_client.payloads[0]["mode"], "draft")

    def test_p1_blocker_stops_before_writing_downstream_artifacts(self) -> None:
        with patch(
            "au_hikari_conditions.p5.build_p1_report",
            return_value={"p1_gate": "not_met"},
        ):
            with self.assertRaisesRegex(P5Error, "P1 source analysis"):
                prepare_monthly_batch(
                    project_root=self.project_root,
                    approved_input_path=self.input_path,
                    p2_result_path=self.p2_path,
                    monthly_comparison_path=self.comparison_path,
                    target_period=PERIOD,
                    batch_id=BATCH_ID,
                    delivery_workbook_path=self.master_path,
                    approval_workbook_path=self.approval_path,
                    candidate_directory=self.candidate_directory,
                    requested_by="anonymous-operator",
                )
        self.assertFalse(self.p2_path.exists())
        self.assertFalse(self.approval_path.exists())
        self.assertFalse(self.candidate_directory.exists())

    def test_changed_recipient_invalidates_approval_snapshot(self) -> None:
        self._prepare()
        self._approve()
        workbook = load_workbook(self.approval_path)
        recipients = workbook["送付先マスター"]
        recipients.cell(2, 4, "changed@example.invalid")
        workbook.save(self.approval_path)
        workbook.close()

        with self.assertRaisesRegex(P5Error, "changed after approval request"):
            self._finalize()
        self.assertFalse(self.final_directory.exists())

    def test_change_in_live_recipient_master_after_request_blocks_finalize(self) -> None:
        self._prepare()
        self._approve()
        workbook = load_workbook(self.master_path)
        workbook["送付先マスター"].cell(2, 4, "changed@example.invalid")
        workbook.save(self.master_path)
        workbook.close()

        with self.assertRaisesRegex(P5Error, "master changed"):
            self._finalize()
        self.assertFalse(self.final_directory.exists())

    def test_change_in_monthly_comparison_after_request_invalidates_snapshot(self) -> None:
        self._prepare()
        self._approve()
        workbook = load_workbook(self.comparison_path)
        worksheet = workbook["上位差分候補"]
        worksheet.cell(2, 6, '["tampered-candidate"]')
        workbook.save(self.comparison_path)
        workbook.close()

        with self.assertRaisesRegex(P5Error, "Approval record no longer matches"):
            self._finalize()
        self.assertFalse(self.final_directory.exists())

    def test_change_in_source_document_after_comparison_blocks_prepare(self) -> None:
        source = next(
            (self.project_root / "01_上位原本" / "2026-11").glob("*.docx")
        )
        source.write_bytes(source.read_bytes() + b"changed")
        with self.assertRaisesRegex(P5Error, "Monthly comparison evidence"):
            self._prepare()
        self.assertFalse(self.p2_path.exists())
        self.assertFalse(self.approval_path.exists())

    def test_candidate_tampering_is_rejected_before_final_output(self) -> None:
        prepared = self._prepare()
        self._approve()
        candidate = prepared.requests[0].candidate_artifacts.spreadsheet_path
        candidate.write_bytes(candidate.read_bytes() + b"tampered")

        with self.assertRaises((P5Error, ValueError)):
            self._finalize()
        self.assertFalse(self.final_directory.exists())


if __name__ == "__main__":
    unittest.main()
