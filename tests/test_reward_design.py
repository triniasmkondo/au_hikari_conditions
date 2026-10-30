from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from au_hikari_conditions.calculation import ApprovalEvidence, ApprovalStatus, LadderMode
from au_hikari_conditions.p2 import (
    P2AgencyResult,
    P2Report,
    calculate_workbook,
)
from au_hikari_conditions.reward_design import (
    AgencyRewardProfile,
    CampaignSchedule,
    CampaignTierInput,
    FTTH_OPENING_FEE_COMPONENT,
    JudgmentReview,
    RewardCategory,
    RewardComponent,
    VolumeScenario,
    connect_reward_design_to_p2,
    design_agency_rewards,
    write_p2_reward_design_candidate,
)
from test_p2_integration import _build_fixture


def _approval(label: str) -> ApprovalEvidence:
    return ApprovalEvidence(
        status=ApprovalStatus.APPROVED,
        approval_id=f"approval-{label}",
        source_document=f"source-{label}.pdf",
        evidence_location="page:1",
        approved_by="approver",
        approved_at=datetime(2026, 10, 29, tzinfo=timezone.utc),
    )


def _profile(
    *,
    agency_id: str = "keioh",
    target: str = "0",
    minimum: str = "0",
    allow_changes: bool = True,
    pay_after_collection: bool | None = True,
) -> AgencyRewardProfile:
    return AgencyRewardProfile(
        agency_id=agency_id,
        target_period="2026-11",
        target_gross_margin=Decimal(target),
        minimum_gross_margin=Decimal(minimum),
        allow_rate_changes=allow_changes,
        pay_only_after_upstream_collection=pay_after_collection,
        collection_timing_approval=_approval("collection"),
    )


def _component(
    item_id: str,
    upstream: str,
    previous: str,
    *,
    target: str | None = None,
    minimum: str | None = None,
    opening_class: str = "5G",
    payment_allowed: bool = True,
    category: RewardCategory = RewardCategory.BASIC,
    payment_component: str = "",
    direct_cost: str = "0",
) -> RewardComponent:
    return RewardComponent(
        component_id=item_id,
        p2_condition_item_id=item_id,
        category=category,
        opening_class=opening_class,
        upstream_unit_reward=Decimal(upstream),
        previous_agency_unit_reward=Decimal(previous),
        direct_cost_per_unit=Decimal(direct_cost),
        payment_component=payment_component,
        payment_allowed=payment_allowed,
        target_gross_margin=Decimal(target) if target is not None else None,
        minimum_gross_margin=Decimal(minimum) if minimum is not None else None,
        upstream_approval=_approval(f"upper-{item_id}"),
        previous_approval=_approval(f"previous-{item_id}"),
        direct_cost_approval=_approval(f"cost-{item_id}"),
    )


def _campaign(
    code: str,
    thresholds: tuple[int, ...],
    upstream: tuple[str, ...],
    previous: tuple[str, ...],
    *,
    target: str,
    opening_classes: tuple[str, ...] = ("5G",),
    mode: LadderMode = LadderMode.REACHED_TIER,
) -> CampaignSchedule:
    tiers = tuple(
        CampaignTierInput(
            threshold=threshold,
            p2_condition_item_id=f"{code}-{threshold}",
            upstream_unit_reward=Decimal(upper),
            previous_agency_unit_reward=Decimal(agency),
            target_gross_margin=Decimal(target),
            upstream_approval=_approval(f"upper-{code}-{threshold}"),
            previous_approval=_approval(f"previous-{code}-{threshold}"),
            direct_cost_approval=_approval(f"cost-{code}-{threshold}"),
        )
        for threshold, upper, agency in zip(
            thresholds, upstream, previous, strict=True
        )
    )
    return CampaignSchedule(
        campaign_code=code,
        counting_scope="commerce-wide",
        payout_scope="keioh-eligible-openings",
        eligible_opening_classes=opening_classes,
        mode=mode,
        tiers=tiers,
    )


class RewardDesignTest(unittest.TestCase):
    def test_reproduces_confirmed_keioh_november_rates_and_margin_differences(self) -> None:
        volume = _campaign(
            "volume-1",
            (500, 750, 1000, 1500, 2000, 2500),
            ("6500", "9000", "11500", "14000", "16500", "17500"),
            ("5500", "8000", "10500", "13000", "15500", "16500"),
            target="1000",
        )
        sales_support = _campaign(
            "sales-support",
            (1800, 2000, 2300),
            ("3000", "5000", "6500"),
            ("1000", "3000", "4500"),
            target="2000",
        )
        plus23 = _campaign(
            "plus-23-volume",
            (220, 440),
            ("3000", "5000"),
            ("1000", "3000"),
            target="2000",
        )
        special = _campaign(
            "plus-3-special",
            (1,),
            ("10000",),
            ("8000",),
            target="2000",
        )
        electricity = _campaign(
            "electricity",
            (30, 50, 80),
            ("3000", "4000", "5000"),
            ("3000", "4000", "5000"),
            target="0",
        )
        result = design_agency_rewards(
            _profile(),
            (
                _component("standard-incentive", "19000", "17000", target="2000"),
                _component(
                    "ftth-opening",
                    "2000",
                    "0",
                    target="2000",
                    opening_class="FTTH",
                    payment_allowed=False,
                    payment_component=FTTH_OPENING_FEE_COMPONENT,
                ),
            ),
            (volume, sales_support, plus23, special, electricity),
        )

        self.assertEqual(result.status, "candidate", result.hold_reasons)
        self.assertEqual(
            {
                proposal.component_id: proposal.proposed_agency_unit_reward
                for proposal in result.component_proposals
            },
            {
                "standard-incentive": Decimal("17000"),
                "ftth-opening": Decimal("0"),
            },
        )
        by_campaign_threshold = {
            (proposal.campaign_code, proposal.threshold):
            proposal.proposed_agency_unit_reward
            for proposal in result.campaign_proposals
        }
        self.assertEqual(by_campaign_threshold[("volume-1", 500)], Decimal("5500"))
        self.assertEqual(
            by_campaign_threshold[("volume-1", 2500)], Decimal("16500")
        )
        self.assertEqual(
            by_campaign_threshold[("sales-support", 2300)], Decimal("4500")
        )
        self.assertEqual(by_campaign_threshold[("plus-23-volume", 440)], Decimal("3000"))
        self.assertEqual(by_campaign_threshold[("plus-3-special", 1)], Decimal("8000"))
        self.assertEqual(by_campaign_threshold[("electricity", 80)], Decimal("5000"))
        self.assertEqual(
            next(
                item.retained_gross_margin_per_unit
                for item in result.component_proposals
                if item.component_id == "standard-incentive"
            ),
            Decimal("2000"),
        )
        self.assertEqual(
            by_campaign_threshold[("volume-1", 500)],
            Decimal("6500") - Decimal("1000"),
        )
        self.assertTrue(
            any(
                scenario.commerce_count == 499
                and scenario.campaign_code == "volume-1"
                and scenario.agency_payment == 0
                for scenario in result.scenarios
            )
        )
        self.assertTrue(
            any(
                scenario.commerce_count == 500
                and scenario.campaign_code == "volume-1"
                and scenario.agency_payment == Decimal("5500")
                for scenario in result.scenarios
            )
        )
        self.assertTrue(
            all(
                scenario.opening_class in ("5G", "FTTH")
                for scenario in result.scenarios
            )
        )
        self.assertEqual(
            sum(
                scenario.campaign_code == "volume-1"
                and scenario.commerce_count in (499, 500, 749, 750, 999, 1000, 1499, 1500, 1999, 2000, 2499, 2500)
                and scenario.upstream_collected
                for scenario in result.scenarios
            ),
            12,
        )

    def test_increasing_target_margin_reduces_rate_only_as_needed(self) -> None:
        result = design_agency_rewards(
            _profile(target="4000", minimum="3000"),
            (_component("standard", "19000", "17000"),),
        )

        self.assertEqual(result.status, "candidate", result.hold_reasons)
        self.assertEqual(
            result.component_proposals[0].proposed_agency_unit_reward,
            Decimal("15000"),
        )

    def test_nonmonotone_campaign_inputs_are_repaired_without_increasing_any_tier(self) -> None:
        campaign = _campaign(
            "nonmonotone",
            (100, 200),
            ("10000", "10000"),
            ("5000", "1000"),
            target="0",
        )
        result = design_agency_rewards(_profile(), (), (campaign,))

        self.assertEqual(result.status, "candidate", result.hold_reasons)
        self.assertEqual(
            tuple(item.proposed_agency_unit_reward for item in result.campaign_proposals),
            (Decimal("1000"), Decimal("1000")),
        )

    def test_target_or_minimum_margin_infeasibility_holds(self) -> None:
        target_unreachable = design_agency_rewards(
            _profile(target="21000", minimum="20000"),
            (_component("standard", "19000", "17000"),),
        )
        minimum_above_target = design_agency_rewards(
            _profile(target="100", minimum="200"),
            (_component("standard", "19000", "17000"),),
        )

        self.assertEqual(target_unreachable.status, "hold")
        self.assertEqual(minimum_above_target.status, "hold")

    def test_ftth_payment_is_forbidden_even_when_previous_price_was_paid(self) -> None:
        result = design_agency_rewards(
            _profile(),
            (
                _component(
                    "ftth-fee",
                    "2000",
                    "2000",
                    target="2000",
                    opening_class="FTTH",
                    payment_allowed=True,
                    payment_component=FTTH_OPENING_FEE_COMPONENT,
                ),
            ),
        )

        self.assertEqual(result.status, "candidate", result.hold_reasons)
        self.assertEqual(
            result.component_proposals[0].proposed_agency_unit_reward,
            Decimal("0"),
        )

    def test_commerce_wide_threshold_applies_unit_rate_to_agency_eligible_count(self) -> None:
        campaign = _campaign(
            "volume",
            (500, 750),
            ("6500", "9000"),
            ("5500", "8000"),
            target="1000",
            opening_classes=("5G", "FTTH"),
        )
        result = design_agency_rewards(
            _profile(),
            (),
            (campaign,),
            volume_scenarios=(
                VolumeScenario(
                    scenario_id="all-channel-500-keioh-3",
                    campaign_code="volume",
                    commerce_count=500,
                    agency_eligible_count=3,
                    opening_class="FTTH",
                    upstream_collected=True,
                    agency_payment_due=True,
                ),
            ),
        )
        outcome = next(
            scenario
            for scenario in result.scenarios
            if scenario.scenario_id == "all-channel-500-keioh-3"
        )

        self.assertEqual(result.status, "candidate", result.hold_reasons)
        self.assertEqual(outcome.commerce_count, 500)
        self.assertEqual(outcome.agency_eligible_count, 3)
        self.assertEqual(outcome.upstream_amount, Decimal("19500"))
        self.assertEqual(outcome.agency_payment, Decimal("16500"))
        self.assertEqual(outcome.retained_gross_margin, Decimal("3000"))

    def test_ftth_opening_fee_amount_must_match_known_contract_rate(self) -> None:
        result = design_agency_rewards(
            _profile(),
            (
                _component(
                    "ftth-fee",
                    "2500",
                    "0",
                    opening_class="FTTH",
                    payment_component=FTTH_OPENING_FEE_COMPONENT,
                ),
            ),
        )

        self.assertEqual(result.status, "hold")
        self.assertFalse(result.component_proposals)
        self.assertTrue(any("2,000円" in reason for reason in result.hold_reasons))

    def test_uncollected_upstream_payment_risk_blocks_candidate(self) -> None:
        result = design_agency_rewards(
            _profile(pay_after_collection=False),
            (_component("standard", "100", "80"),),
        )

        self.assertEqual(result.status, "hold")
        self.assertTrue(any("上位未収" in reason for reason in result.hold_reasons))
        self.assertTrue(
            any(
                scenario.risks
                and not scenario.upstream_collected
                and scenario.agency_payment_due
                for scenario in result.scenarios
            )
        )

    def test_unknown_ladder_mode_and_unresolved_business_judgment_hold(self) -> None:
        campaign = _campaign(
            "unknown-mode",
            (100,),
            ("1000",),
            ("500",),
            target="100",
            mode=LadderMode.UNKNOWN,
        )
        judgment = JudgmentReview(
            issue="対象計数に解約済み案件を含むか",
            source_document="campaign-notice.pdf",
            ai_review_prompt="原本から対象件数の取消・解約条件を要約し、要承認事項を列挙する。",
        )
        result = design_agency_rewards(
            _profile(),
            (),
            (campaign,),
            (judgment,),
        )

        self.assertEqual(result.status, "hold")
        self.assertTrue(any("適用方式" in reason for reason in result.hold_reasons))
        self.assertEqual(result.manual_judgments, (judgment,))

    def test_changes_disabled_preserve_previous_amount_or_hold(self) -> None:
        feasible = design_agency_rewards(
            _profile(target="1000", allow_changes=False),
            (_component("standard", "19000", "17000"),),
        )
        infeasible = design_agency_rewards(
            _profile(target="3000", allow_changes=False),
            (_component("standard", "19000", "17000"),),
        )

        self.assertEqual(feasible.status, "candidate")
        self.assertEqual(
            feasible.component_proposals[0].proposed_agency_unit_reward,
            Decimal("17000"),
        )
        self.assertEqual(infeasible.status, "hold")

    def test_p2_connection_requires_current_source_result_and_never_approves(self) -> None:
        design = design_agency_rewards(
            _profile(target="4000"),
            (_component("standard", "19000", "17000"),),
        )
        p2 = P2Report(
            target_period="2026-11",
            results=(
                P2AgencyResult(
                    target_period="2026-11",
                    agency_id="keioh",
                    condition_item_id="standard",
                    status="proposed",
                    upstream_reward=Decimal("19000"),
                    direct_cost=Decimal("0"),
                    approved_gross_margin=Decimal("2000"),
                    previous_agency_reward=Decimal("17000"),
                    proposed_agency_reward=Decimal("17000"),
                    proposed_difference=Decimal("0"),
                    display_policy="present",
                    hold_reasons=(),
                ),
            ),
        )
        gate = connect_reward_design_to_p2(design, p2)

        self.assertEqual(gate.status, "requires_p2_reapproval")
        self.assertTrue(gate.requires_p2_reapproval)
        self.assertEqual(gate.hold_reasons, ())

        held_p2 = P2Report(
            target_period="2026-11",
            results=(
                P2AgencyResult(
                    target_period="2026-11",
                    agency_id="keioh",
                    condition_item_id="standard",
                    status="hold",
                    upstream_reward=None,
                    direct_cost=None,
                    approved_gross_margin=None,
                    previous_agency_reward=None,
                    proposed_agency_reward=None,
                    proposed_difference=None,
                    display_policy="undecided",
                    hold_reasons=("承認記録がありません。",),
                ),
            ),
        )
        blocked = connect_reward_design_to_p2(design, held_p2)
        self.assertEqual(blocked.status, "hold")
        self.assertTrue(blocked.requires_p2_reapproval)
        self.assertTrue(any("承認記録" in reason for reason in blocked.hold_reasons))

    def test_p2_candidate_copy_changes_margin_and_keeps_hash_gate_closed(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            input_path = Path(temporary_directory) / "approved-input.xlsx"
            candidate_path = Path(temporary_directory) / "p2-design-candidate.xlsx"
            _build_fixture(input_path)
            p2_report = calculate_workbook(input_path, "2026-11")
            design = design_agency_rewards(
                _profile(agency_id="anon-agency-01", target="40"),
                (
                    _component(
                        "anon-item-01",
                        "110",
                        "60",
                        target="40",
                        direct_cost="10",
                        category=RewardCategory.BASIC,
                    ),
                ),
            )

            written = write_p2_reward_design_candidate(
                input_path,
                candidate_path,
                design,
                p2_report,
            )

            self.assertEqual(written.output_path, candidate_path)
            self.assertEqual(
                written.updated_condition_item_ids,
                ("anon-item-01",),
            )
            self.assertEqual(written.post_write_p2_report.hold_count, 1)
            self.assertTrue(
                any(
                    "承認ハッシュ" in reason
                    for reason in written.post_write_p2_report.results[0].hold_reasons
                )
            )
            original = load_workbook(input_path, read_only=True, data_only=True)
            candidate = load_workbook(
                candidate_path, read_only=True, data_only=True
            )
            try:
                margin_column = next(
                    index
                    for index, cell in enumerate(
                        original["粗利原価"][1], start=1
                    )
                    if cell.value == "approved_target_gross_margin"
                )
                item_column = next(
                    index
                    for index, cell in enumerate(
                        original["粗利原価"][1], start=1
                    )
                    if cell.value == "condition_item_id"
                )
                valid_from_column = next(
                    index
                    for index, cell in enumerate(
                        original["粗利原価"][1], start=1
                    )
                    if cell.value == "valid_from"
                )
                item_row = next(
                    row_number
                    for row_number in range(2, original["粗利原価"].max_row + 1)
                    if original["粗利原価"].cell(
                        row_number, item_column
                    ).value == "anon-item-01"
                    and str(
                        original["粗利原価"].cell(
                            row_number, valid_from_column
                        ).value
                    ).startswith("2026-11")
                )
                self.assertEqual(
                    original["粗利原価"].cell(item_row, margin_column).value,
                    "30",
                )
                self.assertEqual(
                    candidate["粗利原価"].cell(item_row, margin_column).value,
                    "40",
                )
            finally:
                original.close()
                candidate.close()


if __name__ == "__main__":
    unittest.main()
