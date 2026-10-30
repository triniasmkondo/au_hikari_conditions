from copy import deepcopy
from dataclasses import replace
from decimal import Decimal
from pathlib import Path
import tempfile
import unittest

from openpyxl import load_workbook
from au_hikari_conditions.condition_ledger import Condition
from au_hikari_conditions.upper_conditions import UpperCondition,extract_upper_pdf
from au_hikari_conditions.phase2_review import (read_phase1,make_proposals,empty_input,
    INPUT_SCHEMAS,evaluate_proposal,build_phase2,read_inputs)
from au_hikari_conditions.phase2_triage import initial_candidate,build_triage

ROOT=Path(__file__).resolve().parents[1]
LEDGER=ROOT/'04_成果物/Phase1/2026-10_11社共通条件台帳_監査修正版.xlsx'


def fixture(agency='架空社'):
    old=Condition(agency,'通常報酬','auひかり','回線','auひかりホーム',amount=22000,
        unit='円/件',tax='税別',period='2026-10-01/2026-10-31',source='old-agency.xlsx',locations=['A1'],digest='old-agent-hash')
    prev=UpperCondition('auひかり','base','net_open',dwelling='ホーム',amount=24000,
        period='2026-10-01/2027-03-31',source='old-upper.pdf',location='page:1',digest='old-upper-hash')
    new=replace(prev,amount=25000,period='2026-11-01/2027-03-31',source='new-upper.pdf',digest='new-upper-hash')
    p=make_proposals([old],[prev],[new],'2026-11')[0]
    records={name:empty_input(headers,p,'snapshot','2026-11') for name,headers in INPUT_SCHEMAS.items()}
    result=evaluate_proposal(p,records,'snapshot','2026-11')
    return p,records,result


class InitialCandidateTest(unittest.TestCase):
    def test_spread_and_conditional_price_are_not_approved_margin_or_zero_cost(self):
        p,inputs,result=fixture();c=initial_candidate(p,p.upper,'2026-10','2026-11')
        self.assertEqual(c.spread,Decimal(2000));self.assertEqual(c.candidate,Decimal(23000))
        self.assertIn('未確認',c.assumptions);self.assertIn('追加受領10',c.missing)
        self.assertIn('実粗利未算出',c.margin_state);self.assertIn('未確定',c.candidate_state)
        self.assertIsNone(result['proposed']);self.assertEqual(result['status'],'HELD')
        self.assertIsNone(inputs['粗利承認入力']['approved_gross_margin'])
        self.assertIsNone(inputs['粗利承認入力']['direct_cost'])
        self.assertIn('old-upper-hash',c.evidence);self.assertIn('new-upper-hash',c.evidence)

    def test_unknown_tax_units_hold_ambiguous_or_dash_never_reverse_margin(self):
        for mutate in (
            lambda p:setattr(p.old,'tax','原本に税区分明示なし'),
            lambda p:setattr(p.old,'unit','円/1ID'),
            lambda p:setattr(p.old,'hold','原本コード未確定'),
            lambda p:setattr(p.previous,'state','DASH'),
            lambda p:setattr(p.previous,'amount',None),
            lambda p:setattr(p.previous,'code','KWE'),
            lambda p:setattr(p,'initial_identity_collision',True),
        ):
            p,_,_=fixture();mutate(p)
            c=initial_candidate(p,p.upper,'2026-10','2026-11')
            self.assertIsNone(c.spread);self.assertIsNone(c.candidate);self.assertTrue(c.exclusion)

    def test_scope_change_or_non_numeric_new_price_keeps_quote_empty(self):
        for mutate in (
            lambda p:setattr(p.upper,'access','事業者変更'),
            lambda p:setattr(p.upper,'type','3'),
            lambda p:setattr(p.upper,'state','DASH'),
            lambda p:setattr(p.upper,'amount',None),
            lambda p:setattr(p.upper,'period','2026-12-01/2027-03-31'),
            lambda p:setattr(p.upper,'amount',1000),
        ):
            p,_,_=fixture();mutate(p)
            c=initial_candidate(p,p.upper,'2026-10','2026-11')
            self.assertEqual(c.spread,Decimal(2000));self.assertIsNone(c.candidate)
            self.assertTrue(c.exclusion)

    def test_previous_rate_change_invalidates_final_fingerprint(self):
        p,inputs,first=fixture();p.previous.amount=26000
        second=evaluate_proposal(p,inputs,'snapshot','2026-11')
        self.assertNotEqual(first['candidate_hash'],second['candidate_hash'])
        self.assertEqual(second['initial_candidate'].candidate,Decimal(21000))
        self.assertEqual(second['status'],'HELD')

    def test_same_axes_multiple_old_rows_are_excluded(self):
        p,_,_=fixture();another=replace(p.old,terms='追加の独自支払条件',amount=21000)
        ps=make_proposals([p.old,another],[p.previous],[p.upper],'2026-11')
        self.assertEqual(sum(p.old is not None for p in ps),2)
        self.assertTrue(all(p.initial_identity_collision for p in ps if p.old))
        self.assertTrue(all(initial_candidate(p,p.upper,'2026-10','2026-11').spread is None for p in ps if p.old))


class TriageCoverageTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        old,_=read_phase1(ROOT,LEDGER);ups=[]
        for month in ('2026-10','2026-11'):
            records=[]
            for path in (ROOT/'01_上位原本'/month).glob('*.pdf'):records+=extract_upper_pdf(path,ROOT)[0]
            ups.append(records)
        ps=make_proposals(old,*ups,'2026-11');cls.results=[]
        for p in ps:
            records={name:empty_input(headers,p,'snapshot','2026-11') for name,headers in INPUT_SCHEMAS.items()}
            cls.results.append(evaluate_proposal(p,records,'snapshot','2026-11'))
        cls.groups,cls.members,cls.stats=build_triage(cls.results,'2026-10','2026-11')

    def test_all_819_review_rows_are_preserved_and_reasons_do_not_claim_resolution(self):
        s=self.stats
        self.assertEqual(s['original_review_rows'],819);self.assertEqual(s['covered_review_rows'],819)
        self.assertEqual(sum(s['primary_reason_counts'].values()),819)
        self.assertEqual(s['unresolved_review_rows'],819);self.assertEqual(s['resolved_by_grouping'],0)
        self.assertGreater(s['question_count_reduction'],0)
        self.assertEqual(s['review_question_count'],s['common_questions']+s['individual_questions'])
        for group in self.groups.values():
            agencies={r['proposal'].agency for r in self.members[group.id]}
            if group.scope=='代理店個別':self.assertEqual(len(agencies),1)
        layout=[i for i in self.groups.values() if i.code=='UPPER_LAYOUT']
        self.assertEqual(len(layout),2)
        self.assertTrue(all(len({r['proposal'].agency for r in self.members[g.id]})==11 for g in layout))
        self.assertTrue(all(r['status']=='HELD' and r['approved_amount'] is None for r in self.results))

    def test_real_arithmetic_429_spreads_407_conditional_prices(self):
        self.assertEqual(self.stats['initial_spread_count'],429)
        self.assertEqual(self.stats['conditional_price_count'],407)
        keioh=[r for r in self.results if r['proposal'].agency=='keioh株式会社']
        support=next(r for r in keioh if r['proposal'].old and r['proposal'].old.item=='販促支援金' and r['proposal'].old.threshold=='2,000件~')
        self.assertEqual(support['initial_candidate'].spread,Decimal(2000))
        self.assertEqual(support['initial_candidate'].candidate,Decimal(3000))
        adjustment=next(r for r in keioh if r['proposal'].old and r['proposal'].old.item=='特別調整金')
        self.assertEqual(adjustment['initial_candidate'].spread,Decimal(2000))
        self.assertIsNone(adjustment['initial_candidate'].candidate)
        self.assertTrue(all(r['initial_candidate'].candidate is None for r in self.results if r['proposal'].old and r['proposal'].old.hold))

    def test_excel_separates_initial_candidates_from_approvals_and_provides_trace(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'keioh.xlsx'
            stats=build_phase2(ROOT,LEDGER,path,agency_filter='keioh株式会社')
            self.assertEqual(stats['triage']['initial_spread_count'],39)
            self.assertEqual(stats['triage']['conditional_price_count'],37)
            self.assertEqual(stats['calculation_proposals'],0)
            book=load_workbook(path,read_only=True)
            self.assertEqual(book['粗利初期候補'].max_row,40)
            self.assertEqual(book['担当者確認一覧'].max_row,151)
            self.assertIn('要確認理由別集計',book.sheetnames)
            header=list(next(book['11月条件案'].values));rows=list(book['11月条件案'].iter_rows(min_row=2,values_only=True))
            self.assertEqual(sum(r[header.index('11月条件付き候補単価')] is not None for r in rows),37)
            self.assertTrue(all(r[header.index('提案単価')] is None for r in rows))
            self.assertTrue(all(r[header.index('承認額')] is None for r in rows))
            review_pids={r[1] for r in rows if r[header.index('分類案')]=='要確認'}
            coverage={r[4] for r in book['論点対象明細'].iter_rows(min_row=2,values_only=True)}
            self.assertTrue(review_pids<=coverage)
            book.close()
            records,_=read_inputs(path)
            self.assertTrue(all(r['approved_gross_margin'] is None and r['direct_cost'] is None for r in records['粗利承認入力'].values()))
