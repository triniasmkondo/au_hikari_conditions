from __future__ import annotations

import sys
import tempfile
import types
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))
sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

from au_hikari_conditions.monthly_bridge import (
    MonthlyBridgeError,
    build_monthly_comparison_workbook,
    classify_agency_document_type,
    discover_agency_document_set,
    validate_monthly_comparison_workbook,
)
from prepare_monthly_comparison import _choose_agency_document


class MonthlyBridgeTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary = tempfile.TemporaryDirectory()
        self.root = Path(self.temporary.name) / "fictional-project"
        self.output = Path(self.temporary.name) / "comparison.xlsx"
        self.previous_upper = self.root / "01_上位原本" / "2026-10"
        self.previous_agency = self.root / "03_条件表" / "2026-10"
        self.current_upper = self.root / "01_上位原本" / "2026-11"
        for folder in (self.previous_upper, self.previous_agency, self.current_upper):
            folder.mkdir(parents=True)

    def tearDown(self) -> None:
        self.temporary.cleanup()

    def _write_docx(self, path: Path, text: str) -> None:
        document = Document()
        document.add_paragraph(text)
        document.save(path)

    def test_builds_held_source_and_dimension_comparison_for_all_three_inputs(self) -> None:
        self._write_docx(
            self.previous_upper / "fictional-previous-upper.docx",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
        )
        self._write_docx(
            self.previous_agency / "fictional-agency-conditions.docx",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
        )
        self._write_docx(
            self.current_upper / "fictional-current-upper.docx",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 120円/件",
        )

        report = build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
        )

        self.assertEqual(report.source_count, 3)
        self.assertGreaterEqual(report.candidate_count, 3)
        self.assertGreater(report.upper_difference_count, 0)
        self.assertGreater(report.unresolved_count, 0)
        self.assertEqual(report.status, "REVIEW_REQUIRED")
        workbook = load_workbook(self.output, read_only=True, data_only=True)
        try:
            self.assertEqual(
                set(workbook.sheetnames),
                {
                    "使い方",
                    "比較情報",
                    "資料一覧",
                    "抽出根拠候補",
                    "上位差分候補",
                    "代理店資料種別",
                    "代理店別上位比較",
                    "代理店資料間照合",
                    "P2項目対応",
                    "不足・保留",
                },
            )
            candidates = workbook["抽出根拠候補"]
            headers = [cell.value for cell in candidates[1]]
            rows = [
                dict(zip(headers, row))
                for row in candidates.iter_rows(min_row=2, values_only=True)
            ]
            agency_rows = [row for row in rows if row["source_type"] == "PREVIOUS_AGENCY"]
            self.assertTrue(agency_rows)
            self.assertTrue(
                all(row["agency_identity_status"] == "MANUAL_AGENCY_ID_REQUIRED" for row in agency_rows)
            )
            self.assertTrue(all(row["approval_status"] == "UNAPPROVED" for row in rows))

            differences = workbook["上位差分候補"]
            diff_headers = [cell.value for cell in differences[1]]
            diff_rows = [
                dict(zip(diff_headers, row))
                for row in differences.iter_rows(min_row=2, values_only=True)
            ]
            amount_differences = [
                row for row in diff_rows if row["item_id"] == "p1.commission_amount"
            ]
            self.assertEqual(len(amount_differences), 1)
            self.assertEqual(amount_differences[0]["candidate_status"], "CHANGED_CANDIDATE")
            self.assertEqual(amount_differences[0]["approval_status"], "UNAPPROVED")

            crosswalk = workbook["P2項目対応"]
            self.assertTrue(
                all(row[4] == "NO" for row in crosswalk.iter_rows(min_row=2, values_only=True))
            )
        finally:
            workbook.close()

    def test_refuses_output_inside_project_and_existing_output(self) -> None:
        with self.assertRaisesRegex(MonthlyBridgeError, "outside the Git repository"):
            build_monthly_comparison_workbook(
                project_root=self.root,
                previous_period="2026-10",
                target_period="2026-11",
                output_path=self.root / "unsafe.xlsx",
            )

        self._write_docx(
            self.previous_upper / "fictional-previous-upper.docx",
            "auひかりプラス BIGLOBE ホーム 1G 開通 報奨金額 100円",
        )
        self._write_docx(
            self.previous_agency / "fictional-agency.docx",
            "auひかりプラス BIGLOBE ホーム 1G 開通",
        )
        self._write_docx(
            self.current_upper / "fictional-current-upper.docx",
            "auひかりプラス BIGLOBE ホーム 1G 開通 報奨金額 110円",
        )
        self.output.write_bytes(b"existing")
        with self.assertRaisesRegex(MonthlyBridgeError, "Refusing to overwrite"):
            build_monthly_comparison_workbook(
                project_root=self.root,
                previous_period="2026-10",
                target_period="2026-11",
                output_path=self.output,
            )

    def test_comparison_validation_binds_periods_and_source_hashes(self) -> None:
        source_paths = (
            self.previous_upper / "fictional-previous-upper.docx",
            self.previous_agency / "fictional-agency-conditions.docx",
            self.current_upper / "fictional-current-upper.docx",
        )
        for index, path in enumerate(source_paths):
            self._write_docx(
                path,
                f"auひかりプラス BIGLOBE ホーム 1G 開通 報奨金額 {100 + index}円",
            )
        build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
        )

        validation = validate_monthly_comparison_workbook(
            self.output,
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
        )
        self.assertEqual(validation.source_count, 3)
        with self.assertRaisesRegex(MonthlyBridgeError, "periods or candidate status"):
            validate_monthly_comparison_workbook(
                self.output,
                project_root=self.root,
                previous_period="2026-09",
                target_period="2026-11",
            )

        source_paths[0].write_bytes(source_paths[0].read_bytes() + b"changed")
        with self.assertRaisesRegex(MonthlyBridgeError, "changed after comparison"):
            validate_monthly_comparison_workbook(
                self.output,
                project_root=self.root,
                previous_period="2026-10",
                target_period="2026-11",
            )

    def test_comparison_candidates_cannot_be_marked_approved(self) -> None:
        for index, folder in enumerate(
            (self.previous_upper, self.previous_agency, self.current_upper)
        ):
            self._write_docx(
                folder / f"fictional-source-{index}.docx",
                f"auひかりプラス BIGLOBE ホーム 1G 開通 報奨金額 {100 + index}円",
            )
        build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
        )
        workbook = load_workbook(self.output)
        worksheet = workbook["抽出根拠候補"]
        headers = [cell.value for cell in worksheet[1]]
        worksheet.cell(2, headers.index("approval_status") + 1, "APPROVED")
        workbook.save(self.output)
        workbook.close()

        with self.assertRaisesRegex(MonthlyBridgeError, "cannot be approved"):
            validate_monthly_comparison_workbook(
                self.output,
                project_root=self.root,
                previous_period="2026-10",
                target_period="2026-11",
            )

    def test_changed_condition_dimension_is_held_as_mismatch(self) -> None:
        self._write_docx(
            self.previous_upper / "fictional-previous-upper.docx",
            "auひかり プラス BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
        )
        self._write_docx(
            self.previous_agency / "fictional-agency-conditions.docx",
            "auひかり プラス BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
        )
        self._write_docx(
            self.current_upper / "fictional-current-upper.docx",
            "auひかり プラス So-net ホーム 1G 新規 開通 報奨金額 120円/件",
        )

        build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
        )
        workbook = load_workbook(self.output, read_only=True, data_only=True)
        try:
            sheet = workbook["上位差分候補"]
            headers = [cell.value for cell in sheet[1]]
            differences = [
                dict(zip(headers, row))
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if row[0] is not None
            ]
            amount_differences = [
                row
                for row in differences
                if row["item_id"] == "p1.commission_amount"
            ]
            self.assertEqual(len(amount_differences), 1)
            self.assertEqual(
                amount_differences[0]["candidate_status"],
                "CONDITION_MISMATCH_REQUIRES_REVIEW",
            )
            self.assertTrue(amount_differences[0]["previous_values_json"])
            self.assertTrue(amount_differences[0]["current_values_json"])
            self.assertEqual(amount_differences[0]["approval_status"], "UNAPPROVED")
        finally:
            workbook.close()

    def test_groups_three_named_condition_document_types_for_one_agency(self) -> None:
        contents = {
            "fictional_agency_auひかり.docx": "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
            "fictional_agency_auひかりプラス.docx": "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 90円/件",
            "fictional_agency_キャンペーン条件.docx": "auひかりプラス BIGLOBE ホーム 1G キャンペーン 特典 20円/件",
        }
        documents = []
        for filename, text in contents.items():
            path = self.previous_agency / filename
            self._write_docx(path, text)
            documents.append(path)

        group = discover_agency_document_set(
            project_root=self.root,
            agency_directory=self.previous_agency,
            selected_document=documents[0],
        )

        self.assertEqual(
            set(group.documents),
            {document.resolve() for document in documents},
        )
        self.assertEqual(
            set(group.type_by_path.values()),
            {"AU_HIKARI", "AU_HIKARI_PLUS", "CAMPAIGN_CONDITIONS"},
        )
        self.assertTrue(
            any("同一代理店" in reason for reason in group.hold_reasons)
        )
        self.assertEqual(classify_agency_document_type(documents[0]), "AU_HIKARI")

    def test_grouping_excludes_other_agency_and_holds_missing_or_unknown_types(self) -> None:
        selected = self.previous_agency / "fictional_agency_auひかり.docx"
        self._write_docx(selected, "auひかり BIGLOBE ホーム 1G 報奨金額 80円/件")
        self._write_docx(
            self.previous_agency / "fictional_agency_auひかりプラス.docx",
            "auひかりプラス BIGLOBE ホーム 1G 報奨金額 90円/件",
        )
        self._write_docx(
            self.previous_agency / "other_agency_キャンペーン条件.docx",
            "auひかりプラス BIGLOBE ホーム 1G キャンペーン 20円/件",
        )

        group = discover_agency_document_set(
            project_root=self.root,
            agency_directory=self.previous_agency,
            selected_document=selected,
        )

        self.assertEqual(len(group.documents), 2)
        self.assertTrue(any("キャンペーン条件" in reason for reason in group.hold_reasons))
        self.assertNotIn(
            "CAMPAIGN_CONDITIONS",
            group.type_by_path.values(),
        )

        unknown = self.previous_agency / "unclassified_agency_terms.docx"
        self._write_docx(unknown, "独自条件")
        unknown_group = discover_agency_document_set(
            project_root=self.root,
            agency_directory=self.previous_agency,
            selected_document=unknown,
        )
        self.assertEqual(unknown_group.type_by_path[unknown.resolve()], "UNKNOWN")
        self.assertTrue(unknown_group.hold_reasons)

    def test_three_document_comparison_keeps_types_separate_and_campaign_unadded(self) -> None:
        self._write_docx(
            self.previous_upper / "fictional-upper-previous.docx",
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
        )
        self._write_docx(
            self.current_upper / "fictional-upper-current.docx",
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 110円/件",
        )
        source_texts = {
            "fictional_agency_auひかり.docx": "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
            "fictional_agency_auひかりプラス.docx": "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 90円/件",
            "fictional_agency_キャンペーン条件.docx": "auひかりプラス BIGLOBE ホーム 1G キャンペーン 特典 報奨金額 20円/件",
        }
        selected_paths = []
        for filename, text in source_texts.items():
            path = self.previous_agency / filename
            self._write_docx(path, text)
            selected_paths.append(path)
        document_set = discover_agency_document_set(
            project_root=self.root,
            agency_directory=self.previous_agency,
            selected_document=selected_paths[0],
        )

        report = build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
            agency_document_set=document_set,
        )

        self.assertEqual(report.source_count, 5)
        self.assertEqual(report.status, "REVIEW_REQUIRED")
        workbook = load_workbook(self.output, read_only=True, data_only=True)
        try:
            roles = workbook["代理店資料種別"]
            role_headers = [cell.value for cell in roles[1]]
            role_rows = [
                dict(zip(role_headers, row))
                for row in roles.iter_rows(min_row=2, values_only=True)
            ]
            self.assertEqual(len(role_rows), 3)
            self.assertEqual(
                {row["document_type_candidate"] for row in role_rows},
                {"AU_HIKARI", "AU_HIKARI_PLUS", "CAMPAIGN_CONDITIONS"},
            )

            comparisons = workbook["代理店別上位比較"]
            comparison_headers = [cell.value for cell in comparisons[1]]
            comparison_rows = [
                dict(zip(comparison_headers, row))
                for row in comparisons.iter_rows(min_row=2, values_only=True)
            ]
            self.assertEqual(
                {row["document_type"] for row in comparison_rows},
                {"AU_HIKARI", "AU_HIKARI_PLUS", "CAMPAIGN_CONDITIONS"},
            )
            campaign_rows = [
                row for row in comparison_rows
                if row["document_type"] == "CAMPAIGN_CONDITIONS"
            ]
            self.assertTrue(campaign_rows)
            self.assertTrue(
                all(row["campaign_product_link_status"] == "PRODUCT_LINK_CANDIDATE" for row in campaign_rows)
            )
            self.assertTrue(
                all(row["candidate_status"] == "REVIEW_REQUIRED" for row in campaign_rows)
            )
            self.assertTrue(
                all(row["review_status"] == "UNAPPROVED" for row in comparison_rows)
            )

            sources = workbook["資料一覧"]
            source_headers = [cell.value for cell in sources[1]]
            source_rows = [
                dict(zip(source_headers, row))
                for row in sources.iter_rows(min_row=2, values_only=True)
            ]
            agency_source_rows = [
                row for row in source_rows if row["source_type"] == "PREVIOUS_AGENCY"
            ]
            self.assertEqual(len(agency_source_rows), 3)
        finally:
            workbook.close()

    def test_duplicate_and_conflicting_same_type_conditions_are_held(self) -> None:
        self._write_docx(
            self.previous_upper / "fictional-upper-previous.docx",
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 100円/件",
        )
        self._write_docx(
            self.current_upper / "fictional-upper-current.docx",
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 110円/件",
        )
        base_name = "fictional_agency_auひかり.docx"
        self._write_docx(
            self.previous_agency / base_name,
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 80円/件",
        )
        duplicate_folder = self.previous_agency / "revision"
        duplicate_folder.mkdir()
        self._write_docx(
            duplicate_folder / base_name,
            "auひかり BIGLOBE ホーム 1G 新規 開通 報奨金額 85円/件",
        )
        self._write_docx(
            self.previous_agency / "fictional_agency_auひかりプラス.docx",
            "auひかりプラス BIGLOBE ホーム 1G 新規 開通 報奨金額 90円/件",
        )
        self._write_docx(
            self.previous_agency / "fictional_agency_キャンペーン条件.docx",
            "auひかりプラス BIGLOBE ホーム 1G キャンペーン 報奨金額 20円/件",
        )

        group = discover_agency_document_set(
            project_root=self.root,
            agency_directory=self.previous_agency,
            selected_document=self.previous_agency / base_name,
        )
        report = build_monthly_comparison_workbook(
            project_root=self.root,
            previous_period="2026-10",
            target_period="2026-11",
            output_path=self.output,
            agency_document_set=group,
        )

        self.assertTrue(group.hold_reasons)
        self.assertGreater(report.unresolved_count, 0)
        workbook = load_workbook(self.output, read_only=True, data_only=True)
        try:
            overlaps = workbook["代理店資料間照合"]
            headers = [cell.value for cell in overlaps[1]]
            rows = [
                dict(zip(headers, row))
                for row in overlaps.iter_rows(min_row=2, values_only=True)
                if row[0] is not None
            ]
            self.assertTrue(
                any(row["candidate_status"] == "CONFLICT_REQUIRES_REVIEW" for row in rows)
            )
            self.assertTrue(
                any(row["candidate_status"] == "DUPLICATE_REQUIRES_REVIEW" for row in rows)
            )
            self.assertTrue(all(row["review_status"] == "UNAPPROVED" for row in rows))
        finally:
            workbook.close()

    def test_interactive_selection_does_not_log_selected_path(self) -> None:
        selected_path = self.previous_agency / "fictional-agency.docx"
        dialog_calls: list[dict[str, object]] = []

        class FakeRoot:
            def withdraw(self) -> None:
                pass

            def destroy(self) -> None:
                pass

        fake_tk = types.ModuleType("tkinter")
        fake_tk.Tk = FakeRoot
        fake_dialog = types.ModuleType("tkinter.filedialog")
        fake_tk.filedialog = fake_dialog
        fake_dialog.askopenfilename = lambda **kwargs: (
            dialog_calls.append(kwargs) or str(selected_path)
        )
        with patch.dict(
            sys.modules,
            {"tkinter": fake_tk, "tkinter.filedialog": fake_dialog},
        ):
            result = _choose_agency_document(self.previous_agency)

        self.assertEqual(result, selected_path)
        self.assertEqual(dialog_calls[0]["initialdir"], str(self.previous_agency))
        self.assertTrue(dialog_calls[0]["filetypes"])


if __name__ == "__main__":
    unittest.main()
