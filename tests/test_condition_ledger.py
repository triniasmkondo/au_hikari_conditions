import hashlib
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from au_hikari_conditions.condition_ledger import (
    agency_candidate, build_ledger, candidate_key, extract_source)

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'03_条件表/2026-10'


class LedgerTest(unittest.TestCase):
    def test_names_and_semantic_keys(self):
        self.assertEqual(agency_candidate(Path('【202610】【会社 御中】条件.xlsx')), '会社')
        self.assertEqual(agency_candidate(Path('【202610】【会社様】条件.xlsx')), '会社')
        self.assertNotEqual(candidate_key('項目',{'施策':'A'},'条件'),candidate_key('項目',{'施策':'B'},'条件'))

    def test_keioh_all_72_cells_and_15_tiers(self):
        paths=list(SOURCE.glob('*keioh*'))
        self.assertEqual(len(paths),3)
        all_rows=[]
        for path in paths:
            before=hashlib.sha256(path.read_bytes()).hexdigest()
            rows,coverage,legal,raw=extract_source(path,ROOT)
            self.assertEqual(before,hashlib.sha256(path.read_bytes()).hexdigest())
            all_rows.extend(rows)
            if path.suffix=='.xlsx':
                plus='プラス' in path.name
                coords=([f'{c}{r}' for r in range(18,21) for c in 'DEF']+
                    [f'{c}{r}' for r in range(24,28) for c in 'DEF']+
                    [f'E{r}' for r in range(32,43)]+[f'E{r}' for r in range(51,63)]+
                    [f'D{r}' for r in range(68,72)]) if plus else (
                    [f'F{r}' for r in [19,20,21,26,27,32,33,34,39,40,50,61,62,63,64,65,66,67,68,76,77]]+
                    ['E83','E85','E87'])
                book=load_workbook(path);sheet=book.worksheets[0]
                for coord in coords:
                    loc=f'sheet:{sheet.title};cell:{coord}'
                    found=[c for c in rows if loc in c.locations]
                    self.assertEqual(len(found),1,loc)
                    c=found[0]
                    self.assertEqual(c.amount,sheet[coord].value,loc)
                    self.assertTrue(c.item and c.section and c.period and c.legal,loc)
                    if c.amount<0:self.assertEqual(c.kind,'控除')
                if plus:
                    net=next(c for c in rows if 'sheet:auひかりプラス;cell:E18' in c.locations)
                    self.assertEqual((net.item,net.type,net.target,net.amount),('ネット開通手数料','タイプ1','ホーム',25000))
                    code=next(c for c in rows if 'sheet:auひかりプラス;cell:E33' in c.locations)
                    self.assertEqual((code.type,code.target,code.code,code.amount),('タイプ1','ホーム(1G)','KCM',500))
                    exclusive=next(c for c in rows if 'sheet:auひかりプラス;cell:E54' in c.locations)
                    self.assertIn('手数料対象外',exclusive.terms)
                    self.assertIn('2027-03-31',net.period)
                else:
                    option=next(c for c in rows if 'sheet:auひかり;cell:F62' in c.locations)
                    self.assertEqual(option.item,'トータル・ ネットセキュリティ')
                    self.assertEqual((option.isp,option.target,option.code),('BIGLOBE','マンション','BDC、NBT、CCB、ACB'))
                    bonus=next(c for c in rows if 'sheet:auひかり;cell:F50' in c.locations)
                    self.assertIn('250件以上',bonus.terms)
                    self.assertTrue(any(c.kind=='違約金・相殺' and c.amount==100000 for c in rows))
                # Every text row, including continuation sentences, is retained on the monetary rows.
                for r in raw:
                    if not any(isinstance(c.value,(int,float)) for c in sheet[int(r[2].split(':')[-1])]):
                        self.assertIn(r[3],legal[0][3])
                book.close()
            else:
                tiers=[c for c in rows if c.kind=='キャンペーン']
                self.assertEqual(len(tiers),15)
                self.assertEqual([c.amount for c in tiers],[5500,8000,10500,13000,15500,16500,1000,2000,3000,1000,3000,3000,3000,4000,5000])
                self.assertEqual([c.threshold for c in tiers[:6]],['500件~','750件~','1,000件~','1,500件~','2,000件~','2,500件~'])
                self.assertTrue(all(c.type and c.terms and c.period for c in tiers))
                self.assertIn('合計数',tiers[0].terms)
                self.assertEqual(sorted(c.amount for c in rows if c.kind=='違約金'),[500000,1000000])
                self.assertTrue(all('6 ヶ月以内' in c.legal for c in tiers))
        self.assertEqual(len(all_rows),90)
        self.assertEqual(len({c.id for c in all_rows}),90)
        zero=[c for c in all_rows if c.amount==0]
        self.assertEqual(len(zero),10)

    def test_workbook_roundtrip_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as folder:
            output=Path(folder)/'out.xlsx'
            result=build_ledger(ROOT,output,agency_filter='keioh株式会社')
            self.assertEqual(result['source_count'],3)
            self.assertEqual(result['monetary_coverage_count'],90)
            self.assertFalse(result['next_phase_allowed'])
            w=load_workbook(output)
            self.assertEqual(w['共通条件台帳'].max_row,91)
            self.assertTrue(all(isinstance(r[11].value,(int,float)) for r in w['共通条件台帳'].iter_rows(min_row=2)))
            w.close()
            with self.assertRaises(FileExistsError):build_ledger(ROOT,output)

    def test_all_32_sources_independent_coverage(self):
        paths=sorted(p for p in SOURCE.iterdir() if p.suffix in ('.xlsx','.docx'))
        self.assertEqual(len(paths),32)
        agency=set()
        for path in paths:
            rows,coverage,legal,raw=extract_source(path,ROOT)
            agency.add(agency_candidate(path))
            self.assertTrue(rows)
            for c in rows:
                self.assertTrue(c.item and c.section and c.period and c.legal)
                self.assertIsInstance(c.amount,(int,float))
            if path.suffix=='.xlsx':
                book=load_workbook(path,data_only=False)
                for sheet in book:
                    for row in sheet:
                        for cell in row:
                            if not isinstance(cell.value,(int,float)) or isinstance(cell.value,bool):continue
                            loc=f'sheet:{sheet.title};cell:{cell.coordinate}'
                            matches=[cv for cv in coverage if cv[2]==loc]
                            self.assertEqual(len(matches),1,(path.name,loc))
                            if cell.row==1 and cell.column>9:
                                self.assertEqual(matches[0][4],'対象外')
                            else:
                                cs=[c for c in rows if loc in c.locations]
                                self.assertEqual(len(cs),1,(path.name,loc))
                                self.assertEqual(cs[0].amount,cell.value)
                book.close()
            else:
                self.assertEqual(len([c for c in rows if c.kind=='キャンペーン']),15)
                self.assertEqual(len([c for c in rows if c.kind=='違約金']),2)
        self.assertEqual(len(agency),11)
