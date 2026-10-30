"""Draft-only delivery of locked notices as PDFs; never regenerate prices or send mail."""
from __future__ import annotations
import hashlib
import importlib.util
import json
import re
import sqlite3
import string
import sys
from contextlib import contextmanager
from datetime import datetime, timezone
from email.utils import getaddresses
from pathlib import Path, PureWindowsPath

SUBJECT_TEMPLATE = '【{target_yyyymm}】【{subject_company_name}】auひかり(AL商流)/auひかりプラス_条件表通知表'


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def data_hash(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def write_json(path: Path, data) -> None:
    with path.open('x', encoding='utf-8') as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2)
        stream.write('\n')


def sibling(name):
    """Load the stdlib-only adapters on Windows without optional Excel dependencies."""
    key = '_locked_delivery_' + name
    if key not in sys.modules:
        spec = importlib.util.spec_from_file_location(key, Path(__file__).with_name(name + '.py'))
        module = importlib.util.module_from_spec(spec)
        sys.modules[key] = module
        spec.loader.exec_module(module)
    return sys.modules[key]


def relative_path(root: Path, value: str) -> Path:
    p = Path(value)
    if p.is_absolute() or '..' in p.parts:
        raise ValueError('Unsafe relative path')
    resolved = (root / p).resolve()
    if not resolved.is_relative_to(root.resolve()):
        raise ValueError('Path outside project')
    return resolved


def render_company_subject(period: str, formal_name: str) -> str:
    if not re.fullmatch(r'\d{4}-(?:0[1-9]|1[0-2])', period):
        raise ValueError('Invalid target period')
    if not isinstance(formal_name, str) or not formal_name.strip() or any(c in formal_name for c in '\r\n'):
        raise ValueError('Formal recipient name is required')
    name = re.sub(r'(?:\s*(?:様|御中))+$', '', formal_name.strip()).rstrip()
    if not name:
        raise ValueError('Formal recipient name is empty')
    return SUBJECT_TEMPLATE.format(target_yyyymm=period.replace('-', ''), subject_company_name=name + '様')


def render_body(template: str, company: str | None, contact: str, period: str) -> str:
    if not isinstance(contact, str) or not contact.strip():
        raise ValueError('mail_contact_name is required')
    company = company or ''
    if not isinstance(company, str):
        raise ValueError('mail_company_name must be text or blank')
    values = dict(mail_company_name=company, mail_contact_name=contact,
                  target_year_month_ja=f'{period[:4]}年{int(period[5:])}月')
    for _, key, spec, conversion in string.Formatter().parse(template):
        if key is not None and (key not in values or spec or conversion):
            raise ValueError('Unsupported body placeholder; legacy mail_addressee is forbidden')
    result = template.format_map(values)
    # Empty company omits its own leading line only. Contact line breaks are untouched.
    if not company and template.startswith('{mail_company_name}\n'):
        result = result[1:]
    return result


def parse_addresses(value, *, required=False) -> list[str]:
    if value in (None, ''):
        if required:
            raise ValueError('TO is required')
        return []
    if not isinstance(value, str) or '\n' in value or '\r' in value:
        raise ValueError('Invalid mail address field')
    addresses = []
    for piece in value.split(';'):
        if not piece.strip():
            continue
        parsed = getaddresses([piece])
        if not parsed:
            raise ValueError('Empty address segment')
        for _, address in parsed:
            if not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?\.[A-Za-z]{2,}", address):
                raise ValueError('Malformed mail address')
            if address.lower() in {a.lower() for a in addresses}:
                raise ValueError('Duplicate mail address')
            addresses.append(address)
    if required and not addresses:
        raise ValueError('TO is required')
    return addresses


def prepare_batch(root: Path, master: Path, template_path: Path, pdf_root: Path, output: Path, period: str='2026-11') -> dict:
    from openpyxl import load_workbook
    from pypdf import PdfReader
    root = root.resolve()
    if not re.fullmatch(r'\d{4}-(?:0[1-9]|1[0-2])',period):
        raise ValueError('Invalid target period')
    master, template_path, pdf_root, output = (p.resolve() for p in (master, template_path, pdf_root, output))
    receipt_path = root / f'docs/{period}-conditions-lock.json'
    receipt = json.loads(receipt_path.read_text(encoding='utf-8'))
    lock_root = root / receipt['lock_directory']
    sibling('artifact_lock').verify_lock(root, lock_root, receipt['manifest_sha256'])
    lock = json.loads((lock_root / 'manifest.json').read_text(encoding='utf-8'))
    if lock['period'] != period:
        raise ValueError('LOCK target month mismatch')
    expected_count=lock['agency_count']
    conversion_path = pdf_root / 'conversion_results.json'
    converted = json.loads(conversion_path.read_text(encoding='utf-8-sig'))
    conversion = {(r['agency'], r['role']): r for r in converted}
    if len(conversion) != expected_count*2 or len(converted) != expected_count*2:
        raise ValueError('Exactly two unique PDF conversions per LOCK agency are required')
    settings = json.loads(template_path.read_text(encoding='utf-8-sig'))
    book = load_workbook(master, read_only=True, data_only=False)
    try:
        sheet = book['配信マスタ']
        rows = list(sheet.iter_rows())
        headers = [c.value for c in rows[0]]
        required = {'agency_id','agency_name','to','cc','bcc','template_id','attach_regular','attach_campaign','enabled','mail_company_name','mail_contact_name'}
        if not required <= set(headers) or len(headers) != len(set(headers)):
            raise ValueError('Missing or duplicated delivery columns')
        records = []
        for cells in rows[1:]:
            if not any(c.value is not None for c in cells):
                continue
            if any(c.data_type == 'f' for c in cells):
                raise ValueError('Delivery formulas are forbidden')
            records.append(dict(zip(headers, [c.value for c in cells])))
    finally:
        book.close()
    agencies = {r['agency'] for r in lock['documents']}
    if len(records) != expected_count or {r['agency_name'] for r in records} != agencies or len({r['agency_id'] for r in records}) != expected_count:
        raise ValueError('Delivery master must match the LOCK agencies exactly')
    if set(lock['excluded_agencies']) & {r['agency_name'] for r in records}:
        raise ValueError('Excluded recipient')
    messages = []
    for record in records:
        agency = record['agency_name']
        if not record['agency_id'] or any(record[k] != 'YES' for k in ('attach_regular','attach_campaign')):
            raise ValueError('Stable agency ID and both PDF attachment flags are required')
        template = settings['templates'][record['template_id']]
        if template['subject'] != SUBJECT_TEMPLATE:
            raise ValueError('Subject template differs from the approved company rule')
        attachments = []
        for role in ('conditions','campaign'):
            documents = [d for d in lock['documents'] if d['agency'] == agency and d['role'] == role]
            if len(documents) != 1:
                raise ValueError('Ambiguous LOCK document pair')
            document = documents[0]
            pdf = pdf_root / agency / (Path(document['path']).stem + '.pdf')
            rendered = conversion[(agency,role)]
            if rendered['source_sha256'] != document['sha256'] or PureWindowsPath(rendered['destination']).name != pdf.name:
                raise ValueError('PDF provenance mismatch')
            if digest(pdf) != rendered['pdf_sha256'] or pdf.read_bytes()[:5] != b'%PDF-':
                raise ValueError('PDF fingerprint or format mismatch')
            reader = PdfReader(pdf)
            text = '\n'.join(page.extract_text() or '' for page in reader.pages)
            normalized = re.sub(r'\s+', '', text)
            if not reader.pages or re.sub(r'\s+', '', agency) not in normalized:
                raise ValueError('PDF page or agency-name validation failed')
            attachments.append(dict(role=role,path=pdf.relative_to(root).as_posix(),file_name=pdf.name,
                                    sha256=digest(pdf),source_word_sha256=document['sha256'],page_count=len(reader.pages)))
        message = dict(agency_id=record['agency_id'],agency_name=agency,enabled=record['enabled'],
                       mail_company_name=record['mail_company_name'] or '', mail_contact_name=record['mail_contact_name'],
                       to=parse_addresses(record['to'],required=True),cc=parse_addresses(record['cc']),bcc=parse_addresses(record['bcc']),
                       subject=render_company_subject(period,agency),
                       body=render_body(template['body'],record['mail_company_name'],record['mail_contact_name'],period),
                       attachments=attachments)
        # Cross-field duplicates are probably a master error; do not silently edit addresses.
        all_addresses = message['to'] + message['cc'] + message['bcc']
        if len({a.lower() for a in all_addresses}) != len(all_addresses):
            raise ValueError('Duplicate addresses across TO/CC/BCC')
        message['message_hash'] = data_hash(message)
        messages.append(message)
    prepared = dict(schema='locked-pdf-drafts-v1',period=period,lock_manifest_sha256=receipt['manifest_sha256'],
        lock_directory=receipt['lock_directory'],receipt_path=receipt_path.relative_to(root).as_posix(),receipt_sha256=digest(receipt_path),
        master_path=master.relative_to(root).as_posix(),master_sha256=digest(master),
        template_path=template_path.relative_to(root).as_posix(),template_sha256=digest(template_path),
        conversion_path=conversion_path.relative_to(root).as_posix(),conversion_sha256=digest(conversion_path),
        messages=messages,agency_count=expected_count,attachment_count=expected_count*2,mode='draft',send_allowed=False)
    output.parent.mkdir(parents=True,exist_ok=True)
    write_json(output,prepared)
    return dict(agency_count=expected_count,attachment_count=expected_count*2,prepared_sha256=digest(output),
                first_agency_id=messages[0]['agency_id'],disabled_count=sum(m['enabled']!='YES' for m in messages))


def choose_sending_account(records, requested_sender: str, configured_id: str | None=None):
    from types import SimpleNamespace
    requested=parse_addresses(requested_sender,required=True)
    if len(requested)!=1:
        raise ValueError('Exactly one sender is required')
    requested_sender=requested[0]
    candidates=[]
    for item in records:
        if item.get('type')!='ZOHO_ACCOUNT' or item.get('enabled') is not True:
            continue
        account_id=str(item.get('accountId') or '')
        if not account_id or configured_id and configured_id!=account_id:
            continue
        registered={v.get('mailId','').lower() for v in item.get('emailAddress',[]) if isinstance(v,dict) and v.get('isConfirmed') is True}
        from_addresses={v.get('fromAddress','').lower() for v in item.get('sendMailDetails',[]) if isinstance(v,dict) and v.get('status') is True}
        if requested_sender.lower() in registered & from_addresses:
            candidates.append(SimpleNamespace(account_id=account_id,primary_email=item.get('primaryEmailAddress',''),sender_address=requested_sender))
    if len(candidates)!=1:
        raise ValueError('Requested sender is not uniquely registered and available in Zoho')
    return candidates[0]


def load_zoho_client(token_file: Path | None=None, requested_sender: str | None=None):
    """Resolve an explicitly requested registered From address, including confirmed aliases."""
    zoho=sibling('zoho_mail')
    provider=zoho.DpapiSecretProvider(token_file=token_file)
    region=provider.get_secret('ZOHO_REGION') or 'us'  # Confirmed US region in repository setup.
    client_id=provider.get_secret('ZOHO_CLIENT_ID');client_secret=provider.get_secret('ZOHO_CLIENT_SECRET')
    token=provider.get_secret('ZOHO_REFRESH_TOKEN');configured=provider.get_secret('ZOHO_ACCOUNT_ID')
    discovery=zoho.ZohoMailClient(zoho.ZohoMailConfig(client_id,client_secret,token,
                                  configured or '__READ_ONLY_ACCOUNT_DISCOVERY__',region))
    status,raw=discovery._request_accounts(discovery._get_access_token())
    discovery._raise_for_status(status,is_token_request=False)
    data=discovery._decode_json(raw,'Zoho account discovery')['data']
    records=data.get('accounts',[data]) if isinstance(data,dict) else data
    if not requested_sender:
        raise ValueError('Explicit requested sender address is required')
    account=choose_sending_account(records,requested_sender,configured)
    client=zoho.ZohoMailClient(zoho.ZohoMailConfig(client_id,client_secret,token,account.account_id,region))
    return client,account


def verify_prepared(root: Path, path: Path, expected_sha256: str) -> dict:
    if digest(path) != expected_sha256:
        raise ValueError('Prepared batch changed')
    data = json.loads(path.read_text(encoding='utf-8'))
    if data.get('schema')!='locked-pdf-drafts-v1' or data.get('mode')!='draft' or data.get('send_allowed') is not False:
        raise ValueError('Only draft preparation is allowed')
    for label in ('receipt','master','template','conversion'):
        if digest(relative_path(root,data[label+'_path']))!=data[label+'_sha256']:
            raise ValueError(label+' changed since preview')
    sibling('artifact_lock').verify_lock(root,relative_path(root,data['lock_directory']),data['lock_manifest_sha256'])
    lock=json.loads((relative_path(root,data['lock_directory'])/'manifest.json').read_text(encoding='utf-8'))
    names={r['agency'] for r in lock['documents']}
    messages=data['messages']
    if len(messages)!=lock['agency_count'] or len({m['agency_id'] for m in messages})!=lock['agency_count'] or {m['agency_name'] for m in messages}!=names:
        raise ValueError('Batch recipient mismatch')
    for message in messages:
        payload={k:v for k,v in message.items() if k!='message_hash'}
        if data_hash(payload)!=message['message_hash'] or message['subject']!=render_company_subject(data['period'],message['agency_name']):
            raise ValueError('Message or subject changed')
        if not message['mail_contact_name'].strip() or not message['to']:
            raise ValueError('Missing contact or TO')
        if len(message['attachments'])!=2 or {a['role'] for a in message['attachments']}!={'conditions','campaign'}:
            raise ValueError('Attachment pair mismatch')
        for attachment in message['attachments']:
            source=next(d for d in lock['documents'] if d['agency']==message['agency_name'] and d['role']==attachment['role'])
            pdf=relative_path(root,attachment['path'])
            if source['sha256']!=attachment['source_word_sha256'] or pdf.name!=attachment['file_name'] or digest(pdf)!=attachment['sha256']:
                raise ValueError('PDF attachment changed or belongs to another agency')
    return data


def _read_original(client, account, message_id):
    from urllib.parse import quote
    status,content=client._request(url=f'{client._config.mail_base_url}/api/accounts/{quote(account.account_id,safe="")}/messages/{quote(message_id,safe="")}/originalmessage',
        method='GET',headers={'Authorization':'Zoho-oauthtoken '+client._get_access_token(),'Accept':'application/json'},body=None)
    client._raise_for_status(status,is_token_request=False)
    response=client._decode_json(content,'Zoho draft readback')
    return response['data']['content']


def verify_mime(original: str, message: dict, sender: str) -> None:
    from email import policy
    from email.parser import BytesParser
    mime=BytesParser(policy=policy.default).parsebytes(original.encode('utf-8'))
    if str(mime['subject'])!=message['subject']:
        raise ValueError('Saved draft subject differs from company rule')
    for header,key in [('to','to'),('cc','cc'),('bcc','bcc')]:
        actual={a.lower() for _,a in getaddresses(mime.get_all(header,[])) if a}
        if actual!={a.lower() for a in message[key]}:
            raise ValueError('Saved draft recipient mismatch')
    actual_sender={a.lower() for _,a in getaddresses(mime.get_all('from',[])) if a}
    if actual_sender!={sender.lower()}:
        raise ValueError('Saved draft sender mismatch')
    body=mime.get_body(preferencelist=('plain',))
    if body is None or body.get_content().replace('\r\n','\n').rstrip('\n')!=message['body'].replace('\r\n','\n').rstrip('\n'):
        raise ValueError('Saved draft body mismatch')
    attachments=list(mime.iter_attachments())
    expected={a['file_name']:a['sha256'] for a in message['attachments']}
    actual={a.get_filename():hashlib.sha256(a.get_payload(decode=True)).hexdigest() for a in attachments}
    if len(attachments)!=2 or actual!=expected:
        raise ValueError('Saved draft PDF attachment mismatch')



@contextmanager
def database(path: Path):
    connection=sqlite3.connect(path)
    try:
        with connection:
            yield connection
    finally:
        connection.close()


def save_one(root: Path, prepared_path: Path, expected_sha256: str, agency_id: str,
             history_path: Path, result_path: Path, *, token_file: Path | None=None,
             authorize_disabled_draft: bool=False, retry_failed_upload: bool=False, skip_readback: bool=False, requested_sender: str | None=None, client=None, account=None) -> dict:
    root=root.resolve()
    prepared=verify_prepared(root,prepared_path,expected_sha256)
    selected=[m for m in prepared['messages'] if m['agency_id']==agency_id]
    if len(selected)!=1:
        raise ValueError('Exactly one explicit agency is required')
    message=selected[0]
    if message['enabled']!='YES' and not authorize_disabled_draft:
        raise ValueError('Master disabled; explicit one-off draft authorization is required')
    if result_path.exists():
        raise FileExistsError('Draft receipt already exists; no automatic retry')
    if client is None:
        client,account=load_zoho_client(token_file,requested_sender)
    sender=getattr(account,'sender_address',account.primary_email)
    if requested_sender and sender.lower()!=requested_sender.lower():
        raise ValueError('Authenticated sender differs from explicitly requested sender')
    if not parse_addresses(sender,required=True):
        raise ValueError('Authenticated sender is invalid')
    history_path.parent.mkdir(parents=True,exist_ok=True)
    request_hash=data_hash(dict(message_hash=message['message_hash'],sender=sender))
    with database(history_path) as db:
        db.execute('CREATE TABLE IF NOT EXISTS locked_pdf_drafts (agency_id TEXT, period TEXT, request_hash TEXT, status TEXT, message_id TEXT, recorded_at TEXT, PRIMARY KEY(agency_id,period))')
        try:
            db.execute('INSERT INTO locked_pdf_drafts VALUES (?,?,?,?,?,?)',(agency_id,prepared['period'],request_hash,'STARTED',None,datetime.now(timezone.utc).isoformat()))
        except sqlite3.IntegrityError:
            previous=db.execute('SELECT status FROM locked_pdf_drafts WHERE agency_id=? AND period=?',(agency_id,prepared['period'])).fetchone()
            if not retry_failed_upload or previous is None or previous[0]!='FAILED_BEFORE_DRAFT':
                raise ValueError('Draft attempt already recorded; inspect Zoho before any retry') from None
            db.execute('CREATE TABLE IF NOT EXISTS upload_retry_audit (agency_id TEXT,period TEXT,previous_status TEXT,reason TEXT,recorded_at TEXT)')
            db.execute('INSERT INTO upload_retry_audit VALUES (?,?,?,?,?)',(agency_id,prepared['period'],previous[0],
                'Explicit retry after confirmed pre-draft attachment failure; no draft POST previously made',datetime.now(timezone.utc).isoformat()))
            db.execute('UPDATE locked_pdf_drafts SET status=?,request_hash=?,recorded_at=? WHERE agency_id=? AND period=?',
                ('STARTED',request_hash,datetime.now(timezone.utc).isoformat(),agency_id,prepared['period']))
    def status(state,message_id=None):
        with database(history_path) as db:
            db.execute('UPDATE locked_pdf_drafts SET status=?,message_id=? WHERE agency_id=? AND period=?',
                       (state,message_id,agency_id,prepared['period']))
    uploaded=[]
    try:
        for attachment in message['attachments']:
            path=relative_path(root,attachment['path'])
            content=path.read_bytes()
            if hashlib.sha256(content).hexdigest()!=attachment['sha256']:
                raise ValueError('Attachment changed before upload')
            response=client.upload_draft_attachment(file_name=attachment['file_name'],content=content,content_type='application/octet-stream')
            from html import unescape
            returned_name=unescape(response.get('attachmentName',''))
            if not returned_name.strip() or not response.get('storeName') or not response.get('attachmentPath'):
                raise ValueError('Uploaded attachment response is incomplete')
            # Explicit user policy: filename changes may proceed to draft-only creation.
            # PDF bytes remain hash-bound; record both names for mandatory manual review.
            uploaded.append(response)
        verify_prepared(root,prepared_path,expected_sha256)
    except Exception:
        status('FAILED_BEFORE_DRAFT');raise
    payload=dict(fromAddress=sender,toAddress=','.join(message['to']),subject=message['subject'],content=message['body'],
                 mailFormat='plaintext',encoding='UTF-8',mode='draft',attachments=uploaded)
    if message['cc']:payload['ccAddress']=','.join(message['cc'])
    if message['bcc']:payload['bccAddress']=','.join(message['bcc'])
    status('POSTING')
    try:
        response=client.save_draft(payload)  # Exactly one POST, always mode=draft.
        data=response.get('data') or {}
        if isinstance(data,list) and len(data)==1:data=data[0]
        message_id=str(data.get('messageId') or '')
        if not message_id:
            raise ValueError('Zoho did not return a draft message ID')
    except Exception:
        status('OUTCOME_UNKNOWN');raise
    status('CREATED',message_id)
    result=dict(status='CREATED',agency_id=agency_id,agency_name=message['agency_name'],period=prepared['period'],
                message_id=message_id,from_address=sender,subject=message['subject'],attachment_count=2,mode='draft',sent=False,
                request_hash=request_hash,prepared_sha256=expected_sha256,
                attachment_name_changed=any(a['file_name']!=u['attachmentName'] for a,u in zip(message['attachments'],uploaded)),
                requires_manual_review=True,
                attachment_name_mapping=[dict(original=a['file_name'],zoho=u['attachmentName'],sha256=a['sha256']) for a,u in zip(message['attachments'],uploaded)],
                one_off_disabled_authorization=message['enabled']!='YES',recorded_at=datetime.now(timezone.utc).isoformat(),
                readback_status='NOT_ATTEMPTED')
    # Persist creation before readback: a later read error must never lead to duplicate creation.
    result_path.parent.mkdir(parents=True,exist_ok=True);write_json(result_path,result)
    if skip_readback:
        result['readback_status']='DEFERRED_BY_USER'
        result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        return result
    try:
        original=_read_original(client,account,message_id)
        verify_mime(original,message,sender)
        result['readback_status']='VERIFIED'
    except Exception as error:
        result['readback_status']='UNVERIFIED'
        result['readback_error_category']=type(error).__name__
    result_path.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    return result


def save_remaining_nine(root: Path, prepared_path: Path, expected_sha256: str, history_path: Path,
                        result_directory: Path, *, token_file: Path, requested_sender: str, retry_failed_upload: bool=False) -> dict:
    """Explicitly authorized remaining nine; exclude the already-created keioh draft."""
    root=root.resolve()
    batch=verify_prepared(root,prepared_path,expected_sha256)
    selected=[m for m in batch['messages'] if m['agency_id']!='keioh']
    if len(selected)!=9 or {m['agency_id'] for m in selected}!={
        'netconnect','an_ocean','centrix','j_and_c','duo_communications',
        'value_ecology','takahashi_daichi','ichimaru','hualan'}:
        raise ValueError('Remaining-nine scope differs from explicit authorization')
    with database(history_path) as db:
        prior=db.execute('SELECT agency_id,status,message_id FROM locked_pdf_drafts').fetchall()
    history={row[0]:row[1:] for row in prior}
    if history.get('keioh',('',''))[0]!='CREATED':
        raise ValueError('The confirmed keioh draft must already be recorded')
    # Do not start a new batch while any previous draft outcome is uncertain.
    for m in selected:
        if m['agency_id'] in history and history[m['agency_id']][0]!='CREATED' and not (retry_failed_upload and history[m['agency_id']][0]=='FAILED_BEFORE_DRAFT'):
            raise ValueError('A prior attempt must be reconciled before resuming')
    client,account=load_zoho_client(token_file,requested_sender)
    created=[];retained=[]
    for m in selected:
        result_path=result_directory/(m['agency_id']+'_draft_result.json')
        if m['agency_id'] in history and history[m['agency_id']][0]=='CREATED':
            existing=json.loads(result_path.read_text(encoding='utf-8'))
            if existing['status']!='CREATED' or existing['prepared_sha256']!=expected_sha256 or existing['from_address']!=requested_sender:
                raise ValueError('Existing receipt differs from current batch')
            retained.append(m['agency_id']);continue
        result=save_one(root,prepared_path,expected_sha256,m['agency_id'],history_path,result_path,
            token_file=token_file,authorize_disabled_draft=True,retry_failed_upload=retry_failed_upload,skip_readback=True,requested_sender=requested_sender,
            client=client,account=account)
        created.append(m['agency_id'])
        print(json.dumps(dict(progress='DRAFT_CREATED',agency_id=m['agency_id'],created_count=len(created),completed_count=len(created)+len(retained),
                              total_to_create=9,from_address=result['from_address'],sent=False),ensure_ascii=False),flush=True)
    summary=dict(status='REMAINING_NINE_CREATED',new_draft_count=len(created),retained_existing_count=len(retained),
                 created_agencies=created,retained_agencies=retained,total_including_keioh=10,sent=False,
                 readback_status='DEFERRED_BY_USER',recorded_at=datetime.now(timezone.utc).isoformat())
    write_json(result_directory/'remaining_nine_result.json',summary)
    return summary
