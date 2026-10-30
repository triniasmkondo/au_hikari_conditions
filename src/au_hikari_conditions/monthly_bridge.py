from __future__ import annotations

import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .document_reader import read_document
from .normalization import NormalizedField, normalize_blocks


SUPPORTED_SUFFIXES = {".pdf", ".docx", ".xlsx"}
DIMENSION_NAMES = (
    "product_family",
    "plus_type",
    "isp",
    "dwelling_type",
    "speed",
    "access_type",
    "trigger",
)
AGENCY_DOCUMENT_TYPES = (
    "AU_HIKARI",
    "AU_HIKARI_PLUS",
    "CAMPAIGN_CONDITIONS",
)
AGENCY_DOCUMENT_TYPE_LABELS = {
    "AU_HIKARI": "auひかり",
    "AU_HIKARI_PLUS": "auひかりプラス",
    "CAMPAIGN_CONDITIONS": "キャンペーン条件",
    "UNKNOWN": "判定不能",
}

P2_CROSSWALK: tuple[tuple[str, str, str, str], ...] = (
    (
        "p1.applicable_period",
        "条件対応.effective_from/effective_until; 粗利原価.valid_from/valid_until",
        "期間候補を照合して人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.isp",
        "条件対応.isp",
        "候補値は参照のみ。代理店ID・サービス条件との対応を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.product_family",
        "条件対応.product_family",
        "商材名称とP2正式値の対応を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.plus_type",
        "条件対応.plus_type",
        "プラン名称・条件の対応を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.dwelling_type",
        "条件対応.dwelling_type",
        "住居区分の定義と適用対象を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.speed",
        "条件対応.speed",
        "速度区分を資料の定義に照合して人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.access_type",
        "条件対応.access_type",
        "申込区分・適用条件の定義を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.condition_dimension.trigger",
        "条件対応.trigger",
        "申込/開通等の成果確定条件を人が確定",
        "MANUAL_REVIEW",
    ),
    (
        "p1.commission_amount",
        "粗利原価.upstream_amount/upstream_amount_state/upstream_tax_basis",
        "報酬候補。税区分・適用条件・金額の承認が必要",
        "MANUAL_REVIEW",
    ),
    (
        "p1.achievement_condition",
        "条件対応.trigger/period_rule/payment_timing/exclusions",
        "原文とP2項目の意味対応を人が確認",
        "MANUAL_REVIEW",
    ),
    (
        "p1.ladder",
        "P2入力スキーマに段階条件専用列なし",
        "ルール定義が追加・承認されるまで計算保留",
        "UNSUPPORTED_HOLD",
    ),
    (
        "p1.clawback_condition",
        "条件対応.reversal_terms",
        "対象・期間・上限・代理店適用を人が構造化して承認",
        "MANUAL_REVIEW",
    ),
    (
        "NOT_IN_UPPER_SOURCE",
        "粗利原価.approved_target_gross_margin; gross_margin_tax_basis",
        "会社承認済みの代理店別粗利データが必要",
        "INTERNAL_APPROVED_SOURCE_REQUIRED",
    ),
    (
        "NOT_IN_UPPER_SOURCE",
        "粗利原価.attributable_direct_cost; direct_cost_tax_basis",
        "会社承認済みの原価内訳データが必要",
        "INTERNAL_APPROVED_SOURCE_REQUIRED",
    ),
    (
        "NOT_IN_UPPER_SOURCE",
        "粗利原価.display_policy; 代理店別例外",
        "掲載可否・例外・告知義務の承認済みルールが必要",
        "INTERNAL_APPROVED_SOURCE_REQUIRED",
    ),
    (
        "NOT_IN_UPPER_SOURCE",
        "前回条件.previous_agency_reward; 条件対応",
        "前月の承認済み支払額・条件対応が必要",
        "PREVIOUS_APPROVED_SOURCE_REQUIRED",
    ),
    (
        "NOT_IN_SOURCE_DOCUMENTS",
        "承認管理",
        "各入力レコードのrecord_hashと承認証跡を人が確定",
        "HUMAN_APPROVAL_REQUIRED",
    ),
)


class MonthlyBridgeError(ValueError):
    """Raised when monthly source comparison cannot be safely prepared."""


@dataclass(frozen=True)
class MonthlyBridgeReport:
    previous_period: str
    target_period: str
    source_count: int
    candidate_count: int
    upper_difference_count: int
    unresolved_count: int
    status: str
    output_path: Path


@dataclass(frozen=True)
class MonthlyComparisonValidation:
    previous_period: str
    target_period: str
    source_count: int
    candidate_count: int
    difference_count: int


@dataclass(frozen=True)
class AgencyDocumentSet:
    agency_key_candidate: str
    documents: tuple[Path, ...]
    type_by_path: dict[Path, str]
    hold_reasons: tuple[str, ...]


@dataclass(frozen=True)
class _Source:
    kind: str
    path: Path
    relative_path: str
    sha256: str
    fields: tuple[NormalizedField, ...]
    block_text: dict[tuple[str, str], str]
    error_category: str | None = None


def build_monthly_comparison_workbook(
    *,
    project_root: Path,
    previous_period: str,
    target_period: str,
    output_path: Path,
    agency_document_set: AgencyDocumentSet | None = None,
) -> MonthlyBridgeReport:
    _validate_period(previous_period)
    _validate_period(target_period)
    if previous_period >= target_period:
        raise MonthlyBridgeError("Previous period must be earlier than target period.")
    _require_external_output(project_root, output_path)
    if output_path.exists():
        raise MonthlyBridgeError("Refusing to overwrite an existing comparison workbook.")

    source_roots = (
        ("PREVIOUS_UPPER", project_root / "01_上位原本" / previous_period),
        ("PREVIOUS_AGENCY", project_root / "03_条件表" / previous_period),
        ("CURRENT_UPPER", project_root / "01_上位原本" / target_period),
    )
    sources: list[_Source] = []
    missing_kinds: list[str] = []
    agency_group_holds: tuple[str, ...] = ()
    for kind, root in source_roots:
        if not root.is_dir():
            missing_kinds.append(kind)
            continue
        files = sorted(
            path
            for path in root.rglob("*")
            if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        )
        if kind == "PREVIOUS_AGENCY" and agency_document_set is not None:
            selected_files = [path.resolve() for path in agency_document_set.documents]
            if not selected_files or len(set(selected_files)) != len(selected_files):
                raise MonthlyBridgeError("Selected agency documents must be distinct and non-empty.")
            for selected in selected_files:
                try:
                    selected.relative_to(root.resolve())
                except ValueError as error:
                    raise MonthlyBridgeError(
                        "Selected agency documents must be inside the previous-period agency folder."
                    ) from error
                if (
                    selected.suffix.lower() not in SUPPORTED_SUFFIXES
                    or not selected.is_file()
                ):
                    raise MonthlyBridgeError(
                        "A selected agency document is missing or unsupported."
                    )
            files = sorted(selected_files)
            agency_group_holds = agency_document_set.hold_reasons
        if not files:
            missing_kinds.append(kind)
        for path in files:
            sources.append(_read_source(path, project_root, kind))

    if not sources:
        raise MonthlyBridgeError("No supported monthly source documents were found.")

    candidates = [
        (source, field)
        for source in sources
        for field in source.fields
    ]
    differences, unresolved_difference_count = _compare_upper_sources(
        sources,
        previous_period=previous_period,
        target_period=target_period,
    )
    agency_sources = [source for source in sources if source.kind == "PREVIOUS_AGENCY"]
    type_by_path = agency_document_set.type_by_path if agency_document_set else {}
    if any(not source.fields for source in agency_sources):
        agency_group_holds = (
            *agency_group_holds,
            "少なくとも1つの代理店資料から比較候補を抽出できません。種別・条件を確認してください。",
        )
    agency_comparisons, agency_comparison_holds = _compare_agency_sources(
        agency_sources,
        sources,
        type_by_path,
        agency_document_set.agency_key_candidate if agency_document_set else "",
        previous_period=previous_period,
        target_period=target_period,
    )
    agency_file_overlaps, agency_overlap_holds = _compare_agency_file_overlaps(
        agency_sources,
        type_by_path,
    )
    agency_file_roles = _agency_file_role_rows(agency_sources, type_by_path)
    extraction_errors = sum(source.error_category is not None for source in sources)
    unresolved_count = (
        unresolved_difference_count
        + agency_comparison_holds
        + agency_overlap_holds
        + len(agency_group_holds)
        + extraction_errors
        + len(missing_kinds)
        + sum(source.kind == "PREVIOUS_AGENCY" for source, _ in candidates)
    )
    status = (
        "REVIEW_REQUIRED"
        if unresolved_count or not differences
        else "CANDIDATES_READY_FOR_HUMAN_REVIEW"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    _write_workbook(
        output_path=output_path,
        previous_period=previous_period,
        target_period=target_period,
        sources=sources,
        candidates=candidates,
        differences=differences,
        agency_comparisons=agency_comparisons,
        agency_file_overlaps=agency_file_overlaps,
        agency_file_roles=agency_file_roles,
        agency_group_key=(
            agency_document_set.agency_key_candidate if agency_document_set else ""
        ),
        agency_group_holds=agency_group_holds,
        agency_overlap_hold_count=agency_overlap_holds,
        agency_comparison_hold_count=agency_comparison_holds,
        missing_kinds=missing_kinds,
        unresolved_difference_count=unresolved_difference_count,
        candidate_status=status,
    )
    return MonthlyBridgeReport(
        previous_period=previous_period,
        target_period=target_period,
        source_count=len(sources),
        candidate_count=len(candidates),
        upper_difference_count=len(differences),
        unresolved_count=unresolved_count,
        status=status,
        output_path=output_path.resolve(),
    )


def classify_agency_document_type(path: Path) -> str:
    stem = re.sub(r"\s+", "", unicodedata.normalize("NFKC", path.stem).casefold())
    if "キャンペーン" in stem:
        return "CAMPAIGN_CONDITIONS"
    if "auひかりプラス" in stem:
        return "AU_HIKARI_PLUS"
    if "auひかり" in stem:
        return "AU_HIKARI"
    return "UNKNOWN"


def _agency_group_key(path: Path) -> str:
    stem = re.sub(r"\s+", "", unicodedata.normalize("NFKC", path.stem).casefold())
    for token in (
        "auひかりプラス",
        "auひかり",
        "キャンペーン条件",
        "キャンペーン",
        "条件表",
        "条件",
    ):
        stem = stem.replace(token, "")
    stem = re.sub(r"20\d{2}[-_.年]?(?:0[1-9]|1[0-2])月?", "", stem)
    return re.sub(r"[\W_]+", "", stem, flags=re.UNICODE)


def discover_agency_document_set(
    *,
    project_root: Path,
    agency_directory: Path,
    selected_document: Path,
) -> AgencyDocumentSet:
    root = agency_directory.resolve()
    selected = selected_document.resolve()
    try:
        selected.relative_to(root)
    except ValueError as error:
        raise MonthlyBridgeError(
            "Selected agency document must be inside the previous-period agency folder."
        ) from error
    if not selected.is_file() or selected.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise MonthlyBridgeError("Selected agency document is missing or unsupported.")

    agency_key = _agency_group_key(selected)
    if not agency_key:
        source = _read_source(selected, project_root, "PREVIOUS_AGENCY")
        document_type = _classify_agency_source_type(source)
        return AgencyDocumentSet(
            agency_key_candidate="",
            documents=(selected,),
            type_by_path={selected: document_type},
            hold_reasons=(
                "代理店識別キーをファイル名から抽出できません。",
                "代理店IDと同一代理店であることを担当者が確認してください。",
                *(f"{AGENCY_DOCUMENT_TYPE_LABELS[kind]}の条件表を確認してください。" for kind in AGENCY_DOCUMENT_TYPES),
            ),
        )

    candidates = sorted(
        path.resolve()
        for path in root.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_SUFFIXES
        and _agency_group_key(path) == agency_key
    )
    if selected not in candidates:
        raise MonthlyBridgeError("Selected agency document could not be grouped.")

    type_by_path = {
        path: _classify_agency_source_type(
            _read_source(path, project_root, "PREVIOUS_AGENCY")
        )
        for path in candidates
    }
    counts = Counter(type_by_path.values())
    holds: list[str] = []
    holds.append(
        "代理店IDと3資料が同一代理店に属することを担当者が確認してください。"
    )
    for document_type in AGENCY_DOCUMENT_TYPES:
        if counts[document_type] == 0:
            holds.append(
                f"{AGENCY_DOCUMENT_TYPE_LABELS[document_type]}の条件表が見つかりません。"
            )
        elif counts[document_type] > 1:
            holds.append(
                f"{AGENCY_DOCUMENT_TYPE_LABELS[document_type]}の候補が重複しています。"
            )
    if counts["UNKNOWN"]:
        holds.append("種別判定できない同一代理店候補ファイルがあります。")
    return AgencyDocumentSet(
        agency_key_candidate=agency_key,
        documents=tuple(candidates),
        type_by_path=type_by_path,
        hold_reasons=tuple(holds),
    )


def _classify_agency_source_type(source: _Source) -> str:
    filename_type = classify_agency_document_type(source.path)
    if filename_type != "UNKNOWN":
        return filename_type

    text = unicodedata.normalize("NFKC", " ".join(source.block_text.values()))
    campaign = "キャンペーン" in text and any(
        marker in text for marker in ("キャンペーン条件", "キャンペーン報酬", "キャンペーン特典")
    )
    if campaign:
        return "CAMPAIGN_CONDITIONS"

    products = {
        value
        for field in source.fields
        for dimension in field.condition_dimensions
        if dimension.name == "product_family"
        for value in dimension.values
    }
    if products == {"auひかり"}:
        return "AU_HIKARI"
    if products == {"auひかりプラス"}:
        return "AU_HIKARI_PLUS"
    return "UNKNOWN"


def validate_monthly_comparison_workbook(
    path: Path,
    *,
    project_root: Path,
    previous_period: str,
    target_period: str,
) -> MonthlyComparisonValidation:
    _validate_period(previous_period)
    _validate_period(target_period)
    if previous_period >= target_period:
        raise MonthlyBridgeError("Previous period must be earlier than target period.")
    if not path.is_file():
        raise MonthlyBridgeError("Monthly comparison workbook is missing.")
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as error:
        raise MonthlyBridgeError("Monthly comparison workbook cannot be opened.") from error
    try:
        required_sheets = {
            "比較情報",
            "資料一覧",
            "抽出根拠候補",
            "上位差分候補",
            "代理店資料種別",
            "代理店別上位比較",
            "代理店資料間照合",
            "P2項目対応",
        }
        if not required_sheets.issubset(workbook.sheetnames):
            raise MonthlyBridgeError("Monthly comparison workbook schema is incomplete.")
        info = {
            row[0]: row[1]
            for row in workbook["比較情報"].iter_rows(min_row=2, values_only=True)
            if len(row) >= 2 and row[0]
        }
        if (
            info.get("previous_period") != previous_period
            or info.get("target_period") != target_period
            or info.get("candidate_status")
            not in {"REVIEW_REQUIRED", "CANDIDATES_READY_FOR_HUMAN_REVIEW"}
        ):
            raise MonthlyBridgeError("Monthly comparison periods or candidate status do not match.")

        source_sheet = workbook["資料一覧"]
        source_rows = list(source_sheet.iter_rows(min_row=2, values_only=True))
        source_headers = [cell.value for cell in source_sheet[1]]
        source_records = [
            dict(zip(source_headers, row))
            for row in source_rows
            if any(value is not None for value in row)
        ]
        required_kinds = {"PREVIOUS_UPPER", "PREVIOUS_AGENCY", "CURRENT_UPPER"}
        actual_kinds = {record.get("source_type") for record in source_records}
        if not required_kinds.issubset(actual_kinds):
            raise MonthlyBridgeError("Monthly comparison is missing one or more source groups.")
        if any(record.get("extraction_status") != "READ" for record in source_records):
            raise MonthlyBridgeError("Monthly comparison contains unreadable source documents.")
        root = project_root.resolve()
        for record in source_records:
            relative_path = record.get("source_path")
            if not isinstance(relative_path, str) or not relative_path:
                raise MonthlyBridgeError("Monthly comparison contains an invalid source path.")
            source_path = (root / Path(relative_path)).resolve()
            try:
                source_path.relative_to(root)
            except ValueError as error:
                raise MonthlyBridgeError("Monthly comparison source path escapes the project root.") from error
            if (
                not source_path.is_file()
                or not isinstance(record.get("sha256"), str)
                or _sha256(source_path) != record["sha256"]
            ):
                raise MonthlyBridgeError("A source document changed after comparison.")

        candidate_sheet = workbook["抽出根拠候補"]
        candidate_rows = list(candidate_sheet.iter_rows(min_row=2, values_only=True))
        candidate_count = sum(any(value is not None for value in row) for row in candidate_rows)
        if candidate_count == 0:
            raise MonthlyBridgeError("Monthly comparison contains no evidence candidates.")
        candidate_headers = [cell.value for cell in candidate_sheet[1]]
        review_column = candidate_headers.index("review_status")
        approval_column = candidate_headers.index("approval_status")
        if any(
            row[review_column] != "NOT_REVIEWED"
            or row[approval_column] != "UNAPPROVED"
            for row in candidate_rows
            if any(value is not None for value in row)
        ):
            raise MonthlyBridgeError("Monthly evidence candidates cannot be approved in this workbook.")

        difference_sheet = workbook["上位差分候補"]
        difference_rows = list(difference_sheet.iter_rows(min_row=2, values_only=True))
        difference_count = sum(any(value is not None for value in row) for row in difference_rows)
        if difference_count == 0:
            raise MonthlyBridgeError("Monthly comparison contains no upper-source comparisons.")
        difference_headers = [cell.value for cell in difference_sheet[1]]
        approval_column = difference_headers.index("approval_status")
        if any(
            row[approval_column] != "UNAPPROVED"
            for row in difference_rows
            if any(value is not None for value in row)
        ):
            raise MonthlyBridgeError("Upper-source comparison candidates cannot be approved in this workbook.")

        for sheet_name in ("代理店別上位比較", "代理店資料間照合"):
            sheet = workbook[sheet_name]
            headers = [cell.value for cell in sheet[1]]
            if "review_status" not in headers:
                raise MonthlyBridgeError("Agency comparison workbook schema is incomplete.")
            review_column = headers.index("review_status")
            if any(
                row[review_column] != "UNAPPROVED"
                for row in sheet.iter_rows(min_row=2, values_only=True)
                if any(value is not None for value in row)
            ):
                raise MonthlyBridgeError("Agency comparison candidates cannot be approved in this workbook.")

        mapping_sheet = workbook["P2項目対応"]
        mapping_headers = [cell.value for cell in mapping_sheet[1]]
        auto_transfer_column = mapping_headers.index("自動転記")
        if any(
            row[auto_transfer_column] != "NO"
            for row in mapping_sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ):
            raise MonthlyBridgeError("P1 candidates cannot be auto-transferred to approved P2 fields.")
        return MonthlyComparisonValidation(
            previous_period=previous_period,
            target_period=target_period,
            source_count=len(source_records),
            candidate_count=candidate_count,
            difference_count=difference_count,
        )
    finally:
        workbook.close()


def _read_source(path: Path, project_root: Path, kind: str) -> _Source:
    relative_path = path.resolve().relative_to(project_root.resolve()).as_posix()
    digest = _sha256(path)
    try:
        blocks = read_document(path, project_root=project_root)
        fields = normalize_blocks(blocks)
        block_text = {
            (block.source_path, block.location): block.text for block in blocks
        }
    except Exception as error:
        return _Source(
            kind=kind,
            path=path,
            relative_path=relative_path,
            sha256=digest,
            fields=(),
            block_text={},
            error_category=type(error).__name__,
        )
    return _Source(
        kind=kind,
        path=path,
        relative_path=relative_path,
        sha256=digest,
        fields=tuple(fields),
        block_text=block_text,
    )


def _compare_upper_sources(
    sources: list[_Source],
    *,
    previous_period: str,
    target_period: str,
) -> tuple[list[dict[str, Any]], int]:
    previous_sources = [source for source in sources if source.kind == "PREVIOUS_UPPER"]
    current_sources = [source for source in sources if source.kind == "CURRENT_UPPER"]
    previous_groups = _group_upper_fields(previous_sources)
    current_groups = _group_upper_fields(current_sources)
    rows: list[dict[str, Any]] = []
    unresolved = 0

    previous_by_item = _group_by_item(previous_groups)
    current_by_item = _group_by_item(current_groups)
    for item_id in sorted(previous_by_item.keys() | current_by_item.keys()):
        previous = previous_by_item.get(item_id, {})
        current = current_by_item.get(item_id, {})
        matched_keys = previous.keys() & current.keys()
        for dimension_key in sorted(matched_keys):
            previous_fields = previous[dimension_key]
            current_fields = current[dimension_key]
            if not dimension_key:
                status = "UNMATCHED_CONDITION_DIMENSIONS"
                unresolved += 1
            elif len(previous_fields) > 1 or len(current_fields) > 1:
                status = "DUPLICATE_REQUIRES_REVIEW"
                unresolved += 1
            else:
                previous_values = _candidate_values(previous_fields)
                current_values = _candidate_values(current_fields)
                status = (
                    "UNCHANGED_CANDIDATE"
                    if previous_values == current_values
                    else "CHANGED_CANDIDATE"
                )
            rows.append(
                _difference_row(
                    item_id,
                    dimension_key,
                    status,
                    previous_fields,
                    current_fields,
                    previous_period,
                    target_period,
                )
            )

        previous_unmatched = [
            (key, previous[key]) for key in previous.keys() - matched_keys
        ]
        current_unmatched = [
            (key, current[key]) for key in current.keys() - matched_keys
        ]
        if previous_unmatched and current_unmatched:
            if any(not key for key, _ in previous_unmatched + current_unmatched):
                status = "UNMATCHED_CONDITION_DIMENSIONS"
            elif len(previous_unmatched) == 1 and len(current_unmatched) == 1:
                status = "CONDITION_MISMATCH_REQUIRES_REVIEW"
            else:
                status = "DUPLICATE_REQUIRES_REVIEW"
            unresolved += 1
            previous_fields = [
                field for _, fields in previous_unmatched for field in fields
            ]
            current_fields = [
                field for _, fields in current_unmatched for field in fields
            ]
            dimension_key = json.dumps(
                {
                    "previous": sorted(key for key, _ in previous_unmatched),
                    "current": sorted(key for key, _ in current_unmatched),
                },
                ensure_ascii=False,
                sort_keys=True,
            )
            rows.append(
                _difference_row(
                    item_id,
                    dimension_key,
                    status,
                    previous_fields,
                    current_fields,
                    previous_period,
                    target_period,
                )
            )
        else:
            for dimension_key, fields in previous_unmatched:
                status = (
                    "UNMATCHED_CONDITION_DIMENSIONS"
                    if not dimension_key
                    else "REMOVED_CANDIDATE"
                )
                unresolved += status == "UNMATCHED_CONDITION_DIMENSIONS"
                rows.append(
                    _difference_row(
                        item_id,
                        dimension_key,
                        status,
                        fields,
                        [],
                        previous_period,
                        target_period,
                    )
                )
            for dimension_key, fields in current_unmatched:
                status = (
                    "UNMATCHED_CONDITION_DIMENSIONS"
                    if not dimension_key
                    else "ADDED_CANDIDATE"
                )
                unresolved += status == "UNMATCHED_CONDITION_DIMENSIONS"
                rows.append(
                    _difference_row(
                        item_id,
                        dimension_key,
                        status,
                        [],
                        fields,
                        previous_period,
                        target_period,
                    )
                )
    return rows, unresolved


def _group_by_item(
    groups: dict[tuple[str, str], list[tuple[_Source, NormalizedField]]],
) -> dict[str, dict[str, list[tuple[_Source, NormalizedField]]]]:
    grouped: dict[str, dict[str, list[tuple[_Source, NormalizedField]]]] = defaultdict(dict)
    for (item_id, dimension_key), fields in groups.items():
        grouped[item_id][dimension_key] = fields
    return grouped


def _difference_row(
    item_id: str,
    dimension_key: str,
    status: str,
    previous: list[tuple[_Source, NormalizedField]],
    current: list[tuple[_Source, NormalizedField]],
    previous_period: str,
    target_period: str,
) -> dict[str, Any]:
    return {
        "previous_period": previous_period,
        "target_period": target_period,
        "item_id": item_id,
        "dimension_key": dimension_key,
        "status": status,
        "previous_values": _candidate_values(previous),
        "current_values": _candidate_values(current),
        "previous_evidence": _evidence_refs(previous),
        "current_evidence": _evidence_refs(current),
        "review_status": "UNAPPROVED",
        "hold_reason": (
            "抽出候補の差分です。正式ID・適用条件・税区分・代理店対応は未承認です。"
        ),
    }


def _compare_agency_sources(
    agency_sources: list[_Source],
    sources: list[_Source],
    type_by_path: dict[Path, str],
    agency_group_key: str,
    *,
    previous_period: str,
    target_period: str,
) -> tuple[list[dict[str, Any]], int]:
    previous_groups = _group_upper_fields(
        [source for source in sources if source.kind == "PREVIOUS_UPPER"]
    )
    current_groups = _group_upper_fields(
        [source for source in sources if source.kind == "CURRENT_UPPER"]
    )
    rows: list[dict[str, Any]] = []
    holds = 0
    expected_products = {
        "AU_HIKARI": "auひかり",
        "AU_HIKARI_PLUS": "auひかりプラス",
    }

    for source in agency_sources:
        document_type = type_by_path.get(
            source.path.resolve(), classify_agency_document_type(source.path)
        )
        agency_groups = _group_upper_fields([source])
        for (item_id, dimensions_key), agency_fields in sorted(agency_groups.items()):
            dimensions = json.loads(dimensions_key) if dimensions_key else {}
            products = dimensions.get("product_family", ())
            previous_fields = previous_groups.get((item_id, dimensions_key), [])
            current_fields = current_groups.get((item_id, dimensions_key), [])
            if document_type == "CAMPAIGN_CONDITIONS":
                previous_fields = [
                    pair for pair in previous_fields if "キャンペーン" in pair[1].raw_value
                ]
                current_fields = [
                    pair for pair in current_fields if "キャンペーン" in pair[1].raw_value
                ]
            reasons: list[str] = []
            if document_type == "UNKNOWN":
                reasons.append("代理店資料の種別をファイル名から判定できません。")
            elif document_type in expected_products:
                if tuple(products) != (expected_products[document_type],):
                    reasons.append("ファイル種別と抽出された商材条件が一致しません。")
            elif document_type == "CAMPAIGN_CONDITIONS" and not products:
                reasons.append("キャンペーン条件の対象商材を特定できません。")
            elif document_type == "CAMPAIGN_CONDITIONS":
                reasons.append(
                    "キャンペーン条件は通常報酬から分離して比較します。上位原本内のキャンペーン固有条件との対応確認が必要です。"
                )
            if not dimensions_key:
                reasons.append("比較に必要な商品・適用条件軸が不足しています。")
            if len(agency_fields) != 1:
                reasons.append("同一ファイル内に同じ条件キーの複数候補があります。")
            if len(previous_fields) > 1 or len(current_fields) > 1:
                reasons.append("上位原本に同一条件キーの複数候補があります。")
            if not previous_fields:
                reasons.append(f"{previous_period}上位原本に完全一致する条件候補がありません。")
            if not current_fields:
                reasons.append(f"{target_period}上位原本に完全一致する条件候補がありません。")

            if reasons:
                status = "REVIEW_REQUIRED"
                holds += 1
            else:
                status = (
                    "UPPER_VALUE_CHANGED_CANDIDATE"
                    if _candidate_values(previous_fields) != _candidate_values(current_fields)
                    else "UPPER_VALUE_UNCHANGED_CANDIDATE"
                )
                if _candidate_values(agency_fields) != _candidate_values(previous_fields):
                    reasons.append(
                        "前月代理店条件と前月上位原本の候補値が異なります。代理店別条件として確認してください。"
                    )
                    holds += 1
            if document_type == "CAMPAIGN_CONDITIONS":
                campaign_link_status = (
                    "PRODUCT_LINK_CANDIDATE"
                    if products
                    else "PRODUCT_LINK_REQUIRED"
                )
            elif document_type in expected_products and tuple(products) == (
                expected_products[document_type],
            ):
                campaign_link_status = "NOT_CAMPAIGN_CONDITION"
            else:
                campaign_link_status = "NOT_CONFIRMED"
            rows.append(
                {
                    "agency_group_key": agency_group_key,
                    "agency_identity_status": "MANUAL_CONFIRMATION_REQUIRED",
                    "agency_source_file": source.path.name,
                    "agency_document_type": document_type,
                    "agency_item_id": item_id,
                    "condition_dimensions": dimensions_key,
                    "agency_values": _candidate_values(agency_fields),
                    "previous_values": _candidate_values(previous_fields),
                    "current_values": _candidate_values(current_fields),
                    "agency_evidence": _evidence_refs(agency_fields),
                    "previous_evidence": _evidence_refs(previous_fields),
                    "current_evidence": _evidence_refs(current_fields),
                    "campaign_product_link_status": campaign_link_status,
                    "status": status,
                    "review_status": "UNAPPROVED",
                    "hold_reason": "; ".join(reasons)
                    or (
                        "条件・単価差分候補です。種別・代理店適用・税区分は未承認です。"
                    ),
                }
            )
    return rows, holds


def _compare_agency_file_overlaps(
    agency_sources: list[_Source],
    type_by_path: dict[Path, str],
) -> tuple[list[dict[str, Any]], int]:
    occurrences: dict[
        tuple[str, str, str],
        list[tuple[_Source, list[tuple[_Source, NormalizedField]]]],
    ] = defaultdict(list)
    for source in agency_sources:
        document_type = type_by_path.get(
            source.path.resolve(), classify_agency_document_type(source.path)
        )
        for key, fields in _group_upper_fields([source]).items():
            occurrences[(document_type, key[0], key[1])].append((source, fields))

    rows: list[dict[str, Any]] = []
    holds = 0
    for (document_type, item_id, dimensions_key), source_entries in sorted(occurrences.items()):
        is_duplicate_within_source = any(len(fields) > 1 for _, fields in source_entries)
        if len(source_entries) < 2 and not is_duplicate_within_source:
            continue
        value_by_source = {
            source.relative_path: _candidate_values(fields)
            for source, fields in source_entries
        }
        if not dimensions_key:
            status = "INCOMPLETE_KEY_REQUIRES_REVIEW"
            reason = "条件軸不足のため、ファイル間で同一条件か判定できません。"
        elif is_duplicate_within_source:
            status = "DUPLICATE_WITHIN_FILE_REQUIRES_REVIEW"
            reason = "同一ファイル内に同じ条件キーが複数あります。"
        elif len(set(value_by_source.values())) > 1:
            status = "CONFLICT_REQUIRES_REVIEW"
            reason = "ファイル間で同一条件キーの候補値が異なります。自動統合しません。"
        else:
            status = "DUPLICATE_REQUIRES_REVIEW"
            reason = "複数ファイルに同一条件キーがあります。自動統合しません。"
        holds += 1
        rows.append(
            {
                "item_id": item_id,
                "document_type": document_type,
                "condition_dimensions": dimensions_key,
                "source_values_json": json.dumps(value_by_source, ensure_ascii=False),
                "source_evidence_json": json.dumps(
                    {
                        source.relative_path: _evidence_refs(fields)
                        for source, fields in source_entries
                    },
                    ensure_ascii=False,
                ),
                "status": status,
                "review_status": "UNAPPROVED",
                "hold_reason": reason,
            }
        )
    return rows, holds


def _agency_file_role_rows(
    agency_sources: list[_Source],
    type_by_path: dict[Path, str],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for source in agency_sources:
        item_counts = Counter(field.item_id for field in source.fields)
        dimensions = sorted(
            {
                dimension.name
                for field in source.fields
                for dimension in field.condition_dimensions
            }
        )
        rows.append(
            {
                "source_file": source.path.name,
                "agency_identity_status": "MANUAL_CONFIRMATION_REQUIRED",
                "document_type": type_by_path.get(
                    source.path.resolve(),
                    classify_agency_document_type(source.path),
                ),
                "candidate_count": len(source.fields),
                "item_id_counts": json.dumps(dict(sorted(item_counts.items())), ensure_ascii=False),
                "condition_dimensions": json.dumps(dimensions, ensure_ascii=False),
                "review_status": "NOT_REVIEWED",
                "hold_reason": "種別はファイル名候補です。agency_idと資料の役割を担当者が確認してください。",
            }
        )
    return rows


def _group_upper_fields(
    sources: list[_Source],
) -> dict[tuple[str, str], list[tuple[_Source, NormalizedField]]]:
    groups: dict[tuple[str, str], list[tuple[_Source, NormalizedField]]] = defaultdict(list)
    for source in sources:
        for field in source.fields:
            dimensions = _dimension_values(field)
            required_dimensions_present = (
                "product_family" in dimensions
                and any(
                    key in dimensions
                    for key in (
                        "isp",
                        "plus_type",
                        "dwelling_type",
                        "speed",
                        "access_type",
                        "trigger",
                    )
                )
            )
            dimension_key = (
                json.dumps(dimensions, ensure_ascii=False, sort_keys=True)
                if required_dimensions_present
                else ""
            )
            groups[(field.item_id, dimension_key)].append((source, field))
    return groups


def _dimension_values(field: NormalizedField) -> dict[str, tuple[str, ...]]:
    return {
        dimension.name: tuple(sorted(set(dimension.values)))
        for dimension in field.condition_dimensions
        if dimension.name in DIMENSION_NAMES
    }


def _candidate_values(
    fields: list[tuple[_Source, NormalizedField]],
) -> tuple[str, ...]:
    return tuple(
        sorted(
            {
                value
                for _, field in fields
                for value in field.normalized_values
            }
        )
    )


def _evidence_refs(
    fields: list[tuple[_Source, NormalizedField]],
) -> tuple[str, ...]:
    return tuple(
        sorted(
            {
                f"{source.relative_path}#{field.location}"
                for source, field in fields
            }
        )
    )


def _write_workbook(
    *,
    output_path: Path,
    previous_period: str,
    target_period: str,
    sources: list[_Source],
    candidates: list[tuple[_Source, NormalizedField]],
    differences: list[dict[str, Any]],
    agency_comparisons: list[dict[str, Any]],
    agency_file_overlaps: list[dict[str, Any]],
    agency_file_roles: list[dict[str, Any]],
    agency_group_key: str,
    agency_group_holds: tuple[str, ...],
    agency_overlap_hold_count: int,
    agency_comparison_hold_count: int,
    missing_kinds: list[str],
    unresolved_difference_count: int,
    candidate_status: str,
) -> None:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "使い方"
    guide.append(("項目", "説明"))
    for row in (
        ("状態", "このブックは比較・抽出候補です。承認済みP2入力、確定計算、配布帳票ではありません。"),
        ("対象", f"{previous_period}上位原本・代理店条件表と{target_period}上位原本。欠落資料は不足一覧に記録します。"),
        ("根拠", "候補には元ファイル相対パス、SHA-256、ページ/シート位置、根拠抜粋を記録します。"),
        ("代理店資料", "auひかり・auひかりプラス・キャンペーン条件を別資料のまま比較します。agency_idと同一代理店であることは人が確認してください。"),
        ("キャンペーン", "対象商材との候補リンクを保持しますが、キャンペーン報酬を通常報酬へ加算しません。"),
        ("重複・矛盾", "同一種別・条件キーの重複または値の矛盾は自動統合せず、代理店資料間照合シートで保留します。"),
        ("未承認", "すべての候補差分は未承認です。候補値を承認済み入力列へそのままコピーしないでください。"),
        ("P2", "粗利・直接原価・税区分・前回報酬・例外・掲載可否は別の承認済み社内根拠が必要です。"),
        ("保存", "原文抜粋と報酬候補を含みます。Git外の会社承認済み保管先に限定してください。"),
    ):
        guide.append(row)

    info_sheet = workbook.create_sheet("比較情報")
    info_sheet.append(("項目", "値"))
    info_sheet.append(("previous_period", previous_period))
    info_sheet.append(("target_period", target_period))
    info_sheet.append(("candidate_status", candidate_status))
    info_sheet.append(("source_count", len(sources)))
    info_sheet.append(("candidate_count", len(candidates)))
    info_sheet.append(("upper_difference_count", len(differences)))
    info_sheet.append(("missing_source_count", len(missing_kinds)))
    info_sheet.append(
        (
            "previous_agency_document_count",
            sum(source.kind == "PREVIOUS_AGENCY" for source in sources),
        )
    )
    info_sheet.append(("agency_comparison_count", len(agency_comparisons)))
    info_sheet.append(("agency_file_overlap_count", len(agency_file_overlaps)))
    info_sheet.append(("agency_group_key_candidate", agency_group_key))
    info_sheet.append(
        (
            "agency_document_type_counts_json",
            json.dumps(
                dict(
                    sorted(
                        Counter(
                            row["document_type"] for row in agency_file_roles
                        ).items()
                    )
                ),
                ensure_ascii=False,
            ),
        )
    )
    info_sheet.append(("agency_group_hold_count", len(agency_group_holds)))
    info_sheet.append(("agency_overlap_hold_count", agency_overlap_hold_count))
    info_sheet.append(("agency_comparison_hold_count", agency_comparison_hold_count))

    source_sheet = workbook.create_sheet("資料一覧")
    source_sheet.append(
        ("source_type", "source_file", "source_path", "sha256", "extraction_status", "error_category")
    )
    for source in sources:
        source_sheet.append(
            (
                source.kind,
                source.path.name,
                source.relative_path,
                source.sha256,
                "READ_ERROR" if source.error_category else "READ",
                source.error_category,
            )
        )

    candidate_sheet = workbook.create_sheet("抽出根拠候補")
    candidate_sheet.append(
        (
            "candidate_id",
            "source_type",
            "source_file",
            "source_path",
            "source_sha256",
            "source_period",
            "agency_identity_status",
            "agency_document_scope",
            "item_id",
            "candidate_values_json",
            "condition_dimensions_json",
            "evidence_location",
            "evidence_excerpt",
            "review_status",
            "approval_status",
            "hold_reason",
        )
    )
    for index, (source, field) in enumerate(candidates, 1):
        source_period = (
            previous_period if source.kind != "CURRENT_UPPER" else target_period
        )
        is_agency_document = source.kind == "PREVIOUS_AGENCY"
        candidate_sheet.append(
            (
                f"{source.sha256[:12]}-{index:05d}",
                source.kind,
                source.path.name,
                source.relative_path,
                source.sha256,
                source_period,
                "MANUAL_AGENCY_ID_REQUIRED" if is_agency_document else "NOT_APPLICABLE",
                source.path.stem if is_agency_document else "",
                field.item_id,
                json.dumps(field.normalized_values, ensure_ascii=False),
                json.dumps(_dimension_values(field), ensure_ascii=False, sort_keys=True),
                field.location,
                source.block_text.get(
                    (field.source_path, field.location), field.raw_value
                ),
                "NOT_REVIEWED",
                "UNAPPROVED",
                "; ".join(field.hold_reasons),
            )
        )

    difference_sheet = workbook.create_sheet("上位差分候補")
    difference_sheet.append(
        (
            "previous_period",
            "target_period",
            "item_id",
            "condition_dimensions_json",
            "candidate_status",
            "previous_values_json",
            "current_values_json",
            "previous_evidence",
            "current_evidence",
            "approval_status",
            "hold_reason",
        )
    )
    for row in differences:
        difference_sheet.append(
            (
                row["previous_period"],
                row["target_period"],
                row["item_id"],
                row["dimension_key"],
                row["status"],
                json.dumps(row["previous_values"], ensure_ascii=False),
                json.dumps(row["current_values"], ensure_ascii=False),
                "; ".join(row["previous_evidence"]),
                "; ".join(row["current_evidence"]),
                row["review_status"],
                row["hold_reason"],
            )
        )

    agency_roles = workbook.create_sheet("代理店資料種別")
    agency_roles.append(
        (
            "agency_group_key_candidate",
            "agency_identity_status",
            "source_file",
            "document_type_candidate",
            "candidate_count",
            "item_id_counts_json",
            "condition_dimensions_json",
            "review_status",
            "hold_reason",
        )
    )
    for row in agency_file_roles:
        agency_roles.append(
            (
                agency_group_key,
                row["agency_identity_status"],
                row["source_file"],
                row["document_type"],
                row["candidate_count"],
                row["item_id_counts"],
                row["condition_dimensions"],
                row["review_status"],
                row["hold_reason"],
            )
        )

    agency_comparison_sheet = workbook.create_sheet("代理店別上位比較")
    agency_comparison_sheet.append(
        (
            "agency_group_key_candidate",
            "agency_identity_status",
            "agency_source_file",
            "document_type",
            "item_id",
            "condition_dimensions_json",
            "agency_values_json",
            "previous_upper_values_json",
            "current_upper_values_json",
            "agency_evidence",
            "previous_upper_evidence",
            "current_upper_evidence",
            "campaign_product_link_status",
            "candidate_status",
            "review_status",
            "hold_reason",
        )
    )
    for row in agency_comparisons:
        agency_comparison_sheet.append(
            (
                row["agency_group_key"],
                row["agency_identity_status"],
                row["agency_source_file"],
                row["agency_document_type"],
                row["agency_item_id"],
                row["condition_dimensions"],
                json.dumps(row["agency_values"], ensure_ascii=False),
                json.dumps(row["previous_values"], ensure_ascii=False),
                json.dumps(row["current_values"], ensure_ascii=False),
                "; ".join(row["agency_evidence"]),
                "; ".join(row["previous_evidence"]),
                "; ".join(row["current_evidence"]),
                row["campaign_product_link_status"],
                row["status"],
                row["review_status"],
                row["hold_reason"],
            )
        )

    agency_overlap_sheet = workbook.create_sheet("代理店資料間照合")
    agency_overlap_sheet.append(
        (
            "document_type",
            "item_id",
            "condition_dimensions_json",
            "source_values_json",
            "source_evidence_json",
            "candidate_status",
            "review_status",
            "hold_reason",
        )
    )
    for row in agency_file_overlaps:
        agency_overlap_sheet.append(
            (
                row["document_type"],
                row["item_id"],
                row["condition_dimensions"],
                row["source_values_json"],
                row["source_evidence_json"],
                row["status"],
                row["review_status"],
                row["hold_reason"],
            )
        )

    crosswalk = workbook.create_sheet("P2項目対応")
    crosswalk.append(
        ("P1項目", "P2入力先", "必要な確認", "状態", "自動転記")
    )
    for source_item, destination, review, status in P2_CROSSWALK:
        crosswalk.append((source_item, destination, review, status, "NO"))

    missing = workbook.create_sheet("不足・保留")
    missing.append(("scope", "status", "reason"))
    for kind in missing_kinds:
        missing.append(
            (
                kind,
                "MISSING_SOURCE",
                "比較入力フォルダーまたは対応する原本がありません。差分判定を確定できません。",
            )
        )
    for source in sources:
        if source.error_category:
            missing.append(
                (
                    source.relative_path,
                    "READ_ERROR",
                    f"読取に失敗しました ({source.error_category})。原本と変換可能性を確認してください。",
                )
            )
    for reason in agency_group_holds:
        missing.append(
            (
                "AGENCY_DOCUMENT_SET",
                "UNRESOLVED",
                reason,
            )
        )
    if agency_overlap_hold_count:
        missing.append(
            (
                "AGENCY_FILE_OVERLAP",
                "UNRESOLVED",
                f"代理店資料間で確認が必要な条件キー: {agency_overlap_hold_count}",
            )
        )
    if agency_comparison_hold_count:
        missing.append(
            (
                "AGENCY_UPPER_COMPARISON",
                "UNRESOLVED",
                f"代理店資料と上位原本の対応確認が必要な行: {agency_comparison_hold_count}",
            )
        )
    missing.append(
        (
            "UPPER_DIFFERENCE",
            "UNRESOLVED",
            f"条件軸不明・重複等によりレビューが必要な差分行: {unresolved_difference_count}",
        )
    )
    missing.append(
        (
            "PREVIOUS_AGENCY_MAPPING",
            "UNRESOLVED",
            "前月代理店帳票からagency_idとP2条件項目への対応は自動確定しません。",
        )
    )
    missing.append(
        (
            "FINANCIAL_RULES",
            "UNRESOLVED",
            "粗利・原価・税区分・ラダー・例外・掲載可否・前回確定額は承認済み社内データが必要です。",
        )
    )

    _style_workbook(workbook)
    workbook.save(output_path)
    workbook.close()


def _style_workbook(workbook: Workbook) -> None:
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in worksheet.columns:
            values = [str(cell.value or "") for cell in column_cells[: min(worksheet.max_row, 30)]]
            width = min(max((len(value) for value in values), default=12) + 2, 72)
            worksheet.column_dimensions[column_cells[0].column_letter].width = max(width, 12)
        table_names = {
            "抽出根拠候補": "MonthlyBridgeEvidence",
            "上位差分候補": "MonthlyBridgeDifferences",
            "代理店資料種別": "MonthlyAgencyRoles",
            "代理店別上位比較": "MonthlyAgencyUpperComparisons",
            "代理店資料間照合": "MonthlyAgencyOverlaps",
            "不足・保留": "MonthlyBridgeHolds",
        }
        if worksheet.title in table_names:
            worksheet.auto_filter.ref = worksheet.dimensions
            table = Table(
                displayName=table_names[worksheet.title],
                ref=worksheet.dimensions,
            )
            table.tableStyleInfo = TableStyleInfo(
                name="TableStyleMedium2",
                showRowStripes=True,
                showColumnStripes=False,
            )
            worksheet.add_table(table)


def _require_external_output(project_root: Path, output_path: Path) -> None:
    try:
        output_path.resolve().relative_to(project_root.resolve())
    except ValueError:
        return
    raise MonthlyBridgeError("Comparison workbook must be stored outside the Git repository.")


def _validate_period(value: str) -> None:
    if len(value) != 7 or value[4] != "-" or not value[:4].isdigit() or not value[5:].isdigit():
        raise MonthlyBridgeError("Periods must use YYYY-MM format.")
    month = int(value[5:])
    if not 1 <= month <= 12:
        raise MonthlyBridgeError("Periods must use a valid month.")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
