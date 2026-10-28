from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass(frozen=True)
class EvidenceBlock:
    source_path: str
    location: str
    text: str


def read_document(path: Path, *, project_root: Path | None = None) -> list[EvidenceBlock]:
    suffix = path.suffix.lower()
    if suffix not in {".pdf", ".xlsx", ".docx"}:
        raise ValueError(f"Unsupported document type: {suffix or '(no extension)'}")

    relative_path = (
        path.resolve().relative_to(project_root.resolve()).as_posix()
        if project_root is not None
        else path.name
    )

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return [
            EvidenceBlock(
                source_path=relative_path,
                location=f"page:{page_number}",
                text=page.extract_text() or "",
            )
            for page_number, page in enumerate(reader.pages, start=1)
        ]

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=True, read_only=True)
        blocks: list[EvidenceBlock] = []
        try:
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    cells = [
                        f"{cell.coordinate}={cell.value}"
                        for cell in row
                        if cell.value is not None
                    ]
                    if cells:
                        blocks.append(
                            EvidenceBlock(
                                source_path=relative_path,
                                location=f"sheet:{worksheet.title};cells:{','.join(cell.split('=', 1)[0] for cell in cells)}",
                                text=" | ".join(cells),
                            )
                        )
        finally:
            workbook.close()
        return blocks

    from docx import Document

    document = Document(path)
    blocks = []
    for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            blocks.append(
                EvidenceBlock(
                    source_path=relative_path,
                    location=f"paragraph:{paragraph_number}",
                    text=paragraph.text,
                )
            )

    for table_number, table in enumerate(document.tables, start=1):
        for row_number, row in enumerate(table.rows, start=1):
            for column_number, cell in enumerate(row.cells, start=1):
                if cell.text.strip():
                    blocks.append(
                        EvidenceBlock(
                            source_path=relative_path,
                            location=(
                                f"table:{table_number};row:{row_number};"
                                f"cell:{column_number}"
                            ),
                            text=cell.text,
                        )
                    )
    return blocks