"""November master policy, separate from financial and final approvals."""
from collections import Counter
import json
from openpyxl import load_workbook


def write_november(book, results, current, groups, members, snapshot, review_input, write_sheet):
    # Import lazily: reuse the existing approval evidence and tamper checks.
    from .phase2_review import APPROVAL_FIELDS, input_hash, approval_errors, old_axes, related_axes
    headers=['issue_id','scope_key','decision','resolution','snapshot_hash']+APPROVAL_FIELDS
    saved={}
    if review_input:
        wb=load_workbook(review_input,read_only=True,data_only=False)
        try:
            if '共通一括判断入力' in wb:
                sheet=wb['共通一括判断入力']
                if list(next(sheet.values))!=headers:raise ValueError('共通一括判断入力の列が不一致')
                for cells in sheet.iter_rows(min_row=2):
                    if any(c.data_type=='f' for c in cells):raise ValueError('共通判断に数式は使用できません')
                    r=dict(zip(headers,[c.value for c in cells]))
                    if r['issue_id'] in saved:raise ValueError('共通論点ID重複')
                    saved[r['issue_id']]=r
        finally:wb.close()
    common={k:g for k,g in groups.items() if g.scope=='共通'}
    if set(saved)-set(common):raise ValueError('対象外の共通論点ID')
    resolved=set();inputs=[];tasks=[]
    for key,g in common.items():
        r=saved.get(key) or dict.fromkeys(headers)
        expected=json.dumps(sorted(x['proposal'].id for x in members[key]),ensure_ascii=False)
        if key in saved and (r['scope_key']!=expected or r['snapshot_hash']!=snapshot):
            raise ValueError('共通判断の対象範囲・資料が不一致')
        r.update(issue_id=key,scope_key=expected,snapshot_hash=snapshot)
        r['record_hash']=input_hash('共通一括判断入力',r)
        valid=r['decision']=='CONFIRMED' and bool(r['resolution']) and not approval_errors('共通一括判断入力',r,snapshot)
        if valid:resolved.add(key)
        inputs.append([r.get(h) for h in headers])
        tasks.append([key,g.label,g.question,len(members[key]),'解消' if valid else '保留',r.get('resolution')])
    write_sheet(book,'共通一括判断入力',headers,inputs,True)
    write_sheet(book,'共通適用条件判断',['論点ID','論点','担当者が判断する内容','対象参照行数','判定','根拠付き判断'],tasks)
    agencies=sorted({r['proposal'].agency for r in results})
    rows=[];remaining={};policy=Counter()
    for agency in agencies:
        old_results=[r for r in results if r['proposal'].agency==agency and r['proposal'].old]
        for u in current:
            exact=[r for r in old_results if old_axes(r['proposal'].old)==u.axes]
            related=[r for r in old_results if related_axes(old_axes(r['proposal'].old),u.axes)]
            matches=exact or related
            classification='同等項目・既存粗利/個別条件引継ぎ' if len(exact)==1 else '11月新規・採用' if not matches else '適用関係保留'
            # Use only this upper row and relevant October references, never another agency.
            references=matches+[r for r in results if r['proposal'].agency==agency and not r['proposal'].old and r['proposal'].upper.id==u.id]
            issues={i.id:i for r in references for i in r['review_issues']}
            pending={k:i for k,i in issues.items() if k not in resolved and i.code!='AGENCY_ADOPTION'}
            if u.state!='VALUE' or u.hold:
                # A resolved common issue may establish applicability, but never creates a price.
                if not resolved.intersection(issues) or any(i.scope=='共通' and i.id not in resolved and i.code in ('UPPER_LAYOUT','UPPER_AMOUNT_STATE','CAMPAIGN_RULE') for i in issues.values()):classification='適用関係保留'
            finance_key='FINANCE:'+agency+':'+u.id
            if not any(r['proposed'] is not None for r in references):remaining[finance_key]=(agency,'代理店別単価・粗利/原価', '既存粗利・個別約定の根拠を引継ぎ、原価・追加報酬を確認。新規は単価と粗利を設定',u.id)
            for k,i in pending.items():remaining[k]=(i.scope,i.label,i.question,k)
            ids=';'.join(r['proposal'].old.id for r in matches)
            amounts=';'.join(str(r['proposal'].old.amount) for r in matches)
            rows.append([agency,u.id,json.dumps(u.axes,ensure_ascii=False),classification,ids,amounts,
                u.amount,u.state,';'.join(pending), '保留（単価・財務確認未完了）',
                '10月条件を参照。実粗利不明は推測しない',u.source+'#'+u.location])
            policy[classification]+=1
        for r in old_results:
            p=r['proposal']
            if any(related_axes(old_axes(p.old),u.axes) for u in current):continue
            individual=bool(p.old.hold or p.old.terms or p.old.legal)
            category='10月のみ・終了'
            policy[category]+=1
            if individual:remaining['OLD:'+p.id]=(agency,'個別約定保留','個別約定の存続・終了を契約根拠で判断',p.old.id)
            rows.append([agency,'',json.dumps(p.axes,ensure_ascii=False),category,p.old.id,p.old.amount,None,
                         '11月対応なし','', '保留' if individual else '終了（方針適用）','個別約定を優先して確認',p.old.source])
    write_sheet(book,'11月正本再分類',['代理店','11月上位ID','条件軸','方針分類','参照10月条件ID','参照10月代理店単価',
        '11月上位単価','上位金額状態','残適用論点ID','単価状態','粗利・個別条件の扱い','正本根拠'],rows)
    write_sheet(book,'一括判定後の残論点',['判断ID','担当範囲','判断区分','担当者が実際に判断する内容','対象ID'],
                [[k,*v] for k,v in sorted(remaining.items())])
    stats=dict(master_rows=len(rows),categories=dict(policy),common_resolved=len(resolved),
               remaining_common=sum(v[0]=='共通' for v in remaining.values()),
               remaining_individual=sum(v[0]!='共通' for v in remaining.values()),remaining_total=len(remaining),
               price_tasks=sum(k.startswith('FINANCE:') for k in remaining))
    write_sheet(book,'11月正本ダッシュボード',['項目','値'],[
        ['方式','11月正本再分類 → 共通適用条件判断 → 一括判定後の残論点。旧確認表は比較・承認用の参照。'],
        ['一括方針','11月新規は採用、10月のみは終了。同等は既存粗利・個別条件を参照。適用不明・個別約定・財務不明は保留。'],
        ['共通判断手順','共通一括判断入力にCONFIRMED・resolutionと根拠を入力し再計算。record_hashを確認してapproved_hashへ転記し既存承認証跡を記入、再計算。代理店単価・最終承認は個別入力。'],
        ['件数定義','共通論点は1判断、単価は代理店×11月上位IDで1判断。244旧論点とは母集団が異なる。共通解消は本方式の適用確認に反映し、旧金額承認には昇格しない。'],
        *[[k,json.dumps(v,ensure_ascii=False)] for k,v in stats.items()]])
    book._sheets.insert(0,book._sheets.pop(book.sheetnames.index('11月正本ダッシュボード')))
    return stats
