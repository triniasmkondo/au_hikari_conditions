from __future__ import annotations

import hashlib
import os
import re
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .p2 import P2Report, calculate_workbook


_REQUIRED_SHEETS = {"代理店別計算", "保留理由"}
_RESULT_HEADERS = (
    "target_period",
    "agency_id",
    "condition_item_id",
    "status",
    "upstream_reward",
    "direct_cost",
    "approved_gross_margin",
    "previous_agency_reward",
    "proposed_agency_reward",
    "proposed_difference",
    "display_policy",
)
_DISPLAY_ATTRIBUTES = (
    ("商材", "product_family"),
    ("プラス種別", "plus_type"),
    ("ISP", "isp"),
    ("住居種別", "dwelling_type"),
    ("速度", "speed"),
    ("アクセス種別", "access_type"),
    ("申込・成果条件", "trigger"),
    ("オプション", "option_code"),
    ("キャンペーン", "campaign_code"),
    ("適用期間", "effective_from/effective_until"),
    ("期間ルール", "period_rule"),
    ("支払時期", "payment_timing"),
    ("除外条件", "exclusions"),
    ("戻入条件", "reversal_terms"),
    ("代理店向け掲載可否", "publication_policy"),
)
_REQUIRED_CONDITION_FIELDS = (
    "product_family",
    "plus_type",
    "isp",
    "dwelling_type",
    "speed",
    "access_type",
    "trigger",
    "effective_from",
    "effective_until",
    "period_rule",
    "payment_timing",
    "exclusions",
    "reversal_terms",
    "option_code",
    "campaign_code",
)
_EXPLICIT_NOT_APPLICABLE = {"N/A", "NA", "NOT_APPLICABLE", "該当なし", "対象外"}
_UNCONFIRMED_VALUES = {
    "UNKNOWN",
    "UNREVIEWED",
    "PENDING",
    "NOT_REVIEWED",
    "REVIEW_REQUIRED",
    "DIFFERENT",
}
_CURRENCY_FORMAT = '#,##0.##;[Red]-#,##0.##'
_NAVY = "17365D"
_BLUE = "D9EAF7"
_WHITE = "FFFFFF"
_DARK_GREY = "666666"


class P3GenerationError(ValueError):
    """Raised when P2 output is not safe to render as an agency condition sheet."""


@dataclass(frozen=True)
class GeneratedAgencyArtifacts:
    agency_id: str
    agency_name: str
    spreadsheet_path: Path
    word_path: Path
    condition_count: int


@dataclass(frozen=True)
class P3GenerationReport:
    target_period: str
    status: str
    artifacts: tuple[GeneratedAgencyArtifacts, ...]


def generate_agency_documents(
    *,
    approved_input_path: Path,
    p2_result_path: Path,
    target_period: str,
    output_directory: Path,
) -> P3GenerationReport:
    return _generate_agency_documents(
        approved_input_path=approved_input_path,
        p2_result_path=p2_result_path,
        target_period=target_period,
        output_directory=output_directory,
        distribution_candidate=False,
    )


def generate_distribution_candidates(
    *,
    approved_input_path: Path,
    p2_result_path: Path,
    target_period: str,
    output_directory: Path,
) -> P3GenerationReport:
    """Create clean-content review candidates; these are not delivery artifacts."""
    return _generate_agency_documents(
        approved_input_path=approved_input_path,
        p2_result_path=p2_result_path,
        target_period=target_period,
        output_directory=output_directory,
        distribution_candidate=True,
    )


def distribution_candidate_paths(
    *, agency_id: str, target_period: str, output_directory: Path
) -> tuple[Path, Path]:
    filename_key = _safe_file_key(agency_id) + "_" + hashlib.sha256(
        agency_id.encode("utf-8")
    ).hexdigest()[:8]
    if not filename_key:
        raise P3GenerationError("Agency ID cannot be safely used for artifact names.")
    stem = f"{target_period}_{filename_key}_approval_candidate"
    return output_directory / f"{stem}.xlsx", output_directory / f"{stem}.docx"


def validate_distribution_candidate(
    *,
    candidate: GeneratedAgencyArtifacts,
    approved_input_path: Path,
    p2_result_path: Path,
    target_period: str,
) -> None:
    recalculated = calculate_workbook(approved_input_path, target_period)
    if not recalculated.results or recalculated.hold_count:
        raise P3GenerationError("P2 approval is not valid for distribution candidates.")
    report_rows = _read_p2_result(p2_result_path, target_period)
    _verify_p2_rows(report_rows, recalculated)
    render_data = _load_render_data(approved_input_path, recalculated, target_period)
    expected_items = render_data.get(candidate.agency_id)
    if (
        expected_items is None
        or not expected_items
        or candidate.condition_count != len(expected_items)
        or not candidate.agency_name
        or expected_items[0]["agency_name"] != candidate.agency_name
    ):
        raise P3GenerationError("Distribution candidate metadata does not match P2.")
    _validate_agency_workbook(
        candidate.spreadsheet_path,
        candidate.agency_name,
        candidate.agency_id,
        target_period,
        expected_items,
        distribution_candidate=True,
    )
    _validate_agency_word(
        candidate.word_path,
        candidate.agency_name,
        candidate.agency_id,
        target_period,
        expected_items,
        distribution_candidate=True,
    )


def _generate_agency_documents(
    *,
    approved_input_path: Path,
    p2_result_path: Path,
    target_period: str,
    output_directory: Path,
    distribution_candidate: bool,
) -> P3GenerationReport:
    """Verify P2 output and create internal forms or isolated review candidates."""
    recalculated = calculate_workbook(approved_input_path, target_period)
    if not recalculated.results:
        raise P3GenerationError("P2 result contains no agency conditions.")
    if recalculated.hold_count:
        raise P3GenerationError(
            f"P2 has {recalculated.hold_count} held condition(s); no documents were generated."
        )
    if recalculated.proposed_count != len(recalculated.results):
        raise P3GenerationError("P2 result includes a non-proposed condition.")

    report_rows = _read_p2_result(p2_result_path, target_period)
    _verify_p2_rows(report_rows, recalculated)
    render_data = _load_render_data(
        approved_input_path, recalculated, target_period
    )
    if output_directory.exists() and not output_directory.is_dir():
        raise P3GenerationError("Output directory path exists and is not a directory.")
    output_directory.mkdir(parents=True, exist_ok=True)

    planned_paths: list[tuple[str, Path, Path, list[dict[str, Any]]]] = []
    for agency_id, items in sorted(render_data.items()):
        if distribution_candidate:
            xlsx_path, docx_path = distribution_candidate_paths(
                agency_id=agency_id,
                target_period=target_period,
                output_directory=output_directory,
            )
        else:
            filename_key = _safe_file_key(agency_id) + "_" + hashlib.sha256(
                agency_id.encode("utf-8")
            ).hexdigest()[:8]
            if not filename_key:
                raise P3GenerationError(
                    "Agency ID cannot be safely used for artifact names."
                )
            stem = f"{target_period}_{filename_key}_conditions"
            xlsx_path = output_directory / f"{stem}.xlsx"
            docx_path = output_directory / f"{stem}.docx"
        planned_paths.append((agency_id, xlsx_path, docx_path, items))
    existing = [
        path
        for _, xlsx_path, docx_path, _ in planned_paths
        for path in (xlsx_path, docx_path)
        if path.exists()
    ]
    if existing:
        raise FileExistsError(
            "Refusing to overwrite existing P3 artifact(s): "
            + ", ".join(str(path) for path in existing)
        )

    staged_artifacts: list[
        tuple[str, str, Path, Path, list[dict[str, Any]]]
    ] = []
    with tempfile.TemporaryDirectory(
        prefix=".p3-staging-", dir=output_directory
    ) as staging_name:
        staging = Path(staging_name)
        for agency_id, xlsx_path, docx_path, items in planned_paths:
            agency_name = items[0]["agency_name"]
            staged_xlsx = staging / xlsx_path.name
            staged_docx = staging / docx_path.name
            _write_agency_workbook(
                staged_xlsx,
                agency_name,
                agency_id,
                target_period,
                items,
                distribution_candidate=distribution_candidate,
            )
            _write_agency_word(
                staged_docx,
                agency_name,
                agency_id,
                target_period,
                items,
                distribution_candidate=distribution_candidate,
            )
            _validate_agency_workbook(
                staged_xlsx,
                agency_name,
                agency_id,
                target_period,
                items,
                distribution_candidate=distribution_candidate,
            )
            _validate_agency_word(
                staged_docx,
                agency_name,
                agency_id,
                target_period,
                items,
                distribution_candidate=distribution_candidate,
            )
            staged_artifacts.append(
                (agency_id, agency_name, staged_xlsx, staged_docx, items)
            )
        for _, _, staged_xlsx, staged_docx, _ in staged_artifacts:
            if (output_directory / staged_xlsx.name).exists() or (
                output_directory / staged_docx.name
            ).exists():
                raise FileExistsError("A P3 artifact appeared during generation; refusing to overwrite.")
        artifacts: list[GeneratedAgencyArtifacts] = []
        for agency_id, agency_name, staged_xlsx, staged_docx, items in staged_artifacts:
            xlsx_path = output_directory / staged_xlsx.name
            docx_path = output_directory / staged_docx.name
            os.replace(staged_xlsx, xlsx_path)
            os.replace(staged_docx, docx_path)
            artifacts.append(
                GeneratedAgencyArtifacts(
                    agency_id=agency_id,
                    agency_name=agency_name,
                    spreadsheet_path=xlsx_path,
                    word_path=docx_path,
                    condition_count=len(items),
                )
            )
    return P3GenerationReport(
        target_period=target_period,
        status=(
            "APPROVAL_CANDIDATES_NOT_APPROVED_FOR_DELIVERY"
            if distribution_candidate
            else "INTERNAL_REVIEW_NOT_APPROVED_FOR_DELIVERY"
        ),
        artifacts=tuple(artifacts),
    )


def _read_p2_result(path: Path, target_period: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = _REQUIRED_SHEETS - set(workbook.sheetnames)
        if missing_sheets:
            raise P3GenerationError(
                f"P2 result workbook is missing sheets: {', '.join(sorted(missing_sheets))}"
            )
        worksheet = workbook["代理店別計算"]
        values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = [str(value).strip() if value is not None else "" for value in values]
        missing_headers = set(_RESULT_HEADERS) - set(headers)
        if missing_headers:
            raise P3GenerationError(
                "P2 result workbook is missing columns: "
                + ", ".join(sorted(missing_headers))
            )
        result_rows: list[dict[str, Any]] = []
        for values in worksheet.iter_rows(min_row=2, values_only=True):
            if not any(value is not None for value in values):
                continue
            row = {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
            if str(row.get("target_period", "")).strip() != target_period:
                raise P3GenerationError("P2 result contains a different target period.")
            if str(row.get("status", "")).strip() != "proposed":
                raise P3GenerationError(
                    "P2 result contains a held or non-final condition; no documents were generated."
                )
            result_rows.append(row)
        hold_sheet = workbook["保留理由"]
        hold_rows = [
            row
            for row in hold_sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
        if hold_rows:
            raise P3GenerationError("P2 result has a non-empty hold-reasons sheet.")
        return result_rows
    finally:
        workbook.close()


def _verify_p2_rows(report_rows: list[dict[str, Any]], recalculated: P2Report) -> None:
    if len(report_rows) != len(recalculated.results):
        raise P3GenerationError(
            "P2 result row count does not match approved-input recalculation."
        )
    by_key = {
        (result.agency_id, result.condition_item_id): result
        for result in recalculated.results
    }
    if len(by_key) != len(recalculated.results):
        raise P3GenerationError("P2 recalculation has duplicate agency/condition keys.")
    seen: set[tuple[str, str]] = set()
    decimal_fields = {
        "upstream_reward": "upstream_reward",
        "direct_cost": "direct_cost",
        "approved_gross_margin": "approved_gross_margin",
        "previous_agency_reward": "previous_agency_reward",
        "proposed_agency_reward": "proposed_agency_reward",
        "proposed_difference": "proposed_difference",
    }
    for row in report_rows:
        key = (
            _as_text(row.get("agency_id")),
            _as_text(row.get("condition_item_id")),
        )
        expected = by_key.get(key)
        if expected is None or key in seen:
            raise P3GenerationError(
                "P2 result has an unexpected or duplicate agency/condition row."
            )
        seen.add(key)
        if _as_text(row.get("status")) != expected.status:
            raise P3GenerationError(f"P2 status mismatch for {key}.")
        if _as_text(row.get("display_policy")) != expected.display_policy:
            raise P3GenerationError(f"P2 display policy mismatch for {key}.")
        for report_field, result_field in decimal_fields.items():
            if _as_decimal(row.get(report_field)) != getattr(expected, result_field):
                raise P3GenerationError(
                    f"P2 {report_field} does not match recalculation for {key}."
                )
    if seen != set(by_key):
        raise P3GenerationError("P2 result is missing one or more agency conditions.")


def _load_render_data(
    approved_input_path: Path, recalculated: P2Report, target_period: str
) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(approved_input_path, read_only=True, data_only=True)
    try:
        gross_rows = _sheet_rows(workbook, "粗利原価")
        mappings = _sheet_rows(workbook, "条件対応")
    finally:
        workbook.close()

    result_by_key = {
        (result.agency_id, result.condition_item_id): result
        for result in recalculated.results
    }
    period_start, period_end = _month_bounds(target_period)
    gross_by_key: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for row in gross_rows:
        key = (
            _as_text(row.get("agency_id")),
            _as_text(row.get("condition_item_id")),
        )
        start = _date_value(row.get("valid_from"))
        end = _date_value(row.get("valid_until"))
        if (
            key in result_by_key
            and start is not None
            and end is not None
            and start <= period_end
            and end >= period_start
        ):
            gross_by_key.setdefault(key, []).append(row)

    mappings_by_id: dict[str, list[dict[str, Any]]] = {}
    for mapping in mappings:
        mappings_by_id.setdefault(_as_text(mapping.get("mapping_id")), []).append(mapping)

    output: dict[str, list[dict[str, Any]]] = {}
    for key, result in result_by_key.items():
        rows = gross_by_key.get(key, [])
        if not rows:
            raise P3GenerationError(f"No approved gross/cost input found for {key}.")
        agency_names = {
            _as_text(row.get("agency_name"))
            for row in rows
            if _as_text(row.get("agency_name"))
        }
        if len(agency_names) != 1:
            raise P3GenerationError(
                f"Agency name is missing or inconsistent for {key}."
            )
        mapping_ids = {_as_text(row.get("mapping_id")) for row in rows}
        if len(mapping_ids) != 1 or "" in mapping_ids:
            raise P3GenerationError(
                f"Approved input rows do not resolve to one condition mapping for {key}."
            )
        mapping_id = next(iter(mapping_ids))
        candidates = mappings_by_id.get(mapping_id, [])
        if len(candidates) != 1:
            raise P3GenerationError(
                f"Condition mapping {mapping_id!r} is missing or ambiguous."
            )
        mapping = candidates[0]
        if (
            _as_text(mapping.get("agency_id")) != result.agency_id
            or _as_text(mapping.get("agency_item_id")) != result.condition_item_id
            or _as_text(mapping.get("approval_status")).upper() != "APPROVED"
            or _as_text(mapping.get("equivalence_review_status")).upper()
            != "REVIEWED_EQUIVALENT"
        ):
            raise P3GenerationError(
                f"Condition mapping {mapping_id!r} is not an approved matching record."
            )
        if result.display_policy not in {"present", "withhold"}:
            raise P3GenerationError(
                f"Condition {key} has an unresolved publication policy."
            )
        missing_fields = [
            field
            for field in _REQUIRED_CONDITION_FIELDS
            if not _is_explicit_value(mapping.get(field))
        ]
        if missing_fields:
            raise P3GenerationError(
                f"Condition {key} has unconfirmed display attributes: "
                + ", ".join(missing_fields)
            )
        item = {
            "agency_name": next(iter(agency_names)),
            "condition_item_id": result.condition_item_id,
            "product_family": _as_text(mapping.get("product_family")),
            "plus_type": _as_text(mapping.get("plus_type")),
            "isp": _as_text(mapping.get("isp")),
            "dwelling_type": _as_text(mapping.get("dwelling_type")),
            "speed": _as_text(mapping.get("speed")),
            "access_type": _as_text(mapping.get("access_type")),
            "trigger": _as_text(mapping.get("trigger")),
            "option_code": _display_optional(mapping.get("option_code")),
            "campaign_code": _display_optional(mapping.get("campaign_code")),
            "effective_from": _date_display(mapping.get("effective_from")),
            "effective_until": _date_display(mapping.get("effective_until")),
            "period_rule": _as_text(mapping.get("period_rule")),
            "payment_timing": _as_text(mapping.get("payment_timing")),
            "exclusions": _as_text(mapping.get("exclusions")),
            "reversal_terms": _as_text(mapping.get("reversal_terms")),
            "proposed_agency_reward": result.proposed_agency_reward,
            "previous_agency_reward": result.previous_agency_reward,
            "proposed_difference": result.proposed_difference,
            "display_policy": result.display_policy,
            "publication_policy": (
                "掲載" if result.display_policy == "present" else "非提示"
            ),
        }
        output.setdefault(result.agency_id, []).append(item)

    for items in output.values():
        items.sort(key=lambda item: item["condition_item_id"])
    return output


def _write_agency_workbook(
    path: Path,
    agency_name: str,
    agency_id: str,
    target_period: str,
    items: list[dict[str, Any]],
    *,
    distribution_candidate: bool = False,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "条件表"
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:E1")
    if not distribution_candidate:
        worksheet["A1"] = "内部確認用・配信用成果物ではありません（未承認）"
        worksheet["A1"].font = Font(name="Meiryo", size=12, bold=True, color="9C0006")
        worksheet["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
        worksheet["A1"].alignment = Alignment(horizontal="center")
    worksheet.merge_cells("A2:E2")
    worksheet["A2"] = f"{agency_name} 御中"
    worksheet["A2"].font = Font(name="Meiryo", size=16, bold=True, color=_NAVY)
    worksheet["A2"].alignment = Alignment(horizontal="center")
    worksheet.merge_cells("A3:E3")
    worksheet["A3"] = f"{target_period} 条件表"
    worksheet["A3"].font = Font(name="Meiryo", size=14, bold=True)
    worksheet["A3"].alignment = Alignment(horizontal="center")
    worksheet.merge_cells("A4:E4")
    worksheet["A4"] = (
        f"代理店ID: {agency_id}"
        if distribution_candidate
        else f"代理店ID: {agency_id} / 状態: INTERNAL_REVIEW_ONLY"
    )
    worksheet["A4"].font = Font(name="Meiryo", size=9, color=_DARK_GREY)
    headers = (
        "成果項目ID",
        "条件項目",
        "適用条件",
        "代理店報酬（税抜）",
        "代理店掲載可否",
    )
    header_row = 6
    for column, header in enumerate(headers, 1):
        cell = worksheet.cell(header_row, column, header)
        cell.font = Font(name="Meiryo", bold=True, color=_WHITE)
        cell.fill = PatternFill("solid", fgColor=_NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row_number = header_row + 1
    thin_grey = Side(style="thin", color="B7C9D6")
    for item in items:
        condition_lines = [
            f"{label}: {_condition_display(item, key)}"
            for label, key in _DISPLAY_ATTRIBUTES
        ]
        values = (
            item["condition_item_id"],
            f"{item['product_family']} / {item['plus_type']}",
            "\n".join(condition_lines),
            _payment_display(item),
            item["publication_policy"],
        )
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row_number, column, value)
            cell.font = Font(name="Meiryo", size=10, color="000000")
            cell.alignment = Alignment(
                vertical="top",
                horizontal="right" if column == 4 and isinstance(value, Decimal) else "left",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin_grey)
        amount_cell = worksheet.cell(row_number, 4)
        amount_cell.number_format = _CURRENCY_FORMAT
        if isinstance(amount_cell.value, Decimal):
            amount_cell.value = float(amount_cell.value)
            amount_cell.number_format = _CURRENCY_FORMAT
        worksheet.cell(row_number, 1).fill = PatternFill("solid", fgColor=_BLUE)
        worksheet.row_dimensions[row_number].height = max(55, len(condition_lines) * 20)
        row_number += 1
    footer_row = row_number + 1
    if not distribution_candidate:
        worksheet.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=5)
        worksheet.cell(footer_row, 1, "この帳票は内部確認用です。承認・QA完了前の配布は禁止です。")
        worksheet.cell(footer_row, 1).font = Font(name="Meiryo", size=9, italic=True, color="9C0006")
        worksheet.cell(footer_row, 1).alignment = Alignment(wrap_text=True)
    for column, width in enumerate((24, 28, 62, 24, 18), 1):
        worksheet.column_dimensions[chr(64 + column)].width = width
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = f"A{header_row}:E{row_number - 1}"
    worksheet.print_title_rows = f"1:{header_row}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.page_setup.orientation = "landscape"
    worksheet.print_area = f"A1:E{footer_row if not distribution_candidate else row_number - 1}"
    worksheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    workbook.save(path)
    workbook.close()


def _write_agency_word(
    path: Path,
    agency_name: str,
    agency_id: str,
    target_period: str,
    items: list[dict[str, Any]],
    *,
    distribution_candidate: bool = False,
) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Meiryo"
    normal.font.size = Pt(9)
    if not distribution_candidate:
        warning = document.add_paragraph()
        warning.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warning.add_run("内部確認用・配信用成果物ではありません（未承認）")
        run.bold = True
        run.font.color.rgb = RGBColor(156, 0, 6)
    title = document.add_heading(f"{target_period} 条件表", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"{agency_name} 御中")
    document.add_paragraph(
        f"代理店ID: {agency_id}"
        if distribution_candidate
        else f"代理店ID: {agency_id} / 状態: INTERNAL_REVIEW_ONLY"
    )
    for item in items:
        document.add_heading(item["condition_item_id"], level=1)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        table.columns[0].width = Inches(1.7)
        table.columns[1].width = Inches(5.3)
        for label, key in _DISPLAY_ATTRIBUTES:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[1].text = _condition_display(item, key)
            row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in row.cells[0].paragraphs[0].runs:
                run.bold = True
            _shade_cell(row.cells[0], _BLUE)
        reward_row = table.add_row()
        reward_row.cells[0].text = "代理店報酬（税抜）"
        reward_row.cells[1].text = _payment_display(item, formatted=True)
        reward_row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        reward_row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in reward_row.cells[0].paragraphs[0].runs:
            run.bold = True
        _shade_cell(reward_row.cells[0], _BLUE)
        document.add_paragraph()
    if not distribution_candidate:
        footer = document.add_paragraph(
            "この帳票は内部確認用です。承認・QA完了前の配布は禁止です。"
        )
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.italic = True
            run.font.color.rgb = RGBColor(156, 0, 6)
    document.save(path)


def _validate_agency_workbook(
    path: Path,
    agency_name: str,
    agency_id: str,
    target_period: str,
    items: list[dict[str, Any]],
    *,
    distribution_candidate: bool = False,
) -> None:
    workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        if workbook.sheetnames != ["条件表"]:
            raise P3GenerationError("Generated Excel has unexpected worksheet names.")
        worksheet = workbook["条件表"]
        if distribution_candidate:
            worksheet_text = "\n".join(
                str(cell.value or "")
                for row in worksheet.iter_rows()
                for cell in row
            )
            if any(
                marker.casefold() in worksheet_text.casefold()
                for marker in (
                    "内部確認用",
                    "未承認",
                    "配信用成果物ではありません",
                    "INTERNAL_REVIEW_ONLY",
                )
            ):
                raise P3GenerationError(
                    "Distribution candidate contains an internal approval marker."
                )
        expected_banner = (
            None
            if distribution_candidate
            else "内部確認用・配信用成果物ではありません（未承認）"
        )
        if worksheet["A1"].value != expected_banner:
            raise P3GenerationError("Generated Excel approval profile is invalid.")
        if worksheet["A2"].value != f"{agency_name} 御中":
            raise P3GenerationError("Generated Excel agency name does not match approved input.")
        if worksheet["A3"].value != f"{target_period} 条件表":
            raise P3GenerationError("Generated Excel target period is incorrect.")
        if str(worksheet["A4"].value) != (
            f"代理店ID: {agency_id}"
            if distribution_candidate
            else f"代理店ID: {agency_id} / 状態: INTERNAL_REVIEW_ONLY"
        ):
            raise P3GenerationError("Generated Excel agency ID does not match.")
        if (
            worksheet["A6"].value != "成果項目ID"
            or worksheet["D6"].value != "代理店報酬（税抜）"
            or worksheet["E6"].value != "代理店掲載可否"
        ):
            raise P3GenerationError("Generated Excel table headers are incorrect.")
        if worksheet.max_row < (7 if distribution_candidate else 8) or worksheet.max_column != 5:
            raise P3GenerationError("Generated Excel table dimensions are invalid.")
        if (
            not worksheet["A6"].font.bold
            or not str(worksheet["A6"].fill.fgColor.rgb).endswith(_NAVY)
        ):
            raise P3GenerationError("Generated Excel header formatting is incorrect.")
        for row_number, item in enumerate(items, 7):
            if worksheet.cell(row_number, 1).value != item["condition_item_id"]:
                raise P3GenerationError("Generated Excel condition ID failed read-back.")
            if worksheet.cell(row_number, 4).value != _excel_payment(item):
                raise P3GenerationError("Generated Excel reward failed read-back.")
            if worksheet.cell(row_number, 4).number_format != _CURRENCY_FORMAT:
                raise P3GenerationError("Generated Excel reward number format is incorrect.")
            if worksheet.cell(row_number, 5).value != item["publication_policy"]:
                raise P3GenerationError("Generated Excel publication policy failed read-back.")
            rendered_conditions = str(worksheet.cell(row_number, 3).value)
            for label, key in _DISPLAY_ATTRIBUTES:
                expected = f"{label}: {_condition_display(item, key)}"
                if expected not in rendered_conditions:
                    raise P3GenerationError(
                        f"Generated Excel condition read-back mismatch: {label}."
                    )
        if not worksheet.freeze_panes or worksheet.page_setup.fitToWidth != 1:
            raise P3GenerationError("Generated Excel layout settings failed read-back.")
        if not worksheet.print_area:
            raise P3GenerationError("Generated Excel has no print area.")
    finally:
        workbook.close()


def _validate_agency_word(
    path: Path,
    agency_name: str,
    agency_id: str,
    target_period: str,
    items: list[dict[str, Any]],
    *,
    distribution_candidate: bool = False,
) -> None:
    document = Document(path)
    paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
    if distribution_candidate:
        all_text = [paragraphs]
        for table in document.tables:
            all_text.extend(cell.text for row in table.rows for cell in row.cells)
        for section in document.sections:
            all_text.extend(
                paragraph.text
                for paragraph in (*section.header.paragraphs, *section.footer.paragraphs)
            )
            for table in (*section.header.tables, *section.footer.tables):
                all_text.extend(cell.text for row in table.rows for cell in row.cells)
        document_text = "\n".join(all_text)
        if any(
            marker.casefold() in document_text.casefold()
            for marker in (
                "内部確認用",
                "未承認",
                "配信用成果物ではありません",
                "INTERNAL_REVIEW_ONLY",
            )
        ):
            raise P3GenerationError(
                "Distribution candidate contains an internal approval marker."
            )
    warning_paragraph = next(
        (
            paragraph
            for paragraph in document.paragraphs
            if "内部確認用・配信用成果物ではありません（未承認）" in paragraph.text
        ),
        None,
    )
    if distribution_candidate:
        if warning_paragraph is not None or any(
            marker in paragraphs
            for marker in ("未承認", "内部確認用", "INTERNAL_REVIEW_ONLY")
        ):
            raise P3GenerationError("Distribution candidate contains an internal marker.")
    else:
        if warning_paragraph is None:
            raise P3GenerationError("Generated Word document is missing its non-delivery warning.")
        if not warning_paragraph.runs or not warning_paragraph.runs[0].bold:
            raise P3GenerationError("Generated Word warning formatting is incorrect.")
    if agency_name not in paragraphs or agency_id not in paragraphs or target_period not in paragraphs:
        raise P3GenerationError("Generated Word agency or target-period data failed read-back.")
    if any(item["condition_item_id"] not in paragraphs for item in items):
        raise P3GenerationError("Generated Word condition IDs failed read-back.")
        warning_color = warning_paragraph.runs[0].font.color.rgb
        if warning_color is None or str(warning_color)[-6:] != "9C0006":
            raise P3GenerationError("Generated Word warning color is incorrect.")
    tables = document.tables
    if len(tables) != len(items):
        raise P3GenerationError("Generated Word condition table count is incorrect.")
    for table, item in zip(tables, items, strict=True):
        if table.style is None or table.style.name != "Table Grid":
            raise P3GenerationError("Generated Word table style is incorrect.")
        values = [
            (row.cells[0].text, row.cells[1].text)
            for row in table.rows
        ]
        if values[0] != ("商材", item["product_family"]):
            raise P3GenerationError("Generated Word condition content failed read-back.")
        if _cell_shading(table.rows[0].cells[0]) != _BLUE:
            raise P3GenerationError("Generated Word condition-label formatting is incorrect.")
        if values[-1] != (
            "代理店報酬（税抜）",
            _payment_display(item, formatted=True),
        ):
            raise P3GenerationError("Generated Word reward failed read-back.")
        for label, key in _DISPLAY_ATTRIBUTES:
            expected = (label, _condition_display(item, key))
            if expected not in values:
                raise P3GenerationError(
                    f"Generated Word condition read-back mismatch: {label}."
                )
    if (
        not document.sections
        or document.sections[0].page_width <= 0
        or document.sections[0].left_margin != Inches(0.7)
        or document.sections[0].right_margin != Inches(0.7)
    ):
        raise P3GenerationError("Generated Word page layout is invalid.")


def _sheet_rows(workbook, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise P3GenerationError(f"Approved input is missing sheet: {sheet_name}")
    worksheet = workbook[sheet_name]
    headers = [
        str(value).strip() if value is not None else ""
        for value in next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True)
        )
    ]
    if len(headers) != len(set(header for header in headers if header)):
        raise P3GenerationError(f"Approved input sheet {sheet_name} has duplicate headers.")
    return [
        {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        for values in worksheet.iter_rows(min_row=2, values_only=True)
        if any(value is not None for value in values)
    ]


def _month_bounds(target_period: str) -> tuple[date, date]:
    try:
        start_datetime = datetime.strptime(target_period, "%Y-%m")
    except ValueError as error:
        raise P3GenerationError("Target period must use YYYY-MM format.") from error
    if start_datetime.strftime("%Y-%m") != target_period:
        raise P3GenerationError("Target period must use YYYY-MM format.")
    start = date(start_datetime.year, start_datetime.month, 1)
    if start.month == 12:
        next_month = date(start.year + 1, 1, 1)
    else:
        next_month = date(start.year, start.month + 1, 1)
    return start, date.fromordinal(next_month.toordinal() - 1)


def _date_value(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def _date_display(value: Any) -> str:
    parsed = _date_value(value)
    return parsed.isoformat() if parsed is not None else _as_text(value)


def _condition_display(item: dict[str, Any], key: str) -> str:
    if key == "effective_from/effective_until":
        return f"{item['effective_from']} ～ {item['effective_until']}"
    if key == "publication_policy":
        return str(item["publication_policy"])
    value = item.get(key)
    return _display_optional(value)


def _display_optional(value: Any) -> str:
    text = _as_text(value)
    if not text or text.upper() in _EXPLICIT_NOT_APPLICABLE:
        return "該当なし"
    return text


def _is_explicit_value(value: Any) -> bool:
    text = _as_text(value)
    return bool(text) and text.upper() not in _UNCONFIRMED_VALUES


def _payment_display(item: dict[str, Any], *, formatted: bool = False) -> Any:
    if item["display_policy"] == "withhold":
        return "非提示（社内条件）"
    amount = item["proposed_agency_reward"]
    if amount is None:
        raise P3GenerationError("Approved proposed reward is missing.")
    if formatted:
        return f"{_format_amount(amount)} 円"
    return amount


def _excel_payment(item: dict[str, Any]) -> float | str | None:
    if item["display_policy"] == "withhold":
        return "非提示（社内条件）"
    amount = item["proposed_agency_reward"]
    return float(amount) if amount is not None else None


def _format_amount(value: Decimal) -> str:
    if not value.is_finite():
        raise P3GenerationError("Reward must be finite.")
    return f"{value:,.2f}".rstrip("0").rstrip(".")


def _as_decimal(value: Any) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        result = Decimal(str(value).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        raise P3GenerationError(f"Invalid monetary value in P2 result: {value!r}") from None
    if not result.is_finite():
        raise P3GenerationError("P2 result contains a non-finite monetary value.")
    return result


def _as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _safe_file_key(agency_id: str) -> str:
    safe = re.sub(r"[^A-Za-z0-9._-]+", "_", agency_id).strip("._-")
    return safe[:48]


def _shade_cell(cell, color: str) -> None:
    shading = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), color)
    shading.append(element)


def _cell_shading(cell) -> str | None:
    from docx.oxml.ns import qn

    properties = cell._tc.tcPr
    shading = properties.find(qn("w:shd")) if properties is not None else None
    return shading.get(qn("w:fill")) if shading is not None else None
