from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from .calculation import (
    ApprovalEvidence,
    ApprovalStatus,
    ApprovedMarginBaseline,
    CalculationCase,
    ConditionSignature,
    DisplayPolicy,
    TaxBasis,
    calculate_agency_reward,
)


GROSS_MARGIN_REQUIRED_HEADERS = (
    "record_id",
    "agency_id",
    "condition_item_id",
    "mapping_id",
    "upstream_amount",
    "upstream_amount_state",
    "upstream_tax_basis",
    "approved_target_gross_margin",
    "gross_margin_state",
    "gross_margin_tax_basis",
    "display_policy",
    "direct_cost_component",
    "attributable_direct_cost",
    "direct_cost_state",
    "direct_cost_tax_basis",
    "minimum_gross_margin",
    "valid_from",
    "valid_until",
    "source_document",
    "evidence_location",
    "approval_status",
    "approved_by",
    "approved_at",
)
MAPPING_REQUIRED_HEADERS = (
    "mapping_id",
    "source_document",
    "source_location",
    "upstream_item_id",
    "agency_id",
    "agency_item_id",
    "product_family",
    "plus_type",
    "isp",
    "dwelling_type",
    "speed",
    "access_type",
    "trigger",
    "option_code",
    "campaign_code",
    "effective_from",
    "effective_until",
    "period_rule",
    "payment_timing",
    "exclusions",
    "reversal_terms",
    "equivalence_review_status",
    "approval_status",
    "approved_by",
    "approved_at",
)
EXCEPTION_REQUIRED_HEADERS = (
    "exception_id",
    "agency_id",
    "condition_item_id_or_rule_group",
    "rule_type",
    "rule_value",
    "display_policy",
    "reason",
    "source_document",
    "evidence_location",
    "approval_status",
    "approved_by",
    "approved_at",
    "valid_from",
    "valid_until",
)
PREVIOUS_REQUIRED_HEADERS = (
    "record_id",
    "target_period",
    "agency_id",
    "condition_item_id",
    "mapping_id",
    "condition_version",
    "previous_agency_reward",
    "amount_state",
    "tax_basis",
    "source_document",
    "evidence_location",
    "approval_status",
    "approved_by",
    "approved_at",
)
APPROVAL_REQUIRED_HEADERS = (
    "approval_id",
    "agency_id",
    "record_type",
    "record_id",
    "record_hash",
    "approval_status",
    "approved_by",
    "approved_at",
    "supersedes_approval_id",
)

_REQUIRED_SHEETS = {
    "粗利原価": GROSS_MARGIN_REQUIRED_HEADERS,
    "条件対応": MAPPING_REQUIRED_HEADERS,
    "代理店別例外": EXCEPTION_REQUIRED_HEADERS,
    "前回条件": PREVIOUS_REQUIRED_HEADERS,
    "承認管理": APPROVAL_REQUIRED_HEADERS,
}
_APPROVAL_FIELDS = {
    "approval_status",
    "approved_by",
    "approved_at",
    "approval_reason",
    "review_reason",
}
_CONDITION_ATTRIBUTES = (
    "upstream_item_id",
    "product_family",
    "plus_type",
    "isp",
    "dwelling_type",
    "speed",
    "access_type",
    "option_code",
    "campaign_code",
)
_SUPPORTED_EXCEPTION_TYPES = {
    "DISPLAY_POLICY",
    "PAYMENT_PROHIBITED",
    "MANDATORY_DISCLOSURE",
    "HOLD",
}


class P2InputError(ValueError):
    """Raised when the workbook cannot be interpreted as the P2 input schema."""


@dataclass(frozen=True)
class P2AgencyResult:
    target_period: str
    agency_id: str
    condition_item_id: str
    status: str
    upstream_reward: Decimal | None
    direct_cost: Decimal | None
    approved_gross_margin: Decimal | None
    previous_agency_reward: Decimal | None
    proposed_agency_reward: Decimal | None
    proposed_difference: Decimal | None
    display_policy: str
    hold_reasons: tuple[str, ...]


@dataclass(frozen=True)
class P2Report:
    target_period: str
    results: tuple[P2AgencyResult, ...]

    @property
    def proposed_count(self) -> int:
        return sum(result.status == "proposed" for result in self.results)

    @property
    def hold_count(self) -> int:
        return sum(result.status == "hold" for result in self.results)


def compute_record_hash(record_type: str, record: dict[str, Any]) -> str:
    """Return the SHA-256 approval fingerprint for a business record."""
    business_values = {
        key: _canonical_value(value)
        for key, value in record.items()
        if key not in _APPROVAL_FIELDS
    }
    payload = json.dumps(
        {"record_type": record_type, "values": business_values},
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def calculate_workbook(input_path: Path, target_period: str) -> P2Report:
    period_start, period_end = _parse_target_period(target_period)
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        sheets = {
            sheet_name: _read_sheet(workbook, sheet_name, headers)
            for sheet_name, headers in _REQUIRED_SHEETS.items()
        }
    finally:
        workbook.close()

    approvals = sheets["承認管理"]
    cost_rows = sheets["粗利原価"]
    mappings = sheets["条件対応"]
    previous_rows = sheets["前回条件"]
    exceptions = sheets["代理店別例外"]
    for sheet_name, rows, key in (
        ("粗利原価", cost_rows, "record_id"),
        ("条件対応", mappings, "mapping_id"),
        ("前回条件", previous_rows, "record_id"),
        ("代理店別例外", exceptions, "exception_id"),
    ):
        _require_unique_ids(sheet_name, rows, key)

    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    invalid_period_rows: list[dict[str, Any]] = []
    for row in cost_rows:
        if not _has_value(row.get("agency_id")) or not _has_value(
            row.get("condition_item_id")
        ):
            invalid_period_rows.append(row)
            continue
        overlaps = _row_overlaps(row, "valid_from", "valid_until", period_start, period_end)
        if overlaps is None:
            invalid_period_rows.append(row)
        elif overlaps:
            key = (_text(row.get("agency_id")), _text(row.get("condition_item_id")))
            groups.setdefault(key, []).append(row)

    results = [
        _calculate_group(
            target_period=target_period,
            period_start=period_start,
            period_end=period_end,
            agency_id=agency_id,
            item_id=item_id,
            rows=rows,
            mappings=mappings,
            previous_rows=previous_rows,
            exceptions=exceptions,
            approvals=approvals,
        )
        for (agency_id, item_id), rows in sorted(groups.items())
    ]
    for row in invalid_period_rows:
        results.append(
            _held_result(
                target_period,
                _text(row.get("agency_id")) or "(agency_id missing)",
                _text(row.get("condition_item_id")) or "(condition_item_id missing)",
                ("粗利原価の代理店・成果項目または有効期間が不足／不正です。対象月を判定できません。",),
            )
        )

    return P2Report(target_period=target_period, results=tuple(results))


def write_report(report: P2Report, output_path: Path) -> None:
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite existing P2 result: {output_path}")
    workbook = Workbook()
    calculations = workbook.active
    calculations.title = "代理店別計算"
    calculations.append(
        (
            "target_period",
            "agency_id",
            "condition_item_id",
            "status",
            "upstream_reward",
            "direct_cost",
            "approved_gross_margin",
            "previous_agency_reward",
            "proposed_agency_reward",
            "proposed_difference",
            "display_policy",
        )
    )
    holds = workbook.create_sheet("保留理由")
    holds.append(("target_period", "agency_id", "condition_item_id", "hold_reason"))
    for result in report.results:
        calculations.append(
            (
                result.target_period,
                result.agency_id,
                result.condition_item_id,
                result.status,
                _excel_amount(result.upstream_reward),
                _excel_amount(result.direct_cost),
                _excel_amount(result.approved_gross_margin),
                _excel_amount(result.previous_agency_reward),
                _excel_amount(result.proposed_agency_reward),
                _excel_amount(result.proposed_difference),
                result.display_policy,
            )
        )
        for reason in result.hold_reasons:
            holds.append(
                (
                    result.target_period,
                    result.agency_id,
                    result.condition_item_id,
                    reason,
                )
            )
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.column_dimensions["A"].width = 18
        worksheet.column_dimensions["B"].width = 24
        worksheet.column_dimensions["C"].width = 32
        worksheet.column_dimensions["D"].width = 28
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def _calculate_group(
    *,
    target_period: str,
    period_start: date,
    period_end: date,
    agency_id: str,
    item_id: str,
    rows: list[dict[str, Any]],
    mappings: list[dict[str, Any]],
    previous_rows: list[dict[str, Any]],
    exceptions: list[dict[str, Any]],
    approvals: list[dict[str, Any]],
) -> P2AgencyResult:
    reasons: list[str] = []
    representative = rows[0]
    current_mapping = _select_mapping(
        mappings,
        _text(representative.get("mapping_id")),
        agency_id,
        item_id,
        period_start,
        period_end,
    )
    previous_start, previous_end = _previous_month(period_start)
    previous_mapping: dict[str, Any] | None = None
    previous_period = f"{previous_start:%Y-%m}"
    prior_candidates = [
        row
        for row in previous_rows
        if _text(row.get("agency_id")) == agency_id
        and _text(row.get("condition_item_id")) == item_id
        and _text(row.get("target_period")) == target_period
    ]
    previous_reward: Decimal | None = None
    if len(prior_candidates) != 1:
        reasons.append(
            "対象月に対応する前回報酬レコードが一意に存在しません。"
        )
    else:
        previous_row = prior_candidates[0]
        previous_mapping = _select_mapping(
            mappings,
            _text(previous_row.get("mapping_id")),
            agency_id,
            item_id,
            previous_start,
            previous_end,
        )
        previous_reward, amount_reason = _read_amount(
            previous_row.get("previous_agency_reward"),
            previous_row.get("amount_state"),
            "前回代理店報酬",
        )
        reasons.extend(amount_reason)
        if _tax_basis(previous_row.get("tax_basis")) != TaxBasis.EXCLUSIVE:
            reasons.append("前回代理店報酬の税区分が不明または税抜ではありません。")
        _, evidence_reasons = _approval_evidence(
            "前回条件", previous_row, approvals
        )
        reasons.extend(evidence_reasons)

    normalized_rows = rows
    common_fields = (
        "upstream_amount",
        "upstream_amount_state",
        "upstream_tax_basis",
        "mapping_id",
        "approved_target_gross_margin",
        "gross_margin_state",
        "gross_margin_tax_basis",
        "display_policy",
        "minimum_gross_margin",
    )
    for field in common_fields:
        values = {_canonical_value(row.get(field)) for row in normalized_rows}
        if len(values) > 1:
            reasons.append(f"粗利原価の複数行で{field}が一致しません。")

    upstream_reward, amount_reasons = _read_amount(
        representative.get("upstream_amount"),
        representative.get("upstream_amount_state"),
        "11月上位報酬",
    )
    reasons.extend(amount_reasons)
    gross_margin, amount_reasons = _read_amount(
        representative.get("approved_target_gross_margin"),
        representative.get("gross_margin_state"),
        "承認済み目標粗利",
    )
    reasons.extend(amount_reasons)
    costs: list[Decimal] = []
    for row in normalized_rows:
        amount, amount_reasons = _read_amount(
            row.get("attributable_direct_cost"),
            row.get("direct_cost_state"),
            "直接原価",
        )
        costs.extend([amount] if amount is not None else [])
        reasons.extend(amount_reasons)
    direct_cost = sum(costs, Decimal("0")) if len(costs) == len(rows) else None
    minimum_gross_margin, minimum_reasons = _optional_amount(
        representative.get("minimum_gross_margin"), "最低粗利"
    )
    reasons.extend(minimum_reasons)
    gross_approvals: list[ApprovalEvidence] = []
    for row in normalized_rows:
        evidence, evidence_reasons = _approval_evidence("粗利原価", row, approvals)
        gross_approvals.append(evidence)
        reasons.extend(evidence_reasons)
    gross_approval = _combine_approvals(gross_approvals)
    if _tax_basis(representative.get("upstream_tax_basis")) != TaxBasis.EXCLUSIVE:
        reasons.append("上位報酬の税区分が不明または税抜ではありません。")
    if _tax_basis(representative.get("gross_margin_tax_basis")) != TaxBasis.EXCLUSIVE:
        reasons.append("目標粗利の税区分が不明または税抜ではありません。")
    if any(
        _tax_basis(row.get("direct_cost_tax_basis")) != TaxBasis.EXCLUSIVE
        for row in normalized_rows
    ):
        reasons.append("直接原価の税区分が不明または税抜ではありません。")

    current_signature: ConditionSignature | None = None
    previous_signature: ConditionSignature | None = None
    current_mapping_approval = ApprovalEvidence()
    if current_mapping is None:
        reasons.append("対象月の承認済み条件対応レコードがありません。")
    else:
        current_signature = _condition_signature(current_mapping)
        current_mapping_approval, mapping_reasons = _mapping_approval(
            current_mapping, approvals, "対象月"
        )
        reasons.extend(mapping_reasons)
    if len(prior_candidates) == 1 and previous_mapping is None:
        reasons.append(f"{previous_period}の前回条件対応レコードがありません。")
    elif previous_mapping is not None:
        previous_signature = _condition_signature(previous_mapping)
        _, mapping_reasons = _mapping_approval(previous_mapping, approvals, "前回")
        reasons.extend(mapping_reasons)

    display_policy, exception_reasons, exception_flags = _resolve_exceptions(
        agency_id,
        item_id,
        period_start,
        period_end,
        exceptions,
        approvals,
        _display_policy(representative.get("display_policy")),
    )
    reasons.extend(exception_reasons)

    if current_signature is not None and previous_signature is not None:
        if not current_signature.matches_terms(previous_signature):
            reasons.append("前回条件と対象月条件の成果・属性・期間ルール等が一致しません。")
    else:
        reasons.append("条件対応レコードから比較可能な条件署名を作成できません。")

    if reasons:
        return _held_result(
            target_period,
            agency_id,
            item_id,
            tuple(dict.fromkeys(reasons)),
            display_policy=display_policy.value,
        )

    assert current_signature is not None
    assert previous_signature is not None
    assert upstream_reward is not None
    assert direct_cost is not None
    assert gross_margin is not None
    baseline = ApprovedMarginBaseline(
        condition_signature=previous_signature,
        tax_basis=_tax_basis(representative.get("gross_margin_tax_basis")),
        target_gross_margin=gross_margin,
        display_policy=display_policy,
        approval=gross_approval,
    )
    calculation = calculate_agency_reward(
        CalculationCase(
            condition_signature=current_signature,
            upstream_tax_basis=_tax_basis(representative.get("upstream_tax_basis")),
            direct_cost_tax_basis=TaxBasis.EXCLUSIVE,
            upstream_reward=upstream_reward,
            direct_cost=direct_cost,
            baseline=baseline,
            upstream_approval=gross_approval,
            direct_cost_approval=_combine_approvals(
                gross_approvals
            ),
            condition_equivalence_approval=current_mapping_approval,
            minimum_gross_margin=minimum_gross_margin,
            exception_reasons=(),
            contract_prohibits_payment=exception_flags["payment_prohibited"],
            mandatory_disclosure=exception_flags["mandatory_disclosure"],
        )
    )
    if calculation.status != "proposed" or calculation.proposed_agency_reward is None:
        return _held_result(
            target_period,
            agency_id,
            item_id,
            calculation.hold_reasons or ("計算エンジンが保留しました。",),
            display_policy=calculation.display_policy.value,
        )
    assert previous_reward is not None
    return P2AgencyResult(
        target_period=target_period,
        agency_id=agency_id,
        condition_item_id=item_id,
        status="proposed",
        upstream_reward=upstream_reward,
        direct_cost=direct_cost,
        approved_gross_margin=gross_margin,
        previous_agency_reward=previous_reward,
        proposed_agency_reward=calculation.proposed_agency_reward,
        proposed_difference=calculation.proposed_agency_reward - previous_reward,
        display_policy=calculation.display_policy.value,
        hold_reasons=(),
    )


def _read_sheet(workbook, sheet_name: str, required_headers: tuple[str, ...]) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise P2InputError(f"Required P2 input sheet is missing: {sheet_name}")
    worksheet = workbook[sheet_name]
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in header_values]
    if len(headers) != len(set(header for header in headers if header)):
        raise P2InputError(f"Sheet {sheet_name} has duplicate column names.")
    missing = [header for header in required_headers if header not in headers]
    if missing:
        raise P2InputError(
            f"Sheet {sheet_name} is missing required columns: {', '.join(missing)}"
        )
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        rows.append(
            {
                header: values[index] if index < len(values) else None
                for index, header in enumerate(headers)
                if header
            }
        )
    return rows


def _require_unique_ids(
    sheet_name: str, rows: list[dict[str, Any]], id_field: str
) -> None:
    ids = [_text(row.get(id_field)) for row in rows if _has_value(row.get(id_field))]
    if len(ids) != len(set(ids)):
        raise P2InputError(f"Sheet {sheet_name} contains duplicate {id_field} values.")


def _approval_evidence(
    record_type: str,
    record: dict[str, Any],
    approval_rows: list[dict[str, Any]],
) -> tuple[ApprovalEvidence, list[str]]:
    record_id = _text(record.get("record_id") or record.get("mapping_id") or record.get("exception_id"))
    reasons: list[str] = []
    if _text(record.get("approval_status")).upper() != "APPROVED":
        reasons.append(f"{record_type} {record_id or '(ID missing)'} は未承認です。")
    matches = [
        row
        for row in approval_rows
        if _text(row.get("record_type")) == record_type
        and _text(row.get("record_id")) == record_id
    ]
    expected_hash = compute_record_hash(record_type, record)
    superseded_ids = {
        _text(row.get("supersedes_approval_id"))
        for row in matches
        if _has_value(row.get("supersedes_approval_id"))
    }
    active = [
        row
        for row in matches
        if _text(row.get("approval_status")).upper() == "APPROVED"
        and _text(row.get("approval_id")) not in superseded_ids
    ]
    if len(active) != 1:
        reasons.append(
            f"{record_type} {record_id or '(ID missing)'} に有効な承認管理レコードが一意にありません "
            f"(期待するrecord_hash: {expected_hash})。"
        )
        return ApprovalEvidence(), reasons
    approval = active[0]
    if _text(approval.get("record_hash")).lower() != expected_hash:
        reasons.append(
            f"{record_type} {record_id} の承認ハッシュが入力内容と一致しません "
            f"(期待値: {expected_hash})。"
        )
    if _text(approval.get("agency_id")) != _text(record.get("agency_id")):
        reasons.append(f"{record_type} {record_id} の承認管理レコードで代理店IDが一致しません。")
    if _text(approval.get("approved_by")) != _text(record.get("approved_by")):
        reasons.append(f"{record_type} {record_id} の承認者が入力と承認管理で一致しません。")
    record_approved_at = _parse_datetime(record.get("approved_at"))
    approval_approved_at = _parse_datetime(approval.get("approved_at"))
    if record_approved_at is None or approval_approved_at is None:
        reasons.append(f"{record_type} {record_id} の承認日時はタイムゾーン付きISO-8601が必要です。")
    elif record_approved_at != approval_approved_at:
        reasons.append(f"{record_type} {record_id} の承認日時が入力と承認管理で一致しません。")
    if _text(approval.get("approval_status")).upper() != "APPROVED":
        reasons.append(f"{record_type} {record_id} は承認管理上で承認済みではありません。")
    evidence = ApprovalEvidence(
        status=ApprovalStatus.APPROVED,
        approval_id=_text(approval.get("approval_id")) or None,
        source_document=_text(record.get("source_document")) or None,
        evidence_location=_text(
            record.get("evidence_location") or record.get("source_location")
        )
        or None,
        approved_by=_text(approval.get("approved_by")) or None,
        approved_at=approval_approved_at,
    )
    if not evidence.is_complete_approval():
        reasons.append(f"{record_type} {record_id} の承認証跡が不完全です。")
    return evidence, reasons


def _mapping_approval(
    mapping: dict[str, Any],
    approvals: list[dict[str, Any]],
    label: str,
) -> tuple[ApprovalEvidence, list[str]]:
    reasons: list[str] = []
    for field in (
        "upstream_item_id",
        "agency_item_id",
        "product_family",
        "plus_type",
        "isp",
        "dwelling_type",
        "speed",
        "access_type",
        "option_code",
        "campaign_code",
        "effective_from",
        "effective_until",
        "period_rule",
        "payment_timing",
        "trigger",
        "exclusions",
        "reversal_terms",
    ):
        value = _text(mapping.get(field)).upper()
        if not value:
            reasons.append(f"{label}条件対応の{field}が空欄です。")
        elif value in {
            "UNKNOWN",
            "UNREVIEWED",
            "PENDING",
            "NOT_REVIEWED",
            "REVIEW_REQUIRED",
            "DIFFERENT",
        }:
            reasons.append(f"{label}条件対応の{field}は未確定値です。")
    review_status = _text(mapping.get("equivalence_review_status")).upper()
    if review_status != "REVIEWED_EQUIVALENT":
        reasons.append(f"{label}条件対応の同等性が承認済み同等としてレビューされていません。")
    evidence, approval_reasons = _approval_evidence("条件対応", mapping, approvals)
    reasons.extend(approval_reasons)
    return evidence, reasons


def _select_mapping(
    mappings: list[dict[str, Any]],
    mapping_id: str,
    agency_id: str,
    item_id: str,
    period_start: date,
    period_end: date,
) -> dict[str, Any] | None:
    candidates: list[dict[str, Any]] = []
    for mapping in mappings:
        if (
            not mapping_id
            or _text(mapping.get("mapping_id")) != mapping_id
            or _text(mapping.get("agency_id")) != agency_id
            or _text(mapping.get("agency_item_id")) != item_id
        ):
            continue
        overlaps = _row_overlaps(
            mapping, "effective_from", "effective_until", period_start, period_end
        )
        if overlaps:
            candidates.append(mapping)
    return candidates[0] if len(candidates) == 1 else None


def _condition_signature(mapping: dict[str, Any]) -> ConditionSignature:
    effective_from = _parse_date(mapping.get("effective_from"))
    effective_until = _parse_date(mapping.get("effective_until"))
    attributes = tuple(
        (key, _text(mapping.get(key)))
        for key in _CONDITION_ATTRIBUTES
        if _has_value(mapping.get(key))
    )
    return ConditionSignature(
        item_key=_text(mapping.get("agency_item_id")),
        attributes=attributes,
        effective_period=(
            f"{effective_from.isoformat()}/{effective_until.isoformat()}"
            if effective_from is not None and effective_until is not None
            else ""
        ),
        period_rule=_text(mapping.get("period_rule")),
        payment_timing=_text(mapping.get("payment_timing")),
        payout_trigger=_text(mapping.get("trigger")),
        exclusions=_split_terms(mapping.get("exclusions")),
        reversal_terms=_split_terms(mapping.get("reversal_terms")),
    )


def _resolve_exceptions(
    agency_id: str,
    item_id: str,
    period_start: date,
    period_end: date,
    exception_rows: list[dict[str, Any]],
    approvals: list[dict[str, Any]],
    default_policy: DisplayPolicy,
) -> tuple[DisplayPolicy, list[str], dict[str, bool]]:
    reasons: list[str] = []
    flags = {"payment_prohibited": False, "mandatory_disclosure": False}
    matching = []
    for row in exception_rows:
        if _text(row.get("agency_id")) != agency_id:
            continue
        scope = _text(row.get("condition_item_id_or_rule_group"))
        if scope not in {item_id, "*"}:
            continue
        overlaps = _row_overlaps(row, "valid_from", "valid_until", period_start, period_end)
        if overlaps is None:
            reasons.append(
                f"例外 { _text(row.get('exception_id')) or '(ID missing)' } の適用期間が不明です。"
            )
        elif overlaps:
            matching.append(row)

    policies: set[DisplayPolicy] = set()
    for row in matching:
        evidence, evidence_reasons = _approval_evidence("代理店別例外", row, approvals)
        reasons.extend(evidence_reasons)
        if not evidence.is_complete_approval():
            continue
        rule_type = _text(row.get("rule_type")).upper()
        if rule_type not in _SUPPORTED_EXCEPTION_TYPES:
            reasons.append(
                f"例外 {_text(row.get('exception_id'))} の規則種別を解釈できません。"
            )
        elif rule_type == "DISPLAY_POLICY":
            value = _text(row.get("display_policy") or row.get("rule_value")).upper()
            policy = _parse_display_policy(value)
            if policy is None or policy == DisplayPolicy.UNDECIDED:
                reasons.append(
                    f"例外 {_text(row.get('exception_id'))} の掲載可否が未確定です。"
                )
            else:
                policies.add(policy)
        elif rule_type == "PAYMENT_PROHIBITED":
            flags["payment_prohibited"] = True
        elif rule_type == "MANDATORY_DISCLOSURE":
            flags["mandatory_disclosure"] = True
        elif rule_type == "HOLD":
            reasons.append(
                f"承認済み例外 {_text(row.get('exception_id'))}: "
                f"{_text(row.get('reason')) or '手動確認が必要です。'}"
            )
    if len(policies) > 1:
        reasons.append("適用例外の掲載可否が競合しています。")
    elif policies:
        default_policy = next(iter(policies))
    return default_policy, reasons, flags


def _read_amount(
    raw_amount: Any, raw_state: Any, label: str
) -> tuple[Decimal | None, list[str]]:
    state = _text(raw_state).upper()
    if state not in {"VALUE", "ZERO"}:
        return None, [f"{label}の状態が承認済み金額ではありません ({state or '空欄'})。"]
    if raw_amount is None or (isinstance(raw_amount, str) and not raw_amount.strip()):
        return None, [f"{label}の金額が空欄です。"]
    try:
        amount = Decimal(str(raw_amount).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None, [f"{label}を数値として解釈できません。"]
    if not amount.is_finite() or amount < 0:
        return None, [f"{label}は有限の0以上である必要があります。"]
    if state == "ZERO" and amount != 0:
        return None, [f"{label}の状態がZEROですが、金額が0ではありません。"]
    if state == "VALUE" and amount == 0:
        return None, [f"{label}が0円の場合、状態をZEROとして明示してください。"]
    return amount, []


def _optional_amount(raw_amount: Any, label: str) -> tuple[Decimal | None, list[str]]:
    if raw_amount is None or (isinstance(raw_amount, str) and not raw_amount.strip()):
        return None, []
    try:
        amount = Decimal(str(raw_amount).replace(",", "").strip())
    except (InvalidOperation, ValueError):
        return None, [f"{label}を数値として解釈できません。"]
    if not amount.is_finite() or amount < 0:
        return None, [f"{label}は有限の0以上である必要があります。"]
    return amount, []


def _combine_approvals(evidences: list[ApprovalEvidence]) -> ApprovalEvidence:
    if not evidences or not all(item.is_complete_approval() for item in evidences):
        return ApprovalEvidence()
    return ApprovalEvidence(
        status=ApprovalStatus.APPROVED,
        approval_id="composite:" + ",".join(item.approval_id or "" for item in evidences),
        source_document=";".join(item.source_document or "" for item in evidences),
        evidence_location=";".join(item.evidence_location or "" for item in evidences),
        approved_by=";".join(item.approved_by or "" for item in evidences),
        approved_at=max(item.approved_at for item in evidences if item.approved_at is not None),
    )


def _row_overlaps(
    row: dict[str, Any],
    start_key: str,
    end_key: str,
    period_start: date,
    period_end: date,
) -> bool | None:
    start = _parse_date(row.get(start_key))
    end = _parse_date(row.get(end_key))
    if start is None or end is None or start > end:
        return None
    return start <= period_end and end >= period_start


def _parse_target_period(value: str) -> tuple[date, date]:
    try:
        parsed = datetime.strptime(value, "%Y-%m")
    except ValueError as error:
        raise ValueError("target_period must use YYYY-MM format") from error
    if parsed.strftime("%Y-%m") != value:
        raise ValueError("target_period must use YYYY-MM format")
    start = date(parsed.year, parsed.month, 1)
    if parsed.month == 12:
        next_month = date(parsed.year + 1, 1, 1)
    else:
        next_month = date(parsed.year, parsed.month + 1, 1)
    return start, date.fromordinal(next_month.toordinal() - 1)


def _previous_month(month_start: date) -> tuple[date, date]:
    previous_end = date.fromordinal(month_start.toordinal() - 1)
    previous_start = date(previous_end.year, previous_end.month, 1)
    return previous_start, previous_end


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str) and value.strip():
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        return None
    return parsed


def _parse_display_policy(value: str) -> DisplayPolicy | None:
    try:
        return DisplayPolicy(value.lower())
    except ValueError:
        return None


def _display_policy(value: Any) -> DisplayPolicy:
    return _parse_display_policy(_text(value)) or DisplayPolicy.UNDECIDED


def _tax_basis(value: Any) -> TaxBasis:
    try:
        return TaxBasis(_text(value).lower())
    except ValueError:
        return TaxBasis.UNKNOWN


def _split_terms(value: Any) -> tuple[str, ...]:
    if not _has_value(value):
        return ()
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    terms = [part.strip() for line in text.split("\n") for part in line.split(";")]
    return tuple(part for part in terms if part)


def _canonical_value(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value.normalize())
    if isinstance(value, float):
        if not Decimal(str(value)).is_finite():
            return str(value)
        return str(Decimal(str(value)).normalize())
    if isinstance(value, str):
        return value.strip() or None
    return value


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _has_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _held_result(
    target_period: str,
    agency_id: str,
    item_id: str,
    reasons: tuple[str, ...],
    *,
    display_policy: str = "undecided",
) -> P2AgencyResult:
    return P2AgencyResult(
        target_period=target_period,
        agency_id=agency_id,
        condition_item_id=item_id,
        status="hold",
        upstream_reward=None,
        direct_cost=None,
        approved_gross_margin=None,
        previous_agency_reward=None,
        proposed_agency_reward=None,
        proposed_difference=None,
        display_policy=display_policy,
        hold_reasons=reasons,
    )


def _excel_amount(value: Decimal | None) -> int | float | None:
    if value is None:
        return None
    if value == value.to_integral_value():
        return int(value)
    return float(value)
