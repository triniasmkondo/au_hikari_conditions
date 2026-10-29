from __future__ import annotations

import calendar
import hashlib
import json
import re
import shutil
import string
import os
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .analysis import build_p1_report
from .monthly_bridge import MonthlyBridgeError, validate_monthly_comparison_workbook
from .p2 import calculate_workbook, write_report
from .p3 import (
    GeneratedAgencyArtifacts,
    P3GenerationError,
    distribution_candidate_paths,
    generate_distribution_candidates,
    validate_distribution_candidate,
)
from .p4b import (
    AgencyMailSettings,
    DraftHistoryStore,
    DraftPreview,
    _file_hash,
    _message_hash,
    _recipient_snapshot_hash,
    _render_message,
    prepare_draft_preview,
    save_draft,
)
from .workflow import AgencyWorkflow, BatchStatus, DeliverySnapshot
from .zoho_mail import ZohoMailClient


_REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
_PERIOD_PATTERN = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")
_SAFE_ID_PATTERN = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]{0,127}$")
_RECIPIENT_HEADERS = {
    "agency_id",
    "agency_name",
    "recipient_status",
    "to_addresses",
    "cc_addresses",
    "bcc_addresses",
    "sender_address",
    "reply_to_address",
    "salutation",
    "approved_by",
    "approved_at",
    "valid_from",
    "valid_until",
    "last_reviewed_at",
}
_TEMPLATE_HEADERS = {
    "template_id",
    "template_version",
    "doc_type",
    "subject_template",
    "body_template",
    "allowed_placeholders",
    "template_hash",
    "approval_status",
    "approved_by",
    "approved_at",
    "valid_from",
    "valid_until",
}
_APPROVAL_HEADERS = {
    "batch_id",
    "target_period",
    "agency_id",
    "approval_status",
    "artifact_hashes",
    "recipient_snapshot_hash",
    "message_hash",
    "template_hash",
    "condition_master_hash",
    "qa_status",
    "requested_by",
    "requested_at",
    "approved_by",
    "approved_at",
    "approval_reason",
}
_TEMPLATE_FIELDS = (
    "template_id",
    "template_version",
    "doc_type",
    "subject_template",
    "body_template",
    "allowed_placeholders",
)


class P5Error(ValueError):
    """Raised when an integrated monthly batch cannot pass a required gate."""


@dataclass(frozen=True)
class P5ApprovalRequest:
    agency_id: str = field(repr=False)
    candidate_artifacts: GeneratedAgencyArtifacts = field(repr=False)
    snapshot: DeliverySnapshot = field(repr=False)


@dataclass(frozen=True)
class P5Preparation:
    p1_status: str
    target_period: str
    agency_count: int
    approval_workbook_path: Path = field(repr=False)
    requests: tuple[P5ApprovalRequest, ...] = field(repr=False)


@dataclass(frozen=True)
class P5ApprovedDraft:
    artifacts: GeneratedAgencyArtifacts = field(repr=False)
    snapshot: DeliverySnapshot = field(repr=False)
    workflow: AgencyWorkflow = field(repr=False)
    preview: DraftPreview = field(repr=False)


@dataclass(frozen=True)
class _ValidatedApproval:
    candidate: GeneratedAgencyArtifacts
    settings: AgencyMailSettings
    snapshot: DeliverySnapshot
    workflow: AgencyWorkflow


def compute_mail_template_hash(template: dict[str, Any]) -> str:
    payload = {key: _canonical_value(template.get(key)) for key in _TEMPLATE_FIELDS}
    return _canonical_hash(payload)


def load_approved_mail_settings(
    workbook_path: Path, *, agency_id: str, target_period: str
) -> AgencyMailSettings:
    _validate_period_and_ids(target_period, agency_id)
    _require_external_path(workbook_path, "Distribution management workbook")
    workbook = _open_workbook(workbook_path)
    try:
        recipients = _read_sheet_records(
            workbook,
            "送付先マスター",
            _RECIPIENT_HEADERS,
            key="agency_id",
        )
        templates = _read_sheet_records(
            workbook,
            "本文テンプレート",
            _TEMPLATE_HEADERS,
            key="template_id",
        )
    finally:
        workbook.close()

    agency_matches = [
        row for row in recipients if _text(row.get("agency_id")) == agency_id
    ]
    if len(agency_matches) != 1:
        raise P5Error("An approved recipient master row is required.")
    recipient = agency_matches[0]
    if _text(recipient.get("reply_to_address")):
        raise P5Error("Reply-To is not supported by the approved Zoho draft adapter.")
    month_start, month_end = _month_range(target_period)
    _require_approved_record(
        recipient,
        status_field="recipient_status",
        date_start="valid_from",
        date_end="valid_until",
        month_start=month_start,
        month_end=month_end,
        review_date="last_reviewed_at",
    )

    template_matches = [
        row
        for row in templates
        if _text(row.get("doc_type")).casefold() == "conditions"
    ]
    if len(template_matches) != 1:
        raise P5Error("Exactly one approved conditions mail template is required.")
    template = template_matches[0]
    _require_approved_record(
        template,
        status_field="approval_status",
        date_start="valid_from",
        date_end="valid_until",
        month_start=month_start,
        month_end=month_end,
    )
    expected_hash = compute_mail_template_hash(template)
    if _text(template.get("template_hash")) != expected_hash:
        raise P5Error("Approved mail template content hash is invalid.")

    allowed = {
        value.strip()
        for value in _text(template.get("allowed_placeholders")).split(",")
        if value.strip()
    }
    used = _template_placeholders(
        _text(template.get("subject_template")),
        _text(template.get("body_template")),
    )
    supported = {"target_period", "agency_name", "salutation"}
    if not used.issubset(supported) or used != allowed:
        raise P5Error("Mail template placeholders do not match the approved allowlist.")

    return AgencyMailSettings(
        agency_id=agency_id,
        from_address=_text(recipient.get("sender_address")),
        to_addresses=_split_addresses(recipient.get("to_addresses"), required=True),
        cc_addresses=_split_addresses(recipient.get("cc_addresses")),
        bcc_addresses=_split_addresses(recipient.get("bcc_addresses")),
        agency_name=_text(recipient.get("agency_name")),
        salutation=_text(recipient.get("salutation")),
        subject_template=_text(template.get("subject_template")),
        body_template=_text(template.get("body_template")),
        template_hash=expected_hash,
    )


def prepare_monthly_batch(
    *,
    project_root: Path,
    approved_input_path: Path,
    p2_result_path: Path,
    monthly_comparison_path: Path,
    target_period: str,
    batch_id: str,
    delivery_workbook_path: Path,
    approval_workbook_path: Path,
    candidate_directory: Path,
    requested_by: str,
) -> P5Preparation:
    _validate_period_and_ids(target_period, batch_id)
    if not _SAFE_ID_PATTERN.fullmatch(requested_by):
        raise P5Error("Requester identity must use a safe ASCII identifier.")
    _require_external_path(approved_input_path, "Approved P2 input")
    _require_external_path(p2_result_path, "P2 result")
    _require_external_path(monthly_comparison_path, "Monthly comparison workbook")
    _require_external_path(delivery_workbook_path, "Distribution management workbook")
    _require_external_path(approval_workbook_path, "Approval workbook")
    _require_external_path(candidate_directory, "Approval candidate directory")
    if p2_result_path.exists() or approval_workbook_path.exists():
        raise P5Error("Refusing to overwrite an existing P2 result or approval workbook.")

    p1_report = build_p1_report(
        project_root,
        previous_period=_previous_period(target_period),
        target_period=target_period,
    )
    if p1_report.get("p1_gate") != "met":
        raise P5Error("P1 source analysis has blockers; downstream processing is stopped.")
    _validate_monthly_comparison(
        monthly_comparison_path, project_root, target_period
    )
    report = calculate_workbook(approved_input_path, target_period)
    write_report(report, p2_result_path)
    if not report.results or report.hold_count:
        raise P5Error(
            "P2 has held or empty results; review the P2 result and correct approvals."
        )
    for agency_id in sorted(
        {result.agency_id for result in report.results if result.status == "proposed"}
    ):
        load_approved_mail_settings(
            delivery_workbook_path,
            agency_id=agency_id,
            target_period=target_period,
        )
    p3_report = generate_distribution_candidates(
        approved_input_path=approved_input_path,
        p2_result_path=p2_result_path,
        target_period=target_period,
        output_directory=candidate_directory,
    )
    if p3_report.status != "APPROVAL_CANDIDATES_NOT_APPROVED_FOR_DELIVERY":
        raise P5Error("P3 did not produce approval-only candidates.")
    source_hash = _monthly_source_snapshot_hash(
        approved_input_path, monthly_comparison_path
    )
    requests: list[P5ApprovalRequest] = []
    for artifacts in p3_report.artifacts:
        settings = load_approved_mail_settings(
            delivery_workbook_path,
            agency_id=artifacts.agency_id,
            target_period=target_period,
        )
        if settings.agency_name != artifacts.agency_name:
            raise P5Error("Approved recipient name does not match P2 agency data.")
        subject, body = _render_message(settings, artifacts.agency_name, target_period)
        snapshot = DeliverySnapshot(
            batch_id=batch_id,
            agency_id=artifacts.agency_id,
            target_period=target_period,
            source_master_hash=source_hash,
            template_hash=settings.template_hash or "",
            artifact_hashes=(
                _file_hash(artifacts.spreadsheet_path),
                _file_hash(artifacts.word_path),
            ),
            recipient_snapshot_hash=_recipient_snapshot_hash(settings),
            message_hash=_message_hash(settings.from_address, subject, body),
        )
        snapshot.fingerprint()
        requests.append(
            P5ApprovalRequest(
                agency_id=artifacts.agency_id,
                candidate_artifacts=artifacts,
                snapshot=snapshot,
            )
        )
    if not requests:
        raise P5Error("P3 did not produce any agency approval candidates.")
    _write_approval_workbook(
        source=delivery_workbook_path,
        destination=approval_workbook_path,
        requests=requests,
        requested_by=requested_by,
        target_period=target_period,
    )
    return P5Preparation(
        p1_status="met",
        target_period=target_period,
        agency_count=len(requests),
        approval_workbook_path=approval_workbook_path.resolve(),
        requests=tuple(requests),
    )


def finalize_approved_monthly_batch(
    *,
    project_root: Path,
    approved_input_path: Path,
    p2_result_path: Path,
    monthly_comparison_path: Path,
    target_period: str,
    batch_id: str,
    delivery_workbook_path: Path,
    approval_workbook_path: Path,
    candidate_directory: Path,
    final_directory: Path,
) -> tuple[P5ApprovedDraft, ...]:
    _validate_period_and_ids(target_period, batch_id)
    for path, label in (
        (approved_input_path, "Approved P2 input"),
        (p2_result_path, "P2 result"),
        (monthly_comparison_path, "Monthly comparison workbook"),
        (delivery_workbook_path, "Distribution management workbook"),
        (approval_workbook_path, "Approval workbook"),
        (candidate_directory, "Approval candidate directory"),
        (final_directory, "Final distribution directory"),
    ):
        _require_external_path(path, label)
    p1_report = build_p1_report(
        project_root,
        previous_period=_previous_period(target_period),
        target_period=target_period,
    )
    if p1_report.get("p1_gate") != "met":
        raise P5Error("P1 source analysis has blockers; finalization is stopped.")
    _validate_monthly_comparison(
        monthly_comparison_path, project_root, target_period
    )
    p2_report = calculate_workbook(approved_input_path, target_period)
    if not p2_report.results or p2_report.hold_count:
        raise P5Error("P2 has held or empty results; finalization is stopped.")

    workbook = _open_workbook(approval_workbook_path)
    try:
        approvals = _read_sheet_records(
            workbook, "月次承認", _APPROVAL_HEADERS, key=None
        )
        all_recipients = _read_sheet_records(
            workbook,
            "送付先マスター",
            _RECIPIENT_HEADERS,
            key="agency_id",
        )
        all_templates = _read_sheet_records(
            workbook,
            "本文テンプレート",
            _TEMPLATE_HEADERS,
            key="template_id",
        )
    finally:
        workbook.close()
    rows = [
        row
        for row in approvals
        if _text(row.get("batch_id")) == batch_id
        and _text(row.get("target_period")) == target_period
    ]
    if not rows:
        raise P5Error("No matching monthly approval records were found.")
    row_agencies = [_text(row.get("agency_id")) for row in rows]
    if any(not agency for agency in row_agencies) or len(row_agencies) != len(
        set(row_agencies)
    ):
        raise P5Error("Monthly approvals contain duplicate or missing agencies.")
    expected_agencies = {
        result.agency_id for result in p2_report.results if result.status == "proposed"
    }
    if set(row_agencies) != expected_agencies:
        raise P5Error("Monthly approvals do not cover exactly the current P2 agencies.")

    validated_approvals: list[_ValidatedApproval] = []
    source_hash = _monthly_source_snapshot_hash(
        approved_input_path, monthly_comparison_path
    )
    for approval in rows:
        agency_id = _text(approval.get("agency_id"))
        _validate_period_and_ids(target_period, agency_id)
        if (
            _text(approval.get("approval_status")) != "APPROVED"
            or _text(approval.get("qa_status")) != "PASSED"
            or not _text(approval.get("approved_by"))
            or _parse_approval_datetime(approval.get("approved_at")) is None
            or not _text(approval.get("approval_reason"))
        ):
            raise P5Error("Every matching agency requires completed QA and approval.")
        settings = _settings_from_records(
            recipients=all_recipients,
            templates=all_templates,
            agency_id=agency_id,
            target_period=target_period,
        )
        current_settings = load_approved_mail_settings(
            delivery_workbook_path,
            agency_id=agency_id,
            target_period=target_period,
        )
        if (
            settings.agency_name != current_settings.agency_name
            or settings.template_hash != current_settings.template_hash
            or _recipient_snapshot_hash(settings)
            != _recipient_snapshot_hash(current_settings)
        ):
            raise P5Error("Recipient or template master changed after approval request.")
        candidate_xlsx, candidate_docx = distribution_candidate_paths(
            agency_id=agency_id,
            target_period=target_period,
            output_directory=candidate_directory,
        )
        candidate = GeneratedAgencyArtifacts(
            agency_id=agency_id,
            agency_name=_agency_name(all_recipients, agency_id),
            spreadsheet_path=candidate_xlsx,
            word_path=candidate_docx,
            condition_count=_condition_count(p2_report.results, agency_id),
        )
        validate_distribution_candidate(
            candidate=candidate,
            approved_input_path=approved_input_path,
            p2_result_path=p2_result_path,
            target_period=target_period,
        )
        subject, body = _render_message(settings, candidate.agency_name, target_period)
        current_subject, current_body = _render_message(
            current_settings,
            candidate.agency_name,
            target_period,
        )
        if (
            _message_hash(settings.from_address, subject, body)
            != _message_hash(
                current_settings.from_address,
                current_subject,
                current_body,
            )
        ):
            raise P5Error("Rendered mail content changed after approval request.")
        expected_hashes = tuple(
            sorted((_file_hash(candidate_xlsx), _file_hash(candidate_docx)))
        )
        recorded_hashes = tuple(
            sorted(
                value.strip()
                for value in _text(approval.get("artifact_hashes")).split(",")
                if value.strip()
            )
        )
        snapshot = DeliverySnapshot(
            batch_id=batch_id,
            agency_id=agency_id,
            target_period=target_period,
            source_master_hash=source_hash,
            template_hash=settings.template_hash or "",
            artifact_hashes=expected_hashes,
            recipient_snapshot_hash=_recipient_snapshot_hash(settings),
            message_hash=_message_hash(settings.from_address, subject, body),
        )
        if (
            recorded_hashes != expected_hashes
            or _text(approval.get("recipient_snapshot_hash"))
            != snapshot.recipient_snapshot_hash
            or _text(approval.get("message_hash")) != snapshot.message_hash
            or _text(approval.get("template_hash")) != snapshot.template_hash
            or _text(approval.get("condition_master_hash")) != snapshot.source_master_hash
        ):
            raise P5Error("Approval record no longer matches the current batch snapshot.")
        workflow = _approved_workflow(
            snapshot=snapshot,
            approver_id=_text(approval.get("approved_by")),
            approved_at=_parse_approval_datetime(approval.get("approved_at")),
        )
        validated_approvals.append(
            _ValidatedApproval(
                candidate=candidate,
                settings=settings,
                snapshot=snapshot,
                workflow=workflow,
            )
        )

    finalized_artifacts: list[tuple[_ValidatedApproval, GeneratedAgencyArtifacts]] = []
    if final_directory.exists():
        for approved in validated_approvals:
            xlsx_target, docx_target = _final_artifact_paths(
                approved.candidate, final_directory
            )
            if not xlsx_target.exists() or not docx_target.exists():
                raise P5Error(
                    "Final output directory contains an incomplete batch; refusing partial recovery."
                )
            finalized = _promote_candidate(
                candidate=approved.candidate,
                target_period=target_period,
                snapshot=approved.snapshot,
                workflow=approved.workflow,
                approved_input_path=approved_input_path,
                p2_result_path=p2_result_path,
                final_directory=final_directory,
            )
            finalized_artifacts.append((approved, finalized))
    else:
        final_directory.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(
            prefix=".p5-finalization-",
            dir=final_directory.parent,
        ) as staging_name:
            staging_directory = Path(staging_name)
            for approved in validated_approvals:
                staged = _promote_candidate(
                    candidate=approved.candidate,
                    target_period=target_period,
                    snapshot=approved.snapshot,
                    workflow=approved.workflow,
                    approved_input_path=approved_input_path,
                    p2_result_path=p2_result_path,
                    final_directory=staging_directory,
                )
                finalized_artifacts.append((approved, staged))
            try:
                os.replace(staging_directory, final_directory)
            except OSError:
                raise P5Error(
                    "Could not atomically publish the approved artifact batch."
                ) from None
            finalized_artifacts = [
                (
                    approved,
                    GeneratedAgencyArtifacts(
                        agency_id=staged.agency_id,
                        agency_name=staged.agency_name,
                        spreadsheet_path=final_directory / staged.spreadsheet_path.name,
                        word_path=final_directory / staged.word_path.name,
                        condition_count=staged.condition_count,
                    ),
                )
                for approved, staged in finalized_artifacts
            ]

    output: list[P5ApprovedDraft] = []
    for approved, finalized in finalized_artifacts:
        preview = prepare_draft_preview(
            artifacts=finalized,
            target_period=target_period,
            settings=approved.settings,
            workflow=approved.workflow,
            snapshot=approved.snapshot,
        )
        output.append(
            P5ApprovedDraft(
                artifacts=finalized,
                snapshot=approved.snapshot,
                workflow=approved.workflow,
                preview=preview,
            )
        )
    return tuple(output)


def create_approved_zoho_drafts(
    *,
    packages: tuple[P5ApprovedDraft, ...],
    mail_client: ZohoMailClient,
    history: DraftHistoryStore,
    company_authorized: bool,
) -> tuple[str, ...]:
    if company_authorized is not True:
        raise P5Error("Company authorization is required before Zoho draft creation.")
    if not packages:
        raise P5Error("At least one approved agency package is required.")
    return tuple(
        save_draft(
            preview=package.preview,
            mail_client=mail_client,
            history=history,
            mode="draft",
        ).status
        for package in packages
    )


def _promote_candidate(
    *,
    candidate: GeneratedAgencyArtifacts,
    target_period: str,
    snapshot: DeliverySnapshot,
    workflow: AgencyWorkflow,
    approved_input_path: Path,
    p2_result_path: Path,
    final_directory: Path,
) -> GeneratedAgencyArtifacts:
    if (
        workflow.status != BatchStatus.APPROVED
        or workflow.approval_fingerprint != snapshot.fingerprint()
        or snapshot.agency_id != candidate.agency_id
        or snapshot.target_period != target_period
    ):
        raise P5Error("Approved workflow does not match the candidate artifact snapshot.")
    if (
        candidate.spreadsheet_path.is_symlink()
        or candidate.word_path.is_symlink()
    ):
        raise P5Error("Approval candidates must be regular files, not symbolic links.")
    validate_distribution_candidate(
        candidate=candidate,
        approved_input_path=approved_input_path,
        p2_result_path=p2_result_path,
        target_period=target_period,
    )
    final_directory.mkdir(parents=True, exist_ok=True)
    xlsx_target, docx_target = _final_artifact_paths(candidate, final_directory)
    if xlsx_target.exists() or docx_target.exists():
        raise P5Error("Refusing to overwrite existing finalized condition sheets.")
    expected = {
        xlsx_target: _file_hash(candidate.spreadsheet_path),
        docx_target: _file_hash(candidate.word_path),
    }
    if tuple(sorted(expected.values())) != tuple(sorted(snapshot.artifact_hashes)):
        raise P5Error("Candidate files do not match approved artifact hashes.")
    if xlsx_target.exists() or docx_target.exists():
        if not xlsx_target.exists() or not docx_target.exists():
            raise P5Error("A partial finalized artifact pair already exists.")
        if any(_file_hash(path) != digest for path, digest in expected.items()):
            raise P5Error("Existing finalized artifacts do not match the approved snapshot.")
        return GeneratedAgencyArtifacts(
            agency_id=candidate.agency_id,
            agency_name=candidate.agency_name,
            spreadsheet_path=xlsx_target,
            word_path=docx_target,
            condition_count=candidate.condition_count,
        )
    created: list[Path] = []
    try:
        for source, target in (
            (candidate.spreadsheet_path, xlsx_target),
            (candidate.word_path, docx_target),
        ):
            with target.open("xb") as target_stream:
                created.append(target)
                with source.open("rb") as source_stream:
                    shutil.copyfileobj(source_stream, target_stream)
            if _file_hash(target) != expected[target]:
                raise P5Error("Finalized artifact hash does not match approved snapshot.")
    except Exception:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return GeneratedAgencyArtifacts(
        agency_id=candidate.agency_id,
        agency_name=candidate.agency_name,
        spreadsheet_path=xlsx_target,
        word_path=docx_target,
        condition_count=candidate.condition_count,
    )


def _final_artifact_paths(
    candidate: GeneratedAgencyArtifacts, final_directory: Path
) -> tuple[Path, Path]:
    xlsx_name = candidate.spreadsheet_path.name
    docx_name = candidate.word_path.name
    if not xlsx_name.endswith("_approval_candidate.xlsx") or not docx_name.endswith(
        "_approval_candidate.docx"
    ):
        raise P5Error("P3 artifact is not an approval candidate.")
    return (
        final_directory / xlsx_name.replace(
            "_approval_candidate.xlsx", "_conditions.xlsx"
        ),
        final_directory / docx_name.replace(
            "_approval_candidate.docx", "_conditions.docx"
        ),
    )


def _approved_workflow(
    *,
    snapshot: DeliverySnapshot,
    approver_id: str,
    approved_at: datetime | None,
) -> AgencyWorkflow:
    if not _SAFE_ID_PATTERN.fullmatch(approver_id) or approved_at is None:
        raise P5Error("A valid approver identity and timestamp are required.")
    workflow = AgencyWorkflow(snapshot.batch_id, snapshot.agency_id)
    for status in (
        BatchStatus.PARSED,
        BatchStatus.VALIDATED,
        BatchStatus.PROPOSED,
        BatchStatus.GENERATED,
        BatchStatus.QA_PASSED,
        BatchStatus.PENDING_APPROVAL,
        BatchStatus.APPROVED,
    ):
        workflow.transition(
            status,
            actor_id=approver_id if status == BatchStatus.APPROVED else "p5-system",
            snapshot=snapshot if status == BatchStatus.APPROVED else None,
            occurred_at=approved_at if status == BatchStatus.APPROVED else None,
        )
    return workflow


def _write_approval_workbook(
    *,
    source: Path,
    destination: Path,
    requests: list[P5ApprovalRequest],
    requested_by: str,
    target_period: str,
) -> None:
    shutil.copy2(source, destination)
    workbook = _open_workbook(destination, data_only=False)
    try:
        if "月次承認" not in workbook.sheetnames:
            raise P5Error("Distribution workbook has no monthly approval sheet.")
        worksheet = workbook["月次承認"]
        headers = _worksheet_headers(worksheet)
        missing = _APPROVAL_HEADERS - set(headers)
        if missing:
            raise P5Error("Monthly approval sheet has an incomplete schema.")
        existing = {
            (_text(row.get("batch_id")), _text(row.get("agency_id")))
            for row in _worksheet_records(worksheet, headers)
        }
        now = datetime.now(timezone.utc).isoformat()
        for request in requests:
            snapshot = request.snapshot
            if (snapshot.batch_id, snapshot.agency_id) in existing:
                raise P5Error("An approval request already exists for this batch and agency.")
            values = {
                "batch_id": snapshot.batch_id,
                "target_period": target_period,
                "agency_id": snapshot.agency_id,
                "approval_status": "PENDING",
                "artifact_hashes": ",".join(sorted(snapshot.artifact_hashes)),
                "recipient_snapshot_hash": snapshot.recipient_snapshot_hash,
                "message_hash": snapshot.message_hash,
                "template_hash": snapshot.template_hash,
                "condition_master_hash": snapshot.source_master_hash,
                "qa_status": "PENDING",
                "requested_by": requested_by,
                "requested_at": now,
                "approved_by": None,
                "approved_at": None,
                "approval_reason": None,
            }
            worksheet.append([values.get(header) for header in headers])
        table_ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
        for table in worksheet.tables.values():
            table.ref = table_ref
            if table.autoFilter is not None:
                table.autoFilter.ref = table_ref
        workbook.save(destination)
    except Exception:
        workbook.close()
        destination.unlink(missing_ok=True)
        raise
    workbook.close()


def _settings_from_records(
    *,
    recipients: list[dict[str, Any]],
    templates: list[dict[str, Any]],
    agency_id: str,
    target_period: str,
) -> AgencyMailSettings:
    matching = [row for row in recipients if _text(row.get("agency_id")) == agency_id]
    template_rows = [
        row
        for row in templates
        if _text(row.get("doc_type")).casefold() == "conditions"
    ]
    if len(matching) != 1 or len(template_rows) != 1:
        raise P5Error("Current recipient or mail template is ambiguous.")
    recipient = matching[0]
    template = template_rows[0]
    if _text(recipient.get("reply_to_address")):
        raise P5Error("Reply-To is not supported by the approved Zoho draft adapter.")
    month_start, month_end = _month_range(target_period)
    _require_approved_record(
        recipient,
        status_field="recipient_status",
        date_start="valid_from",
        date_end="valid_until",
        month_start=month_start,
        month_end=month_end,
        review_date="last_reviewed_at",
    )
    _require_approved_record(
        template,
        status_field="approval_status",
        date_start="valid_from",
        date_end="valid_until",
        month_start=month_start,
        month_end=month_end,
    )
    if _text(template.get("template_hash")) != compute_mail_template_hash(template):
        raise P5Error("Approved mail template content hash is invalid.")
    allowed = {
        value.strip()
        for value in _text(template.get("allowed_placeholders")).split(",")
        if value.strip()
    }
    used = _template_placeholders(
        _text(template.get("subject_template")),
        _text(template.get("body_template")),
    )
    if used != allowed or not used.issubset(
        {"target_period", "agency_name", "salutation"}
    ):
        raise P5Error("Mail template placeholders do not match the approved allowlist.")
    return AgencyMailSettings(
        agency_id=agency_id,
        from_address=_text(recipient.get("sender_address")),
        to_addresses=_split_addresses(recipient.get("to_addresses"), required=True),
        cc_addresses=_split_addresses(recipient.get("cc_addresses")),
        bcc_addresses=_split_addresses(recipient.get("bcc_addresses")),
        agency_name=_text(recipient.get("agency_name")),
        salutation=_text(recipient.get("salutation")),
        subject_template=_text(template.get("subject_template")),
        body_template=_text(template.get("body_template")),
        template_hash=compute_mail_template_hash(template),
    )


def _require_approved_record(
    record: dict[str, Any],
    *,
    status_field: str,
    date_start: str,
    date_end: str,
    month_start: date,
    month_end: date,
    review_date: str | None = None,
) -> None:
    if (
        _text(record.get(status_field)) != "APPROVED"
        or not _text(record.get("approved_by"))
        or not _parse_datetime(record.get("approved_at"))
    ):
        raise P5Error("Recipient or mail template is not approved.")
    valid_from = _parse_date(record.get(date_start))
    valid_until = _parse_date(record.get(date_end))
    if valid_from is None or valid_until is None:
        raise P5Error("Recipient or mail template validity dates are incomplete.")
    if valid_from > month_start or valid_until < month_end:
        raise P5Error("Recipient or mail template is not valid for the full target month.")
    if review_date and not _parse_datetime(record.get(review_date)):
        raise P5Error("Recipient master has not been reviewed.")


def _read_sheet_records(
    workbook, sheet_name: str, required_headers: set[str], *, key: str | None
) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise P5Error(f"Distribution workbook is missing sheet: {sheet_name}.")
    worksheet = workbook[sheet_name]
    headers = _worksheet_headers(worksheet)
    if not required_headers.issubset(headers):
        raise P5Error(f"Distribution workbook has an incomplete {sheet_name} schema.")
    records = _worksheet_records(worksheet, headers)
    if key is not None:
        keys = [_text(record.get(key)) for record in records if _text(record.get(key))]
        if len(keys) != len(set(keys)):
            raise P5Error(f"Distribution workbook has duplicate rows in {sheet_name}.")
    return records


def _worksheet_headers(worksheet) -> list[str]:
    return [
        _text(value)
        for value in next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    ]


def _worksheet_records(worksheet, headers: list[str]) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        values = dict(zip(headers, row, strict=False))
        if any(value is not None for value in values.values()):
            records.append(values)
    return records


def _open_workbook(path: Path, *, data_only: bool = True):
    try:
        return load_workbook(path, read_only=False, data_only=data_only)
    except (OSError, ValueError):
        raise P5Error("Distribution workbook is missing, invalid, or unreadable.") from None


def _template_placeholders(subject: str, body: str) -> set[str]:
    placeholders: set[str] = set()
    try:
        for template in (subject, body):
            for _, field_name, format_spec, conversion in string.Formatter().parse(template):
                if field_name is None:
                    continue
                if format_spec or conversion or not re.fullmatch(r"[a-z_]+", field_name):
                    raise P5Error("Mail template contains an unsupported placeholder.")
                placeholders.add(field_name)
    except ValueError:
        raise P5Error("Mail template syntax is invalid.") from None
    return placeholders


def _split_addresses(value: Any, *, required: bool = False) -> tuple[str, ...]:
    addresses = tuple(
        item.strip()
        for item in re.split(r"[,;\n]", _text(value))
        if item.strip()
    )
    if required and not addresses:
        raise P5Error("An approved To recipient is required.")
    return addresses


def _agency_name(recipients: list[dict[str, Any]], agency_id: str) -> str:
    names = {
        _text(row.get("agency_name"))
        for row in recipients
        if _text(row.get("agency_id")) == agency_id
    }
    if len(names) != 1 or not next(iter(names)):
        raise P5Error("Approved recipient master has no unique agency name.")
    return next(iter(names))


def _condition_count(results, agency_id: str) -> int:
    return sum(result.agency_id == agency_id and result.status == "proposed" for result in results)


def _month_range(target_period: str) -> tuple[date, date]:
    if not _PERIOD_PATTERN.fullmatch(target_period):
        raise P5Error("Target period must use YYYY-MM format.")
    year, month = (int(value) for value in target_period.split("-"))
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


def _validate_monthly_comparison(
    path: Path, project_root: Path, target_period: str
) -> None:
    try:
        validate_monthly_comparison_workbook(
            path,
            project_root=project_root,
            previous_period=_previous_period(target_period),
            target_period=target_period,
        )
    except MonthlyBridgeError as error:
        raise P5Error(
            "Monthly comparison evidence is missing, incomplete, or mismatched."
        ) from error


def _previous_period(target_period: str) -> str:
    target_start, _ = _month_range(target_period)
    if target_start.month == 1:
        return f"{target_start.year - 1:04d}-12"
    return f"{target_start.year:04d}-{target_start.month - 1:02d}"


def _monthly_source_snapshot_hash(
    approved_input_path: Path, monthly_comparison_path: Path
) -> str:
    snapshot = {
        "approved_p2_input_sha256": _file_hash(approved_input_path),
        "monthly_comparison_sha256": _file_hash(monthly_comparison_path),
    }
    return hashlib.sha256(
        json.dumps(snapshot, sort_keys=True, separators=(",", ":")).encode("utf-8")
    ).hexdigest()


def _validate_period_and_ids(target_period: str, *identifiers: str) -> None:
    _month_range(target_period)
    if any(not _SAFE_ID_PATTERN.fullmatch(identifier) for identifier in identifiers):
        raise P5Error("Batch and agency identifiers must use safe ASCII identifiers.")


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def _parse_datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time(), timezone.utc)
    if isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
        except ValueError:
            return None
        if parsed.tzinfo is None:
            return parsed.replace(tzinfo=timezone.utc)
        return parsed
    return None


def _parse_approval_datetime(value: Any) -> datetime | None:
    if not isinstance(value, str):
        return None
    try:
        parsed = datetime.fromisoformat(value.strip().replace("Z", "+00:00"))
    except ValueError:
        return None
    return parsed if parsed.tzinfo is not None else None


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _canonical_value(value: Any) -> str:
    return _text(value)


def _canonical_hash(value: dict[str, str]) -> str:
    encoded = json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _require_external_path(path: Path, label: str) -> None:
    resolved = path.resolve()
    try:
        resolved.relative_to(_REPOSITORY_ROOT)
    except ValueError:
        return
    raise P5Error(f"{label} must be stored outside the Git repository.")
