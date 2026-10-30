from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from enum import Enum
from pathlib import Path

from openpyxl import load_workbook

from .calculation import (
    ApprovalEvidence,
    LadderDefinition,
    LadderMode,
    LadderScenario,
    LadderTier,
    simulate_ladder,
)
from .p2 import P2InputError, P2Report, calculate_workbook


class RewardCategory(str, Enum):
    BASIC = "basic"
    CAMPAIGN = "campaign"
    ISP = "isp"
    OPTION = "option"


FTTH_OPENING_FEE_COMPONENT = "ftth_opening_fee"


@dataclass(frozen=True)
class AgencyRewardProfile:
    agency_id: str
    target_period: str
    target_gross_margin: Decimal
    minimum_gross_margin: Decimal
    allow_rate_changes: bool
    pay_only_after_upstream_collection: bool | None
    collection_timing_approval: ApprovalEvidence = ApprovalEvidence()


@dataclass(frozen=True)
class RewardComponent:
    component_id: str
    p2_condition_item_id: str
    category: RewardCategory
    opening_class: str
    upstream_unit_reward: Decimal
    previous_agency_unit_reward: Decimal
    direct_cost_per_unit: Decimal = Decimal("0")
    payment_component: str = ""
    payment_allowed: bool = True
    target_gross_margin: Decimal | None = None
    minimum_gross_margin: Decimal | None = None
    upstream_approval: ApprovalEvidence = ApprovalEvidence()
    previous_approval: ApprovalEvidence = ApprovalEvidence()
    direct_cost_approval: ApprovalEvidence = ApprovalEvidence()


@dataclass(frozen=True)
class CampaignTierInput:
    threshold: int
    p2_condition_item_id: str
    upstream_unit_reward: Decimal
    previous_agency_unit_reward: Decimal
    direct_cost_per_unit: Decimal = Decimal("0")
    target_gross_margin: Decimal | None = None
    minimum_gross_margin: Decimal | None = None
    upstream_approval: ApprovalEvidence = ApprovalEvidence()
    previous_approval: ApprovalEvidence = ApprovalEvidence()
    direct_cost_approval: ApprovalEvidence = ApprovalEvidence()


@dataclass(frozen=True)
class CampaignSchedule:
    campaign_code: str
    counting_scope: str
    payout_scope: str
    eligible_opening_classes: tuple[str, ...]
    mode: LadderMode
    tiers: tuple[CampaignTierInput, ...]


@dataclass(frozen=True)
class VolumeScenario:
    scenario_id: str
    campaign_code: str
    commerce_count: int
    agency_eligible_count: int
    opening_class: str
    upstream_collected: bool
    agency_payment_due: bool


@dataclass(frozen=True)
class JudgmentReview:
    issue: str
    source_document: str
    ai_review_prompt: str
    resolved: bool = False


@dataclass(frozen=True)
class RewardRateProposal:
    component_id: str
    p2_condition_item_id: str
    category: RewardCategory
    opening_class: str
    upstream_unit_reward: Decimal
    previous_agency_unit_reward: Decimal
    proposed_agency_unit_reward: Decimal
    direct_cost_per_unit: Decimal
    retained_gross_margin_per_unit: Decimal
    payment_component: str
    payment_allowed: bool


@dataclass(frozen=True)
class CampaignTierProposal:
    campaign_code: str
    threshold: int
    p2_condition_item_id: str
    upstream_unit_reward: Decimal
    previous_agency_unit_reward: Decimal
    proposed_agency_unit_reward: Decimal
    direct_cost_per_unit: Decimal
    retained_gross_margin_per_unit: Decimal


@dataclass(frozen=True)
class ScenarioOutcome:
    scenario_id: str
    opening_class: str
    campaign_code: str | None
    commerce_count: int | None
    agency_eligible_count: int
    upstream_collected: bool
    agency_payment_due: bool
    qualified: bool
    upstream_amount: Decimal
    agency_payment: Decimal
    direct_cost: Decimal
    retained_gross_margin: Decimal
    minimum_margin_checked: bool
    risks: tuple[str, ...]


@dataclass(frozen=True)
class RewardDesignResult:
    status: str
    agency_id: str
    target_period: str
    component_proposals: tuple[RewardRateProposal, ...]
    campaign_proposals: tuple[CampaignTierProposal, ...]
    scenarios: tuple[ScenarioOutcome, ...]
    manual_judgments: tuple[JudgmentReview, ...]
    hold_reasons: tuple[str, ...]

    @property
    def total_absolute_change(self) -> Decimal:
        component_change = sum(
            (
                abs(
                    proposal.proposed_agency_unit_reward
                    - proposal.previous_agency_unit_reward
                )
                for proposal in self.component_proposals
            ),
            Decimal("0"),
        )
        campaign_change = sum(
            (
                abs(
                    proposal.proposed_agency_unit_reward
                    - proposal.previous_agency_unit_reward
                )
                for proposal in self.campaign_proposals
            ),
            Decimal("0"),
        )
        return component_change + campaign_change


@dataclass(frozen=True)
class P2RewardDesignGate:
    status: str
    requires_p2_reapproval: bool
    hold_reasons: tuple[str, ...]


@dataclass(frozen=True)
class P2CandidateWorkbookResult:
    output_path: Path
    updated_condition_item_ids: tuple[str, ...]
    post_write_p2_report: P2Report


def design_agency_rewards(
    profile: AgencyRewardProfile,
    components: tuple[RewardComponent, ...],
    campaigns: tuple[CampaignSchedule, ...] = (),
    manual_judgments: tuple[JudgmentReview, ...] = (),
    volume_scenarios: tuple[VolumeScenario, ...] = (),
) -> RewardDesignResult:
    """Propose least-change rates, simulate approved scenarios, and fail closed."""
    hold_reasons: list[str] = []
    proposals: list[RewardRateProposal] = []
    campaign_proposals: list[CampaignTierProposal] = []
    scenarios: list[ScenarioOutcome] = []

    hold_reasons.extend(_validate_profile(profile))
    if not components and not campaigns:
        hold_reasons.append("報酬条件がありません。")

    component_ids = [component.component_id for component in components]
    p2_item_ids = [
        component.p2_condition_item_id for component in components
    ]
    for campaign in campaigns:
        p2_item_ids.extend(tier.p2_condition_item_id for tier in campaign.tiers)
    if len(component_ids) != len(set(component_ids)):
        hold_reasons.append("報酬コンポーネントIDが重複しています。")
    if len(p2_item_ids) != len(set(p2_item_ids)):
        hold_reasons.append("P2成果項目IDが重複しています。")

    for judgment in manual_judgments:
        if not judgment.issue.strip() or not judgment.source_document.strip():
            hold_reasons.append("人手判断事項に内容または根拠資料がありません。")
        if not judgment.resolved:
            hold_reasons.append(
                f"商流・文言の判断を保留しています: {judgment.issue}"
            )

    collection_ready = profile.pay_only_after_upstream_collection is not None
    if not collection_ready:
        hold_reasons.append(
            "上位未収時の代理店支払条件が未定義です。契約文言の確認が必要です。"
        )
    elif not profile.collection_timing_approval.is_complete_approval():
        hold_reasons.append(
            "上位未収時の支払条件に根拠・承認者・承認日時がありません。"
        )
    elif not profile.pay_only_after_upstream_collection:
        hold_reasons.append(
            "上位未収でも代理店支払が発生する条件です。回収不能シナリオの承認が必要です。"
        )

    for component in components:
        input_reasons = _validate_rate_input(
            component.component_id,
            component.p2_condition_item_id,
            component.upstream_unit_reward,
            component.previous_agency_unit_reward,
            component.direct_cost_per_unit,
            component.target_gross_margin,
            component.minimum_gross_margin,
            component.upstream_approval,
            component.previous_approval,
            component.direct_cost_approval,
        )
        hold_reasons.extend(input_reasons)
        if input_reasons or not all(
            _valid_amount(value)
            for value in (
                component.upstream_unit_reward,
                component.previous_agency_unit_reward,
                component.direct_cost_per_unit,
            )
        ):
            continue
        if not component.opening_class.strip():
            hold_reasons.append(
                f"{component.component_id}: 開通区分がありません。"
            )
        payment_allowed = component.payment_allowed
        if component.payment_component == FTTH_OPENING_FEE_COMPONENT:
            if component.opening_class.casefold() != "ftth":
                hold_reasons.append(
                    f"{component.component_id}: FTTH開通手数料の開通区分がFTTHではありません。"
                )
                continue
            if component.upstream_unit_reward != Decimal("2000"):
                hold_reasons.append(
                    f"{component.component_id}: 承認済み上位単価がFTTH開通手数料2,000円と一致しません。"
                )
                continue
            payment_allowed = False

        margin_target = _effective_margin_target(
            profile,
            component.target_gross_margin,
            component.minimum_gross_margin,
        )
        if margin_target is None:
            continue
        maximum_rate = (
            component.upstream_unit_reward
            - component.direct_cost_per_unit
            - margin_target
        )
        if not payment_allowed:
            maximum_rate = Decimal("0")
        if maximum_rate < 0:
            hold_reasons.append(
                f"{component.component_id}: 上位単価から原価と目標粗利を確保できません。"
            )
            continue
        rate = _least_change_rate(
            component.previous_agency_unit_reward,
            maximum_rate,
            profile.allow_rate_changes,
            f"{component.component_id}",
            hold_reasons,
        )
        if rate is None:
            continue
        proposals.append(
            RewardRateProposal(
                component_id=component.component_id,
                p2_condition_item_id=component.p2_condition_item_id,
                category=component.category,
                opening_class=component.opening_class,
                upstream_unit_reward=component.upstream_unit_reward,
                previous_agency_unit_reward=component.previous_agency_unit_reward,
                proposed_agency_unit_reward=rate,
                direct_cost_per_unit=component.direct_cost_per_unit,
                retained_gross_margin_per_unit=(
                    component.upstream_unit_reward
                    - rate
                    - component.direct_cost_per_unit
                ),
                payment_component=component.payment_component,
                payment_allowed=payment_allowed,
            )
        )
        scenarios.extend(
            _component_scenarios(
                profile,
                proposals[-1],
                margin_target,
            )
        )

    for campaign in campaigns:
        schedule_reasons = _validate_campaign(campaign)
        hold_reasons.extend(schedule_reasons)
        if schedule_reasons:
            continue
        if campaign.mode == LadderMode.UNKNOWN:
            hold_reasons.append(
                f"{campaign.campaign_code}: 閾値の適用方式が未確定です。"
            )
        if not campaign.eligible_opening_classes:
            hold_reasons.append(
                f"{campaign.campaign_code}: 支払対象となる開通区分がありません。"
            )

        tier_caps: list[Decimal] = []
        tier_targets: list[Decimal] = []
        for tier in campaign.tiers:
            tier_reasons = _validate_rate_input(
                    f"{campaign.campaign_code}:{tier.threshold}",
                    tier.p2_condition_item_id,
                    tier.upstream_unit_reward,
                    tier.previous_agency_unit_reward,
                    tier.direct_cost_per_unit,
                    tier.target_gross_margin,
                    tier.minimum_gross_margin,
                    tier.upstream_approval,
                    tier.previous_approval,
                    tier.direct_cost_approval,
                )
            hold_reasons.extend(tier_reasons)
            if tier_reasons or not all(
                _valid_amount(value)
                for value in (
                    tier.upstream_unit_reward,
                    tier.previous_agency_unit_reward,
                    tier.direct_cost_per_unit,
                )
            ):
                tier_caps = []
                break
            target = _effective_margin_target(
                profile,
                tier.target_gross_margin,
                tier.minimum_gross_margin,
            )
            if target is None:
                continue
            tier_targets.append(target)
            cap = (
                tier.upstream_unit_reward
                - tier.direct_cost_per_unit
                - target
            )
            if cap < 0 and campaign.mode == LadderMode.REACHED_TIER:
                hold_reasons.append(
                    f"{campaign.campaign_code}:{tier.threshold}: "
                    "上位単価から原価と目標粗利を確保できません。"
                )
            tier_caps.append(cap)

        if len(tier_caps) != len(campaign.tiers):
            continue
        tier_rates = _propose_tier_rates(
            campaign,
            tier_caps,
            tier_targets,
            profile.allow_rate_changes,
            hold_reasons,
        )
        if tier_rates is None:
            continue
        cumulative_upstream = Decimal("0")
        cumulative_cost = Decimal("0")
        cumulative_agency = Decimal("0")
        for tier, rate, target in zip(
            campaign.tiers, tier_rates, tier_targets, strict=True
        ):
            cumulative_upstream += tier.upstream_unit_reward
            cumulative_cost += tier.direct_cost_per_unit
            cumulative_agency += rate
            retained = (
                tier.upstream_unit_reward - rate - tier.direct_cost_per_unit
                if campaign.mode == LadderMode.REACHED_TIER
                else cumulative_upstream - cumulative_agency - cumulative_cost
            )
            campaign_proposals.append(
                CampaignTierProposal(
                    campaign_code=campaign.campaign_code,
                    threshold=tier.threshold,
                    p2_condition_item_id=tier.p2_condition_item_id,
                    upstream_unit_reward=tier.upstream_unit_reward,
                    previous_agency_unit_reward=tier.previous_agency_unit_reward,
                    proposed_agency_unit_reward=rate,
                    direct_cost_per_unit=tier.direct_cost_per_unit,
                    retained_gross_margin_per_unit=retained,
                )
            )
            if campaign_proposals[-1].retained_gross_margin_per_unit < target:
                hold_reasons.append(
                    f"{campaign.campaign_code}:{tier.threshold}: "
                    "目標粗利を下回る提案です。"
                )

        scenarios.extend(
            _campaign_scenarios(
                profile,
                campaign,
                tuple(
                    proposal
                    for proposal in campaign_proposals
                    if proposal.campaign_code == campaign.campaign_code
                ),
            )
        )

    campaign_by_code = {campaign.campaign_code: campaign for campaign in campaigns}
    for scenario in volume_scenarios:
        campaign = campaign_by_code.get(scenario.campaign_code)
        if campaign is None:
            hold_reasons.append(
                f"{scenario.scenario_id}: 対応キャンペーンがありません。"
            )
            continue
        scenario_reasons = _validate_volume_scenario(scenario, campaign, profile)
        hold_reasons.extend(scenario_reasons)
        if scenario_reasons:
            continue
        campaign_proposal_rows = tuple(
            proposal
            for proposal in campaign_proposals
            if proposal.campaign_code == campaign.campaign_code
        )
        if len(campaign_proposal_rows) != len(campaign.tiers):
            continue
        scenarios.append(
            _simulate_campaign_scenario(
                profile,
                campaign,
                campaign_proposal_rows,
                scenario.opening_class,
                scenario.commerce_count,
                scenario.upstream_collected,
                scenario.agency_payment_due,
                agency_eligible_count=scenario.agency_eligible_count,
                scenario_id=scenario.scenario_id,
            )
        )

    for outcome in scenarios:
        if outcome.risks:
            hold_reasons.extend(outcome.risks)
        if (
            outcome.minimum_margin_checked
            and outcome.retained_gross_margin
            < profile.minimum_gross_margin * outcome.agency_eligible_count
        ):
            hold_reasons.append(
                f"{outcome.scenario_id}: 最低粗利"
                f"{profile.minimum_gross_margin}円に対し"
                f"{outcome.retained_gross_margin}円です。"
            )

    deduplicated_reasons = tuple(dict.fromkeys(hold_reasons))
    return RewardDesignResult(
        status="hold" if deduplicated_reasons else "candidate",
        agency_id=profile.agency_id,
        target_period=profile.target_period,
        component_proposals=tuple(proposals),
        campaign_proposals=tuple(campaign_proposals),
        scenarios=tuple(scenarios),
        manual_judgments=manual_judgments,
        hold_reasons=deduplicated_reasons,
    )


def connect_reward_design_to_p2(
    design: RewardDesignResult,
    p2_report: P2Report,
) -> P2RewardDesignGate:
    """Require clean P2 source/mapping results before a candidate enters approval."""
    reasons: list[str] = []
    if design.status != "candidate":
        reasons.extend(design.hold_reasons or ("報酬設計が候補状態ではありません。",))
    if p2_report.target_period != design.target_period:
        reasons.append("報酬設計とP2結果の対象期間が一致しません。")

    expected: dict[str, tuple[Decimal, Decimal]] = {}
    for proposal in design.component_proposals:
        expected[proposal.p2_condition_item_id] = (
            proposal.upstream_unit_reward,
            proposal.direct_cost_per_unit,
        )
    for proposal in design.campaign_proposals:
        expected[proposal.p2_condition_item_id] = (
            proposal.upstream_unit_reward,
            proposal.direct_cost_per_unit,
        )
    if not expected:
        reasons.append("P2へ渡す報酬明細がありません。")

    p2_results = [
        result
        for result in p2_report.results
        if result.agency_id == design.agency_id
    ]
    p2_by_item: dict[str, list] = {}
    for result in p2_results:
        p2_by_item.setdefault(result.condition_item_id, []).append(result)
    requires_reapproval = False
    for item_id, (upstream, direct_cost) in expected.items():
        matches = p2_by_item.get(item_id, [])
        if len(matches) != 1:
            reasons.append(f"P2に対応する成果項目が一意にありません: {item_id}")
            continue
        p2_result = matches[0]
        if p2_result.status != "proposed":
            reasons.extend(
                p2_result.hold_reasons
                or (f"P2成果項目が保留中です: {item_id}",)
            )
            continue
        if p2_result.upstream_reward != upstream:
            reasons.append(f"P2と設計の上位単価が一致しません: {item_id}")
        if p2_result.direct_cost != direct_cost:
            reasons.append(f"P2と設計の直接原価が一致しません: {item_id}")
        proposed_rates = [
            proposal.proposed_agency_unit_reward
            for proposal in design.component_proposals
            if proposal.p2_condition_item_id == item_id
        ] + [
            proposal.proposed_agency_unit_reward
            for proposal in design.campaign_proposals
            if proposal.p2_condition_item_id == item_id
        ]
        if (
            p2_result.proposed_agency_reward is None
            or any(rate != p2_result.proposed_agency_reward for rate in proposed_rates)
        ):
            requires_reapproval = True

    if reasons:
        return P2RewardDesignGate(
            status="hold",
            requires_p2_reapproval=True,
            hold_reasons=tuple(dict.fromkeys(reasons)),
        )
    return P2RewardDesignGate(
        status=(
            "requires_p2_reapproval"
            if requires_reapproval
            else "matches_p2_candidate"
        ),
        requires_p2_reapproval=requires_reapproval,
        hold_reasons=(),
    )


def write_p2_reward_design_candidate(
    input_path: Path,
    output_path: Path,
    design: RewardDesignResult,
    p2_report: P2Report,
) -> P2CandidateWorkbookResult:
    """Write a separate P2 input candidate whose changed hashes require reapproval."""
    gate = connect_reward_design_to_p2(design, p2_report)
    if gate.status == "hold":
        raise P2InputError(
            "報酬設計をP2へ渡せません: " + " / ".join(gate.hold_reasons)
        )
    if not gate.requires_p2_reapproval:
        raise P2InputError("報酬設計は既存P2候補と一致しており、再承認用ブックは不要です。")
    if input_path.resolve() == output_path.resolve():
        raise P2InputError("P2入力ブックを上書きすることはできません。")
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite P2 candidate: {output_path}")

    target_margins: dict[str, Decimal] = {}
    for proposal in design.component_proposals:
        target_margins[proposal.p2_condition_item_id] = (
            proposal.upstream_unit_reward
            - proposal.direct_cost_per_unit
            - proposal.proposed_agency_unit_reward
        )
    for proposal in design.campaign_proposals:
        target_margins[proposal.p2_condition_item_id] = (
            proposal.upstream_unit_reward
            - proposal.direct_cost_per_unit
            - proposal.proposed_agency_unit_reward
        )
    if any(not _valid_amount(value) for value in target_margins.values()):
        raise P2InputError("提案単価からP2へ渡す粗利を算出できません。")

    workbook = load_workbook(input_path)
    try:
        if "粗利原価" not in workbook.sheetnames:
            raise P2InputError("P2入力ブックに粗利原価シートがありません。")
        worksheet = workbook["粗利原価"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        required = ("record_id", "agency_id", "condition_item_id", "valid_from", "valid_until",
                    "approved_target_gross_margin")
        missing = [header for header in required if header not in headers]
        if missing:
            raise P2InputError(
                "P2粗利原価シートに必要な列がありません: " + ", ".join(missing)
            )
        period_start, period_end = _period_bounds(design.target_period)
        row_numbers: dict[str, list[int]] = {item_id: [] for item_id in target_margins}
        for row_number in range(2, worksheet.max_row + 1):
            agency_id = _cell_text(worksheet.cell(row_number, headers["agency_id"]).value)
            item_id = _cell_text(
                worksheet.cell(row_number, headers["condition_item_id"]).value
            )
            if agency_id != design.agency_id or item_id not in row_numbers:
                continue
            valid_from = _cell_date(
                worksheet.cell(row_number, headers["valid_from"]).value
            )
            valid_until = _cell_date(
                worksheet.cell(row_number, headers["valid_until"]).value
            )
            if (
                valid_from is None
                or valid_until is None
                or valid_from > period_end
                or valid_until < period_start
            ):
                continue
            row_numbers[item_id].append(row_number)

        missing_items = [
            item_id for item_id, rows in row_numbers.items() if not rows
        ]
        if missing_items:
            raise P2InputError(
                "P2入力ブックに対象期間の粗利原価行がありません: "
                + ", ".join(missing_items)
            )
        for item_id, rows in row_numbers.items():
            for row_number in rows:
                worksheet.cell(
                    row_number,
                    headers["approved_target_gross_margin"],
                    str(target_margins[item_id]),
                )
    finally:
        workbook.close()

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        with output_path.open("xb") as output_stream:
            workbook.save(output_stream)
    finally:
        workbook.close()

    saved = load_workbook(output_path, read_only=True, data_only=True)
    try:
        worksheet = saved["粗利原価"]
        headers = {
            str(cell.value).strip(): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for item_id, rows in row_numbers.items():
            expected = target_margins[item_id]
            for row_number in rows:
                actual = worksheet.cell(
                    row_number, headers["approved_target_gross_margin"]
                ).value
                if actual is None or Decimal(str(actual)) != expected:
                    raise P2InputError(
                        f"P2候補の保存後読戻し値が一致しません: {item_id}"
                    )
    finally:
        saved.close()

    post_write_report = calculate_workbook(output_path, design.target_period)
    for item_id in target_margins:
        matching = [
            result
            for result in post_write_report.results
            if result.agency_id == design.agency_id
            and result.condition_item_id == item_id
        ]
        if len(matching) != 1 or matching[0].status != "hold":
            raise P2InputError(
                f"P2再計算で承認保留になりませんでした: {item_id}"
            )
        if not any("承認ハッシュ" in reason for reason in matching[0].hold_reasons):
            raise P2InputError(
                f"P2再計算で変更後record_hashの再承認要求を確認できません: {item_id}"
            )
    return P2CandidateWorkbookResult(
        output_path=output_path,
        updated_condition_item_ids=tuple(sorted(target_margins)),
        post_write_p2_report=post_write_report,
    )


def _validate_profile(profile: AgencyRewardProfile) -> list[str]:
    reasons: list[str] = []
    if not profile.agency_id.strip() or not profile.target_period.strip():
        reasons.append("代理店IDまたは対象期間がありません。")
    for label, amount in (
        ("目標粗利", profile.target_gross_margin),
        ("最低粗利", profile.minimum_gross_margin),
    ):
        if not _valid_amount(amount):
            reasons.append(f"{label}は有限の0以上の金額である必要があります。")
    if (
        _valid_amount(profile.target_gross_margin)
        and _valid_amount(profile.minimum_gross_margin)
        and profile.target_gross_margin < profile.minimum_gross_margin
    ):
        reasons.append("目標粗利が最低粗利を下回っています。")
    if not isinstance(profile.allow_rate_changes, bool):
        reasons.append("単価変更可否が明示されていません。")
    if profile.pay_only_after_upstream_collection not in (True, False, None):
        reasons.append("上位未収時の支払条件が不正です。")
    return reasons


def _validate_rate_input(
    label: str,
    p2_item_id: str,
    upstream: Decimal,
    previous: Decimal,
    direct_cost: Decimal,
    target: Decimal | None,
    minimum: Decimal | None,
    upstream_approval: ApprovalEvidence,
    previous_approval: ApprovalEvidence,
    direct_cost_approval: ApprovalEvidence,
) -> list[str]:
    reasons: list[str] = []
    if not p2_item_id.strip():
        reasons.append(f"{label}: P2成果項目IDがありません。")
    for name, amount in (
        ("上位単価", upstream),
        ("前回代理店単価", previous),
        ("直接原価", direct_cost),
    ):
        if not _valid_amount(amount):
            reasons.append(f"{label}: {name}が有限の0以上の金額ではありません。")
    for name, amount in (("目標粗利", target), ("最低粗利", minimum)):
        if amount is not None and not _valid_amount(amount):
            reasons.append(f"{label}: {name}が有限の0以上の金額ではありません。")
    if not upstream_approval.is_complete_approval():
        reasons.append(f"{label}: 上位単価の承認済み根拠がありません。")
    if not previous_approval.is_complete_approval():
        reasons.append(f"{label}: 前回条件の承認済み根拠がありません。")
    if not direct_cost_approval.is_complete_approval():
        reasons.append(f"{label}: 直接原価の承認済み根拠がありません。")
    return reasons


def _effective_margin_target(
    profile: AgencyRewardProfile,
    item_target: Decimal | None,
    item_minimum: Decimal | None,
) -> Decimal | None:
    values = (
        profile.target_gross_margin,
        profile.minimum_gross_margin,
        item_target if item_target is not None else profile.target_gross_margin,
        item_minimum if item_minimum is not None else profile.minimum_gross_margin,
    )
    if any(not _valid_amount(value) for value in values):
        return None
    return max(values)


def _least_change_rate(
    previous: Decimal,
    maximum: Decimal,
    allow_changes: bool,
    label: str,
    hold_reasons: list[str],
) -> Decimal | None:
    if not _valid_amount(previous) or not _valid_amount(maximum):
        return None
    if not allow_changes:
        if previous > maximum:
            hold_reasons.append(
                f"{label}: 単価変更不可ですが、前回単価では目標粗利を確保できません。"
            )
            return None
        return previous
    return min(previous, maximum)


def _validate_campaign(campaign: CampaignSchedule) -> list[str]:
    reasons: list[str] = []
    if not campaign.campaign_code.strip():
        reasons.append("キャンペーンコードがありません。")
    if not campaign.counting_scope.strip() or not campaign.payout_scope.strip():
        reasons.append(
            f"{campaign.campaign_code}: 商流全体の計数範囲または代理店報酬対象範囲がありません。"
        )
    thresholds = [tier.threshold for tier in campaign.tiers]
    if not thresholds:
        reasons.append(f"{campaign.campaign_code}: 閾値がありません。")
    if any(
        isinstance(threshold, bool)
        or not isinstance(threshold, int)
        or threshold <= 0
        for threshold in thresholds
    ):
        reasons.append(f"{campaign.campaign_code}: 閾値は正の整数が必要です。")
    elif thresholds != sorted(thresholds) or len(thresholds) != len(set(thresholds)):
        reasons.append(
            f"{campaign.campaign_code}: 閾値は重複のない昇順で指定してください。"
        )
    if any(not opening_class.strip() for opening_class in campaign.eligible_opening_classes):
        reasons.append(
            f"{campaign.campaign_code}: 空の開通区分が設定されています。"
        )
    if len(campaign.eligible_opening_classes) != len(
        set(campaign.eligible_opening_classes)
    ):
        reasons.append(
            f"{campaign.campaign_code}: 支払対象の開通区分が重複しています。"
        )
    return reasons


def _propose_tier_rates(
    campaign: CampaignSchedule,
    caps: list[Decimal],
    targets: list[Decimal],
    allow_changes: bool,
    hold_reasons: list[str],
) -> tuple[Decimal, ...] | None:
    previous = tuple(tier.previous_agency_unit_reward for tier in campaign.tiers)
    if not all(_valid_amount(amount) for amount in previous):
        return None
    if campaign.mode == LadderMode.REACHED_TIER and any(cap < 0 for cap in caps):
        return None
    if not allow_changes:
        proposed = previous
    elif campaign.mode == LadderMode.REACHED_TIER:
        proposed = _least_change_monotone_rates(previous, caps)
    elif campaign.mode == LadderMode.PROGRESSIVE:
        proposed = _least_change_progressive_rates(
            campaign.tiers, caps, targets
        )
    else:
        proposed = tuple(min(rate, cap) for rate, cap in zip(previous, caps, strict=True))

    if campaign.mode == LadderMode.REACHED_TIER and any(
        left > right for left, right in zip(proposed, proposed[1:])
    ):
        hold_reasons.append(
            f"{campaign.campaign_code}: 閾値上昇に対して代理店単価が逆転しています。"
        )
        return None
    if not allow_changes and any(
        rate > cap for rate, cap in zip(proposed, caps, strict=True)
    ):
        hold_reasons.append(
            f"{campaign.campaign_code}: 単価変更不可では目標粗利を確保できません。"
        )
        return None
    if any(
        rate < 0 or rate > cap
        for rate, cap in zip(proposed, caps, strict=True)
    ) and campaign.mode == LadderMode.REACHED_TIER:
        hold_reasons.append(
            f"{campaign.campaign_code}: 目標粗利を満たさない単価です。"
        )
        return None
    if campaign.mode == LadderMode.PROGRESSIVE:
        cumulative = Decimal("0")
        for index, (rate, tier) in enumerate(
            zip(proposed, campaign.tiers, strict=True)
        ):
            cumulative += rate
            total_cap = sum(
                (
                    value.upstream_unit_reward
                    - value.direct_cost_per_unit
                    for value in campaign.tiers[: index + 1]
                ),
                Decimal("0"),
            ) - targets[index]
            if rate < 0 or cumulative > total_cap:
                hold_reasons.append(
                    f"{campaign.campaign_code}: 累進適用後に目標粗利を確保できません。"
                )
                return None
    return tuple(proposed)


def _least_change_monotone_rates(
    previous: tuple[Decimal, ...],
    caps: list[Decimal],
) -> tuple[Decimal, ...]:
    candidates = sorted({Decimal("0"), *previous, *caps})
    states: dict[Decimal, tuple[Decimal, tuple[Decimal, ...]]] = {}
    for index, old_rate in enumerate(previous):
        allowed = [
            candidate
            for candidate in candidates
            if candidate <= old_rate and candidate <= caps[index]
        ]
        next_states: dict[Decimal, tuple[Decimal, tuple[Decimal, ...]]] = {}
        for candidate in allowed:
            best: tuple[Decimal, tuple[Decimal, ...]] | None = None
            if index == 0:
                best = (Decimal("0"), ())
            elif states:
                compatible = [
                    state
                    for prior_rate, state in states.items()
                    if prior_rate <= candidate
                ]
                if compatible:
                    best = min(compatible, key=lambda state: _state_key(state))
            if best is not None:
                previous_cost, path = best
                next_states[candidate] = (
                    previous_cost + old_rate - candidate,
                    (*path, candidate),
                )
        states = next_states
    if not states:
        return tuple(Decimal("0") for _ in previous)
    return min(states.values(), key=_state_key)[1]


def _state_key(
    state: tuple[Decimal, tuple[Decimal, ...]],
) -> tuple[Decimal, tuple[Decimal, ...]]:
    cost, path = state
    return (
        cost,
        path,
    )


def _least_change_progressive_rates(
    tiers: tuple[CampaignTierInput, ...],
    caps: list[Decimal],
    targets: list[Decimal],
) -> tuple[Decimal, ...]:
    proposed = [tier.previous_agency_unit_reward for tier in tiers]
    for index in range(len(proposed)):
        cumulative = sum(proposed[: index + 1], Decimal("0"))
        cumulative_cap = sum(
            (
                tier.upstream_unit_reward - tier.direct_cost_per_unit
                for tier in tiers[: index + 1]
            ),
            Decimal("0"),
        ) - targets[index]
        excess = max(Decimal("0"), cumulative - cumulative_cap)
        for prior_index in range(index, -1, -1):
            reduction = min(proposed[prior_index], excess)
            proposed[prior_index] -= reduction
            excess -= reduction
            if excess == 0:
                break
    return tuple(proposed)


def _component_scenarios(
    profile: AgencyRewardProfile,
    proposal: RewardRateProposal,
    margin_target: Decimal,
) -> tuple[ScenarioOutcome, ...]:
    outcomes: list[ScenarioOutcome] = []
    for upstream_collected in (True, False):
        payment_due = _payment_due(profile, upstream_collected)
        payout = (
            proposal.proposed_agency_unit_reward
            if payment_due and proposal.opening_class
            else Decimal("0")
        )
        upstream = proposal.upstream_unit_reward if upstream_collected else Decimal("0")
        margin = upstream - payout - proposal.direct_cost_per_unit
        risks = _collection_risks(
            upstream_collected,
            payment_due,
            payout,
            proposal.component_id,
        )
        outcomes.append(
            ScenarioOutcome(
                scenario_id=f"component:{proposal.component_id}:"
                f"{'collected' if upstream_collected else 'uncollected'}",
                opening_class=proposal.opening_class,
                campaign_code=None,
                commerce_count=None,
                agency_eligible_count=1,
                upstream_collected=upstream_collected,
                agency_payment_due=payment_due,
                qualified=True,
                upstream_amount=upstream,
                agency_payment=payout,
                direct_cost=proposal.direct_cost_per_unit,
                retained_gross_margin=margin,
                minimum_margin_checked=upstream_collected,
                risks=risks,
            )
        )
        if upstream_collected and margin < margin_target:
            outcomes[-1] = _add_scenario_risk(
                outcomes[-1],
                f"{proposal.component_id}: 目標粗利{margin_target}円を下回ります。",
            )
    return tuple(outcomes)


def _campaign_scenarios(
    profile: AgencyRewardProfile,
    campaign: CampaignSchedule,
    proposals: tuple[CampaignTierProposal, ...],
) -> tuple[ScenarioOutcome, ...]:
    outcomes: list[ScenarioOutcome] = []
    if len(proposals) != len(campaign.tiers):
        return ()
    for opening_class in campaign.eligible_opening_classes:
        for tier in campaign.tiers:
            for commerce_count in (tier.threshold - 1, tier.threshold):
                for upstream_collected in (True, False):
                    payment_due = _payment_due(profile, upstream_collected)
                    outcome = _simulate_campaign_scenario(
                        profile,
                        campaign,
                        proposals,
                        opening_class,
                        commerce_count,
                        upstream_collected,
                        payment_due,
                        agency_eligible_count=1,
                    )
                    outcomes.append(outcome)
    return tuple(outcomes)


def _simulate_campaign_scenario(
    profile: AgencyRewardProfile,
    campaign: CampaignSchedule,
    proposals: tuple[CampaignTierProposal, ...],
    opening_class: str,
    commerce_count: int,
    upstream_collected: bool,
    payment_due: bool,
    *,
    agency_eligible_count: int,
    scenario_id: str | None = None,
) -> ScenarioOutcome:
    definition = LadderDefinition(
        counting_scope=campaign.counting_scope,
        mode=campaign.mode,
        tiers=tuple(
            LadderTier(tier.threshold, proposal.proposed_agency_unit_reward)
            for tier, proposal in zip(campaign.tiers, proposals, strict=True)
        ),
    )
    ladder = simulate_ladder(
        definition,
        LadderScenario(campaign.counting_scope, commerce_count),
    )
    if ladder.status != "simulated":
        return ScenarioOutcome(
            scenario_id=scenario_id or (
                f"campaign:{campaign.campaign_code}:{commerce_count}:"
                f"{opening_class}"
            ),
            opening_class=opening_class,
            campaign_code=campaign.campaign_code,
            commerce_count=commerce_count,
            agency_eligible_count=agency_eligible_count,
            upstream_collected=upstream_collected,
            agency_payment_due=payment_due,
            qualified=False,
            upstream_amount=Decimal("0"),
            agency_payment=Decimal("0"),
            direct_cost=Decimal("0"),
            retained_gross_margin=Decimal("0"),
            minimum_margin_checked=False,
            risks=ladder.hold_reasons,
        )

    reached = set(ladder.reached_thresholds)
    reached_inputs = [
        (tier, proposal)
        for tier, proposal in zip(campaign.tiers, proposals, strict=True)
        if tier.threshold in reached
    ]
    if campaign.mode == LadderMode.PROGRESSIVE:
        upstream_per_unit = sum(
            (tier.upstream_unit_reward for tier, _ in reached_inputs),
            Decimal("0"),
        )
        cost_per_unit = sum(
            (tier.direct_cost_per_unit for tier, _ in reached_inputs),
            Decimal("0"),
        )
    elif reached_inputs:
        tier, _ = reached_inputs[-1]
        upstream_per_unit = tier.upstream_unit_reward
        cost_per_unit = tier.direct_cost_per_unit
    else:
        upstream_per_unit = Decimal("0")
        cost_per_unit = Decimal("0")
    qualified = bool(reached_inputs)
    payment = (
        ladder.payout_amount * agency_eligible_count
        if payment_due
        and qualified
        and opening_class in campaign.eligible_opening_classes
        else Decimal("0")
    )
    upstream = (
        upstream_per_unit * agency_eligible_count
        if upstream_collected
        else Decimal("0")
    )
    total_cost = cost_per_unit * agency_eligible_count
    retained = upstream - payment - total_cost
    risks = _collection_risks(
        upstream_collected,
        payment_due,
        payment,
        scenario_id or f"{campaign.campaign_code}:{commerce_count}:{opening_class}",
    )
    minimum_checked = qualified and upstream_collected and agency_eligible_count > 0
    return ScenarioOutcome(
        scenario_id=scenario_id or (
            f"campaign:{campaign.campaign_code}:{commerce_count}:"
            f"{opening_class}:"
            f"{'collected' if upstream_collected else 'uncollected'}"
        ),
        opening_class=opening_class,
        campaign_code=campaign.campaign_code,
        commerce_count=commerce_count,
        agency_eligible_count=agency_eligible_count,
        upstream_collected=upstream_collected,
        agency_payment_due=payment_due,
        qualified=qualified,
        upstream_amount=upstream,
        agency_payment=payment,
        direct_cost=total_cost,
        retained_gross_margin=retained,
        minimum_margin_checked=minimum_checked,
        risks=risks,
    )


def _payment_due(
    profile: AgencyRewardProfile,
    upstream_collected: bool,
) -> bool:
    if profile.pay_only_after_upstream_collection is None:
        return not upstream_collected
    return upstream_collected or not profile.pay_only_after_upstream_collection


def _validate_volume_scenario(
    scenario: VolumeScenario,
    campaign: CampaignSchedule,
    profile: AgencyRewardProfile,
) -> list[str]:
    reasons: list[str] = []
    if not scenario.scenario_id.strip():
        reasons.append("出来高シナリオIDがありません。")
    if (
        isinstance(scenario.commerce_count, bool)
        or not isinstance(scenario.commerce_count, int)
        or scenario.commerce_count < 0
    ):
        reasons.append(f"{scenario.scenario_id}: 商流全体件数が不正です。")
    if (
        isinstance(scenario.agency_eligible_count, bool)
        or not isinstance(scenario.agency_eligible_count, int)
        or scenario.agency_eligible_count < 0
    ):
        reasons.append(f"{scenario.scenario_id}: 代理店対象件数が不正です。")
    if scenario.opening_class not in campaign.eligible_opening_classes:
        reasons.append(
            f"{scenario.scenario_id}: 対象開通区分がキャンペーン適用範囲外です。"
        )
    if not isinstance(scenario.upstream_collected, bool) or not isinstance(
        scenario.agency_payment_due, bool
    ):
        reasons.append(f"{scenario.scenario_id}: 入金・支払状態が不正です。")
    if (
        profile.pay_only_after_upstream_collection is True
        and not scenario.upstream_collected
        and scenario.agency_payment_due
    ):
        reasons.append(
            f"{scenario.scenario_id}: 上位未収時に支払うシナリオは承認済み支払条件と矛盾します。"
        )
    return reasons


def _collection_risks(
    upstream_collected: bool,
    payment_due: bool,
    agency_payment: Decimal,
    label: str,
) -> tuple[str, ...]:
    if not upstream_collected and payment_due and agency_payment > 0:
        return (
            f"{label}: 上位未収のまま代理店支払{agency_payment}円が発生します。",
        )
    return ()


def _add_scenario_risk(
    scenario: ScenarioOutcome,
    reason: str,
) -> ScenarioOutcome:
    return ScenarioOutcome(
        scenario_id=scenario.scenario_id,
        opening_class=scenario.opening_class,
        campaign_code=scenario.campaign_code,
        commerce_count=scenario.commerce_count,
        agency_eligible_count=scenario.agency_eligible_count,
        upstream_collected=scenario.upstream_collected,
        agency_payment_due=scenario.agency_payment_due,
        qualified=scenario.qualified,
        upstream_amount=scenario.upstream_amount,
        agency_payment=scenario.agency_payment,
        direct_cost=scenario.direct_cost,
        retained_gross_margin=scenario.retained_gross_margin,
        minimum_margin_checked=scenario.minimum_margin_checked,
        risks=(*scenario.risks, reason),
    )


def _valid_amount(value: Decimal | None) -> bool:
    return (
        isinstance(value, Decimal)
        and value.is_finite()
        and value >= 0
    )


def _period_bounds(target_period: str) -> tuple[date, date]:
    try:
        year_text, month_text = target_period.split("-", maxsplit=1)
        year = int(year_text)
        month = int(month_text)
        start = date(year, month, 1)
    except (ValueError, TypeError) as error:
        raise P2InputError(f"対象期間がYYYY-MM形式ではありません: {target_period}") from error
    if month == 12:
        next_month = date(year + 1, 1, 1)
    else:
        next_month = date(year, month + 1, 1)
    return start, date.fromordinal(next_month.toordinal() - 1)


def _cell_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            return None
    return None


def _cell_text(value: object) -> str:
    return "" if value is None else str(value).strip()
