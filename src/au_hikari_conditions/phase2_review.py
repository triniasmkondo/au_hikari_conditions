"""Phase 2 monthly review, approved-margin calculation and condition-bound decisions.

Nothing here generates deliverables or calls an external service. Review workbooks are
recomputed from source files; editing a displayed result cannot approve a condition.
"""
from __future__ import annotations
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

from .condition_ledger import Condition, HEADERS, normal, candidate_key
from .upper_conditions import (UpperCondition, extract_upper_pdf, compact, types,
                               dwelling, speed, codes, option_name)
from .calculation import (ApprovalEvidence, ApprovalStatus, ApprovedMarginBaseline,
    CalculationCase, ConditionSignature, DisplayPolicy, TaxBasis, calculate_agency_reward)
from .p2 import compute_record_hash, _parse_target_period
from .phase2_triage import (VERSION as TRIAGE_VERSION, PREVIEW_HEADERS, initial_candidate,
    build_triage, write_triage, preview_values)


APPROVAL_FIELDS=['source_document','evidence_location','approval_reason','approval_id',
                 'approval_status','approved_by','approved_at','record_hash','approved_hash']
MAPPING_FIELDS=['proposal_id','agency','condition_id','selected_upper_id','decision',
    'equivalence_status','scope_review','period_rule','payment_timing','payout_trigger',
    'exclusions','reversal_terms','snapshot_hash']+APPROVAL_FIELDS
MARGIN_FIELDS=['proposal_id','agency','condition_id','target_period','approved_gross_margin',
    'gross_margin_tax_basis','direct_cost','direct_cost_tax_basis','minimum_gross_margin',
    'display_policy','snapshot_hash']+APPROVAL_FIELDS
RESOLUTION_FIELDS=['proposal_id','agency','condition_id','resolved_issues','tax_basis',
    'resolution','snapshot_hash']+APPROVAL_FIELDS
DECISION_FIELDS=['proposal_id','agency','condition_id','candidate_hash','approved_amount',
                 'snapshot_hash']+APPROVAL_FIELDS
INPUT_SCHEMAS={'条件対応入力':MAPPING_FIELDS,'粗利承認入力':MARGIN_FIELDS,
               '保留解消入力':RESOLUTION_FIELDS,'最終承認入力':DECISION_FIELDS}
VIEW_HEADERS=['代理店','提案ID','10月条件ID','11月条件ID','区分','商材','成果項目','ISP','タイプ',
    '対象・住居・速度','特典コード','閾値','分類案','分類理由','10月代理店単価','10月上位単価',
    '11月上位単価','上位金額状態','提案単価','承認済み目標粗利','直接原価','提案時粗利',
    '10月適用期間','11月上位適用期間','10月適用条件','11月適用条件','旧原本根拠','新上位原本根拠',
    '上位対応候補','選択上位ID','Phase1保留','税区分未明示','承認への影響・保留理由',
    '計算状態','最終状態','承認額','確定時粗利','承認対象ハッシュ',
    '11月成果項目・条件軸','旧単位','新上位単位','旧税区分','新上位税区分','10月上位原本根拠']+PREVIEW_HEADERS


def file_hash(path):return hashlib.sha256(path.read_bytes()).hexdigest()


def _hash(value):
    return hashlib.sha256(json.dumps(value,sort_keys=True,ensure_ascii=False,default=str,
                                     separators=(',',':')).encode()).hexdigest()


def read_phase1(root: Path, ledger: Path):
    book=load_workbook(ledger,read_only=True,data_only=False)
    try:
        sheet=book['共通条件台帳']
        if [c.value for c in sheet[1]]!=HEADERS:raise ValueError('Phase 1台帳の列定義が不一致')
        result=[]
        for row in sheet.iter_rows(min_row=2):
            if any(c.data_type=='f' for c in row):raise ValueError('Phase 1台帳に数式は使用できません')
            r=[c.value for c in row];r=['' if v is None else v for v in r]
            c=Condition(r[0],r[2],r[3],r[4],r[5],r[6],r[7],r[8],r[9],r[10],r[11],r[12],
                        r[13],r[14],r[15],r[16],r[17],r[18].splitlines(),r[19],r[21])
            if c.id!=r[1]:raise ValueError('Phase 1条件IDと内容が不一致')
            result.append(c)
        if len({c.id for c in result})!=len(result):raise ValueError('Phase 1条件IDが重複')
        sources=[]
        for values in book['資料一覧'].iter_rows(min_row=2,values_only=True):
            _,source,digest,*_=values
            path=(root/source).resolve()
            if not path.is_relative_to(root.resolve()) or file_hash(path)!=digest:
                raise ValueError('Phase 1原本のハッシュが不一致')
            sources.append([source,digest])
        if not result:raise ValueError('Phase 1条件がありません')
        return result,sources
    finally:book.close()


def old_axes(c: Condition):
    """Only explicit source labels and audited name aliases, never price-based matching."""
    product=c.product;item=option_name(c.item);family='option';isp=c.isp;typ=types(c.type)
    dw=dwelling(c.target);sp=speed(c.target);code=codes(c.code);threshold=''
    if c.kind=='キャンペーン':
        family='campaign';item=compact(c.item);threshold=re.sub(r'[,~〜]','',compact(c.threshold)).replace('件','')+'件以上'
        return dict(product=product,family=family,item=item,isp='',type=typ,dwelling='',speed='',access='',code='',threshold=threshold)
    if c.kind=='違約金':
        return dict(product='共通',family='penalty',item=compact(c.item),isp='',type='',dwelling='',speed='',access='',code='',threshold='')
    if c.kind in ('控除','違約金・相殺'):
        item='後確認手数料' if c.item=='後確対応費' else ('消費者センター案件' if c.kind=='違約金・相殺' else item)
        return dict(product='auひかり/auひかりプラス',family='deduction',item=item,isp='',type='',dwelling='',speed='',access='',code='',threshold='1件以上' if item in ('後確認手数料','消費者センター案件') else '')
    if 'プラス' not in product and c.item in ('auひかりホーム','auひかりマンション','電話'):
        family='base';item='phone_open' if c.item=='電話' else 'net_open';dw=dwelling(c.item)
        # Phone is a single common rate in the agency sheet; upstream has two dwelling rows.
        if c.item=='電話':dw=''
    elif 'プラス' in product and c.item in ('5G開通手数料','ネット開通手数料','電話開通手数料'):
        family='base';item={'5G開通手数料':'5g_open','ネット開通手数料':'net_open','電話開通手数料':'phone_open'}[c.item]
    elif c.item in ('Mプラン','Lプラン'):
        product='auでんき';family='electric';isp='';dw=''
    elif isp=='BIGLOBE' and '回線サービス' in c.section:
        family='isp_volume' if 'ボリューム' in c.section else 'isp_line'
        item='ボリュームインセンティブ' if family=='isp_volume' else '回線サービス'
        target=c.target or c.item
        dw='ホームルーター' if 'ホームルーター' in target else dwelling(target)
        sp=speed(target);threshold='250件以上' if family=='isp_volume' else ''
    elif isp=='BIGLOBE':family='isp_option'
    if family in ('option','isp_option'):
        if 'ホームのみ' in c.item:dw='ホーム'
        if item=='So-netID':dw=dwelling(c.item) or dw
    return dict(product=product,family=family,item=item,isp=isp,type=typ,dwelling=dw,speed=sp,
                access='',code=code,threshold=threshold)


def related_axes(a,b):
    if any(a[k]!=b[k] for k in ('product','family','item','isp')):return False
    for key in ('type','dwelling','speed','code','threshold'):
        x,y=a[key],b[key]
        if x==y:continue
        if not x or not y:continue
        if key=='type' and (x=='ALL' or y=='ALL' or set(x.split(','))&set(y.split(','))):continue
        if key=='dwelling' and set(x.split('/'))&set(y.split('/')):continue
        return False
    return True


def _exact_matches(axes,uppers):return [u for u in uppers if u.axes==axes]


def classification(old, previous, current, exact=True):
    """Classification is a review proposal, independently from payment approval."""
    if old is None:return '新規','10月代理店条件に同一の成果項目・条件なし。採用は未承認'
    if current is None:return '要確認','11月上位の一意な対応なし。未掲載だけでは終了を確定しない'
    if current.state=='ENDED':return '終了','上位原本に明示された終了（根拠の承認が必要）'
    if current.state!='VALUE':return '要確認','上位金額は'+current.state+'。ダッシュ・未掲載を0円や終了とみなさない'
    if current.family=='deduction' and old.amount is not None and old.amount>0 and current.amount<0:
        return '変更','請求額の正額表記と上位控除の負額表記が異なる。符号・発生条件の個別承認が必要'
    if not exact:return '変更','対象範囲・特典コード・列構造等に差異。条件対応の個別承認が必要'
    if previous is None:return '要確認','10月上位との対応未確定。代理店単価との差額を粗利と推定しない'
    if previous.state!=current.state or previous.amount!=current.amount:
        return '変更','同じ成果項目・条件の上位単価または金額状態が変更'
    if current.components:return '変更','上位の手数料が複数コンポーネントに分割。合算対応は未承認'
    return '継承','成果項目・条件軸・上位単価が一致する継承候補。期間と支払・除外・戻入条項は別途承認'


META=set(APPROVAL_FIELDS)-{'source_document','evidence_location','approval_reason'}


def input_hash(kind,record):
    return compute_record_hash(kind,{k:v for k,v in record.items() if k not in META})


def approval_errors(kind,record,snapshot):
    if not record:return ['承認入力なし']
    errors=[]
    if record.get('snapshot_hash')!=snapshot:errors.append('入力資料スナップショット不一致')
    if record.get('approval_status')!='APPROVED':errors.append('未承認')
    expected=input_hash(kind,record)
    if record.get('approved_hash')!=expected:errors.append('承認ハッシュ不一致・未記入')
    try:at=datetime.fromisoformat(str(record.get('approved_at') or ''))
    except ValueError:at=None
    evidence=ApprovalEvidence(ApprovalStatus.APPROVED if record.get('approval_status')=='APPROVED' else ApprovalStatus.PENDING,
        record.get('approval_id'),record.get('source_document'),record.get('evidence_location'),record.get('approved_by'),at)
    if not evidence.is_complete_approval():errors.append('承認者・日時（タイムゾーン付き）・根拠・承認ID不足')
    if not record.get('approval_reason'):errors.append('承認理由なし')
    return errors


def approval_evidence(record):
    return ApprovalEvidence(ApprovalStatus.APPROVED,record['approval_id'],record['source_document'],
        record['evidence_location'],record['approved_by'],datetime.fromisoformat(str(record['approved_at'])))


def _decimal(value):
    if value in (None,'') or isinstance(value,bool):return None
    try:
        d=Decimal(str(value))
        return d if d.is_finite() else None
    except InvalidOperation:return None


def read_inputs(path):
    book=load_workbook(path,read_only=True,data_only=False)
    try:
        inputs={}
        for name,headers in INPUT_SCHEMAS.items():
            sheet=book[name]
            if list(next(sheet.values))!=headers:raise ValueError('承認入力の列が不一致: '+name)
            rows={}
            for cells in sheet.iter_rows(min_row=2):
                if any(c.data_type=='f' for c in cells):raise ValueError('承認入力に数式を使用できません')
                values=[c.value for c in cells]
                if not any(v is not None for v in values):continue
                r=dict(zip(headers,values));pid=r['proposal_id']
                if not pid or pid in rows:raise ValueError('承認入力の提案IDが空欄または重複')
                rows[pid]=r
            inputs[name]=rows
        manifest=list(book['資料指紋'].iter_rows(min_row=2,values_only=True))
        return inputs,[list(r) for r in manifest]
    finally:book.close()


@dataclass
class Proposal:
    agency: str
    old: Condition | None
    upper: UpperCondition | None
    candidates: list[UpperCondition]
    previous: UpperCondition | None
    axes: dict
    category: str
    reason: str
    exact: bool
    new_id: str
    id: str
    initial_identity_collision: bool = False


def make_proposals(olds,previous,current,target_period):
    proposals=[];agencies=sorted(set(c.agency for c in olds))
    for agency in agencies:
        used=set();related=set()
        for old in [c for c in olds if c.agency==agency]:
            axes=old_axes(old);exact=_exact_matches(axes,current)
            candidates=exact or [u for u in current if related_axes(axes,u.axes)]
            related.update(u.id for u in candidates)
            upper=exact[0] if len(exact)==1 else (candidates[0] if len(candidates)==1 else None)
            if len(exact)==1:used.add(upper.id)
            prev=_exact_matches(axes,previous);prev=prev[0] if len(prev)==1 else None
            category,reason=classification(old,prev,upper,len(exact)==1)
            if old.hold:category='要確認';reason='Phase 1保留を継承: '+reason
            if len(candidates)>1:category='要確認';reason='上位対応が複数。候補を選択して条件範囲のレビューが必要'
            pid=candidate_key('PROPOSAL',{'agency':agency,'old_id':old.id},target_period)
            newid=candidate_key('NOVEMBER',{'agency':agency,'previous_condition_id':old.id},target_period)
            proposals.append(Proposal(agency,old,upper,candidates,prev,axes,category,reason,len(exact)==1,newid,pid))
        for u in current:
            # All unpaired upper conditions are visible, including ambiguous counterparts.
            if u.id in used:continue
            linked=[c.id for c in olds if c.agency==agency and related_axes(old_axes(c),u.axes)]
            is_related=bool(linked)
            category='要確認' if is_related or u.state not in ('VALUE','ENDED') or u.hold else '新規'
            reason=('既存条件との候補関係あり。重複採用・分割の可否を個別確認: '+','.join(linked)) if is_related else '代理店10月台帳に対応行なし。上位行の新規採用候補'
            pid=candidate_key('NEW_PROPOSAL',{'agency':agency,'upper_id':u.id},target_period)
            newid=candidate_key('NOVEMBER_NEW',{'agency':agency,'upper_id':u.id},target_period)
            prev=_exact_matches(u.axes,previous);prev=prev[0] if len(prev)==1 else None
            if not is_related and prev:reason='10月上位にも存在するが、10月代理店台帳には同条件なし。代理店での新規採用は未承認'
            proposals.append(Proposal(agency,None,u,[u],prev,u.axes,category,reason,not is_related,newid,pid))
    axes_counts=Counter((p.agency,json.dumps(p.axes,sort_keys=True,ensure_ascii=False)) for p in proposals if p.old)
    for p in proposals:
        if p.old:p.initial_identity_collision=axes_counts[(p.agency,json.dumps(p.axes,sort_keys=True,ensure_ascii=False))]>1
    return proposals


def empty_input(headers,p,snapshot,target):
    r={key:None for key in headers}
    r.update(proposal_id=p.id,agency=p.agency,condition_id=p.old.id if p.old else p.new_id,
             snapshot_hash=snapshot,approval_status='PENDING')
    if 'selected_upper_id' in r:r['selected_upper_id']=p.upper.id if p.upper else None
    if 'target_period' in r:r['target_period']=target
    return r


def evaluate_proposal(p: Proposal, inputs, snapshot, target_period, extra_blocks=(), previous_period=None):
    mapping=inputs['条件対応入力'];baseline=inputs['粗利承認入力'];resolution=inputs['保留解消入力'];decision=inputs['最終承認入力']
    by_id={u.id:u for u in p.candidates}
    chosen=by_id.get(mapping.get('selected_upper_id'))
    if p.old:
        p.new_id=candidate_key('NOVEMBER',{'agency':p.agency,'previous_condition_id':p.old.id,
            'upper_id':chosen.id if chosen else 'UNRESOLVED'},target_period)
    blocks=list(extra_blocks)
    if chosen is None:blocks.append('上位対応が未選択・候補外・一意に未確定')
    map_errors=approval_errors('条件対応入力',mapping,snapshot)
    if mapping.get('equivalence_status')!='CONFIRMED':map_errors.append('成果項目・適用条件の同等性未承認')
    if mapping.get('scope_review')!='CONFIRMED':map_errors.append('ISP・タイプ・特典コード・期間・別紙・併給排他の確認未承認')
    for key in ('period_rule','payment_timing','payout_trigger','exclusions','reversal_terms'):
        if not mapping.get(key):map_errors.append(key+'が未確定')
    expected_decision='ADOPT' if p.old is None else 'CONTINUE'
    if chosen and chosen.state=='ENDED':expected_decision='END'
    if mapping.get('decision')!=expected_decision:map_errors.append('対応判断は'+expected_decision+'の個別承認が必要')
    if map_errors:blocks.append('条件対応: '+' / '.join(map_errors))
    issues=[]
    if p.old and p.old.hold:issues.append('PHASE1_HOLD')
    if p.old and p.old.tax!='税別':issues.append('TAX_UNKNOWN')
    if chosen and chosen.tax!='税別':issues.append('UPPER_TAX_UNKNOWN')
    if issues:
        resolved=set(str(resolution.get('resolved_issues') or '').split(';'))
        errors=approval_errors('保留解消入力',resolution,snapshot)
        if not set(issues)<=resolved:errors.append('未解消: '+';'.join(sorted(set(issues)-resolved)))
        if not resolution.get('resolution'):errors.append('解消根拠の具体的説明なし')
        if any('TAX' in i for i in issues) and resolution.get('tax_basis')!='exclusive':errors.append('税抜への確認・換算根拠が未承認')
        if errors:blocks.append('Phase 1引継・税: '+' / '.join(errors))
    if chosen:
        if chosen.state=='VALUE' and _decimal(chosen.amount) is None:
            blocks.append('上位金額の数値が未確定')
        if compact(p.agency).casefold()=='keioh株式会社' and chosen.family=='base' and chosen.item=='ftth_open':
            blocks.append('既存keiohルール: FTTH開通手数料は上位受領のみ・代理店支払禁止。通常報酬に混入不可')
        start,end=_parse_target_period(target_period)
        try:
            first,last=map(date.fromisoformat,chosen.period.split('/'))
            if first>start or last<end:blocks.append('上位適用期間が対象月全体を含まない')
        except (ValueError,AttributeError):blocks.append('上位適用期間が未確定')
        if chosen.state not in ('VALUE','ENDED'):blocks.append('上位金額状態 '+chosen.state+' は単価として確定不可')
        if 'unresolved' in chosen.family:blocks.append('上位の結合条件構造が未確定。原本を個別確認')
    # Extra upper rows related to an old row cannot silently turn into a second adopted payout.
    if p.old is None and not p.exact:blocks.append('既存条件との重複・分割候補。代理店採用範囲の設計が未確定')
    proposed=retained=approved_margin=cost=None
    reward=bool(chosen and chosen.state!='ENDED' and chosen.family not in ('deduction','penalty'))
    margin_errors=[]
    if reward:
        approved_margin=_decimal(baseline.get('approved_gross_margin'));cost=_decimal(baseline.get('direct_cost'))
        margin_errors=approval_errors('粗利承認入力',baseline,snapshot)
        if baseline.get('target_period')!=target_period:margin_errors.append('粗利基準の対象月不一致')
        if approved_margin is None:margin_errors.append('承認済み粗利不明（差額から推測しない）')
        if cost is None:margin_errors.append('承認済み直接原価不明（0円と推測しない）')
        if baseline.get('gross_margin_tax_basis')!='exclusive' or baseline.get('direct_cost_tax_basis')!='exclusive':margin_errors.append('粗利・直接原価の税区分未承認')
        if baseline.get('display_policy') not in ('present','withhold'):margin_errors.append('掲載可否未承認')
        if baseline.get('minimum_gross_margin') not in (None,'') and _decimal(baseline.get('minimum_gross_margin')) is None:margin_errors.append('最低粗利の金額形式不正')
        if margin_errors:blocks.append('粗利・原価: '+' / '.join(margin_errors))
        if not blocks:
            signature=ConditionSignature(p.new_id,tuple(sorted((k,str(v)) for k,v in chosen.axes.items())),
                chosen.period,mapping['period_rule'],mapping['payment_timing'],mapping['payout_trigger'],
                (mapping['exclusions'],),(mapping['reversal_terms'],))
            ev=approval_evidence(mapping);bev=approval_evidence(baseline)
            result=calculate_agency_reward(CalculationCase(signature,TaxBasis.EXCLUSIVE,TaxBasis.EXCLUSIVE,
                Decimal(chosen.amount),cost,ApprovedMarginBaseline(signature,TaxBasis.EXCLUSIVE,
                    approved_margin,DisplayPolicy(baseline['display_policy']),bev),ev,bev,ev,
                    minimum_gross_margin=_decimal(baseline.get('minimum_gross_margin'))))
            if result.status=='proposed':proposed=result.proposed_agency_reward;retained=result.retained_gross_margin
            else:blocks.extend(result.hold_reasons)
    elif chosen and chosen.family in ('deduction','penalty'):
        if not blocks:proposed=Decimal(chosen.amount)
    elif chosen and chosen.state=='ENDED':
        if not blocks:proposed=None
    # Include approvals and all inherited terms in the final fingerprint. Re-approving a mapping
    # or changing a cost, amount, period, source or tax decision invalidates the final decision.
    category,reason=(classification(p.old,p.previous,chosen,chosen.axes==p.axes) if not blocks and chosen else (p.category,p.reason))
    if previous_period is None:
        start,_=_parse_target_period(target_period)
        previous_period=date.fromordinal(start.toordinal()-1).strftime('%Y-%m')
    initial=initial_candidate(p,chosen,previous_period,target_period)
    payload={'triage_version':TRIAGE_VERSION,'initial_candidate':asdict(initial),
        'previous_upper':asdict(p.previous) if p.previous else None,'proposal_id':p.id,'previous_condition_id':p.old.id if p.old else '',
        'new_condition_id':p.new_id,'snapshot':snapshot,'selected_upper':asdict(chosen) if chosen else None,
        'old':asdict(p.old) if p.old else None,'category':category,
        'proposed_amount':str(proposed) if proposed is not None else None,
        'retained_margin':str(retained) if retained is not None else None,'blocks':blocks,
        'inputs':{name:{k:v for k,v in r.items() if k!='record_hash'} for name,r in
                  [('mapping',mapping),('baseline',baseline),('resolution',resolution)]}}
    fingerprint=_hash(payload)
    status='HELD' if blocks else 'PENDING'
    approved_amount=approved_retained=None
    decision_errors=[]
    if decision.get('approval_status')=='REJECTED':status='REJECTED'
    elif decision.get('approval_status')=='APPROVED':
        decision_errors=approval_errors('最終承認入力',decision,snapshot)
        if decision.get('candidate_hash')!=fingerprint:decision_errors.append('条件案ハッシュ不一致・再承認が必要')
        if blocks:decision_errors.append('未解決の承認阻害事項あり')
        amount=_decimal(decision.get('approved_amount'))
        if decision.get('approved_amount') not in (None,'') and amount is None:decision_errors.append('承認単価の金額形式不正')
        if chosen and chosen.state=='ENDED' and decision.get('approved_amount') not in (None,''):decision_errors.append('終了条件に支払額は指定できません')
        final_amount=amount if amount is not None else proposed
        if reward and not blocks:
            if final_amount is None or final_amount<0:decision_errors.append('承認単価不正')
            else:
                margin=Decimal(chosen.amount)-cost-final_amount
                if margin<approved_margin:decision_errors.append('承認額が承認済み粗利を維持しない')
                if not decision_errors:approved_retained=margin
        elif chosen and chosen.family in ('deduction','penalty') and final_amount!=proposed:
            decision_errors.append('控除・違約金額の変更には別の原本根拠が必要')
        if not decision_errors:status='APPROVED';approved_amount=final_amount
        else:status='APPROVAL_INVALID'
    return dict(proposal=p,chosen=chosen,blocks=blocks,proposed=proposed,retained=retained,
                margin=approved_margin,cost=cost,status=status,approved_amount=approved_amount,
                approved_retained=approved_retained,candidate_hash=fingerprint,decision_errors=decision_errors,
                issues=issues,category=category,reason=reason,initial_candidate=initial)


def _upper_ref(u):return u.source+'#'+u.location+' SHA256='+u.digest if u else ''


def view_row(result):
    p=result['proposal'];old=p.old;u=result['chosen'];a=p.axes
    return [p.agency,p.id,old.id if old else '',p.new_id,old.kind if old else ('上位からの新規採用候補'),
        a['product'],old.item if old else a['item'],a['isp'],a['type'],a['dwelling']+' '+a['speed'],a['code'],a['threshold'],
        result['category'],result['reason'],old.amount if old else None,p.previous.amount if p.previous else None,
        u.amount if u else None,u.state if u else 'UNMATCHED',result['proposed'],result['margin'],result['cost'],result['retained'],
        old.period if old else '',u.period if u else '',old.terms if old else '',u.terms if u else '',
        (old.source+'#'+'; '.join(old.locations)+' SHA256='+old.digest) if old else '',_upper_ref(u),
        '\n'.join(x.id+' | '+str(x.amount)+' | '+x.state+' | '+json.dumps(x.axes,ensure_ascii=False)+' | '+_upper_ref(x) for x in p.candidates),
        u.id if u else '',old.hold if old else '',old.tax if old and old.tax!='税別' else '',
        '\n'.join(result['blocks']+result['decision_errors']),
        '保留' if result['blocks'] else ('粗利維持提案' if result['retained'] is not None else '個別条項案'),
        result['status'],result['approved_amount'],result['approved_retained'],result['candidate_hash'],
        json.dumps(u.axes,ensure_ascii=False) if u else '',old.unit if old else '',u.unit if u else '',
        old.tax if old else '',u.tax if u else '',_upper_ref(p.previous)]+preview_values(result)


def _write_sheet(book,name,headers,rows,editable=False):
    sheet=book.create_sheet(name);sheet.append(headers)
    for row in rows:sheet.append([float(v) if isinstance(v,Decimal) else v for v in row])
    sheet.freeze_panes='G2' if len(headers)>10 else 'A2';sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='24476A')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            if isinstance(cell.value,str):
                if len(cell.value)>32767:raise ValueError('Excel文字数上限。原文を切り捨てず停止')
                cell.data_type='s'
            if editable:cell.fill=PatternFill('solid',fgColor='FFF2CC')
        sheet.row_dimensions[row[0].row].height=75
    for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    if editable:
        for field,choices in [('approval_status','PENDING,APPROVED,REJECTED'),('decision','CONTINUE,ADOPT,END,HOLD'),
                              ('equivalence_status','UNREVIEWED,CONFIRMED'),('scope_review','UNREVIEWED,CONFIRMED')]:
            if field in headers:
                col=sheet.cell(1,headers.index(field)+1).column_letter
                val=DataValidation(type='list',formula1='"'+choices+'"');sheet.add_data_validation(val);val.add(f'{col}2:{col}{max(2,sheet.max_row)}')
    return sheet


def find_conflicts(proposals,input_rows,snapshot):
    claims={};conflicts={}
    for p in proposals:
        mapping=input_rows['条件対応入力'].get(p.id)
        if not mapping or approval_errors('条件対応入力',mapping,snapshot) or mapping.get('decision') not in ('CONTINUE','ADOPT'):
            continue
        u=next((u for u in p.candidates if u.id==mapping.get('selected_upper_id')),None)
        if u:
            for uid in (u.id,*u.components):claims.setdefault((p.agency,uid),[]).append(p.id)
    for pids in claims.values():
        if len(set(pids))>1:
            for pid in pids:conflicts.setdefault(pid,[]).append('同一上位条件・合算コンポーネントへの重複採用。分割・併給の個別設計が必要')
    return conflicts


def build_phase2(root: Path, ledger: Path, output: Path, previous_period='2026-10', target_period='2026-11',
                 agency_filter=None, review_input: Path | None=None):
    _parse_target_period(previous_period);_parse_target_period(target_period)
    audit=output.with_suffix('.approvals.json')
    if output.exists() or audit.exists():raise FileExistsError('出力・承認結果は上書きしません')
    olds,phase1_sources=read_phase1(root,ledger)
    manifest=[['Phase1台帳',str(ledger.resolve()),file_hash(ledger)]]+[['Phase1原本',source,digest] for source,digest in phase1_sources]
    previous=[];current=[];raw=[]
    for period,bucket in ((previous_period,previous),(target_period,current)):
        paths=sorted((root/'01_上位原本'/period).glob('*.pdf'))
        if not paths:raise ValueError('上位原本がありません: '+period)
        for path in paths:
            records,pages=extract_upper_pdf(path,root);bucket.extend(records);raw.extend(pages)
            manifest.append(['上位原本',path.relative_to(root).as_posix(),file_hash(path)])
    for bucket in (previous,current):
        if len({u.id for u in bucket})!=len(bucket):raise ValueError('資料間で上位条件IDが重複。版の優先関係を確認してください')
    if agency_filter:olds=[c for c in olds if c.agency==agency_filter]
    if not olds:raise ValueError('対象代理店条件なし')
    snapshot=_hash(dict(manifest=manifest,previous=previous_period,target=target_period,agency=agency_filter))
    input_rows={name:{} for name in INPUT_SCHEMAS}
    if review_input:
        input_rows,saved_manifest=read_inputs(review_input)
        if saved_manifest!=manifest:raise ValueError('原本またはPhase 1台帳の差替え。承認を流用できません')
    proposals=make_proposals(olds,previous,current,target_period)
    allowed={p.id for p in proposals}
    for name,records in input_rows.items():
        if set(records)-allowed:raise ValueError('対象外・不明の提案IDが入力にあります: '+name)
    conflicts=find_conflicts(proposals,input_rows,snapshot)
    results=[]
    for p in proposals:
        records={}
        for name,headers in INPUT_SCHEMAS.items():
            r=input_rows[name].get(p.id) or empty_input(headers,p,snapshot,target_period)
            if r['agency']!=p.agency or r['condition_id']!=(p.old.id if p.old else p.new_id):
                raise ValueError('承認入力の代理店・条件IDが不一致')
            if r['snapshot_hash']!=snapshot:raise ValueError('対象月・代理店範囲・入力資料が変わっています')
            records[name]=r
        result=evaluate_proposal(p,records,snapshot,target_period,conflicts.get(p.id,()),previous_period)
        for name,r in records.items():
            if name=='最終承認入力' and r.get('approval_status') not in ('APPROVED','REJECTED'):
                r['candidate_hash']=result['candidate_hash']
            r['record_hash']=input_hash(name,r)
            input_rows[name][p.id]=r
        results.append(result)
    issue_groups,issue_members,triage_stats=build_triage(results,previous_period,target_period)
    book=Workbook();book.remove(book.active)
    stats=dict(agency_count=len(set(c.agency for c in olds)),phase1_condition_count=len(olds),
        inherited_phase1_holds=sum(bool(c.hold) for c in olds),inherited_unknown_tax=sum(c.tax!='税別' for c in olds),
        proposal_count=len(results),categories=dict(Counter(r['category'] for r in results)),
        calculation_proposals=sum(r['proposed'] is not None for r in results),
        approval_states=dict(Counter(r['status'] for r in results)),target_period=target_period,
        phase='PHASE2_REVIEW',phase3_started=False,triage=triage_stats)
    _write_sheet(book,'利用説明',['項目','内容'],[
        ['状態','Phase 2内部確認用。分類案・計算提案は確定ではありません。最終状態APPROVEDのみ承認済み。'],
        ['編集する場所','黄色の条件対応入力・粗利承認入力・保留解消入力・最終承認入力。表示シートの編集は再計算に反映しません。'],
        ['承認手順','入力後reviewで別Excelへ再計算しrecord_hashを確認。レビュー担当者がapproved_hashへその値を転記し承認者・日時・根拠・理由を記入。APPROVEDを指定して再度review。'],
        ['最終承認','計算・保留解消後、最終承認入力のcandidate_hashを確認。approved_amountは空欄なら提案額、変更時も目標粗利以上が必須。変更後は再計算しrecord_hashの確認・転記・最終承認が必要。'],
        ['Phase1保留','120行の保留・222行の税区分未明示を条件IDで引継。resolved_issuesにPHASE1_HOLD;TAX_UNKNOWN等を指定し根拠付きで個別承認。空欄や状態の書換えだけで解消しません。'],
        ['初期候補','10月の一意に対応する両単価から、費用等控除前の初期粗利候補を算出。条件付き11月候補は未確定で、承認済み粗利・正式提案には自動転記しません。'],
        ['計算','新上位報酬－承認済み直接原価－承認済み目標粗利。上位と旧代理店単価との差額を粗利とみなさず、原価不明を0円としません。'],
        ['分類','継承・変更・新規は条件案。上位からの未掲載だけでは終了としません。ダッシュ・0・未掲載を区別します。継承候補も条項同等性の承認が必要。'],
        ['新規行','上位の条件追加は代理店への採用が未承認。既存条件との分割・重複候補は別途保留。新規の粗利も承認入力なしでは提案しません。'],
        ['必要な確認','scope_reviewはISP・タイプ・特典コード・単位・期間・別紙・集計主体・階層方式・併給排他・支払・適用除外・戻入を確認した場合のみCONFIRMED。'],
        ['非報酬','控除・違約金は粗利維持計算へ混入させず、税・対応の個別承認後に原本金額を条件案とします。'],
        ['上位原文','支払・除外・戻入・別紙をページ単位で全件保持。上位条件一覧は未解析の結合セルも明示して表示。11月keioh完成Wordの金額を逆算根拠に使っていません。'],
        *[[k,json.dumps(v,ensure_ascii=False)] for k,v in stats.items()]])
    write_triage(book,results,issue_groups,issue_members,triage_stats,_write_sheet)
    from .phase2_november import write_november
    stats['november_master']=write_november(book,results,current,issue_groups,issue_members,snapshot,review_input,_write_sheet)
    _write_sheet(book,'代理店別集計',['代理店','10月条件','Phase1保留','税区分未明示','条件案','計算提案','最終承認済'],[
        [agency,sum(c.agency==agency for c in olds),sum(c.agency==agency and bool(c.hold) for c in olds),
         sum(c.agency==agency and c.tax!='税別' for c in olds),sum(r['proposal'].agency==agency for r in results),
         sum(r['proposal'].agency==agency and r['proposed'] is not None for r in results),
         sum(r['proposal'].agency==agency and r['status']=='APPROVED' for r in results)] for agency in sorted(set(c.agency for c in olds))])
    _write_sheet(book,'11月条件案',VIEW_HEADERS,[view_row(r) for r in results])
    for i,agency in enumerate(sorted(set(c.agency for c in olds)),1):
        _write_sheet(book,f'代理店{i:02}',VIEW_HEADERS,[view_row(r) for r in results if r['proposal'].agency==agency])
    for name,headers in INPUT_SCHEMAS.items():
        _write_sheet(book,name,headers,[[r.get(k) for k in headers] for r in input_rows[name].values()],True)
    _write_sheet(book,'Phase1引継',['代理店','条件ID','保留理由','税区分','承認影響'],[
        [c.agency,c.id,c.hold,c.tax,'条件対応・保留解消・税抜確認まで計算と最終承認を禁止'] for c in olds if c.hold or c.tax!='税別'])
    _write_sheet(book,'上位条件一覧',['対象月','上位ID','成果項目・条件軸','金額','金額状態','税区分','単位','期間','条件原文','根拠資料','位置','SHA256','構造上の要確認','コンポーネント'],[
        [period,u.id,json.dumps(u.axes,ensure_ascii=False),u.amount,u.state,u.tax,u.unit,u.period,u.terms,u.source,u.location,u.digest,u.hold,','.join(u.components)]
        for period,bucket in ((previous_period,previous),(target_period,current)) for u in bucket])
    _write_sheet(book,'上位原文',['根拠資料','ページ','原文'],raw)
    _write_sheet(book,'10月支払除外戻入',['代理店','条件ID','原本根拠','条項全文'],[
        [c.agency,c.id,c.source+'#'+'; '.join(c.locations),c.legal] for c in olds])
    _write_sheet(book,'資料指紋',['種類','資料','SHA256'],manifest)
    from .phase2_pricing import write_pricing
    stats['pricing'],pricing_audit=write_pricing(book,results,current,snapshot,review_input,_write_sheet,previous_period,target_period)
    output.parent.mkdir(parents=True,exist_ok=True)
    with output.open('xb') as stream:book.save(stream)
    # Persist the decision result independently of the editable workbook, linked to exact input/output bytes.
    decisions=[dict(proposal_id=r['proposal'].id,agency=r['proposal'].agency,
        previous_condition_id=r['proposal'].old.id if r['proposal'].old else None,
        condition_id=r['proposal'].new_id,candidate_hash=r['candidate_hash'],status=r['status'],
        approved_amount=str(r['approved_amount']) if r['approved_amount'] is not None else None,
        blockers=r['blocks']+r['decision_errors'],initial_candidate=asdict(r['initial_candidate']),
        review_issue_ids=[i.id for i in r['review_issues']],approval=input_rows['最終承認入力'][r['proposal'].id]) for r in results]
    check=load_workbook(output,read_only=True,data_only=False)
    try:
        values=list(check['11月条件案'].iter_rows(min_row=2,values_only=True))
        expected=[[float(v) if isinstance(v,Decimal) else v for v in view_row(r)] for r in results]
        normalize=lambda rows:[[None if x=='' else x for x in row] for row in rows]
        if normalize(values)!=normalize(expected):raise ValueError('Phase 2 Excel保存後の内容不一致')
    finally:check.close()
    with audit.open('x',encoding='utf-8') as stream:
        json.dump(dict(schema='phase2-approval-v1',snapshot_hash=snapshot,workbook_sha256=file_hash(output),
                       target_period=target_period,summary=stats,decisions=decisions,pricing_decisions=pricing_audit),stream,ensure_ascii=False,indent=2,default=str)
    return stats
