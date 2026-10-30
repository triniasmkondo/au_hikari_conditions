"""Evidence-preserving campaign extraction; produces UNREVIEWED rows only."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date
import hashlib
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from .campaign_workbook import CAMPAIGN_HEADERS, CAMPAIGN_SHEET
from .p2 import P2InputError


@dataclass(frozen=True)
class ExtractedCampaign:
    name: str
    terms: str
    tiers: tuple[tuple[int, int, str], ...]
    valid_from: date | None
    valid_until: date | None
    tax_basis: str


@dataclass(frozen=True)
class CampaignCandidates:
    rows: tuple[dict, ...]
    deductions: tuple[dict, ...]
    holds: tuple[str, ...]


def _normal(text: str) -> str:
    return unicodedata.normalize("NFKC", text).strip()


def _key(text: str) -> str:
    return re.sub(r"\s+", "", _normal(text)).lstrip("■")


def _blocks(path: Path) -> list[tuple[str, str]]:
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader
        return [(f"page:{i};line:{j}", line) for i, page in enumerate(PdfReader(path).pages, 1)
                for j, line in enumerate((page.extract_text() or "").splitlines(), 1)]
    if path.suffix.lower() == ".docx":
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(path)
        blocks = []
        paragraph_index = table_index = 0
        for element in doc.element.body:
            if element.tag == qn("w:p"):
                paragraph_index += 1
                blocks.append((f"paragraph:{paragraph_index}", Paragraph(element, doc).text))
            elif element.tag == qn("w:tbl"):
                table_index += 1
                for row_index, row in enumerate(Table(element, doc).rows, 1):
                    blocks.append((f"table:{table_index};row:{row_index}", " ".join(cell.text for cell in row.cells)))
        return blocks
    raise P2InputError("キャンペーン原本はPDFまたはWordを指定してください。")


def extract_campaign_document(path: Path) -> tuple[ExtractedCampaign, ...]:
    """Keep document order, original terms, source locations, and signed prices.

    This recognizes the tabular campaign notices, not arbitrary contracts/OCR.
    Unrecognized or ambiguous data must be reviewed; no fuzzy price matching.
    """
    blocks = _blocks(path)
    whole = "\n".join(text for _, text in blocks)
    period = re.search(r"(\d{4})年(\d{1,2})月(\d{1,2})日[~〜](\d{4})年(\d{1,2})月(\d{1,2})日", _key(whole))
    first = last = None
    if period:
        first = date(*(int(period.group(i)) for i in (1, 2, 3)))
        last = date(*(int(period.group(i)) for i in (4, 5, 6)))
    heading = re.compile(r"^(?:■)?(?:ボリュームインセンティブ|販促支援金|特別調整金|後確認手数料|消費者センター案件)(?:[（(].*[)）])?$")
    price = re.compile(r"([\d,]+)件[~〜]\s*([+−-]?[\d,]+)円/件")
    sections: list[tuple[str, list[tuple[str, str]]]] = []
    general: list[str] = []
    for location, raw in blocks:
        normalized = _normal(raw)
        if heading.fullmatch(normalized):
            sections.append((_key(normalized), [(location, raw)]))
        elif normalized.startswith("【別紙】"):
            break
        elif sections:
            sections[-1][1].append((location, raw))
        else:
            general.append(raw)
    result = []
    names = set()
    for name, section in sections:
        if name in names:
            raise P2InputError("同名のキャンペーン表が重複しています。")
        names.add(name)
        tiers = []
        for location, raw in section:
            match = price.search(re.sub(r"\s+", "", _normal(raw)))
            if match:
                tiers.append((int(match.group(1).replace(",", "")), int(match.group(2).replace(",", "").replace("−", "-")), location))
        if len({threshold for threshold, _, _ in tiers}) != len(tiers):
            raise P2InputError("同一キャンペーンに重複した閾値があります。")
        terms = "\n".join(general + [raw for _, raw in section])
        section_periods = re.findall(r"(\d{4})年(\d{1,2})月(\d{1,2})日[~〜](\d{4})年(\d{1,2})月(\d{1,2})日", _key("\n".join(raw for _, raw in section)))
        if len(set(section_periods)) > 1:
            raise P2InputError("同一キャンペーンに複数の適用期間があります。")
        section_first, section_last = first, last
        if section_periods:
            values = [int(value) for value in section_periods[0]]
            section_first, section_last = date(*values[:3]), date(*values[3:])
        # Each table must explicitly identify its tax basis.
        tax = "EXCLUSIVE" if any("税別" in _normal(raw) for _, raw in section) else "UNKNOWN"
        result.append(ExtractedCampaign(name, terms, tuple(tiers), section_first, section_last, tax))
    if not result:
        raise P2InputError("キャンペーン表を識別できません。原本を確認してください。")
    return tuple(result)


def build_campaign_candidates(
    upstream_path: Path, agency_path: Path, *, agency_id: str, target_period: str,
    previous_agency_path: Path | None = None,
) -> CampaignCandidates:
    from .p2 import _parse_target_period
    expected_start, expected_end = _parse_target_period(target_period)
    if not agency_id.strip():
        raise P2InputError("代理店IDがありません。")
    upstream = extract_campaign_document(upstream_path)
    agency = {item.name: item for item in extract_campaign_document(agency_path)}
    previous = {item.name: item for item in extract_campaign_document(previous_agency_path)} if previous_agency_path else {}
    upper_names = {item.name for item in upstream}
    holds = ["全行は未承認候補です。正式ID・開通区分・原価・支払時期・条件文のレビューと承認が必要です。"]
    for name in agency.keys() - upper_names:
        holds.append(f"上位原本と対応しない代理店キャンペーンがあります: {name}")
    rows, deductions = [], []
    upper_hash = hashlib.sha256(upstream_path.read_bytes()).hexdigest()
    agency_hash = hashlib.sha256(agency_path.read_bytes()).hexdigest()
    previous_hash = hashlib.sha256(previous_agency_path.read_bytes()).hexdigest() if previous_agency_path else ""
    for index, upper in enumerate(upstream, 1):
        current = agency.get(upper.name)
        old = previous.get(upper.name)
        if not upper.tiers:
            holds.append(f"上位の閾値・単価を読めません: {upper.name}")
            continue
        if not current or not current.tiers:
            holds.append(f"代理店の同名キャンペーンを対応付けられません: {upper.name}")
        current_tiers = {t: (a, loc) for t, a, loc in current.tiers} if current else {}
        old_tiers = {t: (a, loc) for t, a, loc in old.tiers} if old else {}
        thresholds = sorted({t for t, _, _ in upper.tiers} | set(current_tiers))
        upper_tiers = {t: (a, loc) for t, a, loc in upper.tiers}
        if set(upper_tiers) != set(current_tiers):
            holds.append(f"上位と代理店の閾値一覧が一致しません: {upper.name}")
        period_ok = upper.valid_from == expected_start and upper.valid_until == expected_end and current is not None and current.valid_from == expected_start and current.valid_until == expected_end
        if not period_ok:
            holds.append(f"原本の適用期間が対象月と一致しません: {upper.name}")
        for threshold in thresholds:
            amount, location = upper_tiers.get(threshold, (None, ""))
            agency_amount, agency_location = current_tiers.get(threshold, (None, ""))
            old_amount, old_location = old_tiers.get(threshold, (None, ""))
            if (amount is not None and amount < 0) or (agency_amount is not None and agency_amount < 0):
                deductions.append(dict(campaign_name=upper.name, threshold_value=threshold,
                    upstream_unit_amount=amount, agency_unit_amount=agency_amount,
                    source_document=str(upstream_path.resolve()), source_sha256=upper_hash,
                    evidence_location=location, agency_source_document=str(agency_path.resolve()),
                    agency_source_sha256=agency_hash, agency_evidence_location=agency_location,
                    approval_status="UNREVIEWED"))
                continue
            row = dict.fromkeys(CAMPAIGN_HEADERS)
            row.update(
                record_id=f"candidate-{upper_hash[:12]}-{agency_hash[:12]}-{index}-{threshold}",
                agency_id=agency_id, campaign_name=upper.name,
                threshold_value=threshold, threshold_values=";".join(str(t) for t in thresholds),
                threshold_metric="登録件数" if "キャリアエントリ" in upper.terms else None,
                threshold_unit="件", threshold_application="reached_tier" if current and "対応する支払いレートを適用" in current.terms else "unknown",
                counting_scope=("商流全体: " + upper.name) if "全体ボリューム" in upper.terms else None,
                payout_scope=(agency_id + "の対象開通案件: " + upper.name) if current and "貴社開通案件" in current.terms else None,
                valid_from=expected_start.isoformat() if period_ok else None, valid_until=expected_end.isoformat() if period_ok else None,
                eligibility_terms="上位条件\n" + upper.terms + "\n代理店条件\n" + (current.terms if current else "未対応"),
                exclusions="\n".join(line for line in (upper.terms + "\n" + (current.terms if current else "")).splitlines() if any(term in line for term in ("取消", "解約", "支払対象外", "適用除外"))) or None, payment_timing=None,
                upstream_unit_amount=amount, upstream_amount_state="VALUE" if amount else ("ZERO" if amount == 0 else "UNKNOWN"),
                upstream_tax_basis=upper.tax_basis,
                agency_unit_amount=agency_amount, agency_amount_state="VALUE" if agency_amount else ("ZERO" if agency_amount == 0 else "UNKNOWN"),
                agency_tax_basis=current.tax_basis if current else "UNKNOWN",
                previous_agency_unit_amount=old_amount,
                previous_amount_state="VALUE" if old_amount else ("ZERO" if old_amount == 0 else "UNKNOWN"),
                previous_tax_basis=old.tax_basis if old else "UNKNOWN",
                previous_period=old.valid_from.strftime("%Y-%m") if old and old.valid_from else None,
                direct_cost_state="UNKNOWN", direct_cost_tax_basis="UNKNOWN",
                source_document=str(upstream_path.resolve()), source_sha256=upper_hash, evidence_location=location,
                agency_source_document=str(agency_path.resolve()), agency_source_sha256=agency_hash,
                agency_evidence_location=agency_location,
                previous_source_document=str(previous_agency_path.resolve()) if old else None,
                previous_source_sha256=previous_hash if old else None, previous_evidence_location=old_location,
                approval_status="UNREVIEWED",
            )
            rows.append(row)
    return CampaignCandidates(tuple(rows), tuple(deductions), tuple(holds))


def write_campaign_candidate_workbook(
    input_path: Path, output_path: Path, candidates: CampaignCandidates,
) -> None:
    """Extend a COPY of v4; preserve the original and all existing approvals."""
    repository = Path(__file__).resolve().parents[2]
    if output_path.resolve().is_relative_to(repository):
        raise P2InputError("実データの出力先はGit外を指定してください。")
    if input_path.resolve() == output_path.resolve() or output_path.exists():
        raise FileExistsError("既存ブックは上書きしません。")
    workbook = load_workbook(input_path)
    try:
        from .p2 import _REQUIRED_SHEETS, _read_sheet
        for name, headers in _REQUIRED_SHEETS.items():
            _read_sheet(workbook, name, headers)
        if CAMPAIGN_SHEET in workbook.sheetnames and any(
            any(value is not None for value in row) for row in workbook[CAMPAIGN_SHEET].iter_rows(min_row=2, values_only=True)
        ):
            raise P2InputError("既存キャンペーン行を含むブックには取り込みません。別の入力を使用してください。")
        sheet = workbook[CAMPAIGN_SHEET] if CAMPAIGN_SHEET in workbook.sheetnames else workbook.create_sheet(CAMPAIGN_SHEET)
        for i, header in enumerate(CAMPAIGN_HEADERS, 1):
            sheet.cell(1, i, header)
        # No production approval or official item ID is created here.
        for row in candidates.rows:
            if row.get("approval_status") != "UNREVIEWED" or row.get("approved_by") or row.get("approved_at"):
                raise P2InputError("取込候補は未承認である必要があります。")
            sheet.append([row.get(header) for header in CAMPAIGN_HEADERS])
        notes = workbook.create_sheet("キャンペーン取込保留")
        notes.append(["hold_reason"])
        for reason in candidates.holds:
            notes.append([reason])
        if candidates.deductions:
            deductions = workbook.create_sheet("キャンペーン控除候補")
            headers = tuple(candidates.deductions[0])
            deductions.append(list(headers))
            for row in candidates.deductions:
                deductions.append([row.get(header) for header in headers])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("xb") as output:
            workbook.save(output)
    finally:
        workbook.close()
    saved = load_workbook(output_path, read_only=True, data_only=False)
    try:
        from .p2 import _read_sheet
        readback = _read_sheet(saved, CAMPAIGN_SHEET, CAMPAIGN_HEADERS)
        if len(readback) != len(candidates.rows):
            raise P2InputError("候補ブックの保存後行数が一致しません。")
        for actual, expected in zip(readback, candidates.rows, strict=True):
            from .p2 import compute_record_hash
            if compute_record_hash(CAMPAIGN_SHEET, actual) != compute_record_hash(CAMPAIGN_SHEET, expected):
                raise P2InputError("候補ブックの保存後内容が一致しません。")
    finally:
        saved.close()
